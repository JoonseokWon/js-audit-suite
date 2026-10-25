"""Formula-free Korean workbooks, explicit confirmation and immutable envelopes."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED
import json
import unittest
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from common import AuditError, digest, dumps
from onboarding import MODULES
from onboarding.demo import CONTEXT, ledger_result
from onboarding.integration import sample, run
from onboarding.confirmation import confirm_input
from onboarding.workbooks import save_input, load_input, save_ledger
from onboarding.workbook_schema import descriptors


class GatewayFixture(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.folder=Path(self.temp.name);self.context=dict(CONTEXT)
        self.ledger=ledger_result(self.context)

    def workbook(self,module='유형자산증감'):
        path=self.folder/(str(len(list(self.folder.glob('*.xlsx'))))+'.xlsx')
        save_input(module,sample(module,self.context),path)
        return path

    def confirm(self,module,payload,originals=()):
        return confirm_input(module,payload,'SYNTHETIC ONLY input reviewer',
            'Synthetic mapping/adjustment fixture only; no audit performance or final approval',
            originals,confirm_adjustments=True)

    def write_json(self,payload,name='confirmed.json'):
        path=self.folder/name;path.write_text(dumps(payload),encoding='utf-8');return path

    @staticmethod
    def mutate_book(path,callback):
        workbook=load_workbook(path);callback(workbook);workbook.save(path);workbook.close()


class GatewayRoundTripTests(GatewayFixture):
    def test_all_seven_korean_templates_roundtrip_and_reconcile_only_after_explicit_confirmation(self):
        for module in MODULES:
            with self.subTest(module=module):
                template=sample(module,self.context);self.assertEqual(template['approval'],{})
                path=self.workbook(module);payload=load_input(module,path,self.context)
                self.assertEqual(payload['approval'],{})
                self.assertEqual(payload['context'],self.context)
                result=run(module,self.write_json(self.confirm(module,payload)),self.context,self.ledger)
                self.assertEqual(result['blocking'],0,result['issues'])
                self.assertTrue(result['detail']['현재원장연결'])
                for desc in descriptors(module,template):
                    self.assertEqual(len(desc['headers']),len(set(desc['headers'])))
                    self.assertNotEqual(desc['sheet'],desc['key'])

    def test_confirmation_is_required_and_does_not_approve_audit_or_mutate_template(self):
        module='유형자산증감';payload=sample(module,self.context);before=deepcopy(payload)
        with self.assertRaises(AuditError):confirm_input(module,payload,'','reason')
        with self.assertRaises(AuditError):confirm_input(module,payload,'actor','')
        confirmed=self.confirm(module,payload)
        self.assertEqual(payload,before)
        self.assertTrue(confirmed['approval']['input_hash'])
        self.assertNotIn('reviewer',confirmed['approval'])
        unconfirmed=run(module,self.write_json(payload),self.context,self.ledger)
        self.assertGreater(unconfirmed['blocking'],0)

    def test_empty_optional_tables_keep_headers_and_accept_new_rows(self):
        path=self.workbook('감사표본선정')
        def add_row(book):
            sheet=book['발견차이'];headers=[c.value for c in sheet[1]]
            values={'행ID':'E1','원본근거':'SYNTHETIC ONLY/evidence','원장행연결':'J01/L1',
                '모집단행ID':'POP-J01-L1','장부금액':100000,'사람입력검사금액':99999,
                '검토입력자':'SYNTHETIC ONLY','사유':'synthetic difference'}
            sheet.append([values[k] for k in headers])
        self.mutate_book(path,add_row)
        payload=load_input('감사표본선정',path,self.context)
        self.assertEqual(payload['errors'][0]['gl_refs'],['J01/L1'])
        result=run('감사표본선정',self.write_json(self.confirm('감사표본선정',payload)),self.context,self.ledger)
        self.assertEqual(result['metrics']['발견차이순액'],-1)
        self.assertEqual(result['metrics']['발견차이절대합'],1)

    def test_existing_workbook_is_not_overwritten(self):
        path=self.workbook();before=path.read_bytes()
        with self.assertRaises(AuditError):save_input('유형자산증감',sample('유형자산증감',self.context),path)
        self.assertEqual(before,path.read_bytes())

    def test_general_input_confirmation_does_not_approve_adjustments_without_explicit_choice(self):
        module='실무자료·원장';payload=sample(module,self.context)
        payload['adjustments']=[dict(adjustment_id='SYNTHETIC-A1',line_id='1',account_code='1200',
            debit=100,credit=0,evidence='SYNTHETIC ONLY',status='승인'),
            dict(adjustment_id='SYNTHETIC-A1',line_id='2',account_code='2000',debit=0,credit=100,
                 evidence='SYNTHETIC ONLY',status='승인')]
        confirmed=confirm_input(module,payload,'SYNTHETIC ONLY','synthetic general mapping')
        self.assertEqual(confirmed['adjustment_approvals'],[])
        explicit=self.confirm(module,payload)
        self.assertEqual(len(explicit['adjustment_approvals']),1)
        self.assertEqual(payload['adjustment_approvals'],[])


class GatewayIntegrityTests(GatewayFixture):
    def test_formula_is_rejected_even_with_true_cached_result(self):
        path=self.workbook();buffer=BytesIO()
        with ZipFile(path) as archive,ZipFile(buffer,'w',ZIP_DEFLATED) as out:
            for item in archive.infolist():
                raw=archive.read(item.filename)
                if item.filename=='xl/worksheets/sheet2.xml':
                    root=ET.fromstring(raw);ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    cell=root.find(".//m:c[@r='B2']",ns);cell.clear();cell.set('r','B2');cell.set('t','b')
                    ET.SubElement(cell,'{'+ns['m']+'}f').text='TRUE()'
                    ET.SubElement(cell,'{'+ns['m']+'}v').text='1';raw=ET.tostring(root)
                out.writestr(item,raw)
        path.write_bytes(buffer.getvalue())
        with self.assertRaisesRegex(AuditError,'수식|오류'):load_input('유형자산증감',path,self.context)

    def test_external_link_parts_are_rejected_without_execution(self):
        path=self.workbook()
        with ZipFile(path,'a',ZIP_DEFLATED) as archive:
            archive.writestr('xl/externalLinks/externalLink1.xml','<externalLink/>')
        with self.assertRaisesRegex(AuditError,'외부연결'):load_input('유형자산증감',path,self.context)

    def test_context_mismatch_is_rejected(self):
        path=self.workbook()
        with self.assertRaises(AuditError):load_input('유형자산증감',path,dict(self.context,회사코드='OTHER'))

    def test_metadata_cannot_redefine_types_or_labels(self):
        for column,value in ((8,'list'),(7,'다른 의미의 열')):
            path=self.workbook()
            def tamper(book):
                meta=book['_양식정보']
                row=next(n for n in range(2,meta.max_row+1) if meta.cell(n,4).value=='assets'
                         and meta.cell(n,6).value=='cost_account')
                old=meta.cell(row,7).value;meta.cell(row,column,value)
                if column==7:
                    sheet=book['자산기초'];col=next(c.column for c in sheet[1] if c.value==old)
                    sheet.cell(1,col,value)
            self.mutate_book(path,tamper)
            with self.subTest(column=column),self.assertRaises(AuditError):load_input('유형자산증감',path,self.context)

    def test_metadata_cannot_drop_a_required_column_from_an_empty_table(self):
        path=self.workbook('감사표본선정')
        def tamper(book):
            meta=book['_양식정보'];row=next(n for n in range(2,meta.max_row+1)
                if meta.cell(n,4).value=='errors' and meta.cell(n,6).value=='reviewer')
            label=meta.cell(row,7).value;meta.delete_rows(row)
            sheet=book['발견차이'];col=next(c.column for c in sheet[1] if c.value==label);sheet.delete_cols(col)
        self.mutate_book(path,tamper)
        with self.assertRaises(AuditError):load_input('감사표본선정',path,self.context)

    def test_metadata_cannot_rename_a_canonical_sheet(self):
        path=self.workbook()
        def tamper(book):
            book['자산기초'].title='다른내용'
            for row in book['_양식정보'].iter_rows(min_row=2):
                if row[3].value=='assets':row[2].value='다른내용'
        self.mutate_book(path,tamper)
        with self.assertRaises(AuditError):load_input('유형자산증감',path,self.context)

    def test_stale_confirmation_and_ledger_basis_are_rejected(self):
        module='유형자산증감';payload=self.confirm(module,sample(module,self.context))
        payload['asset_movements'][0]['source_ref']+=' changed'
        with self.assertRaises(AuditError):run(module,self.write_json(payload),self.context,self.ledger)
        payload=self.confirm(module,sample(module,self.context));changed=deepcopy(self.ledger)
        changed['tables']['입력_총계정원장'][0]['적요']+=' changed'
        result=run(module,self.write_json(payload),self.context,changed)
        self.assertGreater(result['blocking'],0)
        self.assertFalse(result['detail']['현재원장연결'])

    def test_original_source_list_and_envelope_cannot_be_removed_after_confirmation(self):
        module='유형자산증감';source={'id':'SYNTHETIC-SOURCE','sha256':'1'*64}
        original=self.confirm(module,sample(module,self.context),[source])
        for kind in ('remove_sources','empty_sources','remove_both','remove_envelope'):
            payload=deepcopy(original)
            if kind in ('remove_sources','remove_both'):payload.pop('_origin_sources')
            if kind=='empty_sources':payload['_origin_sources']=[]
            if kind in ('remove_both','remove_envelope'):payload['approval'].pop('envelope_hash')
            with self.subTest(kind=kind),self.assertRaises(AuditError):
                run(module,self.write_json(payload),self.context,self.ledger)


if __name__=='__main__':unittest.main()
