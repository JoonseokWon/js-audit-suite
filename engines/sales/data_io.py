"""Bounded, local CSV/XLSX input and formula-safe Excel output."""
from __future__ import annotations
import csv
import io
import zipfile
import re
from datetime import date, datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 1_000
MAX_COLUMNS = 80
FIELDS = {
 'orders': ['order_id','order_line_id','item_code','customer_id','order_date','quantity','unit','unit_price','price_basis','tax_rate','currency'],
 'deliveries': ['delivery_id','delivery_line_id','order_id','order_line_id','item_code','customer_id','delivery_date','acceptance_date','quantity','accepted_quantity','unit'],
 'invoices': ['invoice_id','invoice_line_id','order_id','order_line_id','item_code','customer_id','invoice_date','quantity','unit','unit_price','price_basis','tax_rate','currency','supply_amount','tax_amount','total_amount'],
}
LABELS = {'orders':'주문·계약', 'deliveries':'납품·검수', 'invoices':'청구'}
IMPORT_FIELDS={**FIELDS,
 'ledger':['journal_id','journal_line_id','company_id','account_code','posting_date','invoice_id','invoice_line_id','customer_id','currency','revenue_debit','revenue_credit'],
 'control':['company_id','period_start','period_end','account_codes','expected_rows','debit_total','credit_total','source_reference','extracted_at','prepared_by'],
}

def _cell(value):
    if value is None: return ''
    if isinstance(value, (datetime,date)): return value.strftime('%Y-%m-%d')
    return str(value).strip()

def read_table(data: bytes, filename: str, role: str, sheet_name: str | None = None) -> list[dict]:
    if role not in IMPORT_FIELDS: raise ValueError('알 수 없는 자료 유형입니다')
    if len(data) > MAX_BYTES: raise ValueError('파일당 최대 10MB까지 허용합니다')
    suffix = Path(filename).suffix.lower()
    sheet = 'CSV'
    if suffix == '.xlsx':
        wb = None
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(x.file_size for x in z.infolist()) > 100*1024*1024:
                    raise ValueError('압축 해제 크기가 너무 큽니다')
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            if sheet_name and sheet_name not in wb.sheetnames: raise ValueError(f'시트를 찾을 수 없습니다: {sheet_name}')
            ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
            sheet = ws.title
            # Do not trust producer-supplied dimension metadata. Without this,
            # a small declared range silently hides otherwise valid source rows.
            ws.reset_dimensions()
            raw = []
            positions = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS: raise ValueError('파일당 최대 1,000행까지 허용합니다')
                if len(row) > MAX_COLUMNS: raise ValueError('최대 80개 열까지 허용합니다')
                raw.append(list(row)); positions.append(i+1)
        except (ValueError,): raise
        except Exception as exc: raise ValueError('읽을 수 없는 XLSX 파일입니다') from exc
        finally:
            if wb is not None:
                wb.close()
    elif suffix == '.csv':
        for enc in ('utf-8-sig','cp949'):
            try: decoded=data.decode(enc); break
            except UnicodeDecodeError: pass
        else: raise ValueError('CSV는 UTF-8 또는 CP949 인코딩으로 저장해 주세요')
        try:
            reader = csv.reader(io.StringIO(decoded), strict=True)
            raw=[]
            positions=[]
            previous_end=0
            for i,row in enumerate(reader):
                if i > MAX_ROWS: raise ValueError('파일당 최대 1,000행까지 허용합니다')
                if len(row)>MAX_COLUMNS: raise ValueError('최대 80개 열까지 허용합니다')
                raw.append(row); positions.append(previous_end+1); previous_end=reader.line_num
        except csv.Error as exc: raise ValueError('CSV 형식이 올바르지 않습니다') from exc
    else: raise ValueError('CSV 또는 XLSX 파일만 사용할 수 있습니다')
    if not raw: raise ValueError('빈 파일입니다')
    headers=[_cell(x) for x in raw[0]]
    if len(headers)!=len(set(headers)): raise ValueError('중복된 열 이름이 있습니다')
    missing = [x for x in IMPORT_FIELDS[role] if x not in headers]
    if missing: raise ValueError('필수 열이 없습니다: ' + ', '.join(missing))
    result=[]
    for rowno,row in enumerate(raw[1:],2):
        if not any(_cell(x) for x in row): continue
        if len(row)>len(headers): raise ValueError(f'{rowno}행의 열 수가 제목 행보다 많습니다')
        if any(re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', _cell(x)) for x in row): raise ValueError(f'{positions[rowno-1]}행에 지원하지 않는 제어 문자가 있습니다')
        rec={h:((row[i].isoformat(timespec='seconds')+'+09:00') if h=='extracted_at' and isinstance(row[i],datetime) else _cell(row[i])) if i<len(row) else '' for i,h in enumerate(headers) if h in IMPORT_FIELDS[role]}
        if suffix=='.xlsx':
            idfields={x for x in IMPORT_FIELDS[role] if x.endswith('_id') or x in ('item_code','account_code','account_codes')}
            errors=[f'{h}: XLSX 식별자는 텍스트 셀이어야 합니다 (앞자리 0 손실 가능)' for i,h in enumerate(headers) if h in idfields and i<len(row) and isinstance(row[i],(int,float,bool))]
            if errors: rec['_input_errors']=errors
        rec.update(_source_file=Path(filename).name,_sheet=sheet,_row=positions[rowno-1])
        result.append(rec)
    return result

def template_csv(role: str) -> bytes:
    out=io.StringIO(); csv.writer(out).writerow(IMPORT_FIELDS[role]); return out.getvalue().encode('utf-8-sig')

def _safe(v):
    if v is None: return ''
    if isinstance(v,(list,dict)): import json; v=json.dumps(v,ensure_ascii=False)
    # Prevent Excel formula injection for imported IDs, names and filenames.
    if isinstance(v,str): v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', lambda m: f'\\x{ord(m.group()):02x}',v)
    if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')): return "'" + v
    return v

def export_xlsx(result: dict) -> bytes:
    wb=Workbook(); wb.remove(wb.active)
    sheets=[('검토결과',result['groups']),('사유목록',result['issues']),('수행한검사',result.get('checks',[])),('원본행추적',result['rows']),('실행설정',[result['config']])]
    if 'ledger_reconciliation' in result:sheets += [('원장양방향대사',result['ledger_reconciliation']),('원장통제대사',result['population_checks'])]
    for title,rows in sheets:
        ws=wb.create_sheet(title)
        headers=list(dict.fromkeys(k for row in rows for k in row)) if rows else ['안내']
        ws.append(headers)
        if rows:
            for row in rows: ws.append([_safe(row.get(h,'')) for h in headers])
        else: ws.append(['해당 항목 없음'])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='183D48'); c.alignment=Alignment(wrap_text=True)
        for col in ws.columns:
            letter=col[0].column_letter
            ws.column_dimensions[letter].width=min(55,max(14,max(len(str(c.value or '')) for c in col[:100])+2))
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    ws=wb.create_sheet('읽어주세요')
    for text in [
      '매출 증빙 대사 · 로컬 검토용 프로토타입',
      '결과는 파일에 제공된 증빙 간 기계적 비교입니다. 회계감사 의견·매출 인식·부정 판단을 제공하지 않습니다.',
      'MATCHED=설정 범위에서 수량·금액 일치. REVIEW=추가 검토 필요. ERROR=입력 또는 증빙 식별 오류.',
      'OPEN=일관된 자료에서 미납품·미검수·미청구 잔량이 남은 진행 상태. EXCLUDED=기준일 이후 자료만 있는 그룹.',
      '수량 집계 단위는 주문번호+주문행번호+품목코드입니다. 납품행과 청구행의 개별 배분은 하지 않습니다.',
      '정확히 같은 증빙 ID와 동일 내용은 1회만 합산합니다. 동일 ID의 내용 충돌은 전부 격리합니다.',
      '기준일 이후 자료와 중복·오류·연결 불가 행을 원본행추적에서 함께 확인하세요.',
      '식별자·수치의 정밀도 보존을 위해 결과 수치를 텍스트로 저장합니다. 수식은 포함하지 않습니다.',
    ]: ws.append([text])
    ws.column_dimensions['A'].width=120
    for row in ws:
        row[0].alignment=Alignment(wrap_text=True,vertical='top')
    out=io.BytesIO(); wb.save(out); return out.getvalue()
