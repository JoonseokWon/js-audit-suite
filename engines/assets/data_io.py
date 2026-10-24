"""Bounded XLSX import and atomic static workpaper export for the local program."""
from datetime import date,datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile,BadZipFile
import hashlib,os,tempfile
from .engine import SCHEMA,InputError,analyze

MAX_BYTES=10*1024*1024;MAX_ROWS=1000;MAX_COLS=40
def read_workbook(path):
    from openpyxl import load_workbook
    path=Path(path)
    if path.suffix.lower()!='.xlsx' or path.stat().st_size>MAX_BYTES:raise InputError('10MB 이하 .xlsx 파일만 지원합니다.')
    data=path.read_bytes()
    try:
        with ZipFile(BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist())>50*1024*1024:raise InputError('압축 해제 크기 한도 초과')
            if any(any(t in n.lower() for t in ('externallinks/','vbaproject','embeddings/','connections.xml')) for n in z.namelist()):raise InputError('외부연결·매크로·첨부 개체 미지원')
            from defusedxml.ElementTree import iterparse
            from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
            ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
            for name in z.namelist():
                if not(name.startswith('xl/worksheets/') and name.endswith('.xml')):continue
                previous=0
                for _,element in iterparse(BytesIO(z.read(name)),events=('end',)):
                    if element.tag!=ns+'row':continue
                    n=int(element.attrib.get('r','0'))
                    if n<=previous or n>MAX_ROWS+1:raise InputError('시트 행 순서 오류/1000행 한도 초과')
                    previous=n;previous_col=0
                    for cell in element:
                        if cell.tag!=ns+'c':continue
                        col,rownum=coordinate_from_string(cell.attrib['r']);col=column_index_from_string(col)
                        if rownum!=n or col<=previous_col:raise InputError('셀 좌표 중복/행 순서 오류')
                        previous_col=col
                        if col>MAX_COLS and any(cell.find(ns+t) is not None for t in ('v','is','f')):raise InputError('40열 한도 초과')
                    element.clear()
        w=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            if set(w.sheetnames)!=set(SCHEMA):raise InputError('필수 6개 시트가 정확히 필요합니다: '+', '.join(SCHEMA))
            tables={}
            for name,headers in SCHEMA.items():
                ws=w[name];ws.reset_dimensions();result=[];header_seen=False
                for n,row in enumerate(ws.iter_rows(),1):
                    if n>MAX_ROWS+1:raise InputError(name+': 물리 행 1000개 한도 초과')
                    active=[cell for cell in row if cell.value is not None]
                    if any(cell.column>MAX_COLS for cell in active):raise InputError(name+': 40열 한도 초과')
                    if any(cell.data_type in ('f','e') for cell in active):raise InputError(name+': 수식/오류 셀 미지원 (값으로 제출)')
                    values=[]
                    for cell in row[:MAX_COLS]:
                        v=cell.value
                        if isinstance(v,datetime):v=v.isoformat() if v.time()!=datetime.min.time() else v.date().isoformat()
                        elif isinstance(v,date):v=v.isoformat()
                        if isinstance(v,str) and len(v)>2000:raise InputError(name+': 셀 길이 한도 초과')
                        values.append(v)
                    while values and values[-1] is None:values.pop()
                    if n==1:
                        if tuple(values)!=headers:raise InputError(name+': 열 이름/순서가 표준 양식과 다릅니다.')
                        header_seen=True
                    elif values:
                        if len(values)>len(headers):raise InputError(name+': 정의 밖 입력 열')
                        result.append(dict(zip(headers,values+[None]*(len(headers)-len(values))))|{'_source_row':n})
                if not header_seen:raise InputError(name+': 헤더 없음')
                tables[name]=result
        finally:w.close()
    except (BadZipFile,KeyError) as e:raise InputError('Excel 구조 오류: '+str(e)) from e
    source={'path':str(path.resolve()),'filename':path.name,'sha256':hashlib.sha256(data).hexdigest(),'sheet_rows':{k:len(v) for k,v in tables.items()}}
    return tables,source

def run_file(path,cancel=None):
    tables,source=read_workbook(path)
    result=analyze(tables,source,cancel)
    if hashlib.sha256(Path(path).read_bytes()).hexdigest()!=source['sha256']:raise InputError('실행 도중 입력 파일 변경. 다시 실행하세요.')
    return result

def export_bytes(result):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    w=Workbook();w.remove(w.active)
    def add(name,headers,rows,numbers=()):
        ws=w.create_sheet(name);ws.sheet_view.showGridLines=False;ws.append(list(headers))
        for record in rows:
            ws.append([None if v is None else int(Decimal(v)) if j in numbers else v if isinstance(v,(int,float)) else str(v) for j,v in enumerate(record)])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        for cell in ws[1]:cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17324D')
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
        for j in range(1,len(headers)+1):
            ws.column_dimensions[get_column_letter(j)].width=min(55,max(16,max(len(str(ws.cell(i,j).value or '')) for i in range(1,min(ws.max_row,70)+1))+2))
            for row in ws.iter_rows(min_row=2,min_col=j,max_col=j):
                cell=row[0];cell.font=Font(name='맑은 고딕',size=10);cell.alignment=Alignment(vertical='center')
                if j-1 in numbers:cell.number_format='#,##0;[Red](#,##0);0'
        return ws
    add('비교결과',('상태','자산키','비교','앞금액','뒤금액','차이','허용오차','원본행'),[(f['status'],f['key'],f['check'],f['left'],f['right'],f['difference'],f['tolerance'],f['references']) for f in result['findings']],(3,4,5,6))
    add('예외목록',('중요도','사유','자산키','내용','원본행'),[(i['severity'],i['reason_code'],i['key'],i['message'],i['references']) for i in result['issues']])
    keys=('자산키','월','방법','월초장부','추가원가','상각기준장부','잔존가치','계산전상각','월상각','누계','월말장부','변경근거')
    add('월별계산',keys,[[r[k] for k in keys] for r in result['monthly']],(3,4,5,6,8,9,10))
    keys=('검사','기대','실제','차이','상태','근거');add('원장통제',keys,[[r[k] for k in keys] for r in result['controls']],(1,2,3))
    keys=('비교','건수','차이건수','부호합계','절대합계');add('차이집계',keys,[[r[k] for k in keys] for r in result['difference_summary']],(1,2,3,4))
    headers=('자산키','상태',*('재계산_'+k for k in ('전기말원가','전기말누계','당기상각','기말원가','기말누계','기말장부','처분직전장부','처분손익')))
    add('자산별결과',headers,[(g['key'],g['status'],*[(g['calculated'] or {}).get(k) for k in ('전기말원가','전기말누계','당기상각','기말원가','기말누계','기말장부','처분직전장부','처분손익')]) for g in result['groups']],tuple(range(2,10)))
    movement=('전기말원가','당기취득원가','당기추가원가','당기처분원가','기말원가','전기말누계','당기상각','당기처분누계','기말누계')
    add('원가누계변동',('자산키',*movement),[(g['key'],*[(g['calculated'] or {}).get(k) for k in movement]) for g in result['groups']],tuple(range(1,10)))
    add('실행정보',('항목','값'),list(result['config'].items())+list(result['source'].items())+[('생성시각UTC',result['generated_at_utc']),('범위','KRW 원가모형, 명시적 월할 정책. 계산 일치는 감사결론이 아님.')])
    for name,headers in SCHEMA.items():add('입력_'+name,('원본행',*headers),[(r['row_id'],*[r['raw'].get(k) for k in headers]) for r in result['rows'] if r['sheet']==name])
    b=BytesIO();w.save(b);w.close();return b.getvalue()

def save_result(result,path):
    target=Path(path).resolve()
    if target.suffix.lower()!='.xlsx':raise InputError('.xlsx 파일명 필요')
    if target==Path(result['source']['path']).resolve():raise InputError('입력 원본 덮어쓰기 금지')
    if hashlib.sha256(Path(result['source']['path']).read_bytes()).hexdigest()!=result['source']['sha256']:raise InputError('입력 파일이 변경되었습니다. 다시 실행하세요.')
    data=export_bytes(result);fd,temporary=tempfile.mkstemp(prefix='asset_result_',suffix='.xlsx',dir=target.parent)
    try:
        with os.fdopen(fd,'wb') as stream:stream.write(data)
        os.replace(temporary,target)
    finally:
        if os.path.exists(temporary):os.unlink(temporary)

