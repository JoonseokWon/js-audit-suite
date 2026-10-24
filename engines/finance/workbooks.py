"""Local xlsx templates and static Decimal result snapshots (openpyxl)."""
from __future__ import annotations
from datetime import date,datetime
from decimal import Decimal
from pathlib import Path
import os, tempfile
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from .engine import SCHEMA,CONDITIONS,Result,NOTICE,check_cancel
from .practice_engine import EXTRA_SCHEMA

def append(ws,values):
    ws.append(values)
    for c in ws[ws.max_row]:
        if isinstance(c.value,str):c.data_type='s'

def style(ws,widths=None):
    ws.freeze_panes='A2' if ws.title in ('설정','실행정보') else 'C2'
    ws.auto_filter.ref=ws.dimensions
    ws.sheet_view.showGridLines=False
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor='17324D');c.font=Font(name='맑은 고딕',bold=True,color='FFFFFF',size=11)
        c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height=24
        for c in row:
            c.font=Font(name='맑은 고딕',size=10)
            c.alignment=Alignment(vertical='top',wrap_text=ws.title in ('결과','실행정보'))
            if isinstance(c.value,(date,datetime)):c.number_format='yyyy-mm-dd'
            elif isinstance(c.value,(int,float,Decimal)):c.number_format='#,##0.##########'
    for col in range(1,ws.max_column+1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width=(widths or {}).get(col,18)
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:1'

def save_atomic(wb,path:Path,event=None):
    path=Path(path)
    if path.suffix.lower()!='.xlsx':raise ValueError('저장 확장자는 .xlsx여야 합니다.')
    path.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(prefix='financial_',suffix='.xlsx',dir=path.parent)
    os.close(fd)
    try:
        check_cancel(event)
        wb.save(tmp)
        check_cancel(event)
        os.replace(tmp,path)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)

def base_workbook():
    wb=Workbook();wb.remove(wb.active)
    s=wb.create_sheet('설정');s.append(['항목','값'])
    s.append(['기준일',date(2026,4,11)])
    s.append(['원금허용오차',0]);s.append(['이자허용오차',1])
    for name,headers in SCHEMA.items():
        ws=wb.create_sheet(name);ws.append(headers)
        for col in ('A','B'):ws.column_dimensions[col].number_format='@'
        validations={'D':'"KRW"','E':'"원"'}
        if name=='조건':
            validations.update(F='"단순 정기예금,고정금리 대여,고정금리 차입"',I='"%,소수"',N='"ACT/365F,ACT/360"',O='"포함,제외"',P='"포함,제외"',Q='"아니오,예"',R='"아니오,예"',S='"아니오,예"',T='"아니오,예"')
        for col,formula in validations.items():
            dv=DataValidation(type='list',formula1=formula,allow_blank=False);dv.errorTitle='입력값 확인';dv.error='지원 목록에서 선택하세요.';dv.showErrorMessage=True
            ws.add_data_validation(dv);dv.add(f'{col}2:{col}50001')
    return wb

def demo_rows(key='D001',product='단순 정기예금',principal=1000000,interest=10000):
    common=['BANK01',key,date(2026,4,11),'KRW','원']
    balance=common+[principal,interest]
    condition=common+[product,principal,3.65,'%',date(2026,1,1),date(2026,12,31),date(2026,1,1),date(2026,4,11),'ACT/365F','포함','제외','아니오','아니오','아니오','아니오']
    return balance,condition

def make_samples(folder:Path):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    wb=base_workbook()
    for ws in wb:style(ws)
    save_atomic(wb,folder/'입력템플릿.xlsx')
    wb=base_workbook()
    for key,product,p,i in [('D001','단순 정기예금',1000000,10000),('L001','고정금리 대여',500000,5000),('B001','고정금리 차입',2000000,20000)]:
        balance,condition=demo_rows(key,product,p,i)
        for name in ('명세','원장','회신'):append(wb[name],balance)
        append(wb['조건'],condition)
    for ws in wb:style(ws)
    save_atomic(wb,folder/'더미_정상.xlsx')
    wb=base_workbook()
    for key in ('D001','D002','D003','D004','D005','D006','D007','D008'):
        b,c=demo_rows(key)
        if key=='D007':c[5]='파생상품'
        if key=='D005':c[11]='2026-02-30'
        if key=='D008':b[-2:]=[0,0];c[6]=0;c[7]=0
        append(wb['명세'],b);append(wb['조건'],c);append(wb['원장'],b)
        if key!='D003':
            reply=list(b)
            if key=='D002':reply[5]=999900;reply[6]=9990
            if key=='D006':reply[2]=date(2026,4,10)
            append(wb['회신'],reply)
        if key=='D004':append(wb['원장'],b)
    for ws in wb:style(ws)
    save_atomic(wb,folder/'더미_차이오류.xlsx')

def export_result(result:Result,path:Path,event=None):
    wb=Workbook();ws=wb.active;ws.title='결과'
    append(ws,['상태','기관코드','계약ID','검사','입력행','설명','앞금액(원)','뒤금액(원)','차이(원)','허용오차(원)','계산 근거','검토메모'])
    for f in result.findings:
        check_cancel(event)
        append(ws,[f.status,*(f.key or ('','')),f.check,f.references,f.message,*[int(v) if v is not None else None for v in (f.left,f.right,f.difference,f.tolerance)],f.details,''])
        for c in ws[ws.max_row]:
            if f.status not in ('금액일치','기간확인'):c.fill=PatternFill('solid',fgColor='FFF2E6' if f.status not in ('차이','기간오류') else 'FCE8E6')
    style(ws,{1:14,2:16,3:16,4:24,5:24,6:65,7:22,8:22,9:22,10:22,11:100,12:65})
    for i in range(2,ws.max_row+1):ws.row_dimensions[i].height=50
    meta=wb.create_sheet('실행정보');append(meta,['항목','값'])
    for k,v in [('입력파일',result.input_path),('입력 SHA256',result.input_sha256),('실행시각',result.generated),('범위',NOTICE),('출력','전체 결과 스냅샷. 화면 필터와 무관. 원본 변경 후 다시 실행'),('키','기관코드+계약ID 정확 텍스트 키. 잔액·조건·이자기준은 한 행. 이자구간은 구간ID별 여러 행'),('통화·단위','KRW·원. 차입도 양수. 변환/합산 없음'),('이자 계산','Decimal 정밀도40, ACT/365F 또는 ACT/360, 원 단위 ROUND_HALF_UP'),('차이 방향','앞금액-뒤금액. 명세이자-단순이자. 차이 0원은 일치. 비영차이는 허용범위 이내도 반올림차이로 보존'),*result.settings.items(),*[('건수 '+k,v) for k,v in result.counts.items()]]:append(meta,[k,v if not isinstance(v,Decimal) else int(v)])
    style(meta,{1:24,2:110})
    for i in range(2,meta.max_row+1):meta.row_dimensions[i].height=34
    for title,items in [('기간검사',getattr(result,'period_reports',[])),('비교별차이집계',getattr(result,'difference_summary',[]))]:
        if items:
            view=wb.create_sheet(title);headers=list(items[0]);append(view,headers)
            for item in items:append(view,[int(item[h]) if isinstance(item[h],Decimal) else item[h] for h in headers])
            style(view,{len(headers):100})
    for name,headers in {**SCHEMA,**(EXTRA_SCHEMA if hasattr(result,'period_reports') else {})}.items():
        source=wb.create_sheet('입력_'+name);append(source,['원본행',*headers,'검증오류','미지원'])
        for r in result.rows:
            if r.sheet==name:append(source,[r.row,*[r.raw.get(h) for h in headers],' / '.join(r.errors),' / '.join(r.unsupported)])
        style(source,{1:12,len(headers)+2:70,len(headers)+3:50})
    save_atomic(wb,path,event)
