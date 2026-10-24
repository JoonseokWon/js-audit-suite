"""Deterministic financial snapshot reconciliation. No audit conclusions."""
from __future__ import annotations
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from threading import Event
from typing import Callable
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook

BALANCE = ['기관코드','계약ID','기준일','통화','단위','원금','이자']
CONDITIONS = ['기관코드','계약ID','기준일','통화','단위','상품','원금','연이율','이율단위','계약시작','만기','이자시작','이자끝','DayCount','시작포함','끝포함','원금변동','복리','변동금리','수수료할인할증']
SCHEMA = {'명세':BALANCE,'조건':CONDITIONS,'원장':BALANCE,'회신':BALANCE}
PRODUCTS = ('단순 정기예금','고정금리 대여','고정금리 차입')
MAX_MONEY=Decimal('999999999999999')
MAX_INPUT_ROWS=50_000
NOTICE = '금액일치는 수치 비교 결과입니다. 인식·분류·평가 적정성 및 감사결론은 사람이 판단합니다.'

class InputError(ValueError): pass
class Cancelled(Exception): pass

def check_cancel(event:Event|None):
    if event and event.is_set(): raise Cancelled('사용자가 취소했습니다.')

def as_date(v):
    if isinstance(v,datetime):
        if v.time().isoformat() != '00:00:00': raise InputError('시간이 있는 날짜는 지원하지 않습니다.')
        return v.date()
    if isinstance(v,date): return v
    if isinstance(v,str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}',v):
        try:return date.fromisoformat(v)
        except ValueError: pass
    raise InputError('Excel 날짜 또는 YYYY-MM-DD 형식의 실제 날짜가 필요합니다.')

def number(v):
    if v is None or isinstance(v,bool): raise InputError('숫자 필수값이 없습니다/잘못되었습니다.')
    if isinstance(v,str) and not re.fullmatch(r'(?:0|[1-9]\d*)(?:\.\d+)?',v):
        raise InputError('쉼표·기호·음수·공백 없는 숫자를 입력하세요.')
    try: d=Decimal(str(v))
    except (InvalidOperation,ValueError):raise InputError('숫자가 아닙니다.')
    if not d.is_finite():raise InputError('유한한 숫자가 필요합니다.')
    return d

def money(v):
    d=number(v)
    if not Decimal(0)<=d<=MAX_MONEY or d!=d.to_integral_value():
        raise InputError('금액은 0~999,999,999,999,999원의 정수만 지원합니다.')
    return d

def simple_interest(principal:Decimal,rate:Decimal,start:date,end:date,basis:str,start_in:bool,end_in:bool):
    days=(end-start).days-1+int(start_in)+int(end_in)
    if days<0:raise InputError('포함 규칙에 따른 이자 일수가 음수입니다.')
    denominator={'ACT/365F':365,'ACT/360':360}[basis]
    with localcontext() as ctx:
        ctx.prec=40
        exact=principal*rate*Decimal(days)/Decimal(denominator)
        rounded=exact.quantize(Decimal('1'),rounding=ROUND_HALF_UP)
    return days,denominator,exact,rounded

@dataclass
class SourceRow:
    sheet:str
    row:int
    raw:dict
    key:tuple[str,str]|None
    parsed:dict=field(default_factory=dict)
    errors:list[str]=field(default_factory=list)
    unsupported:list[str]=field(default_factory=list)
    formulas:list[str]=field(default_factory=list)
    @property
    def ref(self):return f'{self.sheet}!{self.row}'

@dataclass
class Finding:
    status:str
    key:tuple[str,str]|None
    check:str
    references:str
    message:str
    left:Decimal|None=None
    right:Decimal|None=None
    difference:Decimal|None=None
    tolerance:Decimal|None=None
    details:str=''

@dataclass
class Result:
    input_path:str
    input_sha256:str
    settings:dict
    rows:list[SourceRow]
    findings:list[Finding]
    generated:str
    @property
    def counts(self):return dict(Counter(f.status for f in self.findings))

def _worksheet_rows(ws, event=None):
    # Producer-supplied dimensions may be missing, inflated or too small.
    # Read actual cells and enforce the physical row limit during iteration.
    ws.reset_dimensions()
    for rowno,cells in enumerate(ws.iter_rows(),1):
        check_cancel(event)
        if rowno>MAX_INPUT_ROWS+1:
            raise InputError(f'{ws.title}: 최대 {MAX_INPUT_ROWS:,}행을 초과합니다.')
        yield cells


def _headers(ws, expected, event=None):
    iterator=_worksheet_rows(ws,event)
    first=next(iterator,None)
    if first is None:raise InputError(f'{ws.title}: 헤더가 없습니다.')
    values=[c.value for c in first]
    while values and values[-1] is None:values.pop()
    if values!=expected:raise InputError(f'{ws.title}: 1행 열 이름/순서가 양식과 다릅니다. 기대: '+', '.join(expected))
    return iterator

def _read(path:Path,event):
    if path.suffix.lower()!='.xlsx':raise InputError('.xlsx만 지원합니다. .xls/.xlsm/CSV는 지원하지 않습니다.')
    if path.stat().st_size>20*1024*1024:raise InputError('입력 파일은 20MB 이하만 지원합니다.')
    data=path.read_bytes()
    if len(data)>20*1024*1024:raise InputError('입력 파일은 20MB 이하만 지원합니다.')
    with ZipFile(BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist())>100*1024*1024:raise InputError('압축 해제 크기가 100MB를 초과합니다.')
        if any(i.filename.endswith('vbaProject.bin') for i in z.infolist()):raise InputError('매크로 포함 파일은 지원하지 않습니다.')
    check_cancel(event)
    wb=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        required={'설정',*SCHEMA}
        if not required<=set(wb.sheetnames) or set(wb.sheetnames)-required-{'이자기준','이자구간'}:raise InputError('설정·명세·조건·원장·회신 및 v2 이자기준·이자구간 양식을 사용하세요.')
        settings={}
        for cells in _headers(wb['설정'],['항목','값'],event):
            vals=[c.value for c in cells]
            if not any(v is not None for v in vals):continue
            if len(vals)>2 and any(v is not None for v in vals[2:]):raise InputError('설정: 양식 밖 값이 있습니다.')
            key=vals[0]
            if key not in ('기준일','원금허용오차','이자허용오차') or key in settings:raise InputError(f'설정: 알 수 없거나 중복된 항목 {key!r}')
            if any(c.data_type=='f' for c in cells):raise InputError('설정: 수식은 지원하지 않습니다. 확정값을 입력하세요.')
            settings[key]=vals[1] if len(vals)>1 else None
        if set(settings)!= {'기준일','원금허용오차','이자허용오차'}:raise InputError('설정: 기준일·원금허용오차·이자허용오차가 필요합니다.')
        settings={'기준일':as_date(settings['기준일']),'원금허용오차':money(settings['원금허용오차']),'이자허용오차':money(settings['이자허용오차'])}
        rows=[]
        for name,headers in SCHEMA.items():
            ws=wb[name]
            for rowno,cells in enumerate(_headers(ws,headers,event),2):
                check_cancel(event)
                vals=[c.value for c in cells]
                if not any(v is not None for v in vals):continue
                if len(vals)>len(headers) and any(v is not None for v in vals[len(headers):]):raise InputError(f'{name}!{rowno}: 양식 밖 열에 값이 있습니다.')
                vals+= [None]*max(0,len(headers)-len(vals))
                raw=dict(zip(headers,vals))
                a,b=raw['기관코드'],raw['계약ID']
                key=(a,b) if all(isinstance(v,str) and v.strip() for v in (a,b)) else None
                r=SourceRow(name,rowno,raw,key)
                r.formulas=[headers[i] for i,c in enumerate(cells[:len(headers)]) if c.data_type=='f']
                rows.append(r)
                _validate(r,settings['기준일'])
        return data,settings,rows
    finally:wb.close()

def _validate(r:SourceRow,asof:date):
    d=r.raw
    if r.key is None:r.errors.append('기관코드·계약ID는 비어 있지 않은 텍스트여야 합니다. 선행0을 보존하세요.')
    if r.formulas:r.errors.append('수식은 지원하지 않습니다: '+', '.join(r.formulas))
    for k in ('기준일',):
        try:r.parsed[k]=as_date(d[k])
        except InputError as e:r.errors.append(f'{k}: {e}')
    if r.parsed.get('기준일') and r.parsed['기준일']!=asof:r.errors.append('기준일 불일치: 설정 '+asof.isoformat())
    for k,want in (('통화','KRW'),('단위','원')):
        if d[k] is None:r.errors.append(f'{k}: 필수값 누락')
        elif d[k]!=want:r.unsupported.append(f'{k}={d[k]!r}, 지원값={want}')
    for k in ('원금','이자') if r.sheet!='조건' else ('원금',):
        try:r.parsed[k]=money(d[k])
        except InputError as e:r.errors.append(f'{k}: {e}')
    if r.sheet!='조건':return
    if d['상품'] is None:r.errors.append('상품: 필수값 누락')
    elif d['상품'] not in PRODUCTS:r.unsupported.append(f'상품={d["상품"]!r}')
    for k in ('원금변동','복리','변동금리','수수료할인할증'):
        if d[k] not in ('아니오','예'):r.errors.append(f'{k}: 예/아니오 필수')
        elif d[k]=='예':r.unsupported.append(k+'=예')
    if d['DayCount'] not in ('ACT/365F','ACT/360'):
        if d['DayCount'] is None:r.errors.append('DayCount: 필수값 누락')
        else:r.unsupported.append(f'DayCount={d["DayCount"]!r}')
    for k in ('시작포함','끝포함'):
        if d[k] not in ('포함','제외'):r.errors.append(f'{k}: 포함/제외 필수')
        else:r.parsed[k]=d[k]=='포함'
    if d['이율단위'] not in ('%','소수'):r.errors.append('이율단위: %/소수 필수')
    try:
        raw_rate=number(d['연이율'])
        if d['이율단위'] in ('%','소수'):
            rate=raw_rate/Decimal(100) if d['이율단위']=='%' else raw_rate
            if not 0<=rate<=1:raise InputError('연이율은 0~100%만 지원합니다.')
            if len(raw_rate.as_tuple().digits)>12:raise InputError('연이율은 유효숫자 12자리 이하입니다.')
            r.parsed['연이율']=rate
    except InputError as e:r.errors.append(f'연이율: {e}')
    for k in ('계약시작','만기','이자시작','이자끝'):
        try:r.parsed[k]=as_date(d[k])
        except InputError as e:r.errors.append(f'{k}: {e}')
    if all(k in r.parsed for k in ('계약시작','만기','이자시작','이자끝')):
        p=r.parsed
        if not p['계약시작']<=p['이자시작']<=p['이자끝']<=asof<=p['만기']:
            r.errors.append('날짜 순서 필요: 계약시작≤이자시작≤이자끝≤설정 기준일≤만기')
        if all(k in p for k in ('시작포함','끝포함')):
            days=(p['이자끝']-p['이자시작']).days-1+int(p['시작포함'])+int(p['끝포함'])
            if days<0:r.errors.append('시작·끝 포함 규칙에 따른 이자 일수가 음수입니다.')

def reconcile(path:Path|str,event:Event|None=None,progress:Callable[[str],None]|None=None)->Result:
    path=Path(path)
    if progress:progress('입력 검증 중')
    try:data,settings,rows=_read(path,event)
    except (BadZipFile,KeyError,OSError,ValueError) as e:
        if isinstance(e,InputError):raise
        raise InputError(f'Excel 파일을 읽을 수 없습니다: {e}') from e
    findings=[]
    def add(status,key,check,refs,msg,**kw):findings.append(Finding(status,key,check,refs,msg,**kw))
    grouped={s:defaultdict(list) for s in SCHEMA}
    blocked=set()
    for r in rows:
        if r.key:grouped[r.sheet][r.key].append(r)
        if r.errors:add('검증오류',r.key,'입력 검증',r.ref,' / '.join(r.errors))
        if r.unsupported:add('미지원',r.key,'지원 범위',r.ref,' / '.join(r.unsupported))
    for sheet,groups in grouped.items():
        for key,items in groups.items():
            if len(items)>1:
                blocked.add(key)
                add('중복',key,'정확키 중복',', '.join(r.ref for r in items),f'{sheet}: {len(items)}행. 합산/첫행 선택 없이 키 전체 계산 차단')
    allkeys=set().union(*(set(g) for g in grouped.values()))
    if not rows:add('검증오류',None,'입력 검증','','입력 자료 4개 시트가 모두 비었습니다.')
    if progress:progress(f'{len(rows):,}행 검증 완료 · {len(allkeys):,}키 대사 중')
    for key in sorted(allkeys):
        check_cancel(event)
        unique={s:g[key][0] for s,g in grouped.items() if len(g.get(key,[]))==1}
        refs=', '.join(r.ref for g in grouped.values() for r in g.get(key,[]))
        for s in SCHEMA:
            if key not in grouped[s]:add('누락',key,'자료 누락',refs,f'{s}에 같은 정확키가 없습니다. 누락은 0으로 치환하지 않습니다.')
        if key in blocked:
            add('계산불가',key,'키 전체',refs,'중복키 때문에 대사·이자 계산을 수행하지 않았습니다.')
            continue
        if any(r.unsupported for r in unique.values()):
            add('계산불가',key,'키 전체',refs,'미지원 조건·통화·단위 때문에 계산하지 않았습니다.')
            continue
        condition=unique.get('조건')
        if not condition or condition.errors:
            add('계산불가',key,'키 전체',refs,'계약조건 누락/검증오류 때문에 계산하지 않았습니다.')
            continue
        def compare(a,b,column,label):
            ra,rb=unique.get(a),unique.get(b)
            tol=settings['원금허용오차' if column=='원금' else '이자허용오차']
            cref=', '.join(r.ref for r in (ra,rb) if r)
            if not ra or not rb or ra.errors or rb.errors:
                add('계산불가',key,label,cref,'관련 자료 누락 또는 검증오류',tolerance=tol)
                return
            left,right=ra.parsed[column],rb.parsed[column]
            delta=left-right
            add('금액일치' if abs(delta)<=tol else '차이',key,label,cref,'허용오차 이내' if abs(delta)<=tol else '허용오차 초과',left=left,right=right,difference=delta,tolerance=tol)
        for a,b in (('명세','원장'),('명세','회신')):
            for column in ('원금','이자'):compare(a,b,column,f'{a}-{b} {column}')
        compare('명세','조건','원금','명세-조건 원금')
        schedule=unique.get('명세')
        if not condition or not schedule or schedule.errors:
            add('계산불가',key,'명세이자-단순이자',refs,'명세 또는 조건 누락/검증오류')
        else:
            p=condition.parsed
            days,denominator,exact,rounded=simple_interest(p['원금'],p['연이율'],p['이자시작'],p['이자끝'],condition.raw['DayCount'],p['시작포함'],p['끝포함'])
            if rounded>MAX_MONEY:
                add('계산불가',key,'명세이자-단순이자',f'{schedule.ref}, {condition.ref}','재계산 이자가 지원 금액 상한을 초과합니다.',details=f'반올림전={exact}; 원 단위 반올림={rounded}')
                continue
            delta=schedule.parsed['이자']-rounded
            tol=settings['이자허용오차']
            details=f'조건원금={p["원금"]}; 연이율(소수)={p["연이율"]}; {p["이자시작"]}~{p["이자끝"]}; 시작{condition.raw["시작포함"]}/끝{condition.raw["끝포함"]}; {condition.raw["DayCount"]}; 일수={days}; 분모={denominator}; 반올림전={exact}; ROUND_HALF_UP(원)'
            add('금액일치' if abs(delta)<=tol else '차이',key,'명세이자-단순이자',f'{schedule.ref}, {condition.ref}','허용오차 이내' if abs(delta)<=tol else '허용오차 초과',left=schedule.parsed['이자'],right=rounded,difference=delta,tolerance=tol,details=details)
    check_cancel(event)
    return Result(str(path.resolve()),sha256(data).hexdigest(),settings,rows,findings,datetime.now().astimezone().isoformat(timespec='seconds'))
