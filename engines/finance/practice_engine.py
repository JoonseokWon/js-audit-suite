"""Mandatory period evidence and conservative difference classification for v2 UI."""
from collections import defaultdict
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from .engine import (reconcile as core_reconcile, InputError, Finding, SourceRow,
                    as_date, _headers, check_cancel, simple_interest)

EXTRA_SCHEMA={
 '이자기준':['기관코드','계약ID','기준일','이자기산일','기산일포함','기준일포함','근거문서'],
 '이자구간':['기관코드','계약ID','구간ID','이자시작','이자끝','시작포함','끝포함'],
}

def bounds(start,end,include_start,include_end):
    return start.toordinal()+(0 if include_start else 1),end.toordinal()+(1 if include_end else 0)

def read_extra(data,event=None):
    wb=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        if not set(EXTRA_SCHEMA)<=set(wb.sheetnames):
            raise InputError('v2 입력에는 이자기준·이자구간 시트가 필요합니다. v2 입력템플릿을 사용하세요.')
        rows=[]
        for name,headers in EXTRA_SCHEMA.items():
            for rowno,cells in enumerate(_headers(wb[name],headers,event),2):
                vals=[c.value for c in cells]
                if not any(v is not None for v in vals):continue
                if any(v is not None for v in vals[len(headers):]):raise InputError(f'{name}!{rowno}: 양식 밖 값')
                vals += [None]*max(0,len(headers)-len(vals));raw=dict(zip(headers,vals))
                a,b=raw['기관코드'],raw['계약ID'];key=(a,b) if all(isinstance(v,str) and v.strip() for v in (a,b)) else None
                r=SourceRow(name,rowno,raw,key);rows.append(r)
                if key is None:r.errors.append('기관코드·계약ID는 텍스트 필수값입니다.')
                if any(c.data_type=='f' for c in cells):r.errors.append('수식은 지원하지 않습니다.')
                dates=('기준일','이자기산일') if name=='이자기준' else ('이자시작','이자끝')
                flags=('기산일포함','기준일포함') if name=='이자기준' else ('시작포함','끝포함')
                for f in dates:
                    try:r.parsed[f]=as_date(raw[f])
                    except InputError as e:r.errors.append(f'{f}: {e}')
                for f in flags:
                    if raw[f] not in ('포함','제외'):r.errors.append(f'{f}: 포함/제외 필수')
                    else:r.parsed[f]=raw[f]=='포함'
                field='근거문서' if name=='이자기준' else '구간ID'
                if not isinstance(raw[field],str) or not raw[field].strip() or len(raw[field])>500:r.errors.append(field+': 비어 있지 않은 텍스트 필수')
        return rows
    finally:wb.close()

def reconcile(path,event=None,progress=None):
    # Core reads immutable bytes, enforces ZIP limits and validates monetary inputs.
    result=core_reconcile(path,event,progress)
    if result.settings['원금허용오차']!=0:
        raise InputError('v2 원금허용오차는 0원입니다. 원금 차이를 허용오차로 제거할 수 없습니다.')
    if result.settings['이자허용오차'] not in (0,1):
        raise InputError('v2 이자허용오차는 0원 또는 1원만 허용합니다. 1원 차이도 반올림차이로 별도 표시합니다.')
    data=Path(path).read_bytes()
    from hashlib import sha256
    if sha256(data).hexdigest()!=result.input_sha256:raise InputError('검사 도중 입력파일이 변경되었습니다. 다시 실행하세요.')
    extra=read_extra(data,event);base_rows=list(result.rows);result.rows+=extra
    groups={s:defaultdict(list) for s in EXTRA_SCHEMA}
    for r in extra:
        if r.key:groups[r.sheet][r.key].append(r)
        if r.errors:result.findings.append(Finding('검증오류',r.key,'이자기간 입력',r.ref,' / '.join(r.errors)))
    conditions=defaultdict(list)
    for r in base_rows:
        if r.sheet=='조건' and r.key:conditions[r.key].append(r)
    allkeys={r.key for r in result.rows if r.key}
    period_reports=[]
    for key in sorted(allkeys):
        check_cancel(event)
        anchors=groups['이자기준'].get(key,[]);segments=groups['이자구간'].get(key,[]);conds=conditions.get(key,[])
        refs=', '.join(r.ref for r in conds+anchors+segments);errors=[];expected_days=None;actual_days=None
        if len(conds)!=1:errors.append('계약조건이 없거나 중복되었습니다.')
        if len(anchors)!=1:errors.append('독립 이자기준이 없거나 중복되었습니다.')
        if not segments:errors.append('이자구간이 없습니다.')
        if any(r.errors or r.unsupported for r in conds+anchors+segments):errors.append('기간 관련 입력에 오류 또는 미지원 조건이 있습니다.')
        ids=[r.raw['구간ID'] for r in segments]
        if len(set(ids))!=len(ids):errors.append('계약 내 구간ID가 중복되었습니다.')
        if not errors:
            c=conds[0];a=anchors[0];p=c.parsed;q=a.parsed;asof=result.settings['기준일']
            if q['기준일']!=asof:errors.append('이자기준의 기준일이 실행 기준일과 다릅니다.')
            if not p['계약시작']<=q['이자기산일']<=asof<=p['만기']:errors.append('이자기산일·기준일이 계약기간 밖입니다.')
            expected=bounds(q['이자기산일'],asof,q['기산일포함'],q['기준일포함'])
            recorded=bounds(p['이자시작'],p['이자끝'],p['시작포함'],p['끝포함'])
            expected_days=expected[1]-expected[0]
            if expected_days<0:errors.append('기준기간 일수가 음수입니다.')
            if recorded!=expected:errors.append('조건의 이자시작·끝/포함규칙이 독립 이자기준부터 기준일까지의 기간과 다릅니다.')
            intervals=[]
            for s in segments:
                t=s.parsed;ab=bounds(t['이자시작'],t['이자끝'],t['시작포함'],t['끝포함'])
                if ab[1]<ab[0] or (ab[0]==ab[1] and expected_days!=0):errors.append(s.ref+': 비어 있거나 역전된 구간입니다.')
                intervals.append((ab[0],ab[1],s.ref))
            cursor=expected[0];actual_days=0
            for start,stop,ref in sorted(intervals):
                if start>cursor:errors.append(f'{ref}: 이자기간 {start-cursor}일 누락')
                if start<cursor:errors.append(f'{ref}: 이자기간 중복 또는 시작 초과')
                if start<expected[0] or stop>expected[1]:errors.append(ref+': 기준기간 밖 구간')
                actual_days+=max(0,stop-start);cursor=max(cursor,stop)
            if cursor<expected[1]:errors.append(f'기준일까지 이자기간 {expected[1]-cursor}일 누락')
            if cursor>expected[1]:errors.append('기준일을 초과하는 이자기간')
        status='기간오류' if errors else '기간확인'
        result.findings.append(Finding(status,key,'이자기간 완전성',refs,' / '.join(errors) if errors else '독립 이자기준·조건·구간의 전체 날짜가 일치합니다.',details=f'기대일수={expected_days}; 입력구간일수={actual_days}'))
        if errors:
            result.findings=[f for f in result.findings if not(f.key==key and f.check=='명세이자-단순이자')]
            result.findings.append(Finding('계산불가',key,'명세이자-단순이자',refs,'이자기간 확인 실패로 이자 재계산 일치 판단을 차단했습니다.'))
        period_reports.append({'기관코드':key[0],'계약ID':key[1],'상태':status,'기대일수':expected_days,'입력구간일수':actual_days,'근거문서':anchors[0].raw['근거문서'] if len(anchors)==1 else '', '입력행':refs,'사유':' / '.join(errors)})
    for f in result.findings:
        if f.status=='금액일치' and f.difference:
            f.status='반올림차이';f.message='0원이 아닌 차이입니다. 반올림 허용범위 이내이며 검토 대상으로 보존합니다.'
        elif f.status=='금액일치':f.message='차이 0원'
    result.settings.update({'검사버전':'2.0','원금정책':'차이 0원만 일치','이자정책':'0원 일치 / 0원 초과 허용범위 이내는 반올림차이 / 초과는 차이','기간근거':'계약·직전 이자정산서에서 확인한 이자기산일과 기준일을 별도 입력'})
    result.period_reports=period_reports
    summary={}
    for f in result.findings:
        if f.difference is not None:
            s=summary.setdefault(f.check,{'비교':f.check,'비교건수':0,'비영차이건수':0,'부호합계':Decimal(0),'절대차이합계':Decimal(0)})
            s['비교건수']+=1;s['비영차이건수']+=int(f.difference!=0);s['부호합계']+=f.difference;s['절대차이합계']+=abs(f.difference)
    result.difference_summary=list(summary.values())
    return result
