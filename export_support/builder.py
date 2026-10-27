"""Identical static-cell contract for normal and write-only Excel workbooks."""
import itertools,json,re
from decimal import Decimal,InvalidOperation
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from common import plain,AuditError
from .text import display_text,invalid_xml_positions
from .layout import split_display,row_height


def create_builder(priority,numeric,rates,quantities,write_only=False,max_rows=1048575):
    book=Workbook(write_only=write_only)
    if not write_only:book.remove(book.active)
    font=Font(name='맑은 고딕',size=10);title=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF')
    fill=PatternFill('solid',fgColor='263F59');indexes=[];finalized=False
    def add(name,rows,headers=None):
        iterator=iter(rows);preview=list(itertools.islice(iterator,35))
        discovery=rows if isinstance(rows,(list,tuple)) else preview
        headers=headers or next((v for k,v in priority.items() if name.endswith(k)),None) or list(dict.fromkeys(k for r in discovery for k in r)) or ['내용']
        widths={key:min(55,max(14,len(key)*1.7+2,max((len(str(r.get(key) or '')) for r in preview),default=0)*1.1)) for key in headers}
        if name.startswith('조서_') and headers==['항목','기록']:widths={'항목':30,'기록':105}
        raw_json=name=='증적JSON'
        def sheet(part):
            label=re.sub(r'[\\/*?:\[\]]','_',name)[:31] or '자료'
            if part>1:label=label[:25]+'_'+str(part)
            base=label;suffix=1
            while label in book.sheetnames:suffix+=1;label=base[:25]+'_'+str(suffix)
            ws=book.create_sheet(label);ws.sheet_view.showGridLines=False;ws.freeze_panes='B2';ws.sheet_format.defaultRowHeight=25
            for j,key in enumerate(headers,1):ws.column_dimensions[get_column_letter(j)].width=widths[key]
            values=[]
            for key in headers:
                value=display_text(key)
                if len(value.encode('utf-16-le'))//2>30000:raise AuditError('출력 항목명이 Excel 표시 한도를 초과했습니다.')
                cell=WriteOnlyCell(ws,value) if write_only else ws.cell(1,len(values)+1,value)
                cell.data_type='s';cell.font=title;cell.fill=fill;cell.alignment=Alignment(wrap_text=True,vertical='top');values.append(cell)
            ws.row_dimensions[1].height=30
            if write_only:ws.append(values);ws.row_dimensions.pop(1,None)
            return ws
        part=1;count=0;ws=sheet(part);first=ws
        for logical,record in enumerate(itertools.chain(preview,iterator),1):
            columns=[];changed=[];splits=[];height=1
            for key in headers:
                value=record.get(key)
                if isinstance(value,(dict,list)):
                    value=json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False,default=plain)
                if key in numeric|rates|quantities|{'전망배수'} and value is not None and isinstance(value,str):
                    try:value=float(Decimal(value)) if '.' in value else int(value)
                    except (ValueError,InvalidOperation):pass
                if isinstance(value,str):
                    if next(invalid_xml_positions(value),None) is not None:changed.append(str(key))
                    rendered=display_text(value)
                    if raw_json:
                        if len(rendered.encode('utf-16-le'))//2>30000:raise AuditError('증적JSON 조각은 30,000 UTF-16 단위 이하로 나누세요.')
                        pieces=[rendered];lines=1
                    else:pieces,lines=split_display(rendered,widths[key])
                    height=max(height,lines)
                    if len(pieces)>1:splits.append(str(key))
                else:pieces=[value]
                columns.append(pieces)
            size=max(map(len,columns));ranges=[]
            for chunk in range(size):
                if count==max_rows:part+=1;count=0;ws=sheet(part)
                physical=count+2
                if not ranges or ranges[-1][0]!=ws.title:ranges.append([ws.title,physical,physical])
                else:ranges[-1][2]=physical
                ws.row_dimensions[physical].height=row_height(height) if not raw_json else 25
                cells=[]
                for j,(key,pieces) in enumerate(zip(headers,columns),1):
                    value=pieces[chunk] if chunk<len(pieces) else None
                    cell=WriteOnlyCell(ws,value) if write_only else ws.cell(physical,j,value)
                    cell.font=font;cell.alignment=Alignment(wrap_text=not raw_json,vertical='top')
                    if isinstance(value,str):cell.data_type='s'
                    if isinstance(value,(int,float)) and not isinstance(value,bool):cell.number_format='0.00####%' if key in rates else '#,##0;[Red](#,##0);0' if isinstance(value,int) else '#,##0.00####;[Red](#,##0.00####);0'
                    cells.append(cell)
                if write_only:
                    ws.append(cells)
                    # write-only already serialized this dimension; retain O(1)
                    # row metadata instead of accumulating one object per GL row.
                    ws.row_dimensions.pop(physical,None)
                count+=1
            if size>1 or changed:
                for label,start,end in ranges:
                    indexes.append({'표':label,'원본행':logical,'시작행':start,'마지막행':end,'분할열':', '.join(splits),
                                    '가시변환열':', '.join(changed),'정책':'계속 행은 남은 문자열만 기록. 숫자는 첫 행에만 기록. 금지문자는 ⟦U+코드⟧ 표시. 원문은 증적JSON에서 복원.'})
            if not write_only:ws.auto_filter.ref=ws.dimensions
        return first
    def finalize():
        nonlocal finalized
        if finalized:return
        finalized=True
        if indexes:add('출력문단연결',list(indexes),['표','원본행','시작행','마지막행','분할열','가시변환열','정책'])
    save=book.save
    def save_complete(target):finalize();return save(target)
    book.save=save_complete;book.audit_finalize=finalize
    return book,add
