import json,tempfile,threading,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from common import read_xlsx,AuditError,dumps
from input_limits import LEDGER_LIMITS,EXCEL_DATA_ROWS
from export_streaming import chunks,create_builder

class LargeLedgerTests(unittest.TestCase):
    def book(self,folder,name='총계정원장',rows=10002):
        path=Path(folder)/'large.xlsx';wb=Workbook(write_only=True);ws=wb.create_sheet(name);ws.append(['코드','금액'])
        for i in range(rows):ws.append(['0001',i])
        wb.save(path);return path
    def test_old_boundary_is_accepted_only_for_the_large_ledger_contract(self):
        with tempfile.TemporaryDirectory() as folder:
            path=self.book(folder)
            with self.assertRaises(AuditError):read_xlsx(path)
            tables,meta=read_xlsx(path,row_limits=LEDGER_LIMITS)
            self.assertEqual(len(tables['총계정원장']),10002);self.assertEqual(tables['총계정원장'][-1]['_row'],10003)
            self.assertEqual(tables['총계정원장'][0]['코드'],'0001')
    def test_large_setting_does_not_relax_other_tables(self):
        with tempfile.TemporaryDirectory() as folder:
            path=self.book(folder,'시산표')
            with self.assertRaises(AuditError):read_xlsx(path,row_limits=LEDGER_LIMITS)
    def test_cancellation_is_checked_during_streaming_preflight(self):
        with tempfile.TemporaryDirectory() as folder:
            path=self.book(folder,rows=10);event=threading.Event();event.set()
            with self.assertRaises(AuditError):read_xlsx(path,event=event,row_limits=LEDGER_LIMITS)
    def test_evidence_chunks_reconstruct_exact_canonical_json(self):
        value={'한글':['은행'*20000,{'account':'0001','debit':123}], 'state':None}
        self.assertEqual(''.join(chunks(value)),dumps(value))
    def test_streamed_rows_codes_and_formula_like_text_are_preserved(self):
        wb,add=create_builder({},set(),set(),set())
        add('원장',[{'코드':'0001','금액':i,'적요':'=1+1'} for i in range(10002)])
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'out.xlsx';wb.save(path);book=load_workbook(path,read_only=True,data_only=False)
            try:
                values=list(book['원장'].values);self.assertEqual(len(values),10003);self.assertEqual(values[-1],('0001',10001,'=1+1'))
                self.assertEqual(next(book['원장'].iter_rows(min_row=2))[2].data_type,'s')
            finally:book.close()
    def test_streaming_keeps_columns_first_seen_after_the_preview(self):
        wb,add=create_builder({},set(),set(),set());rows=[{'금액':i} for i in range(40)];rows[-1]['추가필드']='보존'
        add('원장',rows)
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'late.xlsx';wb.save(path);book=load_workbook(path,read_only=True)
            try:self.assertEqual(list(book['원장'].values)[-1],(39,'보존'))
            finally:book.close()

if __name__=='__main__':unittest.main()
