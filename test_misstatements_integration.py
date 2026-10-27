"""F03 linkage and separately reviewed synthetic workpaper invalidation."""
from copy import deepcopy
from io import BytesIO
import json
from openpyxl import load_workbook
from common import AuditError
from verify import synthetic_procedure
import misstatements as ms
from test_misstatements_support import Case, ACTOR, profile
from test_misstatements_support import ROOT


class IntegrationTests(Case):
    def review_f03(self):
        paper=synthetic_procedure('F03',list(self.ws.sources))
        paper.update(title='SYNTHETIC ONLY 왜곡표시 집계',preparer=ACTOR,
                     performed='SYNTHETIC ONLY 가상 절차 수행 필드 상태전이',
                     conclusion='SYNTHETIC ONLY 가상 판단; 실제 사람 판단·승인 아님')
        self.ws.save_procedure(paper);self.ws.request_review('F03')
        self.ws.review('F03','SYNTHETIC ONLY / 다른 가상검토자','SYNTHETIC ONLY 상태전이 검증')

    def test_f03_is_required_only_after_new_record_and_review_can_follow_current_human_records(self):
        self.assertNotIn('F03',self.ws.required_procedures())
        self.ws.save_workflow_profile(profile(),ACTOR)
        record=ms.save_record(self.ws,self.record_value(),ACTOR)
        self.assertIn('F03',self.ws.required_procedures())
        self.assertIn('MS_F03_REVIEW_PENDING',{r['code'] for r in ms.status(self.ws)['blockers']})
        self.review_f03()
        self.assertEqual(ms.status(self.ws)['blockers'],[])
        snapshot=self.ws.snapshot()
        self.assertEqual(snapshot['misstatements']['records'][0]['id'],record['id'])
        self.assertEqual(snapshot['misstatements']['summary'][0]['gross_amount'],100)
        self.assertEqual(self.ws.procedure_state(self.ws.procedures['F03']),'검토기록완료')

    def test_record_or_group_change_invalidates_workpaper_without_rewriting_human_review(self):
        self.ws.save_workflow_profile(profile(),ACTOR)
        a=ms.save_record(self.ws,self.record_value(),ACTOR)
        self.review_f03();review=deepcopy(self.ws.procedures['F03']);events=deepcopy(self.ws.events)
        b=ms.save_record(self.ws,self.record_value(),ACTOR)
        self.assertEqual(self.ws.procedures['F03'],review)
        self.assertEqual(self.ws.events[:len(events)],events)
        self.assertEqual(self.ws.procedure_state(review),'자료변경·재검토')
        self.review_f03()
        ms.save_group(self.ws,{'member_ids':[a['id'],b['id']],'primary_id':a['id'],
                             'reason':'SYNTHETIC ONLY 동일 송장 확인 상태전이'},ACTOR)
        self.assertEqual(self.ws.procedure_state(self.ws.procedures['F03']),'자료변경·재검토')
        self.assertIn('MS_F03_REVIEW_PENDING',{r['code'] for r in ms.status(self.ws)['blockers']})

    def test_excel_full_register_keeps_gross_numbers_impact_rows_and_manual_drafts(self):
        from exporter import workbook_bytes
        a=ms.save_record(self.ws,self.record_value(description='SYNTHETIC ONLY 한글 받침 수정'),ACTOR)
        b=ms.save_record(self.ws,self.record_value(amount=-100),ACTOR)
        # An unfinished assessment must export as a draft, not crash or imply tax 0.
        c=ms.save_record(self.ws,self.record_value(amount=7,tax_effect=None,tax_basis=''),ACTOR)
        snapshot=self.ws.snapshot()
        catalog=json.loads((ROOT/'procedure_catalog.json').read_text(encoding='utf-8'))
        book=load_workbook(BytesIO(workbook_bytes(snapshot,catalog)),read_only=True,data_only=False)
        try:
            for name in ('왜곡표시검토대장','왜곡표시상세기록','왜곡표시금액집계','왜곡표시계정영향'):
                self.assertIn(name,book.sheetnames)
            rows=list(book['왜곡표시금액집계'].values);header=list(rows[0]);row=dict(zip(header,rows[1]))
            self.assertEqual(row['절대금액 합계'],207)
            self.assertEqual(row['양수 금액'],107);self.assertEqual(row['음수 절대금액'],100)
            self.assertEqual(book['왜곡표시계정영향'].max_row,7)
            details=list(book['왜곡표시상세기록'].values)
            self.assertTrue(any('한글 받침 수정' in str(cell) for row in details for cell in row))
            self.assertTrue(any(c['id'] in str(cell) for row in details for cell in row))
            self.assertEqual(snapshot['misstatements']['summary'][0]['tax_unknown_count'],1)
        finally:book.close()
