"""Streaming workbook builder preserves result rows and the full evidence JSON."""
import itertools,json,re
from decimal import Decimal,InvalidOperation
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from common import plain
from input_limits import EXCEL_DATA_ROWS

def chunks(value,size=28000):
    encoder=json.JSONEncoder(ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False,default=plain)
    buffer=''
    for token in encoder.iterencode(value):
        buffer+=token
        while len(buffer)>=size:yield buffer[:size];buffer=buffer[size:]
    if buffer:yield buffer

def bounded_json(value):
    parts=[];length=0
    encoder=json.JSONEncoder(ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False,default=plain)
    for token in encoder.iterencode(value):
        parts.append(token);length+=len(token)
        if length>30000:break
    return ''.join(parts)

def create_builder(priority,numeric,rates,quantities):
    workbook=Workbook(write_only=True);font=Font(name='맑은 고딕',size=10)
    title=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF');fill=PatternFill('solid',fgColor='263F59');align=Alignment(vertical='center')
    def add(name,rows,headers=None):
        iterator=iter(rows);preview=list(itertools.islice(iterator,35))
        discovery=rows if isinstance(rows,(list,tuple)) else preview
        headers=headers or next((v for k,v in priority.items() if name.endswith(k)),None) or list(dict.fromkeys(k for r in discovery for k in r)) or ['내용']
        def sheet(part):
            label=re.sub(r'[\\/*?:\[\]]','_',name)[:31] or '자료'
            if part>1:label=label[:25]+'_'+str(part)
            base=label;suffix=1
            while label in workbook.sheetnames:suffix+=1;label=base[:25]+'_'+str(suffix)
            ws=workbook.create_sheet(label);ws.sheet_view.showGridLines=False;ws.freeze_panes='B2'
            for j,key in enumerate(headers,1):
                width=max((len(str(r.get(key) or '')) for r in preview),default=0)
                ws.column_dimensions[get_column_letter(j)].width=min(55,max(14,len(key)*1.7+2,width*1.1))
            cells=[]
            for key in headers:
                cell=WriteOnlyCell(ws,key);cell.data_type='s';cell.font=title;cell.fill=fill;cells.append(cell)
            ws.append(cells);return ws
        part=1;count=0;ws=sheet(part);first=ws
        for record in itertools.chain(preview,iterator):
            if count==EXCEL_DATA_ROWS:part+=1;count=0;ws=sheet(part)
            values=[]
            for key in headers:
                value=record.get(key)
                if isinstance(value,(dict,list)):value=bounded_json(value)
                if key in numeric|rates|quantities|{'전망배수'} and value is not None and isinstance(value,str):
                    try:value=float(Decimal(value)) if '.' in value else int(value)
                    except (ValueError,InvalidOperation):pass
                if isinstance(value,str) and len(value)>30000:value=value[:29970]+' [전체 내용은 증적JSON 참조]'
                cell=WriteOnlyCell(ws,value);cell.font=font;cell.alignment=align
                if isinstance(value,str):cell.data_type='s'
                if isinstance(value,(int,float)):cell.number_format='0.00####%' if key in rates else '#,##0;[Red](#,##0);0' if isinstance(value,int) else '#,##0.00####;[Red](#,##0.00####);0'
                values.append(cell)
            ws.append(values);count+=1
        return first
    return workbook,add
