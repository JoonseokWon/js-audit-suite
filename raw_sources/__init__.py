"""Bounded, reproducible source packages; original XLSX bytes remain evidence.

The catalog is recalculated from the packaged workbooks on EVERY import. SHA256
detects changes relative to this package, not authenticity or independence.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED, ZIP_STORED, BadZipFile
import hashlib
import json
import math
import re
import stat
import unicodedata

from common import AuditError, atomic, cancelled, dumps, day

PACK_LIMIT = 40 * 1024**2
UNPACK_LIMIT = 100 * 1024**2
FILE_LIMIT = 20 * 1024**2
JSON_LIMIT = 20 * 1024**2
MEMBER_LIMIT = 128
RATIO_LIMIT = 200
ROW_LIMIT = 10000
COL_LIMIT = 80
TABLE_LIMIT = 200
CONTEXT_KEYS = ('회사코드', '시작일', '기준일', '통화', '연결구분')
SELECTED_FOLDERS = ('01_현재프로그램입력', '02_절차별원천자료')
POLICY = {
    'timestamp': 'ISO8601 타임스탬프 앞 작은따옴표 1개만 제거; 원문/변환값/셀 위치 기록',
    'null_row': '물리적으로 존재하는 모든 셀 null 행은 카탈로그 데이터행에서 제외; 원문 행번호/값 기록',
    'dates': 'Excel date/datetime을 ISO8601 문자열로 직렬화; 자료형/원문/변환값 기록',
    'formulas': '수식·오류 셀 차단; 캐시값을 계산 결과로 신뢰하지 않음',
    'context_defaults': {'통화': 'KRW', '연결구분': '별도'},
    'selection': '현행 spec books 중 01_현재프로그램입력·02_절차별원천자료의 전체 시트; CSV/오류/행정책자 제외',
}
_TS = re.compile(r"^'\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})$")
_CONTROL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _error(message):
    raise AuditError('원천자료 패키지: ' + message)


def _hash(data):
    return hashlib.sha256(data).hexdigest()


def _path(name):
    if (not isinstance(name, str) or not name or len(name) > 260 or '\\' in name
            or ':' in name or re.search(r'[\x00-\x1f]', name) or name.startswith('/')
            or any(p in ('', '.', '..') or p.endswith((' ', '.')) for p in name.split('/'))):
        _error('허용하지 않는 파일 경로')
    if PurePosixPath(name).is_absolute():
        _error('절대경로 차단')
    return name


def _json(data, label):
    if len(data) > JSON_LIMIT:
        _error(label + ': JSON 크기 한도 초과')

    def pairs(items):
        out = {}
        for k, v in items:
            if k in out:
                _error(label + ': JSON 중복 키 ' + k)
            out[k] = v
        return out

    def constant(value):
        _error(label + ': 유한 숫자 필수 ' + value)

    try:
        obj = json.loads(data.decode('utf-8'), object_pairs_hook=pairs, parse_constant=constant)
        todo = [(obj, 0)]
        nodes = 0
        while todo:
            value, depth = todo.pop()
            nodes += 1
            if depth > 24 or nodes > 600000:
                _error(label + ': JSON 구조 한도 초과')
            if isinstance(value, dict):
                todo.extend((v, depth + 1) for v in value.values())
            elif isinstance(value, list):
                todo.extend((v, depth + 1) for v in value)
            elif isinstance(value, float) and not math.isfinite(value):
                _error(label + ': 유한 숫자 필수')
        return obj
    except (UnicodeError, json.JSONDecodeError, RecursionError, OverflowError) as exc:
        raise AuditError('원천자료 패키지: ' + label + ': JSON 형식 오류') from exc


def _context(obj, defaults=False):
    if not isinstance(obj, dict):
        _error('회사·기간 문맥 필수')
    result = dict(obj)
    if defaults:
        for k, v in POLICY['context_defaults'].items():
            result.setdefault(k, v)
    for key in CONTEXT_KEYS:
        if not isinstance(result.get(key), str) or not result[key].strip():
            _error('문맥 누락: ' + key)
    if day(result['시작일']) > day(result['기준일']):
        _error('시작일이 기준일보다 늦음')
    if result['시작일'] != day(result['시작일']).isoformat() or result['기준일'] != day(result['기준일']).isoformat():
        _error('기간은 YYYY-MM-DD 필수')
    if result['통화'] != 'KRW' or result['연결구분'] not in ('별도', '연결'):
        _error('지원 문맥은 KRW 및 별도/연결')
    return {key: result[key] for key in CONTEXT_KEYS}


def _zip_members(z, max_members=MEMBER_LIMIT, total_limit=UNPACK_LIMIT):
    infos = z.infolist()
    if not infos or len(infos) > max_members:
        _error('ZIP 파일 개수 한도 초과/빈 파일')
    seen = set()
    total = 0
    for info in infos:
        _path(info.filename)
        identity = unicodedata.normalize('NFC', info.filename).casefold()
        if identity in seen:
            _error('ZIP 중복 파일 경로')
        seen.add(identity)
        mode = (info.external_attr >> 16) & 0o170000
        if info.is_dir() or mode not in (0, stat.S_IFREG) or info.flag_bits & 1:
            _error('폴더·symlink·특수파일·암호 ZIP 차단')
        if info.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
            _error('지원하지 않는 ZIP 압축 방식')
        if info.file_size < 0 or info.file_size > FILE_LIMIT:
            _error('ZIP 개별 파일 크기 한도 초과')
        total += info.file_size
        if total > total_limit or info.file_size > max(1, info.compress_size) * RATIO_LIMIT:
            _error('ZIP 압축해제 크기/압축률 한도 초과')
    return infos


def _read_member(z, info):
    with z.open(info) as stream:
        data = stream.read(min(FILE_LIMIT, info.file_size) + 1)
    if len(data) != info.file_size:
        _error('ZIP 파일 실제 크기 불일치')
    return data


def _workbook(data, name, event=None):
    """Preflight nested OOXML before bounded, read-only value extraction."""
    from defusedxml.ElementTree import fromstring
    from openpyxl import load_workbook
    try:
        with ZipFile(BytesIO(data)) as z:
            infos = _zip_members(z, max_members=600)
            for info in infos:
                cancelled(event)
                lower = info.filename.lower()
                if any(x in lower for x in ('externallinks/', 'vbaproject', 'embeddings/', 'connections.xml', 'macrosheets/', 'dialogsheet')):
                    _error(name + ': 외부연결·매크로·첨부 차단')
                if lower.endswith(('.xml', '.rels')):
                    root = fromstring(_read_member(z, info))
                    if lower.endswith('.rels') and any(el.get('TargetMode') == 'External' for el in root):
                        _error(name + ': 외부 참조 차단')
                    if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', lower):
                        last = 0
                        for row in root.iter():
                            if row.tag.rsplit('}', 1)[-1] != 'row':
                                continue
                            r = int(row.get('r', '0'))
                            if r <= last or r > ROW_LIMIT + 1:
                                _error(name + ': 물리 행번호/행수 한도 오류')
                            last = r
                            col = 0
                            for cell in row:
                                if cell.tag.rsplit('}', 1)[-1] != 'c':
                                    continue
                                m = re.fullmatch(r'([A-Z]+)(\d+)', cell.get('r', ''))
                                if not m or int(m[2]) != r:
                                    _error(name + ': 잘못된 셀 좌표')
                                c = 0
                                for ch in m[1]:
                                    c = 26 * c + ord(ch) - 64
                                if c <= col or c > COL_LIMIT:
                                    _error(name + ': 중복/역순/범위밖 열')
                                col = c
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        out = {}
        try:
            if len(book.sheetnames) > TABLE_LIMIT:
                _error(name + ': 시트 개수 한도 초과')
            for sheet in book:
                sheet.reset_dimensions()
                rows = []
                for n, cells in enumerate(sheet.iter_rows(), 1):
                    cancelled(event)
                    if n > ROW_LIMIT + 1 or len(cells) > COL_LIMIT:
                        _error(name + ': 시트 크기 한도 초과')
                    vals = []
                    for cell in cells:
                        if cell.data_type in ('f', 'e'):
                            _error(f'{name}/{sheet.title}!{n}: 수식·오류 셀 차단')
                        value = cell.value
                        if isinstance(value, str) and (len(value) > 4000 or _CONTROL.search(value)):
                            _error(name + ': 셀 텍스트 한도/제어문자 오류')
                        if isinstance(value, float) and not math.isfinite(value):
                            _error(name + ': 유한 숫자 필수')
                        if value is not None and not isinstance(value, (str, int, float, bool, date, datetime)):
                            _error(name + ': 지원하지 않는 셀 자료형')
                        vals.append(value)
                    rows.append((n, vals))
                out[sheet.title] = rows
        finally:
            book.close()
        return out
    except AuditError:
        raise
    except Exception as exc:
        raise AuditError('원천자료 패키지: ' + name + ': Excel 구조/값 읽기 실패') from exc


def _selection(spec):
    if not isinstance(spec, dict) or not isinstance(spec.get('books'), list) or not isinstance(spec.get('datasets'), list):
        _error('spec books/datasets 필수')
    books = []
    names = set()
    for book in spec['books']:
        if not isinstance(book, dict):
            _error('spec book 형식 오류')
        if book.get('folder') not in SELECTED_FOLDERS:
            continue
        name = book.get('name')
        if not isinstance(name, str) or not name or '/' in name or '\\' in name:
            _error('spec book 이름 오류')
        file = _path(book['folder'] + '/' + name + '.xlsx')
        if file in names or not isinstance(book.get('headers'), dict) or not book['headers']:
            _error('spec 중복 book/헤더 없음')
        names.add(file)
        books.append((file, book))
    if not books or len(books) > 32:
        _error('선택할 현행 XLSX가 없음/개수 초과')
    mapping = {}
    ids = set()
    for item in spec['datasets']:
        if not isinstance(item, dict) or not re.fullmatch(r'D\d{3}', str(item.get('id', ''))):
            _error('spec dataset ID 형식 오류')
        if not isinstance(item.get('file'), str) or not isinstance(item.get('name'), str):
            _error('spec dataset 원천 파일/시트 형식 오류')
        key = (item['file'], item['name'])
        if item['id'] in ids or key in mapping:
            _error('spec 중복 dataset ID/원천 시트')
        ids.add(item['id'])
        if key[0] not in names or not isinstance(key[1], str):
            _error('spec dataset의 현행 파일/시트 연결 없음')
        mapping[key] = item['id']
    return books, mapping


def _catalog(spec, source_bytes, event=None):
    context = _context(spec.get('context'), defaults=True)
    books, mapping = _selection(spec)
    tables, transformations = [], []
    used_ids = set()
    used_mappings = set()
    for file, book in books:
        cancelled(event)
        source_file = 'sources/' + file
        if source_file not in source_bytes:
            _error('원문 XLSX 누락: ' + file)
        sheets = _workbook(source_bytes[source_file], file, event)
        if set(sheets) != set(book['headers']):
            _error(file + ': 원문/spec 시트 목록 불일치')
        for sheet, expected_headers in book['headers'].items():
            key = (file, sheet)
            dataset_id = mapping.get(key)
            if dataset_id:
                used_mappings.add(key)
            else:
                prefix = re.match(r'\d{2}', book['name'])
                if not book.get('native') or not prefix:
                    _error(file + ': dataset ID 미등록')
                dataset_id = 'N' + prefix[0] + '_' + sheet
            if dataset_id in used_ids:
                _error('중복 카탈로그 dataset ID')
            used_ids.add(dataset_id)
            raw = sheets[sheet]
            if not raw:
                _error(file + '/' + sheet + ': 빈 시트')
            headers = list(raw[0][1])
            while headers and headers[-1] is None:
                headers.pop()
            if (not headers or any(not isinstance(h, str) or not h.strip() for h in headers)
                    or len(set(headers)) != len(headers) or headers != expected_headers):
                _error(file + '/' + sheet + ': 헤더 이름/순서/중복 오류')
            records = []
            for row, values in raw[1:]:
                if len(values) > len(headers) and any(v is not None for v in values[len(headers):]):
                    _error(file + '/' + sheet + ': 헤더 밖 데이터')
                values = values[:len(headers)] + [None] * max(0, len(headers) - len(values))
                if all(v is None for v in values):
                    transformations.append({'dataset_id': dataset_id, 'row': row, 'kind': 'all_null_formatted_row', 'original': values, 'normalized': None})
                    continue
                out = []
                for col, value in enumerate(values):
                    transformed = value
                    kind = None
                    if isinstance(value, (date, datetime)):
                        transformed = value.isoformat()
                        kind = 'excel_datetime' if isinstance(value, datetime) else 'excel_date'
                    elif isinstance(value, str) and _TS.fullmatch(value):
                        try:
                            datetime.fromisoformat(value[1:].replace('Z', '+00:00'))
                        except ValueError as exc:
                            raise AuditError('원천자료 패키지: 작은따옴표 날짜 형식 오류') from exc
                        transformed = value[1:]
                        kind = 'quoted_iso_timestamp'
                    if kind:
                        transformations.append({'dataset_id': dataset_id, 'row': row, 'column': headers[col], 'kind': kind, 'original': value.isoformat() if isinstance(value, (date, datetime)) else value, 'normalized': transformed})
                    out.append(transformed)
                records.append({'row': row, 'values': out})
            tables.append({'dataset_id': dataset_id, 'source_file': source_file, 'sheet': sheet, 'headers': headers, 'rows': records})
    if set(mapping) != used_mappings:
        _error('spec dataset 시트를 모두 찾지 못함')
    if len(tables) > TABLE_LIMIT:
        _error('카탈로그 테이블 개수 초과')
    return {'schema_version': 'raw-catalog-1', 'context': context, 'tables': tables, 'transformations': transformations}


def _links(data, context):
    if (not isinstance(data, dict) or data.get('schema_version') != 'raw-links-1'
            or _context(data.get('context')) != context
            or not isinstance(data.get('populations'), list)
            or not isinstance(data.get('allocations'), list)):
        _error('links 문맥/스키마 오류')
    return data


def _encode(value):
    try:
        return dumps(value).encode('utf-8')
    except (TypeError, ValueError, RecursionError) as exc:
        raise AuditError('원천자료 패키지: JSON 직렬화 오류') from exc


def build_pack(target, source_root, links=None):
    """Build a deterministic package from current spec and original workbooks."""
    root = Path(source_root).resolve()
    spec_path = root / 'spec.json'
    if not spec_path.is_file() or spec_path.stat().st_size > FILE_LIMIT:
        _error('source_root/spec.json 없음/한도 초과')
    spec_data = spec_path.read_bytes()
    spec = _json(spec_data, 'spec.json')
    books, _ = _selection(spec)
    source_bytes = {'sources/spec.json': spec_data}
    for file, _ in books:
        path = (root / '배포' / Path(file)).resolve()
        if not path.is_relative_to(root) or not path.is_file() or path.stat().st_size > FILE_LIMIT:
            _error('원본 XLSX 경로 없음/범위/크기 오류: ' + file)
        source_bytes['sources/' + file] = path.read_bytes()
    catalog = _catalog(spec, source_bytes)
    if links is None:
        from raw_accounts import build_links
        links = build_links(catalog)
    links = _links(links, catalog['context'])
    payload = source_bytes | {'catalog.json': _encode(catalog), 'links.json': _encode(links)}
    manifest = {'schema_version': 'auditpack-1', 'context': catalog['context'], 'normalization_policy': POLICY,
                'selection': {'source_spec': 'sources/spec.json', 'workbooks': ['sources/' + f for f, _ in books]},
                'files': [{'path': k, 'sha256': _hash(v), 'bytes': len(v)} for k, v in sorted(payload.items())]}
    payload['manifest.json'] = _encode(manifest)
    buf = BytesIO()
    with ZipFile(buf, 'w', compression=ZIP_DEFLATED, compresslevel=6) as z:
        for name, data in sorted(payload.items()):
            info = ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = (stat.S_IFREG | 0o600) << 16
            z.writestr(info, data, compresslevel=6)
    output = Path(target)
    if output.suffix.lower() != '.auditpack' or len(buf.getvalue()) > PACK_LIMIT:
        _error('대상은 40MB 이하 .auditpack')
    # Validate the in-memory bytes through the exact import path before commit.
    _load(buf.getvalue(), catalog['context'])
    atomic(output, buf.getvalue())
    return output


def _load(data, context, event=None):
    expected_context = _context(context)
    try:
        with ZipFile(BytesIO(data)) as z:
            infos = _zip_members(z)
            by_name = {i.filename: i for i in infos}
            if not {'manifest.json', 'catalog.json', 'links.json', 'sources/spec.json'} <= set(by_name):
                _error('필수 파일 누락')
            for name in by_name:
                if name not in ('manifest.json', 'catalog.json', 'links.json', 'sources/spec.json') and not (name.startswith('sources/') and name.endswith('.xlsx')):
                    _error('미지원 파일: ' + name)
            manifest = _json(_read_member(z, by_name['manifest.json']), 'manifest.json')
            if (not isinstance(manifest, dict) or manifest.get('schema_version') != 'auditpack-1'
                    or _context(manifest.get('context')) != expected_context
                    or manifest.get('normalization_policy') != POLICY
                    or not isinstance(manifest.get('files'), list)):
                _error('manifest 문맥/정책/스키마 오류')
            declared = {}
            for item in manifest['files']:
                if not isinstance(item, dict) or set(item) != {'path', 'sha256', 'bytes'}:
                    _error('manifest 파일 선언 형식 오류')
                name = _path(item['path'])
                if name in declared or name == 'manifest.json' or not re.fullmatch(r'[0-9a-f]{64}', str(item['sha256'])) or type(item['bytes']) is not int or not 0 <= item['bytes'] <= FILE_LIMIT:
                    _error('manifest 중복/해시/크기 오류')
                declared[name] = item
            if set(declared) != set(by_name) - {'manifest.json'}:
                _error('미등록/누락 파일 또는 manifest 불일치')
            payload = {}
            for name, item in declared.items():
                cancelled(event)
                binary = _read_member(z, by_name[name])
                if len(binary) != item['bytes'] or _hash(binary) != item['sha256']:
                    _error('파일 SHA256/크기 불일치: ' + name)
                payload[name] = binary
            spec = _json(payload['sources/spec.json'], 'sources/spec.json')
            books, _ = _selection(spec)
            selection = {'source_spec': 'sources/spec.json', 'workbooks': ['sources/' + f for f, _ in books]}
            source_set = {'sources/spec.json'} | set(selection['workbooks'])
            if manifest.get('selection') != selection or {n for n in payload if n.startswith('sources/')} != source_set:
                _error('현행 spec 원천 파일 선택 불일치')
            catalog = _json(payload['catalog.json'], 'catalog.json')
            recalculated = _catalog(spec, payload, event)
            if recalculated['context'] != expected_context or _encode(catalog) != _encode(recalculated):
                _error('카탈로그와 원본 XLSX 헤더·행·값·변환기록 불일치')
            links = _links(_json(payload['links.json'], 'links.json'), expected_context)
        return {'catalog': catalog, 'links': links, 'manifest': manifest}
    except AuditError:
        raise
    except (BadZipFile, OSError, RuntimeError, KeyError, ValueError, OverflowError) as exc:
        raise AuditError('원천자료 패키지: ZIP/입력 구조 오류') from exc


def load_pack(path, context, event=None):
    """Strictly import without extracting files or executing document contents."""
    p = Path(path)
    if p.suffix.lower() != '.auditpack' or not p.is_file() or not 0 < p.stat().st_size <= PACK_LIMIT:
        _error('40MB 이하 .auditpack 파일 필수')
    cancelled(event)
    data = p.read_bytes()
    if len(data) > PACK_LIMIT:
        _error('파일 크기 한도 초과')
    result = _load(data, context, event)
    result['source'] = {'filename': p.name, 'bytes': len(data), 'sha256': _hash(data)}
    cancelled(event)
    if p.stat().st_size != len(data) or _hash(p.read_bytes()) != result['source']['sha256']:
        _error('가져오는 중 원본 패키지 변경')
    return result
