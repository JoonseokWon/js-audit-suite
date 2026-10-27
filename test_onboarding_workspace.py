"""Temporary-workspace end-to-end evidence lifecycle; no model or live company."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import unittest
from openpyxl import load_workbook
from common import AuditError, digest, dumps
from onboarding import MODULES, PROCEDURES
from onboarding.demo import ledger_tables
from onboarding.integration import sample
from onboarding.workbooks import save_input, load_input, save_ledger
from procedures import task_definitions
from workspace import Workspace
import exporter
from ai_workpapers import FIELDS, PROMPT_VERSION
from test_onboarding_gateway import GatewayFixture

CATALOG=json.loads((Path(__file__).resolve().parent/'procedure_catalog.json').read_text(encoding='utf-8'))


def synthetic_response(packet):
    """Fixed mock response; explicitly no inference, provider or network call."""
    ref=next(r['id'] for r in packet['references'] if r['kind']=='run_summary')
    fields={key:'SYNTHETIC ONLY: 프로그램 자료·대사 근거의 모의 초안이며 사람이 실제 수행·판단하지 않았습니다. ['+ref+']'
            for key in FIELDS}
    fields['performed']='자동계산·대사 근거의 모의 기록입니다. 사람의 실사·조회·증빙 확인은 미수행이며 별도 추가확인이 필요합니다. ['+ref+']'
    fields['conclusion']='사람 검토 전 모의 초안입니다. 차이·미확인 항목은 추가확인 및 사람의 판단이 필요합니다. ['+ref+']'
    citations={key:[ref] for key in FIELDS}
    exceptions=[r for r in packet['references'] if r['kind']=='exception']
    if exceptions:
        fields['performed']+=' 예외코드: '+','.join(str(r['data'].get('code','')) for r in exceptions)
        fields['performed']+=' '+''.join('['+r['id']+']' for r in exceptions)
        citations['performed'] += [r['id'] for r in exceptions]
    context={'company_code':packet['engagement']['company_code'],
             'procedure_id':packet['procedure']['id'],'basis':packet['basis']}
    from test_review_ai_helpers import response_from_fields
    return response_from_fields(packet,fields,citations,model='SYNTHETIC-ONLY-no-inference')



class WorkspaceFixture(GatewayFixture):
    def setUp(self):
        super().setUp()
        engagement=dict(self.context,회사명='SYNTHETIC ONLY JS전자',회계체계='합성 입력 테스트',
                        감사기준='합성 상태검사·준수 인증 아님',작성자='SYNTHETIC ONLY',더미업무=True)
        self.ws=Workspace(self.folder/'engagement',create=engagement)
        self.ledger_path=self.folder/'ledger.xlsx';save_ledger(ledger_tables(self.context),self.ledger_path)
        self.ledger_run=self.ws.execute('원장·리드',self.ledger_path)
        self.assertEqual(self.ledger_run['result']['blocking'],0)
        self.originals={};self.payloads={};self.inputs={}

    def execute_module(self,module):
        template=self.workbook(module);sid=self.ws.add_source(template,'SYNTHETIC ONLY 원본 '+module)
        payload=load_input(module,template,self.context)
        confirmed=self.confirm(module,payload,[self.ws.sources[sid]])
        path=self.write_json(confirmed,module+'.json');self.inputs[module]=path
        self.originals[module]=sid;self.payloads[module]=confirmed
        run=self.ws.execute(module,path)
        self.assertEqual(run['result']['blocking'],0,run['result']['issues'])
        return run

    def manual_procedure(self,pid,module,evidence):
        self.ws.save_procedure(dict(id=pid,title=module,preparer='SYNTHETIC ONLY',
            purpose='가상 검토상태 검사',risk='가상 증거누락 반례',scope='가상 자료 범위',
            performed='SYNTHETIC ONLY 상태검사 입력; 실제 감사수행 아님',
            conclusion='실제 결론 아님·검토필요',performed_date=self.context['기준일'],evidence_ids=evidence))


class OnboardingWorkspaceTests(WorkspaceFixture):
    def test_all_seven_modules_reach_current_ledger_tasks_ai_drafts_and_formula_free_export(self):
        runs={module:self.execute_module(module) for module in MODULES}
        cross={r['모듈']:r for r in self.ws.cross_checks() if r['모듈'] in MODULES}
        self.assertEqual(set(cross),set(MODULES))
        self.assertTrue(all(r['상태']=='일치' for r in cross.values()))
        tasks={r['id']:r for r in task_definitions(self.ws,CATALOG)}
        self.assertTrue(set(PROCEDURES) <= set(tasks))
        self.assertEqual(self.ws.procedures,{})
        for pid,module in PROCEDURES.items():
            refs=self.ws.calculation_refs(pid)
            self.assertEqual(refs['실행ID'],runs[module]['id'])
            self.assertIn(self.originals[module],refs['원본자료ID'])
            self.assertIn(self.originals[module],tasks[pid]['evidence_ids'])
            packet=self.ws.ai_packet(pid,CATALOG)
            selected=next(r for r in packet['runs'] if r['module']==module)
            self.assertEqual(selected['result_hash'],runs[module]['result_hash'])
            kinds={r['kind'] for r in packet['references']}
            self.assertTrue({'run_summary','ledger_reconciliation','calculation_row'} <= kinds, (pid,kinds))
            self.assertTrue(any(r['kind']=='calculation_row' and r.get('run_id')==runs[module]['id']
                                for r in packet['references']), pid)
            input_ids={runs[module]['source_id'],self.originals[module]}
            self.assertTrue(any(r['kind'] in ('raw_json','raw_xlsx_row') and r.get('source_id') in input_ids
                                for r in packet['references']), pid)
            before=dict(self.ws.procedures)
            did=self.ws.record_ai_draft(packet,synthetic_response(packet),'SYNTHETIC ONLY')
            self.assertEqual(before,self.ws.procedures)
            procedure=self.ws.apply_ai_draft(did,'SYNTHETIC ONLY')
            self.assertEqual(procedure['state'],'작성중')
            self.assertEqual(procedure['reviewer'],'')
            self.assertEqual(procedure['performed_date'],'')
            self.assertFalse(self.ws.ai_status(pid)['approved'])
        self.assertFalse(self.ws.ai_approvals);self.assertFalse(self.ws.requests)
        self.assertTrue(set(PROCEDURES) <= self.ws.required_procedures())
        snapshot=self.ws.snapshot(False)
        self.assertEqual(snapshot['export_state'],'검토중 초안')
        self.assertTrue(all(p['current_state']=='작성중' for p in snapshot['procedures']))
        data=exporter.workbook_bytes(snapshot,CATALOG)
        workbook=load_workbook(BytesIO(data),data_only=False)
        try:
            self.assertIn('AI근거참조',workbook.sheetnames)
            self.assertIn('검사실행목록',workbook.sheetnames)
            self.assertTrue(all(c.data_type!='f' for sheet in workbook for row in sheet for c in row))
            listed={row[0] for row in workbook['검사실행목록'].iter_rows(min_row=2,values_only=True)}
            self.assertTrue(set(MODULES) <= listed)
        finally:workbook.close()
        with self.assertRaises(AuditError):self.ws.snapshot(True)

    def test_partial_confirmations_classification_pending_and_asset_reversals_are_preserved(self):
        for module in ('유형자산증감','금융기관조회','분류검토'):self.execute_module(module)
        latest=self.ws.latest()
        pairs=latest['유형자산증감']['result']['tables']['대체취소쌍검사']
        self.assertEqual({r['쌍ID'] for r in pairs},{'T1','C1'})
        self.assertEqual(latest['금융기관조회']['result']['metrics']['후속요청항목'],25)
        classification=latest['분류검토']['result']['tables']['분류AsIsToBe검토'][0]
        self.assertEqual(classification['검토상태'],'근거입력완비·사람판단필요')
        self.assertFalse(classification['프로그램적정판정'])
        self.assertFalse(self.ws.procedures);self.assertFalse(self.ws.requests)

    def test_rerunning_identical_ledger_requires_new_module_run_and_rejects_stale_review(self):
        module='유형자산증감';run=self.execute_module(module)
        newer=self.ws.execute('원장·리드',self.ledger_path)
        self.assertNotEqual(newer['id'],self.ledger_run['id'])
        row=next(r for r in self.ws.cross_checks() if r['모듈']==module)
        self.assertEqual(row['상태'],'미확인')
        self.manual_procedure('D02',module,[run['source_id'],newer['source_id'],self.originals[module]])
        with self.assertRaisesRegex(AuditError,'현재 원장'):self.ws.request_review('D02')
        rerun=self.ws.execute(module,self.inputs[module])
        self.assertEqual(rerun['result']['blocking'],0)
        self.assertEqual(next(r for r in self.ws.cross_checks() if r['모듈']==module)['상태'],'일치')
        self.assertEqual(self.ws.procedure_state(self.ws.procedures['D02']),'자료변경·재검토')

    def test_original_evidence_is_required_in_addition_to_calculation_and_ledger(self):
        module='유형자산증감';run=self.execute_module(module)
        self.manual_procedure('D02',module,[run['source_id'],self.ledger_run['source_id']])
        with self.assertRaisesRegex(AuditError,'원본 자료'):self.ws.request_review('D02')
        self.assertEqual(self.ws.procedures['D02']['state'],'작성중')
        self.assertFalse(self.ws.ai_approvals)

    def test_changed_or_deleted_registered_original_blocks_packet_run_and_export(self):
        module='유형자산증감';self.execute_module(module)
        source=self.ws.sources[self.originals[module]]
        (self.ws.root/source['file']).write_bytes(b'changed synthetic original')
        before=self.ws.head
        with self.assertRaises(AuditError):self.ws.ai_packet('D02',CATALOG)
        with self.assertRaises(AuditError):self.ws.execute(module,self.inputs[module])
        with self.assertRaises(AuditError):self.ws.snapshot(False)
        self.assertEqual(before,self.ws.head)

    def test_invalid_origin_id_or_hash_cannot_become_an_accepted_run(self):
        module='유형자산증감';payload=sample(module,self.context)
        raw=self.workbook(module);sid=self.ws.add_source(raw,'SYNTHETIC ONLY original')
        for source in ({'id':'missing','sha256':'1'*64},{'id':sid,'sha256':'1'*64}):
            wrong=self.confirm(module,payload,[source])
            run=self.ws.execute(module,self.write_json(wrong))
            self.assertGreater(run['result']['blocking'],0)
            self.assertIn('RUN_FAILED',{r['code'] for r in run['result']['issues']})
            self.assertEqual(next(r for r in self.ws.cross_checks() if r['모듈']==module)['상태'],'미확인')
        self.assertFalse(self.ws.procedures)

    def test_missing_required_procedures_never_create_review_or_final_records(self):
        self.execute_module('유형자산증감')
        self.assertIn('D02',self.ws.required_procedures())
        before=self.ws.head
        with self.assertRaises(AuditError):self.ws.snapshot(True)
        self.assertEqual(before,self.ws.head)
        self.assertEqual(self.ws.procedures,{})
        self.assertEqual(self.ws.ai_approvals,{})


if __name__=='__main__':unittest.main()
