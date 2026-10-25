"""Synthetic lifecycle integration shared by source and frozen verification.

All performance and review declarations below are explicitly synthetic test
inputs. This code does not exercise a model or approve a user's workpaper.
"""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import sys
import traceback
import zipfile
import runtime
ROOT=runtime.configure()
from common import AuditError, VERSION, digest, dumps
from workspace import Workspace
from procedures import task_definitions
import exporter
from audit_archive import create_archive,verify_archive
from audit_forms import load_registry,forms_for_procedure

CONTEXT={'회사코드':'SYNTH01','회사명':'SYNTHETIC ONLY / v1.9 흐름 검증',
         '시작일':'2026-01-01','기준일':'2026-12-31','회계체계':'K-IFRS',
         '감사기준':'SYNTHETIC ONLY / 실제 감사 수행 아님','작성자':'합성작성자','더미업무':True}


def synthetic_procedure(pid,ids):
    record={'id':pid,'title':'합성 검증 '+pid,'preparer':'합성작성자','purpose':'합성 입력·연결 검증',
            'risk':'SYNTHETIC ONLY / 실제 위험 판단 아님','scope':'선택한 합성 입력',
            'performed':'SYNTHETIC ONLY / 검증 코드의 모의 수행 입력',
            'conclusion':'실제 감사 결론 없음. 프로그램 상태 검증용.',
            'performed_date':'2026-10-25','evidence_ids':ids,'excluded':False}
    if pid in ('F01','FS01'):record['purpose_code']='subsequent_events' if pid=='F01' else 'financial_statement_reconciliation'
    return record


def exercise_core(output,resource_root=None):
    resource_root=Path(resource_root or ROOT);output=Path(output)
    if output.exists() and any(output.iterdir()):raise AuditError('기존 검증 결과를 보존하고 새 폴더를 사용하세요.')
    output.mkdir(parents=True,exist_ok=True);checks=[]
    def check(value,name):
        if not value:raise AssertionError(name)
        checks.append(name)
    def rejected(operation,name):
        try:operation()
        except AuditError:check(True,name)
        else:raise AssertionError(name+' did not reject')
    ws=Workspace(output/'업무',create=CONTEXT)
    for module,name in [('원장·리드','통합_01_원장정상.xlsx'),('금융상품','통합업무_금융정상.xlsx'),
                        ('매출','통합업무_매출정상.xlsx'),('감가상각','통합업무_감가상각정상.xlsx')]:
        ws.execute(module,resource_root/'samples/통합'/name)
    check(not any(run['result']['blocking'] for run in ws.latest().values()),'four_existing_modules_no_blocking')
    check(all(row['상태']=='일치' for row in ws.cross_checks()),'independent_ledger_calculation_links_match')
    tasks=task_definitions(ws,json.loads((resource_root/'procedure_catalog.json').read_text(encoding='utf-8')))
    check(len({t['id'] for t in tasks})==len(tasks),'all_workpaper_definitions_have_unique_ids')
    byid={task['id']:task for task in tasks}
    check(byid['F01']['purpose_code']=='subsequent_events' and byid['FS01']['purpose_code']=='financial_statement_reconciliation','F01_and_FS01_are_distinct_purposes')
    registry=load_registry()
    check(len(registry['forms'])==116 and sum(len(f['sheet_names']) for f in registry['forms'])==435,'reference_inventory_116_files_435_sheets')
    check({s for form in forms_for_procedure('FS01','KIFRS') for s in form['linked_sheets']}=={'3650','8400'},'FS01_declared_reference_sheet_names')
    pids=sorted(ws.required_procedures())
    profile={'audit_framework':'KSA_200_720','audit_framework_reason':'SYNTHETIC ONLY / 모의 사람 선택',
             'form_framework':'KIFRS','materiality':{'benchmark_name':'합성 기준금액','benchmark_amount':100000000,
                 'materiality_rate':'0.05','performance_rate':'0.75','trivial_rate':'0'},
             'scope':[{'procedure_id':pid,'applicable':True,'reason':'SYNTHETIC ONLY / 검증 범위'} for pid in pids]}
    saved=ws.save_workflow_profile(profile,'합성계획자')
    check((saved['materiality']['overall_materiality'],saved['materiality']['performance_materiality'],saved['materiality']['clearly_trivial_threshold'])==(5000000,3750000,0),'materiality_recalculates_explicit_human_rates')
    check(ws.engagement['회계체계']=='K-IFRS' and saved['form_framework']=='KIFRS','reporting_framework_and_form_family_stay_separate')
    risk={'id':'SYNTH-R1','title':'합성 자료 누락 검증','accounts':['선택 합성 계정'],'assertions':['완전성'],
          'procedure_ids':pids,'basis':ws.data_basis(),'evidence_ids':list(ws.sources),
          'nature':'합성 증빙·원장 대사','timing':'합성 기준일 이후','extent':'제공된 합성 범위','assessment':'SYNTHETIC ONLY / 모의 판단'}
    ws.save_workflow_risk(risk,'합성위험기록자')
    links={}
    for pid in pids:
        link={'procedure_id':pid,'risk_ids':['SYNTH-R1'],'evidence_ids':list(ws.sources),'basis':ws.data_basis(),
              'nature':'합성 대사','timing':'합성 기준일','extent':'합성 자료','plan':'SYNTHETIC ONLY / 모의 수행 계획'}
        ws.save_workflow_procedure(link,'합성계획자');links[pid]=link
    check(any(row['code']=='PROCEDURE_NOT_PERFORMED' for row in ws.workflow_status()['human_pending']),'plans_do_not_create_actual_performance')
    check(not ws.procedures and not ws.ai_approvals,'planning_does_not_create_workpaper_approval')
    rejected(lambda:ws.snapshot(True),'final_output_blocks_missing_human_performance_and_review')
    before=ws.head;draft=create_archive(ws,output/'검토중취합.zip','합성보관자')
    check(draft['valid'] and draft['mode']=='draft' and ws.head==before,'draft_archive_marks_scope_without_mutating_history')
    check(verify_archive(output/'검토중취합.zip')['valid'],'draft_archive_full_verification')
    for pid,link in links.items():
        ws.save_workflow_procedure({**link,'performed':'SYNTHETIC ONLY / 테스트의 모의 수행 기록',
            'performed_date':'2026-10-25','performer':'합성수행자'},'합성계획자')
    for pid in pids:
        ws.save_procedure(synthetic_procedure(pid,list(ws.sources)))
        ws.request_review(pid);ws.review(pid,'합성별도검토자','SYNTHETIC ONLY / 모의 검토')
    check(not ws.workflow_status()['blockers'],'complete_synthetic_records_satisfy_mechanical_gates')
    snap=ws.snapshot(True)
    bookraw=exporter.workbook_bytes(snap,json.loads((resource_root/'procedure_catalog.json').read_text(encoding='utf-8')))
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(bookraw),read_only=True,data_only=False)
    try:check({'감사계획·중요성','위험주장대응','절차계획·수행연결','완료점검표','기준서식참고'}<=set(book.sheetnames),'workbook_contains_planning_links_completion_and_references')
    finally:book.close()
    ws.finalize_workflow('합성보관자','SYNTHETIC ONLY / 모의 검토 상태 잠금')
    check(ws.workflow_status()['locked'],'explicit_finalization_locks_current_records')
    locked_head=ws.head
    rejected(lambda:ws.add_note('잠긴 상태 변경','차단 검증','합성작성자'),'locked_workspace_rejects_edits')
    check(ws.head==locked_head,'rejected_edit_preserves_frozen_head')
    final=create_archive(ws,output/'검토완료취합.zip','합성보관자',final=True)
    check(final['valid'] and final['mode']=='final' and verify_archive(output/'검토완료취합.zip')['head']['hash']==ws.head,'final_archive_reconstructs_same_reviewed_snapshot')
    rejected(lambda:create_archive(ws,output/'검토완료취합.zip','합성보관자',final=True),'archive_never_overwrites_existing_file')
    raw=(output/'검토완료취합.zip').read_bytes();before_hash=hashlib.sha256(raw).hexdigest()
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        damaged={entry.filename:archive.read(entry) for entry in archive.infolist()}
    source_name=next(name for name in damaged if name.startswith('원본/'))
    damaged[source_name]+=b'CORRUPT'
    with zipfile.ZipFile(output/'변조취합.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for name,value in damaged.items():archive.writestr(name,value)
    rejected(lambda:verify_archive(output/'변조취합.zip'),'tampered_archive_source_is_rejected')
    ws.reopen_workflow('합성재개자','SYNTHETIC ONLY / 수정 재개 검증')
    check(not ws.workflow_status()['locked'],'explicit_reopen_records_reason')
    ws.save_workflow_risk({**risk,'assessment':'SYNTHETIC ONLY / 근거 재검토 입력'},'합성위험기록자')
    check(any(row['code']=='WORKPAPER_REVIEW_PENDING' for row in ws.workflow_status()['human_pending']),'planning_change_requires_fresh_review')
    check(hashlib.sha256((output/'검토완료취합.zip').read_bytes()).hexdigest()==before_hash,'reopen_and_change_preserve_prior_archive_bytes')
    old=Workspace(output/'구형업무',create=CONTEXT)
    record=synthetic_procedure('F01',[]);record.pop('purpose_code')
    record.update(basis=old.basis(),state='작성중',reviewer='',review_note='',reviewed_at='',edit_revision='legacy-test')
    old._append('procedure',record);old.refresh();events=deepcopy(old.events)
    check(any(row['code']=='LEGACY_F01_UNCLASSIFIED' for row in old.workflow_status()['blockers']),'legacy_ambiguous_F01_is_not_guessed')
    old.classify_legacy_f01('financial_statement_reconciliation','합성분류자','SYNTHETIC ONLY / 과거 목적 확인')
    check(old.events[:len(events)]==events and old.legacy_f01_documents==[record],'legacy_classification_preserves_original_events')
    for pid in ('F01','FS01'):old.save_procedure(synthetic_procedure(pid,[]))
    check(set(old.procedures)=={'F01','FS01'},'legacy_resolution_allows_new_independent_workpapers')
    return {'ok':True,'version':VERSION,'checks':checks,'workflow_workspace':str(ws.root),
            'fixture_scope':'Synthetic performance/review inputs only; no real audit or model call.',
            'current_actual_model_calls':0,'archive_validated':True,'legacy_f01_preserved':True}


def self_test(report_path):
    report_path=Path(report_path).resolve();out=report_path.parent
    if out.exists() and any(out.iterdir()):raise AuditError('새 검증 출력 폴더를 사용하세요.')
    out.mkdir(parents=True,exist_ok=True);root=None;app=None
    try:
        record=exercise_core(out/'lifecycle');checks=record['checks']
        from unittest.mock import patch
        import tkinter as tk
        from app import App,CATALOG
        root=tk.Tk();root.withdraw();app=App(root,sector_home=out/'isolated-sectors')
        app.open_workspace(record['workflow_workspace']);app.open_workflow_view()
        def check(value,name):
            if not value:raise AssertionError(name)
            checks.append(name)
        check(root.state()=='withdrawn','test_root_stays_hidden')
        check(app.nb.select()==str(app.tabs['계획·절차·취합']),'desktop_workflow_tab_is_reachable')
        check(len(app._wf_tasks())>=54,'all_planning_procedures_available_before_execution')
        d=app.open_workflow_profile()
        check(d.win.state()=='withdrawn','test_plan_dialog_stays_hidden')
        check(d.value('form_framework')=='KIFRS','desktop_reads_separate_form_framework')
        d.close(force=True)
        app.load_procedure('A04');narrative='SYNTHETIC ONLY / 한글 조서 수정\n받침 검토 · 증거 연결 · 😀\n사람의 실제 감사 판단은 별도입니다.'
        app.proc_text['conclusion'].delete('1.0','end');app.proc_text['conclusion'].insert('1.0',narrative);app.save_procedure()
        check(app.ws.procedures['A04']['conclusion']==narrative,'native_unicode_and_newlines_survive_save')
        path=out/'UI검토중조서.xlsx';exporter.export(app.ws,path,CATALOG)
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True)
        try:
            proof=''.join(str(row[1] or '') for row in book['증적JSON'].iter_rows(min_row=2,values_only=True))
            saved=json.loads(proof);actual=next(p['conclusion'] for p in saved['procedures'] if p['id']=='A04')
            check(actual==narrative,'native_unicode_matches_Excel_evidence_JSON')
        finally:book.close()
        check(not app.ws.ai_approvals,'desktop_does_not_invent_AI_approval')
        checks.append('human_pending_remains_visible_after_edit') if app.wf_state['human_pending'] else (_ for _ in ()).throw(AssertionError('pending missing'))
        record.update(ok=True,frozen=bool(getattr(sys,'frozen',False)),unicode_roundtrip_verified=True,
                      checks=checks,uses_hidden_native_windows=True,current_actual_key_test=False)
    except Exception as exc:
        record={'ok':False,'version':VERSION,'frozen':bool(getattr(sys,'frozen',False)),
                'error':str(exc),'traceback':traceback.format_exc()}
    finally:
        if app is not None:app.shutdown_ai()
        if root is not None:
            from window_capture import destroy_test_window
            destroy_test_window(root)
    report_path.write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(record,ensure_ascii=False));return 0 if record['ok'] else 1
