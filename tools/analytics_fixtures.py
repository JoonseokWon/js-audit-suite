"""Write reproducible synthetic analyzer inputs through the established XLSX format."""
from pathlib import Path
from copy import deepcopy
import sys
ROOT=Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
from common import read_xlsx
import ledger

def load_base():
    tables,_=read_xlsx(ROOT/'samples/통합/통합_v1.3_원장_리드.xlsx',ledger.SCHEMA)
    return deepcopy(tables)

def write_input(path,extra_schema,extra_tables,base=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    schemas={**ledger.SCHEMA,**extra_schema};tables={**(load_base() if base is None else deepcopy(base)),**deepcopy(extra_tables)}
    w=Workbook();w.remove(w.active)
    for name,headers in schemas.items():
        s=w.create_sheet(name);s.append(headers)
        for row in tables.get(name,[]):
            s.append([row.get(h) for h in headers])
            for c in s[s.max_row]:
                if isinstance(c.value,str):c.data_type='s'
        for c in s[1]:c.font=Font(name='맑은 고딕',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='263F59')
        for i,h in enumerate(headers,1):s.column_dimensions[get_column_letter(i)].width=min(45,max(17,len(h)*2+3))
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions;s.sheet_view.showGridLines=False
    w.save(path);w.close();return path
