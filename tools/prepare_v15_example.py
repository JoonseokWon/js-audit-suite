from pathlib import Path
import sys,json,shutil,hashlib
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from app import App,CATALOG
from analytics_validation import BASE_INPUTS,CONTEXT
from workspace import Workspace
from common import AuditError
import exporter
out=ROOT/'outputs/01a16140-43b4-7e72-8585-04010bd5e6c9';folder=out/'예시업무_미검토'
assert not folder.exists(),'Preserve earlier examples'
ws=Workspace(folder,create=CONTEXT)
for module,name in BASE_INPUTS:
    r=ws.execute(module,ROOT/'samples/통합'/name);assert r['result']['blocking']==0,(module,r['result']['issues'])
assert len(ws.latest())==15 and len(ws.cross_checks())==35
for task in App.task_list(type('Context',(),{'ws':ws})()):
    if task['id'] in ws.required_procedures():ws.save_procedure({**task,'preparer':'합성작성자'})
row=ws.risk_register()[0];draft=ws.candidate_request_draft(row['후보키'])
ws.save_request({**draft,'requested_from':'가상 담당자 / SYNTHETIC ONLY'})
path=out/'통합감사_미검토초안_v1.5.xlsx';snap=exporter.export(ws,path,CATALOG)
assert all(not p['performed'] and not p['conclusion'] and not p['reviewer'] for p in snap['procedures'])
try:ws.snapshot(True)
except AuditError:pass
else:raise AssertionError('Unreviewed example bypass')
from review_support import save_context
save_context(ws,out/'검토맥락_예시.json')
from openpyxl import load_workbook
w=load_workbook(path,data_only=False)
assert all(c.data_type not in ('f','e') for s in w for row in s for c in row)
record={'ok':True,'modules':15,'current_links':35,'candidates':len(snap['risk_register']),'procedures':len(snap['procedures']),'human_fields_blank':True,'final_export_blocked':True,'xlsx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column} for s in w]};w.close()
(out/'예시검증.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
dist=ROOT/'배포';(dist/'예시조서').mkdir(exist_ok=True)
for p in (path,out/'검토맥락_예시.json'):shutil.copy2(p,dist/'예시조서'/p.name)
shutil.copytree(folder,dist/'예시업무_미검토')
print(json.dumps({k:v for k,v in record.items() if k!='sheets'},ensure_ascii=False))
