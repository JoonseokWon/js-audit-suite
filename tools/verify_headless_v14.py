"""Deterministic file-case and workspace checks without importing Tk or App.

Run from any directory, supplying a new, explicit --output-dir. Expected values
are assertions after the normal APIs run; none are passed into calculators.
The 270 figure counts fixture cases, not UI assertions or unit test methods.
"""
from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path
import hashlib
import json
import platform
import sys
import time
import traceback

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import runtime
runtime.configure()
from common import AuditError, atomic, digest, dumps, sha
import adapters
import exporter
from workspace import Workspace
from verify import check_book, regression

CONTEXT = {'회사코드':'SYNTH01', '시작일':'2026-01-01', '기준일':'2026-12-31'}
# Kept explicit to avoid importing app.py/Tk just to discover the demonstration.
NORMAL_INPUTS = (
    ('원장·리드','통합_v1.3_원장_리드.xlsx'), ('금융상품','통합_v1.3_금융상품.xlsx'),
    ('매출','통합_v1.3_매출.xlsx'), ('감가상각','통합_v1.3_감가상각.xlsx'),
    ('확장검사','통합_v1.3_확장검사.xlsx'), ('리스','통합_v1.3_리스.xlsx'),
    ('대손충당금','통합_v1.3_대손충당금.xlsx'), ('재고평가','통합_v1.3_재고평가.xlsx'),
    ('종업원급여','통합_v1.3_종업원급여.xlsx'), ('미기록부채·비용이연','비용이연_01_정상.xlsx'),
    ('원천증빙','통합_v1.4_원천증빙.auditpack'), ('재무제표·주석','통합_v1.4_재무제표.json'),
)
COMMON_EXPECTED = {
    'normal': (False, False, None),
    'balanced_journal_missing': (True, False, 'RAW_REFERENCE'),
    'duplicate_allocation': (True, False, 'RAW_ALLOC_DUP'),
    'partial_collection': (False, True, 'RAW_UNALLOCATED'),
    'split_payment': (False, False, None),
    'wrong_company': (True, False, 'RAW_COMPANY'),
    'wrong_period': (True, False, 'RAW_PERIOD'),
    'blank_source_amount': (True, False, 'RAW_AMOUNT'),
    'wrong_sign': (True, False, 'RAW_SIGN'),
    'overallocated': (True, False, 'RAW_OVERALLOC'),
}


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def exported_rows_match(book, sheet, expected):
    """Reconstruct exact logical cells; physical continuation rows are not records.

    This verification helper neither changes the exporter nor treats a larger
    physical row count as sufficient. Every physical row must belong to exactly
    one expected record, and scalar/numeric cells occur only on its first row.
    """
    from export_support.text import display_text
    matrix = list(sheet.iter_rows(values_only=True))
    if not matrix or not expected:
        return False
    headers = matrix[0]
    keys = {key for record in expected for key in record}
    if (any(not isinstance(key, str) for key in headers) or len(headers) != len(set(headers))
            or set(headers) != keys):
        return False
    mappings = {}
    if '출력문단연결' in book.sheetnames:
        index = list(book['출력문단연결'].iter_rows(values_only=True))
        required = ('표', '원본행', '시작행', '마지막행', '분할열')
        if not index or any(key not in index[0] for key in required):
            return False
        positions = [index[0].index(key) for key in required]
        for row in index[1:]:
            name, logical, start, end, split = [row[i] for i in positions]
            if name != sheet.title:
                continue
            if (any(type(value) is not int for value in (logical, start, end)) or
                    not 1 <= logical <= len(expected) or logical in mappings or
                    not 2 <= start <= end <= len(matrix)):
                return False
            columns = split.split(', ') if isinstance(split, str) and split else []
            if len(columns) != len(set(columns)) or not set(columns) <= set(headers):
                return False
            mappings[logical] = (start, end, set(columns))
    cursor = 2
    for logical, record in enumerate(expected, 1):
        start, end, splits = mappings.get(logical, (cursor, cursor, set()))
        if start != cursor or end > len(matrix):
            return False
        fragments = matrix[start-1:end]
        for column, key in enumerate(headers):
            original = record.get(key)
            if isinstance(original, (dict, list)):
                original = json.dumps(original, ensure_ascii=False, sort_keys=True,
                                      separators=(',', ':'), allow_nan=False)
            if isinstance(original, str):
                original = display_text(original)
            values = [row[column] for row in fragments]
            if key in splits:
                if not isinstance(original, str) or any(value is not None and not isinstance(value, str) for value in values):
                    return False
                actual = ''.join(value or '' for value in values)
            else:
                if any(value is not None for value in values[1:]):
                    return False
                actual = values[0]
            original = None if original == '' else original
            if isinstance(original, bool):
                if type(actual) is not bool or actual != original:
                    return False
            elif isinstance(original, (int, float)):
                if isinstance(actual, bool) or not isinstance(actual, (int, float)) or actual != original:
                    return False
            elif actual != original:
                return False
        cursor = end + 1
    return cursor == len(matrix) + 1


def fresh(folder):
    return Workspace(folder,create={**CONTEXT, '회사명':'SYNTHETIC ONLY / headless 합성업무',
        '회계체계':'K-IFRS / 제한된 합성 사례', '감사기준':'학습용 / 실제 감사 수행 아님',
        '작성자':'합성작성자', '더미업무':True})


def observed_case(case, module, actual):
    return {'case':case, 'module':module, 'blocking':actual['blocking'], 'review':actual['review'],
            'codes':[x['code'] for x in actual['issues']], 'metrics':actual['metrics']}


def tracked_inputs():
    """Fingerprint input fixtures and immutable application code, excluding output."""
    paths = [p for p in (ROOT/'samples').rglob('*') if p.is_file() and p.suffix.lower() in ('.xlsx','.json','.auditpack')]
    paths += [ROOT/p for p in ('samples_specs.json','new_specs.json','v12_specs.json','v13_specs.json',
                              'worker_cases.json','procedure_catalog.json')]
    paths += list(ROOT.glob('*.py'))
    for package in ('engines','financial_bridge','raw_accounts','raw_sources','worker_library'):
        paths += [p for p in (ROOT/package).rglob('*.py') if '__pycache__' not in p.parts]
    return {p.relative_to(ROOT).as_posix():sha(p) for p in sorted(set(paths))}


def run(output):
    started = time.monotonic()
    checks, cases = [], []
    report = {'ok':False, 'mode':'headless', 'python':platform.python_version(), 'platform':platform.system(),
              'checks':checks, 'cases':cases, 'limitations':[
                  'Tk/App/screenshots/OS dialogs are not exercised; native EXE verification is a separate run.',
                  'All inputs are synthetic; no independent external evidence or actual audit conclusion.',
                  'The 33 account cases are independent exercises and never summed into the common TB.',
                  'The raw package indexes 114 tables but tests only explicitly selected allocations/scopes.',
              ]}
    def check(value, name):
        if not value: raise AssertionError(name)
        checks.append(name)
    def rejected(fn, name):
        try: fn()
        except AuditError: check(True, name)
        else: raise AssertionError(name + ': expected AuditError')
    try:
        before = tracked_inputs()
        check('tkinter' not in sys.modules and 'app' not in sys.modules, 'no_gui_import_at_start')

        # Existing 55 engine fixtures plus 158 integrated fixtures = 213.
        old = regression(ROOT)
        check(len(old)==55 and all(x['ok'] for x in old), 'legacy_engine_55_cases')
        cases.extend({'group':'legacy','case':x['name'],'module':x['module'],'expected_verified':x['ok']} for x in old)
        legacy_count = len(old)
        for spec_name, expected_count in [('samples_specs.json',28),('new_specs.json',36),('v12_specs.json',44),('v13_specs.json',50)]:
            spec = read_json(ROOT/spec_name)
            books = [b for b in spec['books'] if not b['expected'].get('template')]
            check(len(books)==expected_count, spec_name+'_fixture_count')
            for book in books:
                path = ROOT/'samples/통합'/(book['name']+'.xlsx')
                if 'input_error' in book['expected']:
                    try: adapters.run(book['module'],path,spec['context'])
                    except AuditError as exc: check(book['expected']['input_error'] in str(exc), book['name']+'_expected_input_rejection')
                    else: raise AssertionError(book['name']+': expected input rejection')
                    cases.append({'group':'legacy','case':book['name'],'module':book['module'],'expected_rejection':True})
                else:
                    actual = check_book(book,path,spec['context'])
                    check(True,book['name']+'_pinned_expectations')
                    cases.append({'group':'legacy',**observed_case(book['name'],book['module'],actual)})
                legacy_count += 1
        check(legacy_count==213,'all_213_existing_file_cases')

        # Normal 12-module workspace. The data APIs and immutable source copies
        # are exactly those used by the desktop application; no App import.
        ws = fresh(output/'normal_workspace')
        for module, filename in NORMAL_INPUTS:
            run_record = ws.execute(module,ROOT/'samples/통합'/filename)
            check(run_record['result']['blocking']==0,module+'_normal_not_blocked')
            check(run_record['result']['review']==0,module+'_normal_no_difference')
        check(len(ws.latest())==12,'normal_twelve_modules')
        cross = ws.cross_checks()
        check(len(cross)==32 and all(x['상태']=='일치' for x in cross),'normal_32_current_ledger_reconciliations')
        check(ws.verify_sources()==[],'normal_source_copies_intact')
        latest = ws.latest()
        lead = latest['원장·리드']['result']['tables']['리드스케줄']
        check(len(lead)==59,'normal_59_lead_accounts')
        financial = latest['재무제표·주석']['result']
        check(financial['metrics']=={'대사항목수':94,'차이항목수':0,'매핑계정수':59,'필수주석수':8},'financial_94_checks_59_accounts_8_notes')
        amounts = {x['항목ID']:x['보고금액'] for x in financial['tables']['재무제표대사'] if x['코드']=='FS_TB'}
        for item, value in {'A_CURRENT':4002752,'A_NONCURRENT':1435817,'ASSETS':5438569,'LIABILITIES':5283551,
                            'EQUITY':155018,'REVENUE':146400,'EXPENSE':6730382,'PROFIT':-6583982}.items():
            check(amounts[item]==value,'financial_hand_oracle_'+item)
        check(amounts['ASSETS']==amounts['LIABILITIES']+amounts['EQUITY'],'financial_hand_balance_equation')
        raw = latest['원천증빙']['result']
        check(len(raw['tables']['현재원장연결'])==141,'raw_141_actual_workspace_ledger_rows')
        check(raw['metrics']['카탈로그테이블수']==114 and raw['metrics']['배부건수']==32,'raw_114_indexed_tables_32_selected_allocations')
        check(raw['metrics']['미확인범위수']==8,'raw_eight_unconfirmed_scopes_preserved')
        for module in ('원천증빙','재무제표·주석'):
            d=latest[module]['result']['detail']
            check(d['ledger_run_id']==latest['원장·리드']['id'] and d['ledger_result_hash']==latest['원장·리드']['result_hash'],module+'_pinned_to_actual_ledger_run')

        # A blank draft must neither invent human execution/conclusions nor
        # permit a reviewed export. Export/reopen checks the actual workbook.
        for pid in ('FS01','R01'):
            ws.save_procedure({'purpose_code':'financial_statement_reconciliation' if pid=='FS01' else pid,'id':pid,'title':'합성 검토 초안 '+pid,'preparer':'합성작성자',
                'purpose':'등록 원천과 수치 대사 범위 검토','risk':'합성 자료·선택 범위 제한',
                'scope':'현재 원장 연결 자동검사','performed':'','conclusion':'','performed_date':'',
                'evidence_ids':list(ws.sources)})
            check(ws.procedures[pid]['state']=='작성중' and not ws.procedures[pid]['reviewer'],pid+'_draft_not_auto_reviewed')
            rejected(lambda pid=pid:ws.request_review(pid),pid+'_blank_performance_cannot_request_review')
        rejected(lambda:ws.snapshot(True),'normal_unreviewed_export_blocked')
        catalog=read_json(ROOT/'procedure_catalog.json')
        draft=output/'normal_draft.xlsx'
        exporter.export(ws,draft,catalog,False)
        from openpyxl import load_workbook
        book=load_workbook(draft,read_only=True,data_only=False)
        try:
            check('조서_FS01' in book.sheetnames and '조서_R01' in book.sheetnames,'draft_FS01_R01_sheets')
            fs_sheet=next(s for s in book if s.title.endswith('_재무제표대사'))
            raw_sheet=next(s for s in book if s.title.endswith('_현재원장연결'))
            check(len(financial['tables']['재무제표대사'])==94 and len(raw['tables']['현재원장연결'])==141
                  and exported_rows_match(book,fs_sheet,financial['tables']['재무제표대사'])
                  and exported_rows_match(book,raw_sheet,raw['tables']['현재원장연결']),
                  'draft_all_FS_and_current_ledger_rows_exported')
            proof=json.loads(''.join(row[1] for row in book['증적JSON'].iter_rows(min_row=2,values_only=True)))
            check(len(proof['latest'])==12 and len(proof['cross'])==32,'draft_full_proof_12_modules_32_links')
            check(all(p['current_state']=='작성중' and not p['performed'] and not p['conclusion'] and not p['reviewer'] for p in proof['procedures']),'draft_proof_has_no_human_conclusion')
            check(all(c.data_type!='f' for sheet in book for row in sheet for c in row),'draft_no_executable_excel_formulas')
        finally: book.close()
        normal_summary={'modules':12,'reconciliations':32,'lead_accounts':59,'financial_comparisons':94,
                        'raw_current_ledger_rows':141,'draft_file':draft.name,'draft_sha256':sha(draft)}

        # Independent workers: API reads each input file and removes any embedded
        # expected data before the calculator is called. The oracle stays here.
        entries=read_json(ROOT/'worker_cases.json')
        check(len(entries)==33,'independent_33_case_files')
        account_checks=0
        for entry in entries:
            isolated=fresh(output/'independent_accounts'/entry['case_id'])
            actual=isolated.execute('계정별증빙연습',ROOT/entry['file'])['result']
            rows=actual['tables']['계정별검사']; indexed={r['검사ID']:r for r in rows}
            expected=entry['expected']
            check(len(indexed)==len(rows) and set(indexed)=={e['check_id'] for e in expected},entry['case_id']+'_exact_check_scope')
            for e in expected:
                for field,key in [('상태','status'),('기대','expected_value'),('실제','actual_value'),('차이','difference')]:
                    check(indexed[e['check_id']][field]==e[key],entry['case_id']+'_'+e['check_id']+'_'+field)
                account_checks+=1
            check(actual['detail']['현재원장연결'] is False and actual['detail']['입력기대값사용'] is False,entry['case_id']+'_independent_no_expected_injection')
            check(set(isolated.latest())=={'계정별증빙연습'},entry['case_id']+'_not_merged_into_TB')
            cases.append({'group':'account',**observed_case(entry['case_id'],'계정별증빙연습',actual)})
        check(account_checks==300,'account_300_pinned_check_results')

        # Normal financial case was already exercised on the 12-module ledger.
        cases.append({'group':'financial',**observed_case('normal','재무제표·주석',financial)})
        fs_root=ROOT/'samples/재무제표'
        fs_cases=read_json(fs_root/'financial_cases/cases.json')['cases']
        check(len(fs_cases)==13,'financial_13_altered_files_plus_normal')
        fs_ws=fresh(output/'financial_variants')
        fs_ws.execute('원장·리드',ROOT/'samples/통합/통합_v1.3_원장_리드.xlsx')
        for entry in fs_cases:
            actual=fs_ws.execute('재무제표·주석',fs_root/'financial_cases'/entry['file'])['result']
            for key in ('blocking','review'):
                if key in entry:check(actual[key]==entry[key],entry['name']+'_'+key)
            if 'required_code' in entry:check(entry['required_code'] in [x['code'] for x in actual['issues']],entry['name']+'_required_code')
            if entry['blocking']:
                check(all(not rows for rows in actual['tables'].values()),entry['name']+'_no_comparisons_on_invalid_input')
            if entry['name']=='02_주석차이':
                differences=[x for x in actual['tables']['재무제표대사'] if x['차이']]
                check(len(differences)==1 and differences[0]['코드']=='FS_NOTE' and differences[0]['차이']==700,'financial_note_only_700_won')
            if entry['name']=='10_허용범위내차이':
                differences=[x for x in actual['tables']['재무제표대사'] if x['차이']]
                check(bool(differences) and all(x['상태']=='허용범위내차이' for x in differences),'financial_one_won_difference_remains_open')
            if entry['name']=='12_표시부호오류':
                row=next(x for x in actual['tables']['재무제표대사'] if x['코드']=='FS_TB' and x['항목ID']=='B_1390')
                check(row['차이']==-21296,'financial_signed_allowance_difference')
            if entry['name']=='13_합계는맞지만TB불일치':
                check({x['code'] for x in actual['issues']}=={'FS_TB'},'financial_totals_and_notes_pass_but_actual_TB_fails')
                cash=next(x for x in actual['tables']['재무제표대사'] if x['코드']=='FS_TB' and x['항목ID']=='B_1000')
                check(cash['차이']==1200,'financial_coherent_but_wrong_cash_1200_won')
            cases.append({'group':'financial',**observed_case(entry['name'],'재무제표·주석',actual)})

        # Ten shared-transaction packages: these genuinely share one fixed small
        # ledger. Fixed hand values are assertions, never balancing plugs.
        from raw_accounts.samples import CASES
        check(set(CASES)==set(COMMON_EXPECTED),'common_ten_fixed_scenarios')
        for name in CASES:
            folder=ROOT/'samples/공통거래'/name
            oracle=read_json(folder/'검증결과.json')
            common_ws=fresh(output/'common_transactions'/name)
            gl=common_ws.execute('원장·리드',folder/'원장_공통거래.xlsx')['result']
            actual=common_ws.execute('원천증빙',folder/('공통거래_'+name+'.auditpack'))['result']
            blocked, difference, required_code=COMMON_EXPECTED[name]
            check(bool(actual['blocking'])==blocked,name+'_fixed_workspace_blocking')
            if not blocked:check(bool(actual['review'])==difference,name+'_fixed_workspace_difference')
            if required_code:check(required_code in [x['code'] for x in actual['issues']],name+'_fixed_required_code')
            check(gl['blocking']==oracle['ledger_blocking'],name+'_pinned_ledger_blocking')
            check(actual['metrics']==oracle['raw_metrics'],name+'_pinned_raw_metrics')
            for code in oracle['raw_codes']:check(code in [x['code'] for x in actual['issues']],name+'_pinned_raw_code_'+code)
            tb={x['계정코드']:x for x in gl['tables']['리드스케줄']}
            check(tb['1000']['조정전']==130 and tb['1400']['조정전']==150 and tb['4000']['조정전']==-280,name+'_hand_TB_130_150_minus280')
            inputs=gl['tables']['입력_총계정원장']
            check(len(inputs)==(18 if name=='balanced_journal_missing' else 20),name+'_full_journal_row_count')
            check(len({r['전표ID'] for r in inputs})==(9 if name=='balanced_journal_missing' else 10),name+'_journal_population_count')
            if name=='balanced_journal_missing':
                check('LEDGER_CONTROL' in [x['code'] for x in gl['issues']] and 'RAW_LEDGER_REQUIRED' in [x['code'] for x in actual['issues']],name+'_balanced_missing_detected_before_conclusion')
            if name=='partial_collection':
                check(actual['metrics']['미대사행수']==2 and actual['metrics']['배부건수']==13,name+'_partial_receipt_kept_open')
            if name in ('normal','split_payment'):
                check(actual['metrics']['미대사행수']==0 and actual['metrics']['배부건수']==14,name+'_many_to_many_allocations_complete')
            cases.append({'group':'common',**observed_case(name,'원천증빙',actual)})

        # Re-execution must invalidate both FS and raw bindings even when a new
        # run happens to use identical ledger bytes. Retained history stays put.
        old_ids={m:ws.latest()[m]['id'] for m in ('원천증빙','재무제표·주석')}
        ws.execute('원장·리드',ROOT/'samples/통합/통합_v1.3_원장_리드.xlsx')
        changed=ws.cross_checks()
        for module,pid in [('재무제표·주석','FS01'),('원천증빙','R01')]:
            check(next(x['상태'] for x in changed if x['모듈']==module)=='미확인',module+'_old_binding_invalid_after_ledger_rerun')
            check(ws.latest()[module]['id']==old_ids[module],module+'_old_result_preserved_in_history')
            ws.save_procedure({'purpose_code':'financial_statement_reconciliation' if pid=='FS01' else pid,'id':pid,'title':'합성 연결 재검토 '+pid,'preparer':'합성작성자',
                'purpose':'연결 통제','risk':'자료 버전 변경','scope':'합성 회귀검사',
                'performed':'SYNTHETIC ONLY / 검사 상태 연습','conclusion':'실제 감사 결론 없음',
                'performed_date':'2026-12-31','evidence_ids':list(ws.sources)})
            rejected(lambda pid=pid:ws.request_review(pid),pid+'_stale_binding_review_request_blocked')
        rejected(lambda:ws.snapshot(True),'stale_bindings_completed_export_blocked')

        counts={g:sum(x['group']==g for x in cases) for g in ('legacy','account','financial','common')}
        check(counts=={'legacy':213,'account':33,'financial':14,'common':10},'exact_270_file_case_breakdown')
        check(len(cases)==270,'all_270_file_cases_passed')
        check('tkinter' not in sys.modules and 'app' not in sys.modules,'no_gui_import_at_completion')
        check(tracked_inputs()==before,'all_source_code_and_fixture_hashes_unchanged')
        report.update(ok=True, fixture_cases=270, counts=counts, account_expectations=account_checks,
                      normal_workspace=normal_summary, input_inventory_sha256=digest(before),
                      tracked_input_files=len(before), assertions=len(checks))
    except Exception as exc:
        report.update(error=str(exc),traceback=traceback.format_exc(),assertions=len(checks))
    report['elapsed_seconds']=round(time.monotonic()-started,3)
    atomic(output/'report.json',json.dumps(report,ensure_ascii=False,indent=2).encode('utf-8'))
    print(json.dumps({k:report[k] for k in ('ok','fixture_cases','counts','assertions','elapsed_seconds','error') if k in report},ensure_ascii=False))
    return 0 if report['ok'] else 1


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',required=True,help='New or empty output directory for reports, draft workbook and isolated workspaces.')
    args=parser.parse_args()
    output=Path(args.output_dir).expanduser().resolve()
    protected=[ROOT/p for p in ('samples','templates','engines','financial_bridge','raw_sources','raw_accounts','worker_library','tools')]
    if output==ROOT or output in ROOT.parents or any(output==p or output.is_relative_to(p) for p in protected):
        parser.error('Output must not be the project/source/input directory or its parent.')
    if output.exists() and (not output.is_dir() or any(output.iterdir())):
        parser.error('Output directory must be new or empty; prior evidence is never overwritten.')
    output.mkdir(parents=True,exist_ok=True)
    return run(output)


if __name__=='__main__':raise SystemExit(main())
