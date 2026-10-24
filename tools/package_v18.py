"""Package verified v1.8.0 artifacts. No work, network or application startup on import.

Run only after the PM has frozen build inputs, actual-model evidence and visual
review. Every input is explicitly selected; portable AI folders are never walked.
Original local reports and workbooks remain unchanged. Public reading copies
redact development-PC paths and credential-shaped values and say so explicitly.
"""
from __future__ import annotations

import copy
import hashlib
import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import re
import stat
import sys
import unicodedata
import uuid
import zipfile

ROOT = Path(__file__).resolve().parents[1]
VERSION = '1.8.0'
SOURCE_PREFIX = 'js-audit-suite/'
WINDOWS_PREFIX = 'JS-Audit-Suite-v1.8.0/'
EXE_NAME = '통합 회계감사 v1.8.0.exe'
SOURCE_NAME = 'js-audit-suite-v1.8.0-source.zip'
WINDOWS_NAME = 'JS-Audit-Suite-v1.8.0-windows.zip'
VALIDATION_NAME = 'validation-report.json'
MODEL = 'qwen3:4b-instruct-2507-q4_K_M'
MODEL_DIGEST = '0edcdef34593eac1aa2be9c7d06c432dcf81945adca5eca2f27662c18f168ba0'

# Preserve all previous public paths, including old tools/docs. New additions
# require an explicit name here as well as the release helper's public policy.
NEW_REQUIRED = {
    'ai_desktop.py', 'ai_provider.py', 'ai_workpapers.py', 'local_ai_runtime.py',
    'test_ai_desktop.py', 'test_ai_provider.py', 'test_ai_workpapers.py',
    'test_local_ai_runtime.py', 'verify_ai_actual.py', 'verify_v18.py',
    '문서/AI응답계약.md', '문서/v1.8_사용안내.md', '문서/로컬AI_설치안내.md',
    'ci/audit-suite-v18.yml', 'tools/build_v18.py', 'tools/package_v18.py',
    'tools/github_release_v18.py',
}
LEGACY_MUTABLE = {
    'AGENTS.md', 'README.md', 'app.py', 'common.py', 'exporter.py', 'workspace.py',
    'windows.manifest', 'AuditSuite.spec', 'test_runtime_gui.py',
    '문서/기능별_작동설명_노션원고.md',
}
SECRET = re.compile(r'(?i)(?:\bsk-[a-z0-9_-]{12,}|\bBearer\s+[a-z0-9._-]{12,}|(?:api[_ -]?key|access[_ -]?token|password|secret)\s*[=:]\s*[\"\']?[^\s\"\',}]{6,})')
ABSOLUTE = re.compile(r'(?i)(?<![a-z0-9])(?:[a-z]:[\\/][^\r\n\t\"\'<>|]*|\\\\[^\r\n\t\"\'<>|]+|file:///[^\r\n\t\"\'<>|]+|/(?:Users|home|tmp|mnt|root|private|var|opt)/[^\r\n\t\"\'<>|]*)')
SENSITIVE_KEY = re.compile(r'(?i)^(?:api[_-]?key|access[_-]?token|password|secret|authorization)$')
PUBLIC_NOTE = ('공개 열람용 사본: 개발 PC의 절대경로·인증값 형태를 제거했습니다. '
               '숫자·재무계산 결과는 변경하지 않았습니다. 원본 증적은 로컬에 보존합니다. '
               '경로를 제거한 JSON/이력은 원본 바이트·이벤트 체인 재검증용이 아닙니다. '
               '사람 승인 기록은 합성 상태전이 시험이며 실제 감사 승인·감사의견이 아닙니다.')


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def sha(data):
    return hashlib.sha256(data).hexdigest()


def json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n').encode('utf-8')


def json_digest(value):
    return sha(json.dumps(value, ensure_ascii=False, sort_keys=True,
                          separators=(',', ':'), allow_nan=False).encode('utf-8'))


def _pairs(items):
    result = {}
    for key, value in items:
        require(key not in result, 'JSON 중복 키는 배포 근거로 사용할 수 없습니다.')
        result[key] = value
    return result


def read_json(path):
    require(path.is_file() and path.stat().st_size <= 32 * 1024**2,
            '필수 검증 JSON 누락/크기 오류: ' + path.name)
    return json.loads(path.read_text(encoding='utf-8-sig'), object_pairs_hook=_pairs,
                      parse_constant=lambda _: (_ for _ in ()).throw(RuntimeError('비유한 JSON 값')))


def safe_file(root, relative):
    require(isinstance(relative, str) and relative and '\\' not in relative and ':' not in relative,
            '상대 파일 경로가 필요합니다.')
    parts = relative.split('/')
    require(not relative.startswith('/') and all(p not in ('', '.', '..') for p in parts),
            '배포 경로 이탈을 차단했습니다.')
    path = root.joinpath(*parts)
    require(path.resolve().is_relative_to(root.resolve()), '배포 파일 경로 이탈')
    for parent in (path, *path.parents):
        if parent.exists() or parent.is_symlink():
            attrs = getattr(parent.lstat(), 'st_file_attributes', 0)
            require(not parent.is_symlink() and not attrs & stat.FILE_ATTRIBUTE_REPARSE_POINT,
                    '공개 입력의 링크·재분석 지점은 허용하지 않습니다.')
        if parent == root:
            break
    return path


def load_helper(root):
    path = safe_file(root, 'tools/github_release_v18.py')
    spec = importlib.util.spec_from_file_location('_js_audit_public_policy_v18', path)
    require(spec is not None and spec.loader is not None, '공개 정책 도구를 읽을 수 없습니다.')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # This helper has an explicit __main__ guard.
    require(module.TAG == 'v1.8.0' and module.SOURCE_NAME == SOURCE_NAME,
            '공개 정책의 버전/첨부 이름이 다릅니다.')
    return module


def redact_text(value):
    # Known workspace roots are replaced first so whitespace in Windows paths
    # cannot leave a username or a partial personal directory behind.
    for base in (ROOT, ROOT.parent):
        for form in (str(base), base.as_posix(), str(base).replace('\\', '\\\\')):
            value = value.replace(form, '[개발PC경로 제외]')
    return ABSOLUTE.sub('[절대경로 제외]', SECRET.sub('[인증값 제외]', value))


def sanitize(value):
    if isinstance(value, dict):
        result = {}
        for key, item in value.items():
            clean_key = redact_text(str(key))
            require(clean_key not in result, '공개 경로 제거 후 JSON 키가 중복됩니다.')
            result[clean_key] = '[인증값 제외]' if SENSITIVE_KEY.fullmatch(str(key)) else sanitize(item)
        return result
    if isinstance(value, list):
        return [sanitize(item) for item in value]
    if isinstance(value, str):
        return redact_text(value)
    return value


def no_leaks(value):
    if isinstance(value, dict):
        return all(no_leaks(str(key)) and no_leaks(item) for key, item in value.items())
    if isinstance(value, list):
        return all(no_leaks(item) for item in value)
    return not isinstance(value, str) or not (ABSOLUTE.search(value) or SECRET.search(value))


def public_json(value):
    clean = sanitize(value)
    require(no_leaks(clean), '공개 JSON의 경로/인증값 제거 검증 실패')
    return json_bytes(clean)


def public_xlsx(data):
    """Redact a synthetic exported reading copy, preserving financial numbers."""
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(data), keep_links=False, data_only=False)
    try:
        require('증적JSON' in book.sheetnames and '조서표지' in book.sheetnames,
                '예시 조서의 전체 증적JSON/표지가 없습니다.')
        proof = book['증적JSON']
        require(proof.cell(1, 2).value == 'JSON조각', '예시 증적JSON 구조 오류')
        full = ''.join(str(proof.cell(row, 2).value or '') for row in range(2, proof.max_row + 1))
        snapshot = json.loads(full, object_pairs_hook=_pairs)
        require(snapshot.get('engagement', {}).get('더미업무') is True,
                '합성 업무가 아닌 조서는 공개 예시로 포함할 수 없습니다.')
        before_numbers = {(sheet.title, cell.coordinate): cell.value
                          for sheet in book if sheet.title != '증적JSON'
                          for row in sheet for cell in row
                          if isinstance(cell.value, (int, float, bool))}
        for sheet in book:
            for row in sheet:
                for cell in row:
                    require(cell.data_type != 'f', '공개 예시에는 실행 수식을 포함하지 않습니다.')
                    if isinstance(cell.value, str):
                        cell.value = redact_text(cell.value)
                        cell.data_type = 's'
                    if cell.comment:
                        cell.comment.text = redact_text(cell.comment.text)
                        cell.comment.author = redact_text(cell.comment.author)
                    if cell.hyperlink and cell.hyperlink.target and not cell.hyperlink.target.startswith('#'):
                        cell.hyperlink = None
        # Redact the complete proof before splitting, including paths that
        # happened to span an original Excel-cell boundary.
        clean = json.dumps(sanitize(snapshot), ensure_ascii=False, sort_keys=True,
                           separators=(',', ':'), allow_nan=False)
        require(no_leaks(json.loads(clean)), '공개 증적JSON에 경로/인증값이 남아 있습니다.')
        style = [copy.copy(proof.cell(2, column)._style) for column in (1, 2)]
        if proof.max_row > 1:
            proof.delete_rows(2, proof.max_row - 1)
        for index, start in enumerate(range(0, len(clean), 28000), 2):
            for column, value in ((1, index - 1), (2, clean[start:start + 28000])):
                cell = proof.cell(index, column, value)
                cell._style = copy.copy(style[column - 1])
                if isinstance(value, str):
                    cell.data_type = 's'
        book['조서표지'].append(['공개 사본 안내', PUBLIC_NOTE])
        notice = book.create_sheet('공개사본안내')
        notice.append(['범위', PUBLIC_NOTE])
        notice.column_dimensions['A'].width = 20
        notice.column_dimensions['B'].width = 100
        for key in ('creator', 'lastModifiedBy', 'title', 'subject', 'description', 'keywords', 'category'):
            value = getattr(book.properties, key, None)
            if isinstance(value, str):
                setattr(book.properties, key, redact_text(value))
        after_numbers = {(sheet.title, cell.coordinate): cell.value
                         for sheet in book if sheet.title not in ('증적JSON', '공개사본안내')
                         for row in sheet for cell in row
                         if isinstance(cell.value, (int, float, bool))}
        require(before_numbers == after_numbers, '공개 예시의 숫자/재무 결과가 변경됐습니다.')
        require(json.loads(''.join(str(proof.cell(row, 2).value or '')
                                  for row in range(2, proof.max_row + 1))) == sanitize(snapshot),
                '공개 조서의 증적JSON 재구성 오류')
        out = BytesIO()
        book.save(out)
    finally:
        book.close()
    result = out.getvalue()
    with zipfile.ZipFile(BytesIO(result)) as archive:
        require(archive.testzip() is None, '공개 조서 ZIP CRC 오류')
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')):
                text = archive.read(name).decode('utf-8')
                require(not re.search(r'(?i)(?<![a-z0-9])[a-z]:[\\/]|file:///|/(?:Users|home)/', text)
                        and not SECRET.search(text),
                        '공개 조서에 절대경로/인증값이 남아 있습니다.')
    return result


def source_payloads(root, helper, build):
    baseline = root.parent / 'v1.7.1_화면선명도_조서안내_2026-10-24/배포/audit-suite-source-v1.7.1.zip'
    require(baseline.is_file(), '기존 1,708개 소스 ZIP이 필요합니다.')
    old = {}
    with zipfile.ZipFile(baseline) as archive:
        require(archive.testzip() is None, '기존 소스 ZIP CRC 오류')
        for info in archive.infolist():
            require(not info.is_dir() and info.filename.startswith(SOURCE_PREFIX), '기존 소스 prefix 오류')
            relative = info.filename[len(SOURCE_PREFIX):]
            helper.allowed_public_path(relative)
            require(relative not in old, '기존 소스 경로 중복')
            old[relative] = archive.read(info)
    require(len(old) == 1708, '기존 공개 소스 경로 1,708개 보존 기준 오류')
    selected = set(old) | NEW_REQUIRED
    manifest = {entry['file']: entry['sha256'] for entry in build['source_manifest']}
    require(len(manifest) == len(build['source_manifest']), '빌드 소스 manifest 중복')
    output = {}
    for relative in sorted(selected):
        helper.allowed_public_path(relative)
        file = safe_file(root, relative)
        require(file.is_file(), '공개 필수 소스 누락: ' + relative)
        payload = file.read_bytes()
        if relative in old and relative not in LEGACY_MUTABLE:
            require(payload == old[relative], '보존 대상의 기존 공개 파일 변경: ' + relative)
        if relative in NEW_REQUIRED or payload != old.get(relative):
            require(manifest.get(relative) == sha(payload), '신규/변경 소스가 빌드 manifest에 고정되지 않았습니다: ' + relative)
        output[relative] = payload
    require(output['ci/audit-suite-v18.yml'] == helper.DEFAULT_WORKFLOW.encode('utf-8'), '공개 CI 설정 불일치')
    return output


def archive_bytes(target, files, prefix, helper, extra_check=None):
    require(not target.exists(), '기존 배포 ZIP은 덮어쓰지 않습니다.')
    normalized = set()
    with zipfile.ZipFile(target, 'x', zipfile.ZIP_DEFLATED, compresslevel=8) as archive:
        for relative, payload in sorted(files.items()):
            helper.safe_source_path(SOURCE_PREFIX + relative)
            key = unicodedata.normalize('NFC', relative).casefold()
            require(key not in normalized, 'ZIP의 정규화 경로 중복')
            normalized.add(key)
            info = zipfile.ZipInfo(prefix + relative, date_time=(2026, 10, 24, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = (stat.S_IFREG | 0o644) << 16
            archive.writestr(info, payload, compresslevel=8)
    with zipfile.ZipFile(target) as archive:
        require(archive.testzip() is None, '배포 ZIP CRC 오류')
        require(set(archive.namelist()) == {prefix + name for name in files}, 'ZIP 전체 파일목록 불일치')
        require(all(sha(archive.read(prefix + name)) == sha(payload) for name, payload in files.items()),
                'ZIP 내부 파일의 SHA256 불일치')
    if extra_check:
        extra_check(target)


def actual_evidence(root, path, frozen):
    report = read_json(path)
    require(report.get('status') == 'passed' and report.get('version') == VERSION
            and report.get('actual_local_inference') is True and report.get('frozen') is frozen,
            '정식 실제 AI 소스/EXE 검증이 필요합니다.')
    require(type(report.get('actual_openai_calls')) is int and report['actual_openai_calls'] == 0,
            '실제 OpenAI 미호출 범위가 다릅니다.')
    require(report.get('model') == MODEL and report.get('model_digest') == MODEL_DIGEST,
            '실제 모델 이름/digest가 배포 계약과 다릅니다.')
    require(report.get('checks') and 'synthetic' in report.get('human_approval_scope', '').casefold(),
            '실제 AI 검증의 합성 승인 범위 기록이 필요합니다.')
    packet = read_json(path.parent / '근거패킷.json')
    response = read_json(path.parent / '실제응답.json')
    require(report['packet_hash'] == response['packet_hash'] == json_digest(packet), '실제 AI 패킷 해시 불일치')
    require(response['provider'] == 'ollama' and response['model'] == MODEL
            and response['fields'] == report['fields'] and response['citations'] == report['citations'],
            '실제 AI 응답과 최종 보고서가 다릅니다.')
    context = {'company_code': packet['engagement']['company_code'], 'procedure_id': packet['procedure']['id'], 'basis': packet['basis']}
    raw = json.loads(response['raw_response'], object_pairs_hook=_pairs)
    require(raw == {'fields': response['fields'], 'citations': response['citations'], 'context': context},
            '실제 AI 원문/적용 필드/근거 문맥 불일치')
    # Pure offline revalidation of saved output, including the current provider's
    # exception-code/citation safety contract. This never invokes generate_draft.
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    from ai_provider import validate_draft_response
    validate_draft_response(packet, response['fields'], response['citations'])
    workbook = path.parent / '실제모델_합성_미검토조서.xlsx'
    require(workbook.is_file(), '실제 AI의 정식 미검토 출력 조서가 필요합니다.')
    return report, packet, response, workbook


def main():
    root = ROOT
    qa, dist = root / '검증', root / '배포'
    helper = load_helper(root)
    build = read_json(qa / '빌드결과.json')
    require(build.get('status') == 'passed' and build.get('version') == VERSION, 'v1.8.0 빌드 검증이 필요합니다.')
    unit = (qa / '단위검사_최종.txt').read_text(encoding='utf-8-sig')
    matches = re.findall(r'(?m)^Ran (\d+) tests? in [0-9.]+s\r?$', unit)
    require(len(matches) == 1 and unit.rstrip().endswith('OK') and 'skipped=' not in unit,
            '최종 단위검사 완료 로그가 필요합니다.')
    unit_tests = int(matches[0])
    require(unit_tests > 0 and build['unit_tests'] == unit_tests, '단위검사 수와 빌드 기록 불일치')
    for entry in build['source_manifest']:
        file = safe_file(root, entry['file'])
        require(file.is_file() and sha(file.read_bytes()) == entry['sha256'], '고정한 빌드 입력이 변경됐습니다: ' + entry['file'])
    reports = {key: safe_file(qa, build[key]) for key in ('source_report', 'frozen_report', 'headless_report', 'actual_ai_source_report', 'actual_ai_frozen_report')}
    source, frozen, headless = (read_json(reports[key]) for key in ('source_report', 'frozen_report', 'headless_report'))
    require(source.get('ok') is True and source.get('frozen') is False and frozen.get('ok') is True
            and frozen.get('frozen') is True and source.get('version') == frozen.get('version') == VERSION,
            '소스/독립 EXE 네이티브 검증이 필요합니다.')
    for key in ('observations', 'sector_observations', 'display_observations', 'native_checks', 'ai_native_checks'):
        require(source[key] == frozen[key], '소스/EXE 검증 불일치: ' + key)
    require('mock' in source.get('ai_transport', '').casefold()
            and 'mock' in frozen.get('ai_transport', '').casefold(), '네이티브 AI 통신의 모의 표시가 필요합니다.')
    require(headless.get('ok') is True and (headless['fixture_cases'], headless['legacy_assertions'], headless['sector_assertions']) == (332, 1949, 1152),
            '기존 더미/화면 없는 통합검증 범위 불일치')
    require((headless['normal_module_executions'], headless['current_ledger_links'],
             headless['cross_company_negative_checks'], headless['parent_cross_company_negative_checks']) == (61, 140, 72, 4),
            '회사별 원장 연결/혼입 검증 범위 불일치')
    sector_path = safe_file(reports['headless_report'].parent, headless['sector_report'])
    sector = read_json(sector_path)
    require(sector.get('status') == 'passed', '업종별 headless 검증이 필요합니다.')
    visual = read_json(qa / 'v18_시각검토.json')
    require(visual.get('status') == 'passed', 'v18 시각검토가 완료되지 않았습니다.')
    actual_source = actual_evidence(root, reports['actual_ai_source_report'], False)
    actual_frozen = actual_evidence(root, reports['actual_ai_frozen_report'], True)
    exe_path = safe_file(dist, EXE_NAME)
    require(exe_path.is_file(), '배포 EXE가 없습니다.')
    exe = exe_path.read_bytes()
    require(len(exe) == build['exe_bytes'] and sha(exe) == build['exe_sha256'], '배포 EXE 크기/해시 불일치')
    require(build.get('actual_local_inference') is True and build.get('actual_openai_calls') == 0,
            '빌드의 실제 AI 범위 불일치')
    sources = source_payloads(root, helper, build)
    windows = {EXE_NAME: exe, 'README.md': sources['README.md']}
    for name, payload in sources.items():
        if name.startswith('sector_resources/'):
            windows['업종별자료/' + name.removeprefix('sector_resources/')] = payload
        elif name.startswith(('문서/', '라이선스/')):
            if name.startswith('라이선스/'):
                windows[name] = payload
            elif name.endswith('.json'):
                windows[name] = public_json(json.loads(payload.decode('utf-8-sig'), object_pairs_hook=_pairs))
            else:
                windows[name] = redact_text(payload.decode('utf-8-sig')).encode('utf-8')
    public_records = {
        '빌드결과.json': build, '소스UI.json': source, 'EXE_UI.json': frozen,
        'headless.json': headless, '업종별_대사검증.json': sector, '시각검토.json': visual,
        '실제AI_소스/report.json': actual_source[0], '실제AI_소스/근거패킷.json': actual_source[1], '실제AI_소스/실제응답.json': actual_source[2],
        '실제AI_EXE/report.json': actual_frozen[0], '실제AI_EXE/근거패킷.json': actual_frozen[1], '실제AI_EXE/실제응답.json': actual_frozen[2],
    }
    for name, value in public_records.items():
        windows['검증기록/' + name] = public_json(value)
    windows['검증기록/단위검사.txt'] = redact_text(unit).encode('utf-8')
    windows['검증기록/공개사본안내.txt'] = PUBLIC_NOTE.encode('utf-8')
    baseline_folder = safe_file(qa, build['source_baseline_report']).parent / 'baseline_v17'
    examples = {code + '_미검토초안.xlsx': baseline_folder / (code + '_미검토초안.xlsx')
                for code in ('JS', 'JS_ELEC', 'JS_AUTO', 'JS_SEMI', 'JS_BIO')}
    examples.update({'모의UI_AI내용승인_초안.xlsx': reports['source_report'].parent / 'AI조서_내용승인_초안.xlsx',
                     '실제AI_소스_합성_미검토조서.xlsx': actual_source[3], '실제AI_EXE_합성_미검토조서.xlsx': actual_frozen[3]})
    example_records = []
    for name, path in examples.items():
        require(path.is_file(), '필수 합성 예시 조서 누락: ' + name)
        original = path.read_bytes()
        public = public_xlsx(original)
        windows['예시조서/' + name] = public
        example_records.append({'file': name, 'original_sha256': sha(original), 'public_sha256': sha(public),
                                'financial_numbers_unchanged': True, 'public_reading_copy': True})
    names = (WINDOWS_NAME, SOURCE_NAME, VALIDATION_NAME, 'SHA256SUMS.txt')
    dist.mkdir(exist_ok=True)
    require(not any((dist / name).exists() for name in names), '기존 v1.8.0 첨부는 덮어쓰지 않습니다.')
    require(not (qa / '최종배포검증.json').exists(), '기존 최종 배포 기록을 보존하세요.')
    stage = safe_file(dist, 'package-v18-' + uuid.uuid4().hex)
    stage.mkdir()
    source_zip = stage / SOURCE_NAME
    def check_source_zip(path):
        actual = helper.publication_files(path)
        require(actual == sources, '공개 도구가 읽은 소스 ZIP 원문 불일치')
    archive_bytes(source_zip, sources, SOURCE_PREFIX, helper, check_source_zip)
    validation = {
        'version': VERSION, 'status': 'passed', 'unit_tests': unit_tests,
        'legacy_fixture_cases': 332, 'legacy_headless_assertions': 1949, 'sector_headless_assertions': 1152,
        'normal_module_executions': 61, 'current_ledger_links': 140, 'cross_company_negative_checks': 72,
        'parent_cross_company_negative_checks': 4, 'readable_raw_workbooks': headless['extracted_workbooks'],
        'native_assertions': len(frozen['native_checks']), 'ai_native_assertions': len(frozen['ai_native_checks']),
        'display_assertions': len(frozen['display_checks']), 'baseline_native_assertions': len(frozen['baseline_native_checks']),
        'ai_transport': frozen['ai_transport'], 'native_ai_transport_is_mock': True,
        'actual_local_inference': True, 'uses_ai_inference': True, 'actual_openai_calls': 0,
        'actual_inference': {'provider': 'ollama', 'model': MODEL, 'model_digest': MODEL_DIGEST,
                             'source_checks': len(actual_source[0]['checks']), 'frozen_checks': len(actual_frozen[0]['checks']),
                             'human_approval_scope': 'synthetic state-transition test, not actual audit approval'},
        'source_frozen_observations_identical': True, 'embedded_manifest_dpi_verified': build['embedded_manifest_dpi_verified'],
        'dpi_policy': 'system', 'display_environment': sanitize(frozen['display_environment']), 'app_zoom_verified': [100, 125, 150],
        'resource_manifest_sha256': frozen['resource_manifest_sha256'],
        'exe': {'name': EXE_NAME, 'bytes': len(exe), 'sha256': sha(exe)},
        'source': {'name': SOURCE_NAME, 'files': len(sources), 'sha256': sha(source_zip.read_bytes()), 'previous_public_paths_preserved': 1708},
        'company_names': ['JS', 'JS전자', 'JS자동차', 'JS반도체', 'JS바이오'],
        'balanced_omission_control_failures': headless['balanced_omission']['independent_control_failures'],
        'core_calculation_workpaper_links': ['A02', 'A04', 'A12'], 'manual_industry_topics': 20, 'document_templates': 60,
        'unreviewed_final_export_blocked': True, 'synthetic_examples': example_records,
        'portable_ai_runtime_models_profiles_excluded': True, 'screenshots_not_bundled': True,
        'public_reports_and_xlsx_are_redacted_reading_copies': True,
        'limitations': ['Same synthetic generator; not independent audit evidence.',
                       'Parent JS has no industry; cash/investments/capital only, no investment valuation or consolidation.',
                       'AI quotes limited evidence excerpts; reference structure does not prove every statement or accounting judgment.',
                       'Human content approval is separate from independent workpaper review; synthetic tests are not real audit approvals.',
                       'Native UI AI transport is mocked. Saved actual local source/EXE inference is checked separately. OpenAI real calls untested.',
                       'Models, Ollama runtime/cache/private profile/API keys and actual customer workspaces are not distributed.',
                       'Public report/workbook path redaction changes reading-copy bytes; original local evidence remains preserved.',
                       'Windows development PC/EXE-only folder; no clean OS or other-PC claim. System DPI monitor changes may require restart.'],
        'github_verification': 'See the v1.8.0 release and exact-commit Python 3.12/3.14 Actions evidence; packaging alone is not publication.',
    }
    validation_data = public_json(validation)
    windows[VALIDATION_NAME] = validation_data
    file_manifest = [{'file': name, 'bytes': len(payload), 'sha256': sha(payload)} for name, payload in sorted(windows.items())]
    windows['파일목록_SHA256.json'] = json_bytes(file_manifest)
    archive_bytes(stage / WINDOWS_NAME, windows, WINDOWS_PREFIX, helper)
    (stage / VALIDATION_NAME).write_bytes(validation_data)
    assets = [stage / name for name in (WINDOWS_NAME, SOURCE_NAME, VALIDATION_NAME)]
    (stage / 'SHA256SUMS.txt').write_text(''.join(sha(path.read_bytes()) + '  ' + path.name + '\n' for path in assets), encoding='utf-8')
    assets.append(stage / 'SHA256SUMS.txt')
    # Recheck original build inputs after workbook processing and ZIP creation.
    require(all(sha(safe_file(root, row['file']).read_bytes()) == row['sha256'] for row in build['source_manifest']),
            '패키징 도중 고정 소스가 변경됐습니다.')
    require(sha(exe_path.read_bytes()) == build['exe_sha256'], '패키징 도중 EXE가 변경됐습니다.')
    require(not any((dist / name).exists() for name in names), '완료 직전 기존 첨부 충돌')
    result = {'status': 'passed', 'version': VERSION, 'unit_tests': unit_tests, 'source_files': len(sources), 'windows_files': len(windows),
              'previous_public_paths_preserved': 1708, 'new_source_allowlist': sorted(NEW_REQUIRED), 'source_manifest_unchanged': True,
              'zip_crc_and_all_file_hashes_verified': True, 'runtime_models_profiles_excluded': True,
              'public_reading_copies_sanitized': True, 'actual_local_inference': True, 'actual_openai_calls': 0,
              'native_ai_transport': 'mock; separate actual source/EXE inference required',
              'assets': [{'name': path.name, 'bytes': path.stat().st_size, 'sha256': sha(path.read_bytes())} for path in assets]}
    for path in assets:
        os.replace(path, dist / path.name)
    stage.rmdir()  # Empty, task-created directory only; never recursive deletion.
    target = qa / '최종배포검증.json'
    with target.open('xb') as file:
        file.write(json_bytes(result))
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(str(exc) if isinstance(exc, RuntimeError) else '배포 입력/실행 오류. 원본과 기존 첨부는 보존했습니다.', file=sys.stderr)
        raise SystemExit(1)
