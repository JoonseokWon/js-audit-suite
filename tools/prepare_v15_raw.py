"""Split approved synthetic analyzer inputs without creating new evidence.

Existing workbooks are copied with cell values, types, row positions and styles
preserved. This is a reproducible source-layout exercise, not an ERP/bank fetch.
"""
from __future__ import annotations
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import ledger
import gla
import fia
import cash_analytics

OUT = ROOT / 'raw_analytics'
ANALYSIS_INPUTS = {
    'GLA': ROOT / 'samples/통합/통합_v1.5_GLA.xlsx',
    'FIA': ROOT / 'samples/통합/통합_v1.5_FIA.xlsx',
    '현금분석': ROOT / 'samples/통합/통합_v1.5_현금분석.xlsx',
}
GL_INPUT = ANALYSIS_INPUTS['GLA']
BASE_REFERENCE = ROOT / 'samples/통합/통합_v1.3_원장_리드.xlsx'
SPECS = [
    ('01_GL_TB_합성원천.xlsx', GL_INPUT, ledger.SCHEMA, ['GLA', 'FIA', '현금분석']),
    ('02_ERP전표로그_합성.xlsx', ANALYSIS_INPUTS['FIA'], fia.SCHEMA, ['FIA']),
    ('03_은행현금_합성원천.xlsx', ANALYSIS_INPUTS['현금분석'], cash_analytics.SCHEMA, ['현금분석']),
    ('04_GLA_정책.xlsx', ANALYSIS_INPUTS['GLA'], gla.SCHEMA, ['GLA']),
]


def encode(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False).encode('utf-8')


def checksum(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def typed(value):
    if value is None:
        return {'type': 'null', 'value': None}
    if isinstance(value, datetime):
        return {'type': 'datetime', 'value': value.isoformat()}
    if isinstance(value, date):
        return {'type': 'date', 'value': value.isoformat()}
    if type(value) not in (str, int, float, bool):
        raise ValueError('지원하지 않는 셀 값 형식')
    return {'type': type(value).__name__, 'value': value}


def snapshot(path, expected=None):
    book = load_workbook(path, data_only=False, keep_links=False)
    out = {}
    try:
        if expected is not None and set(book.sheetnames) != set(expected):
            raise ValueError('선택 시트 목록이 예상과 다릅니다: ' + Path(path).name)
        for sheet in book:
            headers = [c.value for c in sheet[1]]
            while headers and headers[-1] is None:
                headers.pop()
            if expected is not None and headers != list(expected[sheet.title]):
                raise ValueError('원본 헤더가 스키마와 다릅니다: ' + sheet.title)
            rows = []
            for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=len(headers)):
                values = []
                for cell in row:
                    if cell.data_type in ('f', 'e'):
                        raise ValueError('확정값 원천에 수식/오류 셀이 있습니다: ' + sheet.title + '!' + cell.coordinate)
                    values.append({'excel_type': cell.data_type, **typed(cell.value)})
                rows.append(values)
            out[sheet.title] = {'headers': headers, 'rows': rows}
    finally:
        book.close()
    return out


def summaries(snap):
    result = []
    for name, data in snap.items():
        rows = data['rows']
        result.append({'sheet': name, 'headers': data['headers'], 'columns': len(data['headers']),
                       'physical_rows_including_header': len(rows), 'data_rows': len(rows)-1,
                       'header_sha256': hashlib.sha256(encode(rows[0])).hexdigest(),
                       'row_values_sha256': hashlib.sha256(encode(rows[1:])).hexdigest(),
                       'row_hashes': [{'excel_row': i+2, 'sha256': hashlib.sha256(encode(row)).hexdigest()}
                                      for i, row in enumerate(rows[1:])]})
    return result


def extract(source, target, schema):
    book = load_workbook(source, data_only=False, keep_links=False)
    try:
        for sheet in list(book):
            if sheet.title not in schema:
                book.remove(sheet)
        if book.sheetnames != list(schema):
            raise ValueError('원본 시트 순서가 예상과 다릅니다.')
        book.save(target)
    finally:
        book.close()


PBC_HEADERS = ['자료ID', '자료·범위', '필수키', '날짜·시간', '통제', '제출형태', '연결 분석', '유의사항']
PBC_ROWS = [
    ['GL01', '회사·회계기간·통화·추출범위 설정', '항목 고유, 회사코드', '시작일·기준일 YYYY-MM-DD', 'KRW 및 대상계정 명시', '01_GL_TB_합성원천 / 설정', 'GLA·FIA·현금', '조건 누락·회사 혼입을 0 또는 전체 범위로 해석하지 않음'],
    ['GL02', '전체 TB. 기초·당기 차변·대변·기말', '회사코드+계정코드', '설정의 동일 회계기간', '기초+차변-대변=기말. 차대 합계', '01_GL_TB_합성원천 / 시산표', 'GLA·FIA·현금', '전기 이월·미제출 계정·폐쇄 여부와 잔액 부호를 확인'],
    ['GL03', '선언 계정의 전체 전표행. 필터 포함', '회사코드+전표ID+행ID', '전기일 YYYY-MM-DD', '전표 차대 및 계정별 행수·총차변·총대변', '01_GL_TB_합성원천 / 총계정원장', 'GLA·FIA·현금', '차대 상계 후 0만 확인하면 균형 전표 통째 누락을 놓칠 수 있음'],
    ['GL04', '계정코드와 리드·BS/PL 매핑', '계정코드', '현재 회계기간에 적용한 버전', '중복·미매핑·리드 분류 충돌', '01_GL_TB_합성원천 / 계정매핑', 'GLA·FIA·현금', '계정명·번호만 보고 수동/자동·현금 성격을 추정하지 않음'],
    ['GL05', 'ERP 별도 추출 통제', '회사코드+계정코드', '기간·시간대 포함 추출시각', '행수·총차변·총대변·기초·기말·근거', '01_GL_TB_합성원천 / 독립통제', 'GLA·FIA·현금', '본 더미 통제도 동일 생성기 자료. 실제 독립 통제로 주장하지 않음'],
    ['GL06', '감사 조정 전체. 없으면 헤더만 제출', '조정ID+행ID', '업무 기준일 및 승인상태', '조정별 차대·상태 일관성', '01_GL_TB_합성원천 / 감사조정', 'GLA·FIA·현금', '초안과 승인을 구분. 분석은 현재 업무 조정전 GL 기준'],
    ['GL07', '계정과 계산 모듈 지표 연결', '모듈+지표+계정코드', '동일 회계기간', '비교기준·표시 부호·중복', '01_GL_TB_합성원천 / 모듈연결', 'GLA·FIA·현금', '연결 누락을 계산값 0으로 대체하지 않음'],
    ['FI01', '마감·메타 추출·근무시간 등 정책', '항목 고유·규칙버전·정책명', '시간대 포함 ISO8601 시각', '마감·추출시각 및 정책 설정 완전성', '02_ERP전표로그_합성 / FIA설정', 'FIA', '마감은 회계기준일과 다름. 휴일·근무시간을 임의 추정하지 않음'],
    ['FI02', '전체 GL행의 ERP 메타·승인 로그', '회사코드+전표ID+행ID', '입력·승인시각 ISO8601+offset', '전표 전체 행의 일관성·중복·고아·시각', '02_ERP전표로그_합성 / 전표메타', 'FIA', '작성자·수동자동·승인자·부서·시스템은 실제 로그로 별도 확인'],
    ['CA01', '선택 현금계정과 검토 규칙 설정', '항목 고유', '반복·변경 검색일수', '대상계정·집중 최소건수·비율', '03_은행현금_합성원천 / 분석설정', '현금분석', '선택 범위가 회사 전체 계좌목록이라는 뜻이 아님'],
    ['CA02', '은행계좌와 GL계정·기초/기말', '회사코드+계좌ID', '설정과 동일 기간', 'KRW·계좌목록·기초/기말·TB 대사', '03_은행현금_합성원천 / 현금계좌', '현금분석', '완전 미제출 계좌는 독립 계좌목록·금융기관조회로 추가 확인'],
    ['CA03', '계좌별 전체 은행 입출금', '회사코드+은행거래ID', '거래일 YYYY-MM-DD', '입금·출금 별도 정수원·중복·기간', '03_은행현금_합성원천 / 은행거래', '현금분석', '개인 실명·실제 계좌번호 없음. 모든 상대방·계좌는 합성'],
    ['CA04', '은행 별도 추출 행수·입출금', '회사코드+계좌ID', '기간·시간대 포함 추출시각', '계좌별 행수·입출금 합계. 무거래 0', '03_은행현금_합성원천 / 은행독립통제', '현금분석', '은행거래 자체를 더한 값만으로 실제 추출 완전성을 입증하지 못함'],
    ['CA05', '선택 계정의 GL 전체 행과 지급 메타', '회사코드+전표ID+행ID', '현재 GL 전기일', '한 번씩 전체 선택·계좌/계정 일치', '03_은행현금_합성원천 / 현금원장선택', '현금분석', '거래처·상대계좌·근거·유형이 은행 쪽 명시값과 모순되면 차단'],
    ['CA06', '은행과 원장행의 명시적 배부', '배부ID·은행ID·회사+전표+행', '동일 업무 기간의 양쪽 원문', '중복쌍·부호·과배부·각각 잔여', '03_은행현금_합성원천 / 현금배부', '현금분석', '같은 금액이라는 이유로 임의 매칭하지 않음. 다대다·부분 배부 명시'],
    ['CA07', '거래처 계좌와 소유 확인자료', '회사코드+거래처ID+상대계좌', '확인일·효력 범위는 별도 확보', '외부/자사/미확인 구분·근거', '03_은행현금_합성원천 / 거래처계좌', '현금분석', '동일 계좌·공동 수납·집중 자체는 부정 확정이 아님'],
    ['CA08', '거래처 계좌 변경·승인 이력', '변경ID·회사+거래처', '효력일·승인일 YYYY-MM-DD', '이전/신규 계좌·중복·근거·범위', '03_은행현금_합성원천 / 계좌변경', '현금분석', '변경 이력 미제출은 변경 없음이 아님. 전체성도 별도 확인'],
    ['CA09', '승인 지급·송장/배치·정기/분할 근거', '회사코드+지급근거ID+거래처', '승인 적용 기간은 별도 확보', 'KRW·승인총액·일회/분할/정기/이체', '03_은행현금_합성원천 / 지급근거', '현금분석', '분할·정기 선언의 진위와 승인총액의 실질을 사람이 검토'],
    ['GA01', '주말·월말·고액·반복 검토 정책', '항목 고유·규칙버전', '월말 범위·기준 회계기간', '사용/미사용·임계금액·최소횟수', '04_GLA_정책 / GLA설정', 'GLA', '오프라인 규칙 기반. 외부 AI 전송·AI 추론·부정 확정 없음'],
]


def write_pbc(path):
    book = Workbook(); sheet = book.active; sheet.title = '자료요청'
    for values in [PBC_HEADERS] + PBC_ROWS:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Preserve literal =,+,-,@ text without evaluating.
            cell.font = Font(name='맑은 고딕', size=10)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='263F59')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    widths = [12, 35, 33, 32, 42, 41, 20, 67]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.row_dimensions[1].height = 28
    for i in range(2, len(PBC_ROWS)+2):
        sheet.row_dimensions[i].height = 46
    sheet.freeze_panes = 'C2'; sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    book.save(path); book.close()


def prepare():
    OUT.mkdir(parents=True, exist_ok=True)
    mother_paths = [BASE_REFERENCE] + list(ANALYSIS_INPUTS.values())
    mother_hashes = {str(p.relative_to(ROOT)).replace('\\', '/'): checksum(p) for p in mother_paths}
    targets = {name: snapshot(path, {**ledger.SCHEMA, **module.SCHEMA}) for name, path, module in [
        ('GLA', ANALYSIS_INPUTS['GLA'], gla), ('FIA', ANALYSIS_INPUTS['FIA'], fia),
        ('현금분석', ANALYSIS_INPUTS['현금분석'], cash_analytics)]}
    gl_snap = {name: targets['GLA'][name] for name in ledger.SCHEMA}
    reference = snapshot(BASE_REFERENCE, ledger.SCHEMA)
    blank_row_differences = []
    for sheet in ledger.SCHEMA:
        def nonblank_rows(table):
            return [row for row in table['rows'] if any(cell['type'] != 'null' for cell in row)]
        if reference[sheet]['headers'] != gl_snap[sheet]['headers'] or nonblank_rows(reference[sheet]) != nonblank_rows(gl_snap[sheet]):
            raise ValueError('구 기준원장과 정상 분석 입력의 실제 값이 다릅니다: ' + sheet)
        if reference[sheet] != gl_snap[sheet]:
            blank_row_differences.append({'sheet': sheet,
                'baseline_blank_physical_rows': [i+1 for i, row in enumerate(reference[sheet]['rows']) if all(c['type'] == 'null' for c in row)],
                'analysis_blank_physical_rows': [i+1 for i, row in enumerate(gl_snap[sheet]['rows']) if all(c['type'] == 'null' for c in row)],
                'policy': '기존 원본은 변경하지 않음. 분해본은 정상 GLA 결합입력의 공통 7시트를 그대로 보존하여 분석 입력과 정확 재구성.'})
    for name, snap in targets.items():
        for sheet in ledger.SCHEMA:
            if snap[sheet] != gl_snap[sheet]:
                raise ValueError('기준 원장과 분석 입력 공통행/타입/순서가 다릅니다: ' + name + '/' + sheet)
    files, raw_snaps, checks = [], {}, []
    for filename, mother, schema, consumers in SPECS:
        source_all = snapshot(mother)
        expected = {name: source_all[name] for name in schema}
        extract(mother, OUT/filename, schema)
        actual = snapshot(OUT/filename, schema)
        if actual != expected:
            raise ValueError('분해본 값/행/타입/순서 보존 실패: ' + filename)
        raw_snaps[filename] = actual
        sheets = summaries(actual)
        files.append({'file': filename, 'sha256': checksum(OUT/filename), 'bytes': (OUT/filename).stat().st_size,
                      'mother_input': str(mother.relative_to(ROOT)).replace('\\', '/'), 'mother_sha256': checksum(mother),
                      'analysis_targets': [{'module': name, 'file': str(ANALYSIS_INPUTS[name].relative_to(ROOT)).replace('\\', '/'),
                                            'sha256': checksum(ANALYSIS_INPUTS[name])} for name in consumers], 'sheets': sheets})
        checks.append({'file': filename, 'sheets': len(sheets), 'data_rows': sum(s['data_rows'] for s in sheets),
                       'exact_values_types_row_order': True, 'missing_rows': 0, 'changed_cells': 0})
    for name, extra_filename in [('GLA', '04_GLA_정책.xlsx'), ('FIA', '02_ERP전표로그_합성.xlsx'), ('현금분석', '03_은행현금_합성원천.xlsx')]:
        combined = {**raw_snaps['01_GL_TB_합성원천.xlsx'], **raw_snaps[extra_filename]}
        if combined != targets[name]:
            raise ValueError('원천 조합과 분석용 XLSX 전체행 재구성 대조 실패: ' + name)
        checks.append({'module': name, 'reconstructed_sheets': len(combined), 'all_values_types_row_order_exact': True,
                       'missing_rows': 0, 'changed_cells': 0, 'analysis_input_sha256': checksum(ANALYSIS_INPUTS[name])})
    pbc = OUT/'00_분석자료요청목록.xlsx'; write_pbc(pbc)
    pbc_snap = snapshot(pbc, {'자료요청': PBC_HEADERS})
    if len(pbc_snap['자료요청']['rows']) != len(PBC_ROWS)+1:
        raise ValueError('PBC 행수 오류')
    for rel, expected in mother_hashes.items():
        if checksum(ROOT/rel) != expected:
            raise ValueError('원본 분석 입력이 변경되었습니다.')
    manifest = {'schema_version': 'raw-analytics-1', 'context': {'회사코드': 'SYNTH01', '시작일': '2026-01-01', '기준일': '2026-12-31', '통화': 'KRW'},
                'kind': 'synthetic-input-decomposition', 'independent_evidence': False,
                'source_statement': '기존 정상 분석용 합성 XLSX를 분해한 재배치 자료. ERP·은행에서 새로 수집한 원천 또는 독립 감사증거가 아님.',
                'baseline_reference': {'file': str(BASE_REFERENCE.relative_to(ROOT)).replace('\\', '/'),
                    'sha256': checksum(BASE_REFERENCE), 'nonblank_cells_types_row_order_match': True,
                    'format_only_blank_row_differences': blank_row_differences},
                'hash_policy': 'SHA256. 각 행의 Excel 셀형식·Python 값타입·값과 물리행순서를 포함한 UTF-8 canonical JSON. 헤더 별도, 데이터행 Excel 2행부터.',
                'files': files,
                'pbc': {'file': pbc.name, 'sha256': checksum(pbc), 'bytes': pbc.stat().st_size, 'data_rows': len(PBC_ROWS)},
                'analysis_inputs': [{'module': name, 'file': str(path.relative_to(ROOT)).replace('\\', '/'), 'sha256': checksum(path)}
                                    for name, path in ANALYSIS_INPUTS.items()],
                'verification': {'source_workbooks': 4, 'source_sheets': sum(len(s['sheets']) for s in files),
                                 'data_rows': sum(s['data_rows'] for f in files for s in f['sheets']),
                                 'reconstructed_analysis_workbooks': 3, 'missing_rows': 0, 'changed_cells': 0,
                                 'original_analysis_inputs_unchanged': True, 'checks': checks}}
    (OUT/'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    (OUT/'verification.json').write_text(json.dumps(manifest['verification'], ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    return manifest['verification']


if __name__ == '__main__':
    print(json.dumps(prepare(), ensure_ascii=False, indent=2))
