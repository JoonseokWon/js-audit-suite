"""Provider-based five-company regression, independent of Tk and frozen UI.

Writes fresh isolated sessions below the explicit output directory. Existing
source releases, resources and previously completed sessions are never edited.
"""
from pathlib import Path
import argparse
import hashlib
import json
import sys
import uuid
from copy import deepcopy

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from common import AuditError
from workspace import Workspace
import adapters
import exporter

COMPANIES = {'JS': 'JS', 'JS_ELEC': 'JS전자', 'JS_AUTO': 'JS자동차', 'JS_SEMI': 'JS반도체', 'JS_BIO': 'JS바이오'}
MODULES = {'원장·리드', '금융상품', '매출', '감가상각', '확장검사', '리스', '대손충당금', '재고평가', '종업원급여',
           '미기록부채·비용이연', '원천증빙', '재무제표·주석', 'GLA', 'FIA', '현금분석'}
SUBSIDIARIES = tuple(code for code in COMPANIES if code != 'JS')
PARENT_PROCEDURES = {'C02', 'P01', 'C04', 'C05', 'P07'}
COMPANY_BOUND_MODULES = ('원장·리드', '원천증빙', '재무제표·주석', 'GLA', 'FIA', '현금분석')
# Existing pack preserves these separate manual consolidation examples. Their
# SYNTH_SUB rows are not part of the selected company's automated GL/TB scope.
MANUAL_SUBSIDIARY_EXAMPLES = {'D001', 'D031', 'D032', 'D038', 'D039'}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def fingerprint(folder):
    """Relative regular-file bytes, excluding links; no reads outside folder."""
    root = Path(folder).resolve()
    return {p.relative_to(root).as_posix(): sha(p) for p in sorted(root.rglob('*'))
            if p.is_file() and not p.is_symlink() and p.resolve().is_relative_to(root)}


def original_files():
    """Optional local pre/post protection, portable CI reports absent originals."""
    paths = set()
    for version in ('v1.5_GLA_FIA_현금분석_2026-10-23', 'v1.6_업종선택통합_2026-10-23'):
        old = ROOT.parent/version
        build = old/'검증/빌드결과.json'
        if build.is_file():
            paths.add(build)
            data = json.loads(build.read_text(encoding='utf-8'))
            for row in data.get('source_manifest', []):
                candidate = old/row['file']
                if candidate.is_file():
                    paths.add(candidate)
            paths.update((old/'배포').glob('*.exe'))
    js = ROOT.parent/'JS_업종별_더미업무_2026-10-23'
    if js.is_dir():
        paths.update(p for p in js.rglob('*') if p.is_file() and not p.is_symlink()
                     and p.resolve().is_relative_to(js.resolve()))
    return {str(p.resolve()): sha(p) for p in paths}


def validate_workspace(ws, code, checks):
    def check(condition, label):
        if not condition:
            raise AssertionError(code + ': ' + label)
        checks.append(code + ':' + label)
    ws.refresh()
    check(ws.engagement['회사코드'] == code, 'engagement_company')
    check(ws.engagement['회사명'] == COMPANIES[code], 'engagement_name')
    check(ws.engagement.get('더미업무') is True, 'synthetic_engagement')
    check(not ws.verify_sources(), 'immutable_source_copies')
    latest = ws.latest()
    expected_modules = {'원장·리드'} if code == 'JS' else MODULES
    check(set(latest) == expected_modules, 'parent_one_module' if code == 'JS' else 'subsidiary_fifteen_modules')
    for module, run in latest.items():
        check(run['result']['blocking'] == 0, module + ':no_blocking')
        check(run['result']['context'].get('회사코드') == code, module + ':result_company')
        check(run['source_id'] in ws.sources, module + ':registered_input')
        source = ws.sources[run['source_id']]
        check(sha(ws.root/source['file']) == source['sha256'], module + ':source_sha')
    active = latest['원장·리드']
    for module in (() if code == 'JS' else ('원천증빙', '재무제표·주석', 'GLA', 'FIA', '현금분석')):
        detail = latest[module]['result']['detail']
        check(detail.get('ledger_run_id') == active['id'], module + ':pinned_current_ledger')
        check(detail.get('ledger_result_hash') == active['result_hash'], module + ':current_ledger_hash')
    for name in ('입력_총계정원장', '입력_시산표'):
        check(all(r['회사코드'] == code for r in active['result']['tables'][name]), name + ':row_company')
    subsidiary_examples = []
    if code != 'JS':
        # Inspect the pinned source payload as well as the adapter-assigned context.
        from raw_sources import load_pack
        raw_source = ws.sources[latest['원천증빙']['source_id']]
        pack = load_pack(ws.root/raw_source['file'], {**ws.context, '통화': 'KRW', '연결구분': '별도'})
        check(pack['catalog']['context']['회사코드'] == code and pack['links']['context']['회사코드'] == code,
              'raw_catalog_and_allocation_company')
        check(len(pack['catalog']['tables']) == 114, 'raw_source_tables_preserved')
        for table in pack['catalog']['tables']:
            if '회사코드' in table['headers']:
                index = table['headers'].index('회사코드')
                allowed = {code, 'SYNTH_SUB'} if table['dataset_id'] in MANUAL_SUBSIDIARY_EXAMPLES else {code}
                check(all(r['values'][index] in allowed for r in table['rows']),
                      'raw_row_company:' + table['dataset_id'])
                foreign = [r for r in table['rows'] if r['values'][index] != code]
                if foreign:
                    subsidiary_examples.append({'dataset_id': table['dataset_id'], 'company_code': 'SYNTH_SUB',
                                                'rows': len(foreign), 'scope': 'preserved manual example, excluded from selected company GL/TB'})
            if table['headers'] == ['항목', '값']:
                check(all(r['values'][1] == code for r in table['rows'] if r['values'][0] == '회사코드'),
                      'raw_settings_company:' + table['dataset_id'])
    cross = ws.cross_checks()
    check(len(cross) == (0 if code == 'JS' else 35), 'parent_no_calculation_claim' if code == 'JS' else 'thirty_five_reconciliations')
    if code == 'JS':
        check(set(ws.procedures) == PARENT_PROCEDURES, 'parent_five_scoped_drafts')
        leads = active['result']['tables']['리드스케줄']
        check(sum(r['조정전'] for r in leads) == 0, 'parent_signed_tb_zero')
        check(sum(r['조정전'] for r in leads if r['조정전'] < 0) == -10000000, 'parent_capital_ten_million')
        check(sum(r['조정전'] for r in leads if '현금' in r['계정명']) == 2000000, 'parent_cash_two_million')
        investments = [r for r in leads if '투자' in r['계정명']]
        check(len(investments) == 4 and all(r['조정전'] == 2000000 for r in investments), 'parent_four_two_million_investments')
    check(all(r['상태'] == '일치' for r in cross), 'reconciliations_equal')
    check(bool(ws.procedures), 'workpaper_drafts_exist')
    check(all(not p.get('performed') and not p.get('conclusion') and not p.get('reviewer') for p in ws.procedures.values()), 'human_fields_blank')
    check(not ws.requests, 'no_request_automatically_sent_or_saved')
    for p in ws.procedures.values():
        check(p['title'].startswith(COMPANIES[code] + ' '), 'workpaper_company_title:' + p['id'])
        check(all(sid in ws.sources for sid in p.get('evidence_ids', [])), 'workpaper_registered_evidence:' + p['id'])
        refs = ws.calculation_refs(p['id'])
        if refs:
            check({refs['자료ID'], active['source_id']} <= set(p.get('evidence_ids', [])),
                  'workpaper_own_calculation_and_ledger:' + p['id'])
    metadata = [s for s in ws.sources.values() if s['role'] in ('합성 업종설정', '합성 회사설정', '미수행 검토계획')]
    check(len(metadata) == 3, 'three_company_metadata_sources')
    for source in metadata:
        data = json.loads((ws.root/source['file']).read_text(encoding='utf-8'))
        check(data['company_code'] == code, 'metadata_company:' + source['role'])
        if code == 'JS' and data.get('kind') == 'parent' and 'operating_data_provided' in data:
            check(data.get('industry') is None and data['operating_data_provided'] is False,
                  'parent_industry_and_operating_data_not_invented')
            check(data.get('ownership_percentage') is None, 'parent_ownership_not_assumed')
        if source['role'] == '미수행 검토계획':
            check(len(data['audit_topics']) == (0 if code == 'JS' else 5), 'conditional_industry_topic_count')
            for topic in data['audit_topics']:
                for pid in [topic['workpaper_id']] + topic.get('additional_workpaper_ids', []):
                    check(topic['title'] in ws.procedures[pid]['scope'], 'manual_topic_link:' + topic['id'] + ':' + pid)
    try:
        ws.snapshot(True)
    except AuditError:
        check(True, 'unreviewed_final_export_blocked')
    else:
        raise AssertionError('미검토 조서의 최종본이 허용되었습니다: ' + code)
    return {'company_code': code, 'company_name': ws.engagement['회사명'], 'modules': len(latest), 'current_ledger_links': len(cross),
            'procedures': len(ws.procedures), 'review_candidates': len(ws.risk_register()),
            'human_fields_blank': True, 'unreviewed_final_export_blocked': True,
            'preserved_manual_subsidiary_examples': subsidiary_examples}


def reject_foreign(module, path, context, current):
    try:
        result = adapters.run(module, path, context, ledger_result=current)
    except AuditError:
        return 'input_rejected'
    if not result['blocking']:
        raise AssertionError('다른 회사 자료가 통과했습니다: ' + context['회사코드'] + '/' + module)
    return 'blocking_result'


def validate_data_folder(ws, code, checks):
    """Readable copies match immutable packaged bytes; opening adds no evidence/event."""
    import sector_demo
    before_head, before_basis, before_sources = ws.head, ws.basis(), set(ws.sources)
    folder = sector_demo.data_folder(ws)
    if not isinstance(folder, Path) or folder != ws.root.parent:
        raise AssertionError('자료 보기 위치는 현재 업무 세션이어야 합니다.')
    inputs = folder/'inputs'
    expected_count = 1 if code == 'JS' else 15
    if not inputs.is_dir() or len([p for p in inputs.iterdir() if p.is_file()]) != expected_count:
        raise AssertionError('현재 회사 입력 사본 개수 오류: ' + code)
    extracted = folder/'원천엑셀'
    if code == 'JS':
        if extracted.exists():
            raise AssertionError('업종 없는 모회사에 자회사 원천팩이 복제되었습니다.')
        books = 0
    else:
        from raw_sources import load_pack
        source = ws.sources[ws.latest()['원천증빙']['source_id']]
        pack = load_pack(ws.root/source['file'], {**ws.context, '통화': 'KRW', '연결구분': '별도'})
        expected = {r['path'][len('sources/'):]: r['sha256'] for r in pack['manifest']['files']
                    if r['path'].startswith('sources/') and r['path'].endswith('.xlsx')}
        actual = fingerprint(extracted)
        if len(expected) != 16 or actual != expected:
            raise AssertionError('원천엑셀 16권의 목록/원문 SHA256이 다릅니다: ' + code)
        books = len(actual)
    ws.refresh()
    if ws.head != before_head or ws.basis() != before_basis or set(ws.sources) != before_sources:
        raise AssertionError('자료 열람이 증거/이력을 새로 등록했습니다.')
    checks.extend([code + ':current_session_data_folder', code + ':data_view_preserves_basis'])
    return {'company_code': code, 'input_files': expected_count, 'extracted_workbooks': books,
            'immutable_pack_bytes_equal': True, 'evidence_and_event_head_unchanged': True}


def balanced_omission(resource_root):
    """Removing an entire balanced JE must still fail unchanged gross controls/TB."""
    from common import read_xlsx
    import ledger
    root = Path(resource_root)
    manifest = json.loads((root/'manifest.json').read_text(encoding='utf-8'))
    entry = next(e for e in manifest['companies'] if e['code'] == 'JS_ELEC')
    path = root/next(r['path'] for r in entry['inputs'] if r['module'] == '원장·리드')
    tables, _ = read_xlsx(path, ledger.SCHEMA)
    normal = ledger.analyze(tables)
    if normal['blocking']:
        raise AssertionError('균형 누락 반례의 모본 원장이 차단 상태입니다.')
    changed = deepcopy(tables)
    removed = [r for r in changed['총계정원장'] if r['전표ID'] == 'J-PAY']
    changed['총계정원장'] = [r for r in changed['총계정원장'] if r['전표ID'] != 'J-PAY']
    if not removed or sum(r['차변'] - r['대변'] for r in removed) != 0:
        raise AssertionError('실제 차대가 일치하는 전표 전체를 제거해야 합니다.')
    actual_balance = sum(r['차변'] - r['대변'] for r in changed['총계정원장'])
    result = ledger.analyze(changed)
    failures = [r for r in result['tables']['원장통제'] if r['상태'] == '불일치']
    issue_codes = {r['code'] for r in result['issues']}
    if actual_balance != 0 or not result['blocking'] or not failures or 'LEDGER_CONTROL' not in issue_codes or 'JE_UNBALANCED' in issue_codes:
        raise AssertionError('균형 전표 누락이 독립 통제 대사로 차단되지 않았습니다.')
    return {'company_code': 'JS_ELEC', 'omitted_journal': 'J-PAY', 'omitted_rows': len(removed),
            'remaining_debit_minus_credit': actual_balance, 'independent_control_failures': len(failures),
            'blocking_issues': result['blocking'], 'debit_credit_balance_alone_would_pass': True,
            'blocked_by_independent_controls': True}


def exercise_saved_history(ws, checks):
    """Explicit saves append request/procedure revisions and human review records."""
    ws.refresh()
    initial_head, initial_count = ws.head, len(ws.events)
    pid = 'C02'
    draft = deepcopy(ws.procedures[pid])
    draft['performed'] = '아직 저장하지 않은 합성 메모'
    fresh = Workspace(ws.root)
    if fresh.head != initial_head or fresh.procedures[pid].get('performed'):
        raise AssertionError('명시 저장 전의 메모가 DB 이력에 기록되었습니다.')
    sid = ws.latest()['원장·리드']['source_id']
    request = ws.save_request({'title': '합성 추가자료 요청', 'requested_from': '합성 자료담당',
                              'note': '프로그램의 로컬 이력 검증용, 실제 발송 아님',
                              'due_date': '2027-01-15', 'state': '요청', 'source_ids': []})
    request['state'] = '수령'; request['source_ids'] = [sid]
    ws.save_request(request)
    request['state'] = '검토완료'; ws.save_request(request)
    draft.update(preparer='합성 작성담당', performed='현재 입력의 독립 통제·TB 연결을 확인한 합성 수행 메모',
                 conclusion='합성 자료의 제한된 검사 결과만 확인', performed_date='2026-12-31',
                 purpose='저장 시점 이력 확인', risk='입력 이력 누락 위험', scope='JS전자 검증용 단일 조서',
                 evidence_ids=[sid])
    ws.save_procedure(draft)
    ws.request_review(pid)
    ws.review(pid, '합성 검토담당', '합성 수행 메모 및 연결 원본 확인; 로컬 자유입력 기록')
    saved = Workspace(ws.root)
    added = saved.events[initial_count:]
    kinds = [e['kind'] for e in added]
    if kinds != ['evidence_request']*3 + ['procedure']*3:
        raise AssertionError('자료요청·조서·검토의 저장 이력 순서가 다릅니다.')
    if [e['data']['state'] for e in added[:3]] != ['요청', '수령', '검토완료']:
        raise AssertionError('자료요청의 과거 상태가 보존되지 않았습니다.')
    if [e['data']['state'] for e in added[3:]] != ['작성중', '검토요청', '검토기록완료']:
        raise AssertionError('조서 작성·검토의 과거 상태가 보존되지 않았습니다.')
    if saved.procedures[pid]['reviewer'] != '합성 검토담당' or saved.requests[request['id']]['state'] != '검토완료':
        raise AssertionError('재개한 업무의 검토자/요청 상태가 다릅니다.')
    checks.extend(['JS_ELEC:unsaved_note_not_persisted', 'JS_ELEC:six_explicit_save_events',
                   'JS_ELEC:request_revision_history', 'JS_ELEC:procedure_and_review_revision_history'])
    return {'events_added': len(added), 'request_versions': 3, 'procedure_versions': 3,
            'unsaved_note_persisted': False, 'reopened_head_equal': saved.head == ws.head,
            'reviewer_is_free_text_not_authenticated': True, 'no_email_sent': True}


def exercise(resource_root, output_dir, *, negatives=True, check_originals=True):
    import sector_demo
    resource_root, output_dir = Path(resource_root).resolve(), Path(output_dir).resolve()
    if output_dir == resource_root or output_dir.is_relative_to(resource_root):
        raise AuditError('내장 자료 폴더를 검증 출력 위치로 사용할 수 없습니다.')
    output_dir.mkdir(parents=True, exist_ok=True)
    run_root = output_dir/('run-' + uuid.uuid4().hex[:12])
    run_root.mkdir()
    store = run_root/'company-store'
    before_resources = fingerprint(resource_root)
    before_originals = original_files() if check_originals else {}
    checks, company_rows, returned, workspaces, wrong_company = [], [], {}, {}, []
    parent_wrong_company = []
    raw_view_checks = []
    entries = sector_demo.list_companies(resource_root)
    if {r['code'] for r in entries} != set(COMPANIES):
        raise AssertionError('회사 선택 목록이 고정 5개와 다릅니다.')
    checks.append('five_registered_company_choices')
    catalog = json.loads((ROOT/'procedure_catalog.json').read_text(encoding='utf-8'))
    for code in COMPANIES:
        outcome = sector_demo.open_company(code, store, resource_root)
        if outcome['reused'] or outcome['company_code'] != code or not outcome['workspace_id']:
            raise AssertionError('새 회사 선택 provider 상태가 다릅니다: ' + code)
        ws = Workspace(outcome['folder'])
        if not ws.root.is_relative_to(store.resolve()):
            raise AssertionError('업무가 지정 저장 위치 밖에 생성되었습니다.')
        company_rows.append(validate_workspace(ws, code, checks))
        raw_view_checks.append(validate_data_folder(ws, code, checks))
        returned[code], workspaces[code] = outcome, ws
        exports = run_root/'exports'; exports.mkdir(exist_ok=True)
        path = exports/(code + '_조서초안.xlsx')
        snap = exporter.export(ws, path, catalog)
        if snap['engagement']['회사코드'] != code or snap['export_state'] != '검토중 초안':
            raise AssertionError('조서 증적 회사/상태가 다릅니다.')
        from openpyxl import load_workbook
        book = load_workbook(path, data_only=False)
        try:
            cover = dict(book['조서표지'].iter_rows(min_row=2, values_only=True))
            if cover['회사코드'] != code or cover['회사명'] != COMPANIES[code]:
                raise AssertionError('실제 저장 조서 표지의 회사가 다릅니다.')
            if any(c.data_type in ('f', 'e') for sheet in book for row in sheet for c in row):
                raise AssertionError('조서에 수식 또는 오류 셀이 있습니다.')
            proof = json.loads(''.join(row[1] for row in book['증적JSON'].iter_rows(min_row=2, values_only=True)))
            if proof['engagement']['회사코드'] != code or proof['head'] != ws.head:
                raise AssertionError('조서의 원문 증적/이력 연결이 다릅니다.')
        finally:
            book.close()
        checks.extend([code + ':saved_workpaper_company', code + ':saved_workpaper_evidence_hash'])
    all_source_ids = [set(ws.sources) for ws in workspaces.values()]
    all_run_ids = [set(ws.runs) for ws in workspaces.values()]
    for i in range(5):
        for j in range(i+1, 5):
            if all_source_ids[i] & all_source_ids[j] or all_run_ids[i] & all_run_ids[j]:
                raise AssertionError('서로 다른 회사의 등록 자료/실행 ID가 공유됩니다.')
    checks.append('source_and_run_identity_isolated_across_five_companies')
    # Reopening after switching companies must preserve the complete event head.
    for code, ws in workspaces.items():
        old_head, old_files = ws.head, fingerprint(ws.root)
        again = sector_demo.open_company(code, store, resource_root)
        reopened = Workspace(again['folder'])
        if not again['reused'] or again['workspace_id'] != returned[code]['workspace_id'] or reopened.root != ws.root or reopened.head != old_head or fingerprint(ws.root) != old_files:
            raise AssertionError('회사 재선택이 기존 업무/이력을 바꿨습니다: ' + code)
        checks.append(code + ':reselection_preserves_history_bytes')
    if negatives:
        for code in SUBSIDIARIES:
            ws = workspaces[code]
            current = ws.latest()['원장·리드']['result']
            for other_code in SUBSIDIARIES:
                if code == other_code:
                    continue
                other = workspaces[other_code]
                for module in COMPANY_BOUND_MODULES:
                    source = other.sources[other.latest()[module]['source_id']]
                    path = other.root/source['file']
                    mode = reject_foreign(module, path, ws.context, current)
                    wrong_company.append({'target': code, 'source': other_code, 'module': module, 'status': mode})
        if len(wrong_company) != 72:
            raise AssertionError('자회사 간 혼입 반례 개수가 다릅니다.')
        parent = workspaces['JS']
        for other_code in SUBSIDIARIES:
            other = workspaces[other_code]
            source = other.sources[other.latest()['원장·리드']['source_id']]
            mode = reject_foreign('원장·리드', other.root/source['file'], parent.context,
                                  parent.latest()['원장·리드']['result'])
            parent_wrong_company.append({'target': 'JS', 'source': other_code, 'module': '원장·리드', 'status': mode})
        if len(parent_wrong_company) != 4:
            raise AssertionError('모회사 원장 혼입 반례 개수가 다릅니다.')
    balanced = balanced_omission(resource_root)
    checks.append('balanced_je_omission_blocked_by_independent_controls')
    history = exercise_saved_history(workspaces['JS_ELEC'], checks)
    if fingerprint(resource_root) != before_resources:
        raise AssertionError('내장 원천 모본이 변경되었습니다.')
    if any(not Path(path).is_file() or sha(path) != value for path, value in before_originals.items()):
        raise AssertionError('구버전 원본 파일이 변경되었습니다.')
    result = {'status': 'passed', 'kind': 'source provider headless verification, not native/frozen UI',
              'normal_company_count': 5, 'normal_module_executions': sum(r['modules'] for r in company_rows),
              'current_ledger_links': sum(r['current_ledger_links'] for r in company_rows),
              'cross_company_negative_checks': len(wrong_company), 'parent_cross_company_negative_checks': len(parent_wrong_company),
              'companies': company_rows, 'checks': checks,
              'negative_cases': wrong_company, 'parent_negative_cases': parent_wrong_company,
              'balanced_omission': balanced, 'saved_history': history, 'raw_data_views': raw_view_checks, 'resource_files_unchanged': len(before_resources),
              'original_protection': {'status': 'checked_pre_post' if before_originals else 'not_available_or_not_requested', 'files': len(before_originals)},
              'workspace_root': run_root.name, 'initial_human_fields_blank': True, 'unreviewed_final_export_blocked': True,
              'limitations': ['Parent JS has one separate cash/investment/capital ledger only; four subsidiaries have industry demonstrations; no consolidated group', 'Same synthetic source generator, not independent audit evidence',
                             'Raw D001/D031/D032/D038/D039 preserve separate SYNTH_SUB manual examples; excluded from selected company automated GL/TB',
                             'Company-axis negative checks cover the six modules with actual company/current-ledger binding', 'No native EXE or UI claims from this script']}
    (output_dir/'headless_sectors.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    return result


def main():
    parser = argparse.ArgumentParser(description='5회사 선택·원장·원천·조서 provider 검증')
    parser.add_argument('--resource-root', type=Path, default=ROOT/'sector_resources')
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--no-negatives', action='store_true')
    parser.add_argument('--no-originals', action='store_true')
    args = parser.parse_args()
    result = exercise(args.resource_root, args.output_dir, negatives=not args.no_negatives, check_originals=not args.no_originals)
    print(json.dumps({k: v for k, v in result.items() if k not in ('checks', 'negative_cases', 'companies')}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
