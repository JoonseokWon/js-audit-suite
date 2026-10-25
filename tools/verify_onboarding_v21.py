"""Synthetic source/frozen lifecycle. Never call a model or approve real work."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import json,sys,traceback


def self_test(report_path):
    from tools.verify_bento_v20 import self_test as previous
    report_path=Path(report_path).resolve()
    code=previous(report_path)
    if code:return code
    record=json.loads(report_path.read_text(encoding='utf-8'))
    try:
        from common import AuditError,VERSION,digest,dumps
        from onboarding import MODULES,PROCEDURES
        from onboarding.coverage import COVERAGE
        from onboarding.fixtures import create_demo
        from onboarding.integration import sample,engine
        from onboarding.confirmation import confirm_input
        from onboarding.workbooks import save_input,load_input
        from onboarding.demo import ledger_result
        import exporter,runtime
        from openpyxl import load_workbook
        def check(value,label):
            if not value:raise AssertionError(label)
            record['checks'].append(label)
        base=report_path.parent/'onboarding';base.mkdir()
        ws,runs=create_demo(base/'업무',confirmed=True,actor='SYNTHETIC ONLY 가상확인자',reason='명시적 시험 입력 확인 · 실제 감사승인 아님')
        check(all(not r['result']['blocking'] and r['result']['detail']['현재원장연결'] for r in runs.values()),'v21_all_seven_modules_mechanically_match_current_ledger')
        check(all(x['상태']=='일치' for x in ws.cross_checks()),'v21_cross_module_current_run_pins_match')
        check(len(COVERAGE)==16 and set(COVERAGE)=={f'AU{i:02}' for i in range(1,9)}|{f'CF{i:02}' for i in range(1,9)},'v21_sixteen_review_topics_have_implemented_paths')
        metrics=runs['현금흐름표연결']['result']['metrics']
        check({k:metrics[k] for k in ('net_profit','CFO','CFI','CFF','cash_closing')}==
            {'net_profit':137000,'CFO':103000,'CFI':-55000,'CFF':40000,'cash_closing':1088000},'v21_cashflow_matches_independent_fixed_numbers')
        check(all(p['state']=='작성중' and not p['performed'] and not p['performed_date'] and not p['conclusion']
            and not p['reviewer'] for p in ws.procedures.values()),'v21_human_performance_conclusion_and_review_remain_empty')
        for pid,module in PROCEDURES.items():
            refs=ws.calculation_refs(pid);run=runs[module]
            check(refs['실행ID']==run['id'] and refs['자료ID']==run['source_id'], 'v21_procedure_run_'+pid)
            packet=ws.ai_packet(pid)
            check(any(r['kind']=='calculation_row' and r.get('run_id')==run['id'] for r in packet['references'])
                and any(r['kind'] in ('raw_json','raw_xlsx_row') and r.get('source_id')==run['source_id'] for r in packet['references']),
                'v21_ai_packet_has_own_calculation_and_own_input_'+pid)
            template=sample(module,ws.context)
            check(not template.get('approval') or not template['approval'].get('actor'),'v21_templates_do_not_supply_person_approval_'+pid)
            path=base/(pid+'_양식.xlsx');save_input(module,template,path)
            restored=load_input(module,path,ws.context)
            def canonical(value):
                if isinstance(value,dict):return {k:canonical(v) for k,v in value.items() if k not in ('area','approval')}
                if isinstance(value,list):return [canonical(v) for v in value]
                return None if value=='' else value
            check(canonical(restored)==canonical(template) and not restored['approval'],'v21_business_excel_roundtrip_'+pid)
        # A changed original-input link or missing confirmation envelope must fail.
        from onboarding.integration import run as execute
        payload=confirm_input(MODULES[-1],sample(MODULES[-1],ws.context),'SYNTHETIC','negative confirmation test')
        payload['approval'].pop('envelope_hash')
        bad=base/'removed-envelope.json';bad.write_text(dumps(payload),encoding='utf-8')
        try:execute(MODULES[-1],bad,ws.context,ledger_result(ws.context))
        except AuditError:blocked=True
        else:blocked=False
        check(blocked,'v21_removed_confirmation_envelope_is_rejected')
        # Formula and stale cached TRUE cannot become independently recomputed data.
        formula=base/'formula.xlsx';save_input(MODULES[2],sample(MODULES[2],ws.context),formula)
        workbook=load_workbook(formula);first=next(s for s in workbook if s.title!='_양식정보');first.cell(2,2,'=TRUE()');workbook.save(formula);workbook.close()
        try:load_input(MODULES[2],formula,ws.context)
        except AuditError:blocked=True
        else:blocked=False
        check(blocked,'v21_formula_and_cache_are_not_new_verification')
        catalog=json.loads((runtime.configure()/'procedure_catalog.json').read_text(encoding='utf-8'))
        data=exporter.workbook_bytes(ws.snapshot(False),catalog)
        output=base/'연결조서_검토중.xlsx';output.write_bytes(data)
        workbook=load_workbook(BytesIO(data),read_only=True,data_only=False)
        try:
            text=str([[cell.value for row in s for cell in row] for s in workbook])
            check(all(pid in text for pid in PROCEDURES),'v21_all_workpaper_drafts_export_with_calculation_links')
            check(not any(c.data_type=='f' for s in workbook for row in s for c in row),'v21_export_contains_no_formula_execution')
        finally:workbook.close()
        old_run=ws.latest()['원장·리드']
        ws.execute('원장·리드',ws.root/ws.sources[old_run['source_id']]['file'])
        check(all(x['상태']=='미확인' for x in ws.cross_checks() if x['모듈'] in MODULES),'v21_new_ledger_run_invalidates_old_module_links')
        check(not ws.ai_approvals and all(not p['reviewer'] for p in ws.procedures.values()),'v21_no_final_content_or_review_approval_created')
        record.update(ok=True,version=VERSION,frozen=bool(getattr(sys,'frozen',False)),
            coverage=sorted(COVERAGE),gates={'human_approval_preserved':True,'input_security':True,'ledger_pin':True,'workpaper_export':True},
            current_actual_model_calls=0,onboarding_workspace=str(ws.root),onboarding_export=str(output))
    except Exception as exc:record.update(ok=False,error=str(exc),traceback=traceback.format_exc())
    report_path.write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'ok':record['ok'],'version':record.get('version'),'checks':len(record['checks']),'error':record.get('error')},ensure_ascii=False))
    return 0 if record['ok'] else 1
