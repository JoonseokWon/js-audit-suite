"""Build the v1.7 parent-only JS fixture; preserve all four subsidiary inputs byte for byte."""
from __future__ import annotations
import argparse,copy,hashlib,io,json,sys,zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from common import AuditError,atomic,read_xlsx
import ledger

DEFAULT_SOURCE=ROOT.parent/'v1.6_업종선택통합_2026-10-23'/'sector_resources'
DEFAULT_OUTPUT=ROOT/'sector_resources'
PARENT_FILE='모회사_JS_원장_리드.xlsx'
PARENT_PROCEDURES=['C02','C04','C05','P01','P07']
SUBS=('JS_ELEC','JS_AUTO','JS_SEMI','JS_BIO')
LIMITS=['실제 독립 감사증거가 아닌 가상 선택 시나리오',
        '가상 자본납입 10000000원, 가상 투자 2000000원씩 4건, 가상 기말현금 2000000원',
        '실제 지분율·지배력·투자취득 근거·투자주식 평가·금융기관 조회·연결계산은 미검사',
        'JS 모회사의 업종·매출·재고·제조·유형자산 자료는 생성하지 않음. 미제공을 0으로 대체하지 않음']

def encoded(x):return (json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n').encode('utf-8')
def record(path,data):return {'path':path,'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data)}

def parent_book(context):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    rows={k:[] for k in ledger.SCHEMA}
    accounts=[('1000','현금및예금'),('1801','JS전자 투자주식'),('1802','JS자동차 투자주식'),('1803','JS반도체 투자주식'),('1804','JS바이오 투자주식'),('3000','자본')]
    code=context['회사코드'];start=context['시작일'];cut=context['기준일']
    rows['설정']=[['회사코드',code],['시작일',start],['기준일',cut],['통화','KRW'],['추출범위',';'.join(a for a,_ in accounts)]]
    def post(jid,line,account,debit,credit,description,date):
        rows['총계정원장'].append([code,jid,line,date,account,debit,credit,'SYNTHETIC ONLY / '+description,'합성작성자'])
    post('JS-CAPITAL','1','1000',10000000,0,'가상 자본납입 선택 시나리오','2026-01-02')
    post('JS-CAPITAL','2','3000',0,10000000,'가상 자본납입 선택 시나리오','2026-01-02')
    for number,(_,label) in enumerate(accounts[1:5],1):
        account='180'+str(number);jid='JS-INVEST-'+str(number)
        description=label+' 가상 취득 선택 시나리오 · 지분율/평가 미설정'
        post(jid,'1',account,2000000,0,description,'2026-02-02')
        post(jid,'2','1000',0,2000000,description,'2026-02-02')
    sums=defaultdict(lambda:[0,0,0])
    for r in rows['총계정원장']:
        sums[r[4]][0]+=1;sums[r[4]][1]+=r[5];sums[r[4]][2]+=r[6]
    for account,label in accounts:
        count,debit,credit=sums[account];balance=debit-credit
        rows['시산표'].append([code,account,label,0,debit,credit,balance])
        lead='현금및예금' if account=='1000' else '자본' if account=='3000' else '투자주식'
        rows['계정매핑'].append([account,lead,lead,'BS'])
        rows['독립통제'].append([code,start,cut,account,count,debit,credit,0,balance,'SYNTHETIC ONLY / 동일 생성기 가상 통제값; 실제 독립 추출증거 아님','2027-01-04T09:00:00+09:00','합성작성자'])
    wb=Workbook();wb.remove(wb.active)
    wb.properties.created=wb.properties.modified=datetime(2026,10,24)
    for sheet,headers in ledger.SCHEMA.items():
        ws=wb.create_sheet(sheet);ws.append(headers)
        for row in rows[sheet]:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=22
    buffer=io.BytesIO();wb.save(buffer);wb.close()
    # Stable OOXML member timestamps allow the fixture to be regenerated identically.
    target=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as original,zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
        for name in original.namelist():
            info=zipfile.ZipInfo(name,(2026,10,24,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,original.read(name))
    return target.getvalue()

def prepare_resources(source_dir=DEFAULT_SOURCE,output_dir=DEFAULT_OUTPUT,run_engines=False):
    source=Path(source_dir).resolve();output=Path(output_dir).resolve()
    if output!=DEFAULT_OUTPUT.resolve() or source==output:raise AuditError('생성 위치는 이 v1.7의 sector_resources 폴더여야 합니다.')
    old=json.loads((source/'manifest.json').read_text(encoding='utf-8'))
    if old.get('engine_contract')!='1.6':raise AuditError('불변 v1.6 모본이 필요합니다.')
    payload={}
    for entry in old['files']:
        name=entry['path'];data=(source/name).read_bytes()
        if record(name,data)!=entry:raise AuditError('원본 v1.6 모본 해시 불일치: '+name)
        if not name.startswith('companies/JS/'):payload[name]=data
    companies=copy.deepcopy(old['companies'])
    parent=companies[0]
    parent.update(kind='parent',sector='업종 없음 · 모회사',business_description='JS 모회사의 현금·투자주식·자본만 있는 별도 합성 원장. 제조·유통 업종을 부여하지 않음.',
                  industry_topics=[],industry_topic_count=0,procedure_ids=PARENT_PROCEDURES,source_baseline=None,
                  unperformed_scope=LIMITS,scenario_is_arbitrary=True,ownership_percentage=None,ledger_links_expected=0)
    data=parent_book(parent['context']);name='companies/JS/inputs/'+PARENT_FILE;payload[name]=data
    parent['inputs']=[{'module':'원장·리드',**record(name,data)}]
    investments=[{'company_code':c['code'],'company_name':c['name'],'scenario_amount':2000000,'ownership_percentage':None,'control_assessment':None} for c in companies[1:]]
    objects={
      'master':{'company_code':'JS','company_name':'JS','kind':'parent','sector':parent['sector'],'items':[],'assets':[],
                'investments':investments,'currency':'KRW','profile_file':'profile.json','scenario_is_arbitrary':True,'limitations':LIMITS},
      'profile':{'company_code':'JS','company_name':'JS','kind':'parent','sector':parent['sector'],'ownership_percentage':None,
                 'scenario_is_arbitrary':True,'capital_contribution':10000000,'investment_amount_each':2000000,'closing_cash':2000000,
                 'scenario_amounts_are_user_selectable_examples':True,'industry':None,'operating_data_provided':False,'limitations':LIMITS},
      'plan':{'company_code':'JS','company_name':'JS','kind':'parent','state':'계획·미검토','audit_topics':[],
              'procedure_ids':PARENT_PROCEDURES,'performed':'','conclusion':'','reviewer':'','limitations':LIMITS}}
    parent['metadata']={}
    for key,value in objects.items():
        name='companies/JS/'+key+'.json';data=encoded(value);payload[name]=data;parent['metadata'][key]=record(name,data)
    for c in companies[1:]:c.update(kind='subsidiary',ledger_links_expected=35)
    payload['README.md']=('''# v1.7 회사별 내장 합성자료

JS는 업종 없는 모회사이며 현금·4개 투자주식·자본의 가상 선택 시나리오 원장 1권만 제공합니다.
자본납입 10,000,000원, 투자 2,000,000원씩 4건, 기말현금 2,000,000원은 가상 예시입니다.
실제 지분율·지배력·투자취득 근거·평가·연결계산·은행확인은 미검사입니다.
매출·재고·제조·유형자산 자료를 모회사에 만들거나 미제공 자료를 0으로 대체하지 않습니다.

JS전자·JS자동차·JS반도체·JS바이오는 각각 별도 합성업무입니다.
네 자회사의 60개 계산 입력은 v1.6 바이트를 그대로 보존했습니다. 실제 독립 감사증거는 아닙니다.
20개 업종별 주제는 미수행 수동 계획이며 지분율·내부거래·연결제거는 설정하지 않았습니다.
reference/의 과거 항목표·변환설계는 역사적 참고자료이며 현행 JS 모회사의 입력이 아닙니다.

앱은 선택한 회사 자료를 전용 사용자 폴더에 복사하고 실제 계산·원장·조서 연결을 수행합니다.
자회사는 엄격 검증한 auditpack의 원천 엑셀 16권을 해당 연습의 원천엑셀 폴더에 풉니다.
열람사본은 신규 수령 증거로 다시 등록하지 않으며 누락·변조 시 자동 복구하지 않고 차단합니다.
해시와 로컬 이력은 실제 문서의 진위·본인인증·독립성을 입증하지 않습니다.
생성·검증 도구: tools/prepare_sector_resources_v17.py. 실제 AI 추론·외부 전송은 없습니다.
''').encode('utf-8')
    m={**old,'engine_contract':'1.7','scope':'one industry-free parent ledger and four separate subsidiary engagements; no consolidation or completed human judgment',
       'source_bundle':'v1.6 subsidiary inputs preserved; JS parent newly generated arbitrary scenario','companies':companies,
       'files':[record(n,b) for n,b in sorted(payload.items())]}
    prior={r['resource_path']:r for r in old['provenance']}
    m['provenance']=[]
    for name,data in sorted(payload.items()):
        if name=='README.md':continue
        if name.startswith('companies/JS/'):
            m['provenance'].append({'source_root':'generated-parent-scenario','source_path':'tools/prepare_sector_resources_v17.py',
                'source_sha256':None,'resource_path':name,'resource_sha256':hashlib.sha256(data).hexdigest(),'normalization':'new arbitrary synthetic parent-only fixture; no external evidence'})
        else:m['provenance'].append(prior[name])
    # Remove only the prior declared JS fixture files in this new version, never other versions or user workspaces.
    for entry in old['files']:
        if entry['path'].startswith('companies/JS/') and entry['path'] not in payload:
            path=(output/entry['path']).resolve()
            if not path.is_relative_to(output/'companies'/'JS'):raise AuditError('모회사 교체 경로 오류')
            if path.is_file():path.unlink()
    for name,data in payload.items():atomic(output/name,data)
    atomic(output/'manifest.json',encoded(m))
    return validate_resources(output,run_engines)

def validate_resources(resource_dir=DEFAULT_OUTPUT,run_engines=False):
    from sector_demo import _catalog
    import adapters
    from raw_sources import load_pack
    root,m=_catalog(resource_dir)
    actual={p.relative_to(root).as_posix() for p in root.rglob('*') if p.is_file()}
    if actual!={r['path'] for r in m['files']}|{'manifest.json'}:raise AuditError('미등록·누락 리소스 파일')
    observations=[]
    for c in m['companies']:
        by={r['module']:root/r['path'] for r in c['inputs']};tables,_=read_xlsx(by['원장·리드'],ledger.SCHEMA);lr=ledger.analyze(tables)
        if lr['blocking']:raise AuditError('원장 입력 차단: '+c['code'])
        if c['kind']=='parent':
            if set(r['계정코드'] for r in tables['시산표'])!={'1000','1801','1802','1803','1804','3000'} or tables['모듈연결']:raise AuditError('모회사 범위 불일치')
            if tables['시산표'][0]['기말잔액']!=2000000:raise AuditError('모회사 예시 잔액 불일치')
        else:load_pack(by['원천증빙'],c['context'])
        ran=1
        if run_engines:
            for mod,path in by.items():
                if mod=='원장·리드':continue
                if adapters.run(mod,path,c['context'],ledger_result=lr)['blocking']:raise AuditError('계산 차단: '+c['code']+'/'+mod)
                ran+=1
        observations.append({'code':c['code'],'kind':c['kind'],'inputs':len(by),'ledger_rows':len(tables['총계정원장']),'tb_accounts':len(tables['시산표']),'engines_executed':ran})
    return {'status':'passed','engine_contract':'1.7','companies':5,'input_files':61,'subsidiary_inputs_preserved':60,'manual_industry_topics':20,'observations':observations}

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--source-dir',type=Path,default=DEFAULT_SOURCE);p.add_argument('--output-dir',type=Path,default=DEFAULT_OUTPUT)
    p.add_argument('--validate-only',action='store_true');p.add_argument('--run-engines',action='store_true');a=p.parse_args()
    result=validate_resources(a.output_dir,a.run_engines) if a.validate_only else prepare_resources(a.source_dir,a.output_dir,a.run_engines)
    print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
