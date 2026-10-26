"""Exercise an unsplit synthetic ledger through actual source/frozen Workspace and Excel."""
import hashlib,json,sys,time,traceback
from pathlib import Path

def self_test(report_path,input_path,rows,total):
    from common import VERSION
    from workspace import Workspace
    from exporter import export
    from xlsx_reader import file_hash
    from openpyxl import load_workbook
    report_path=Path(report_path).resolve();folder=report_path.parent
    if folder.exists() and any(folder.iterdir()):raise FileExistsError('새 검증 폴더를 사용하세요.')
    folder.mkdir(parents=True,exist_ok=True);started=time.perf_counter();timings={}
    def phase(name):
        (folder/'progress.json').write_text(json.dumps({'phase':name,'seconds':round(time.perf_counter()-started,2)}),encoding='utf-8')
        return time.perf_counter()
    try:
        source_hash=file_hash(input_path);stamp=phase('workspace_execute')
        scope={'회사코드':'LARGE_DEMO','회사명':'큰 원장 입력 시험','시작일':'2026-01-01','기준일':'2026-01-31',
               '회계체계':'명시 시험 계정체계','감사기준':'입력 호환성 시험','작성자':'SYNTHETIC ONLY','더미업무':True}
        work=folder/'workspace';task=Workspace(work,create=scope);run=task.execute('원장·리드',str(input_path));result=run['result']
        assert not result['blocking'],result['issues']
        gl=result['tables']['입력_총계정원장'];controls=result['tables']['원장통제']
        assert len(gl)==rows and sum(r['차변'] for r in gl)==sum(r['대변'] for r in gl)==total
        assert gl[0]['계정코드']=='0001' and gl[-1]['행ID']=='002'
        assert all(r['상태']=='일치' for r in controls)
        checks=len(controls);result_hash=run['result_hash'];timings['execute']=round(time.perf_counter()-stamp,3)
        del result,gl,run
        stamp=phase('reopen_and_verify');task=Workspace(work)
        assert task.latest()['원장·리드']['result_hash']==result_hash and not task.verify_sources()
        assert len(task.latest()['원장·리드']['result']['tables']['입력_총계정원장'])==rows
        timings['reopen']=round(time.perf_counter()-stamp,3)
        stamp=phase('excel_export');root=Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parents[1]))
        catalog=json.loads((root/'procedure_catalog.json').read_text(encoding='utf-8'))
        output=folder/'large-review.xlsx';export(task,output,catalog,False);timings['export']=round(time.perf_counter()-stamp,3)
        stamp=phase('saved_excel_verify');wb=load_workbook(output,read_only=True,data_only=False)
        try:
            name=next(n for n in wb.sheetnames if n.endswith('입력_총계정원장'));sheet=wb[name];count=0;debit=credit=0
            headers=list(next(sheet.values));di=headers.index('차변');ci=headers.index('대변');account=headers.index('계정코드');line=headers.index('행ID')
            for row in sheet.iter_rows(min_row=2,values_only=True):
                count+=1;debit+=row[di];credit+=row[ci]
                assert row[account] in ('0001','0002') and row[line] in ('001','002')
            assert count==rows and debit==credit==total
            proof_rows=sum(1 for _ in wb['증적JSON'].iter_rows(min_row=2,values_only=True));assert proof_rows>1
        finally:wb.close()
        timings['saved_excel_verify']=round(time.perf_counter()-stamp,3)
        stamp=phase('desktop_first_last_pages')
        import runtime;runtime.configure_tk_libraries()
        import tkinter as tk
        from app import App
        root=tk.Tk();root.withdraw()
        try:
            app=App(root);app.open_workspace(str(work));app.result_module.set('원장·리드');app.choose_result_module()
            app.result_table.set('입력_총계정원장');app.load_table();root.update()
            assert len(app.table_rows)==rows and len(app.visible_rows)==200
            app.turn_page(10**9);root.update();assert app.visible_rows[-1]['전표ID']==f'L-{rows//2-1:07d}'
            assert len(app.result_tree.get_children())<=200
        finally:root.destroy()
        timings['desktop_pages']=round(time.perf_counter()-stamp,3)
        assert file_hash(input_path)==source_hash
        report={'ok':True,'version':VERSION,'frozen':bool(getattr(sys,'frozen',False)),'rows':rows,'journals':rows//2,
                'debitTotal':total,'creditTotal':total,'controlChecks':checks,'matchedChecks':checks,'blockingIssues':0,
                'workspaceReopened':True,'sourceIntegrity':True,'excelRows':count,'proofChunks':proof_rows,
                'leadingZerosPreserved':True,'desktopFirstLastPages':True,'desktopPageRows':200,
                'inputSha256':source_hash,'resultHash':result_hash,'seconds':round(time.perf_counter()-started,3),'timings':timings}
    except Exception as exc:report={'ok':False,'version':VERSION,'error':str(exc),'traceback':traceback.format_exc(),'timings':timings}
    report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');phase('complete' if report['ok'] else 'failed')
    print(json.dumps(report,ensure_ascii=False));return 0 if report['ok'] else 1

if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--report',required=True);parser.add_argument('--input',required=True)
    parser.add_argument('--rows',type=int,required=True);parser.add_argument('--total',type=int,required=True);args=parser.parse_args()
    raise SystemExit(self_test(args.report,args.input,args.rows,args.total))
