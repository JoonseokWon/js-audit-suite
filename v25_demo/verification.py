"""Deterministic mock route only. Never evidence of actual model inference."""
from pathlib import Path
from unittest.mock import patch
import json
import sys
from common import digest
from .factory import create_demo, load_extension
from .case import actual_case
from .flow import complete_actual


def mock_response(packet):
    from ai_contract.generation import generation_plan
    from test_review_ai_helpers import response_from_fields
    plan = generation_plan(packet)
    sentence = 'SYNTHETIC ONLY / 모델을 호출하지 않은 고정 응답입니다. 인용 근거의 의미와 실제 감사판단은 미검증이며 사람이 별도로 확인해야 합니다.'
    proposals = {field: plan['prefixes'].get(field, '') + sentence for field in ('purpose', 'risk', 'scope', 'performed', 'conclusion')}
    return response_from_fields(packet, proposals, plan['citations'], model='SYNTHETIC_ONLY_NO_MODEL_CALL')


def run_demo_verification(folder, *, include_hidden_ui=True):
    from demo_case import load_manifest
    from openpyxl import load_workbook
    import audit_links, misstatements
    base = Path(folder)
    if base.exists() and any(base.iterdir()):
        raise ValueError('검증은 새 빈 폴더에서 실행하며 이전 시도를 보존합니다.')
    base.mkdir(parents=True, exist_ok=True); checks = []
    def check(ok, name):
        if not ok: raise AssertionError(name)
        if name in checks: raise AssertionError('duplicate_check_name: ' + name)
        checks.append(name)
    with patch('ai_provider.generate_draft', side_effect=AssertionError('deterministic path must never call a model')) as provider:
        user, manifest = create_demo(base/'사용자용_JS전자_빈기록')
        original = load_manifest(user.root); extension = load_extension(user.root)
        check(digest(original['inputs']) == digest(manifest['inputs']), 'v25_original_sealed_demo_manifest_and_input_contract_remain_compatible')
        check(not user.procedures and not user.ai_approvals and not user.ai_drafts and not user.workflow_profile,
              'v25_user_walkthrough_contains_no_human_procedures_assessment_approval_or_model_response')
        check(not audit_links.status(user)['active'] and not misstatements.list_records(user),
              'v25_user_walkthrough_never_invents_risk_or_misstatement_judgments')
        book = load_workbook(user.root/extension['inputs'][0]['path'], data_only=False)
        try:
            row = list(book['검토표_빈양식'].iter_rows(values_only=True))[1]
            check(all(value in (None, '') for value in row[2:]), 'v25_user_review_template_leaves_all_human_choice_and_actor_cells_blank')
            check(not any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row),
                  'v25_user_supplementary_inputs_contain_no_formula_cells')
        finally: book.close()
        if include_hidden_ui:
            if sys.platform != 'win32': raise RuntimeError('hidden native UI verification requires Windows')
            from .ui import verify_user_ui
            verify_user_ui(user, check)
        ws, _ = actual_case(base/'SYNTHETIC_ONLY_통합상태전이', check)
        packet = ws.ai_packet('A04', limits={'max_chars': 24000, 'max_references': 70, 'rows_per_table': 3, 'max_tables': 8, 'cell_chars': 300})
        response = mock_response(packet)
        check(response['model'] == 'SYNTHETIC_ONLY_NO_MODEL_CALL', 'v25_mock_response_is_explicitly_labelled_and_separate_from_actual_inference')
        flow = complete_actual(ws, packet, response, base/'SYNTHETIC_ONLY_통합상태전이', check)
        check(provider.call_count == 0, 'v25_deterministic_flow_makes_zero_provider_or_model_calls')
    return {'ok': True, 'schema': 'v25-deterministic-only-1', 'checks': checks, 'flow': flow,
            'actual_model_calls': 0, 'current_actual_model_calls': 0, 'actual_local_inference': False,
            'actual_human_approvals': 0, 'synthetic_approval_events': flow['synthetic_approval_events'],
            'user_workspace': str(user.root), 'hidden_UI_verified': include_hidden_ui,
            'physical_IME_verified': False, 'physical_DPI_verified': False, 'actual_Office_render_verified': False,
            'scope': 'Mock response, synthetic state transitions and static OOXML only. Actual inference must be validated separately in current source and EXE.'}

