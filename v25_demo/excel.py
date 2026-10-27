"""Static OOXML and exact evidence-JSON checks; not physical Office rendering."""
from pathlib import Path
import json
from common import digest


def verify_export(ws, did, original, out, check):
    import exporter, runtime
    from openpyxl import load_workbook
    target = Path(out)/'JS전자_v25_위험_왜곡표시_AI조서_합성검증.xlsx'
    catalog = json.loads((runtime.configure()/'procedure_catalog.json').read_text(encoding='utf-8'))
    expected = exporter.export(ws, target, catalog, False)
    book = load_workbook(target, data_only=False)
    try:
        required = ('왜곡표시검토대장', '왜곡표시상세기록', '왜곡표시금액집계', '동일오류그룹', '왜곡표시계정영향',
                    '위험절차연결점검', '경영진주장위험평가', '위험대응절차', '사람수행기록', '위험대응조서연결',
                    'AI프로그램근거', 'AI미검증제안', 'AI구조화주장', 'AI최종내용승인', 'AI근거참조', '조서_A04', '조서_F03')
        check(all(name in book.sheetnames for name in required), 'v25_Excel_has_separate_misstatement_risk_human_AI_and_workpaper_tables')
        rows = list(book['왜곡표시금액집계'].iter_rows(values_only=True)); headers = rows[0]
        gross = sum(row[headers.index('절대금액 합계')] or 0 for row in rows[1:])
        count = sum(row[headers.index('대표 건수')] or 0 for row in rows[1:])
        check(gross == 7 and count == 1, 'v25_Excel_gross_seven_won_and_one_representative_not_doubled')
        impacts = list(book['왜곡표시계정영향'].iter_rows(values_only=True))
        check(len(impacts) >= 5, 'v25_Excel_keeps_at_least_four_member_account_impacts')
        check(not any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row), 'v25_OOXML_has_zero_executable_formula_cells')
        check(all((dim.height or 0) <= 409 for sheet in book for dim in sheet.row_dimensions.values()), 'v25_OOXML_row_heights_stay_within_409_points')
        restored = json.loads(''.join(row[1] for row in book['증적JSON'].iter_rows(min_row=2, values_only=True)))
        draft = next(d for d in restored['ai_drafts'] if d['id'] == did)
        check(draft['raw_response'] == original, 'v25_evidence_JSON_restores_exact_received_model_original')
        check(digest(restored) == digest(expected), 'v25_evidence_JSON_restores_complete_snapshot_without_record_loss')
        check(restored['export_state'] == '검토중 초안', 'v25_Excel_does_not_claim_overall_audit_completion')
        check(len(restored['misstatements']['records']) == 2 and len(restored['misstatements']['account_impacts']) >= 4,
              'v25_evidence_JSON_keeps_two_records_and_all_account_impacts')
        response = next(iter(restored['audit_links']['responses'].values()))
        check(bool(response['sources']) and bool(response['runs']) and all(r['result_hash'] for r in response['runs']),
              'v25_evidence_JSON_keeps_risk_assertions_original_locations_and_execution_hashes')
    finally:
        book.close()
    return target
