"""Explicit, non-overwriting engagement archives; verification never extracts files.

An archive records an existing event head. Creating or checking one never appends
an event, grants approval, locks an engagement, or rewrites its anchor.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path, PurePosixPath
import copy
import hashlib
import json
import os
import re
import sqlite3
import stat
import tempfile
import unicodedata
import zipfile

from common import AuditError, dumps, now, text
from exporter import workbook_bytes
from workspace import Workspace

SCHEMA = 'audit-suite-archive-1'
MANIFEST = 'manifest.json'
DATABASE = '업무.sqlite3'
ANCHOR = '이력헤드.json'
SNAPSHOT = '감사스냅샷.json'
WORKBOOK = '감사조서.xlsx'
SCOPE = '보관범위.json'
CORE = {DATABASE, ANCHOR, SNAPSHOT, WORKBOOK, SCOPE}
MAX_FILES = 10000
MAX_FILE_BYTES = 512 * 1024**2
MAX_TOTAL_BYTES = 1024**3
from input_limits import MAX_SOURCE_BYTES
MAX_MANIFEST_BYTES = 8 * 1024**2
_HEX = re.compile(r'[0-9a-f]{64}\Z')
_SOURCE_ID = re.compile(r'[0-9a-f]{32}\Z')
_PRIVATE_SUFFIXES = {'.gguf', '.safetensors', '.bin', '.pth', '.pt', '.ckpt',
                     '.pem', '.key', '.p12', '.pfx', '.sqlite', '.sqlite3',
                     '.db', '.exe', '.dll', '.ps1', '.bat', '.cmd', '.py'}
_PRIVATE_NAMES = {'credentials.json', 'token.json', 'tokens.json', 'secrets.json',
                  'api_keys.json', 'api-key.json', '.env', 'id_rsa', 'id_ed25519'}


def _hash(raw):
    return hashlib.sha256(raw).hexdigest()


def _json(raw, label):
    def unique(pairs):
        out = {}
        for key, value in pairs:
            if key in out:
                raise AuditError(label + ': 중복 JSON 키')
            out[key] = value
        return out
    try:
        return json.loads(raw, object_pairs_hook=unique,
                          parse_constant=lambda value: (_ for _ in ()).throw(ValueError(value)))
    except (ValueError, TypeError, UnicodeError, RecursionError) as exc:
        raise AuditError(label + ': JSON 구조 오류') from exc


def _safe_name(name):
    """Use the same unambiguous, relative path on Windows and ZIP readers."""
    if not isinstance(name, str) or not name or len(name) > 240:
        raise AuditError('보관 파일 경로 오류')
    if name != unicodedata.normalize('NFC', name) or '\\' in name or '\x00' in name:
        raise AuditError('보관 파일 경로 정규화 오류')
    p = PurePosixPath(name)
    if p.is_absolute() or p.as_posix() != name or any(part in ('', '.', '..') for part in p.parts):
        raise AuditError('보관 파일의 경로 탈출/상대경로 오류')
    for part in p.parts:
        if any(ch in part for ch in ':<>"|?*') or part.endswith((' ', '.')) or any(ord(ch) < 32 for ch in part):
            raise AuditError('보관 파일 경로 오류')
        if re.fullmatch(r'(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?', part, re.I):
            raise AuditError('예약된 보관 파일 경로')
    return name


def _path_without_links(path):
    """Check lexical ancestors before resolve, including Windows junctions."""
    p = Path(os.path.abspath(os.fspath(path)))
    for current in (p, *p.parents):
        try:
            info = current.lstat()
        except FileNotFoundError:
            continue
        if stat.S_ISLNK(info.st_mode) or getattr(info, 'st_file_attributes', 0) & getattr(stat, 'FILE_ATTRIBUTE_REPARSE_POINT', 0x400):
            raise AuditError('심볼릭 링크/재분석 지점 경로는 보관에 사용할 수 없습니다.')
    return p


def _identity(info):
    return (info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns)


def _read_regular(path, limit=MAX_FILE_BYTES):
    try:
        return _read_regular_checked(path, limit)
    except OSError as exc:
        raise AuditError('보관 파일을 읽을 수 없거나 변경/누락되었습니다: ' + Path(path).name) from exc


def _read_regular_checked(path, limit):
    p = _path_without_links(path)
    before = p.lstat()
    if not stat.S_ISREG(before.st_mode) or before.st_size > limit:
        raise AuditError('보관 파일 형식/크기 한도 오류: ' + p.name)
    flags = os.O_RDONLY | getattr(os, 'O_BINARY', 0) | getattr(os, 'O_NOFOLLOW', 0)
    fd = os.open(p, flags)
    with os.fdopen(fd, 'rb') as handle:
        opened = os.fstat(handle.fileno())
        if _identity(before) != _identity(opened):
            raise AuditError('보관 파일이 읽기 전에 변경되었습니다.')
        raw = handle.read(limit + 1)
        after = os.fstat(handle.fileno())
    _path_without_links(p)
    if len(raw) > limit or _identity(before) != _identity(after) or _identity(before) != _identity(p.lstat()):
        raise AuditError('보관 파일이 읽는 도중 변경되었습니다.')
    return raw, _identity(after)


def _zip_name(info):
    original = _safe_name(info.orig_filename)
    if original != info.filename:
        raise AuditError('보관 ZIP 원문 경로와 정규화 경로가 다릅니다.')
    return original


def _source_records(sources):
    records = {}
    names = set()
    for source in sources:
        if not isinstance(source, dict):
            raise AuditError('보관 원본 등록부 구조 오류')
        sid = source.get('id', '')
        name = _safe_name(source.get('file'))
        p = PurePosixPath(name)
        if not isinstance(sid, str) or not _SOURCE_ID.fullmatch(sid) or sid in records:
            raise AuditError('보관 원본 ID 중복/오류')
        if len(p.parts) != 2 or p.parts[0] != '원본' or p.stem != sid or p.suffix != p.suffix.lower():
            raise AuditError('다른 업무/고아 원본 경로는 보관할 수 없습니다.')
        original_name = str(source.get('original', '')).replace('\\', '/').rsplit('/', 1)[-1].lower()
        if p.suffix in _PRIVATE_SUFFIXES or original_name in _PRIVATE_NAMES or original_name.endswith('.env'):
            raise AuditError('AI 모델·키·설정·실행파일은 증거 취합 범위가 아닙니다.')
        if not _HEX.fullmatch(str(source.get('sha256', ''))) or type(source.get('bytes')) is not int or not 0 <= source['bytes'] <= MAX_SOURCE_BYTES:
            raise AuditError('보관 원본의 SHA256/크기 오류')
        if name.casefold() in names:
            raise AuditError('중복 원본 경로')
        names.add(name.casefold())
        records[sid] = copy.deepcopy(source)
    return records


def _workflow(ws):
    method = getattr(ws, 'workflow_status', None)
    status = method() if callable(method) else {'active': False, 'locked': False, 'blockers': [], 'human_pending': []}
    if not isinstance(status, dict):
        raise AuditError('감사 흐름 상태 구조 오류')
    return copy.deepcopy(status)


def _pin(ws):
    ws.refresh()
    root = _path_without_links(ws.root).resolve(strict=True)
    _path_without_links(ws.dbpath)
    anchor_raw, anchor_identity = _read_regular(ws.anchor, 65536)
    head = {'seq': len(ws.events), 'hash': ws.head}
    if _json(anchor_raw, ANCHOR) != head:
        raise AuditError('보관 준비 중 이력 헤드 변경')
    basis = ws.basis()
    records = _source_records(ws.sources.values())
    expected_paths = {s['file'] for s in records.values()}
    folder = root / '원본'
    if folder.exists():
        _path_without_links(folder)
        if not folder.is_dir():
            raise AuditError('원본 폴더 구조 오류')
        actual = set()
        for child in folder.iterdir():
            _path_without_links(child)
            if not child.is_file():
                raise AuditError('원본 폴더의 하위폴더/고아 경로는 보관할 수 없습니다.')
            actual.add('원본/' + child.name)
        if actual != expected_paths:
            raise AuditError('등록되지 않은 고아 원본/누락 원본을 먼저 확인하세요.')
    elif records:
        raise AuditError('등록 원본 폴더가 없습니다.')
    source_state = {}
    for sid, source in sorted(records.items()):
        p = root / source['file']
        if p.resolve() != p or not p.is_relative_to(root):
            raise AuditError('다른 업무의 원본/경로 탈출')
        raw, identity = _read_regular(p, MAX_SOURCE_BYTES)
        if len(raw) != source['bytes'] or _hash(raw) != source['sha256']:
            raise AuditError('등록 원본 사본의 크기/해시 변조: ' + source['name'])
        source_state[sid] = identity
    workflow = _workflow(ws)
    ws.refresh()
    if ws.head != head['hash'] or len(ws.events) != head['seq'] or ws.basis() != basis:
        raise AuditError('보관 준비 중 자료·재분석·위험·이력 변경')
    return {'root': str(root), 'head': head, 'basis': basis, 'anchor': anchor_raw,
            'anchor_identity': anchor_identity, 'sources': records, 'source_state': source_state,
            'workflow': workflow, 'engagement_id': ws.engagement['id']}


def _assert_pin(ws, pin):
    if _pin(ws) != pin:
        raise AuditError('보관 도중 자료·재분석·위험·이력 또는 원본 변경. 새 경로로 다시 취합하세요.')


def _final_gate(ws, workflow):
    if not workflow.get('active') or not workflow.get('locked'):
        raise AuditError('최종 취합은 사람이 완료 점검 후 업무 보관 잠금을 먼저 기록해야 합니다.')
    if workflow.get('blockers') or workflow.get('human_pending'):
        raise AuditError('기계적 차단/사람 판단 미완료 상태에서는 최종 취합할 수 없습니다.')
    return ws.snapshot(final=True)


def _backup_database(ws, path, pin):
    source = sqlite3.connect(ws.dbpath.as_uri() + '?mode=ro', uri=True, timeout=5)
    target = sqlite3.connect(path)
    try:
        source.execute('PRAGMA trusted_schema=OFF')
        source.execute('BEGIN')
        events, head = ws._read(source)
        if head != pin['head']['hash'] or len(events) != pin['head']['seq']:
            raise AuditError('DB 스냅샷과 고정한 이력 헤드 불일치')
        source.backup(target)
        target.execute('PRAGMA journal_mode=DELETE')
        target.commit()
    finally:
        target.close()
        source.close()


def _catalog():
    import runtime
    path = runtime.configure() / 'procedure_catalog.json'
    return _json(path.read_bytes(), '고정 절차목록')


def _scope_record(mode, actor, created, engagement_id, head, basis, sources, workflow):
    return {'schema': SCHEMA, 'mode': mode,
            'state': '잠금된 선택범위 최종 취합' if mode == 'final' else '미완료 상태를 포함한 검토중 초안',
            'created_utc': created, 'actor': actor, 'engagement_id': engagement_id,
            'head': head, 'basis': basis, 'source_ids': sorted(sources),
            'workflow_locked': bool(workflow.get('locked')),
            'blockers': workflow.get('blockers', []), 'human_pending': workflow.get('human_pending', []),
            'included': '현재 업무 DB·이력헤드·등록 원본 사본 전부·같은 스냅샷의 정적 XLSX·스냅샷JSON·범위·manifest',
            'excluded': ['등록 전 사용자 원본 경로의 파일', '다른 업무·사용자 폴더', '임의 첨부·미등록 파일', 'AI 모델·API 키·설정·런타임'],
            'meaning': '기록한 범위의 보관이며 전체 감사완료·기준 준수 인증·감사의견·본인인증 서명이 아님.',
            'history_policy': '취합은 이력 이벤트/승인/잠금을 추가하거나 헤드를 다시 쓰지 않음.'}


def create_archive(ws, path, actor, final=False):
    """Create a new ZIP at an explicit absolute path, or leave no new archive.

    ``final=True`` requires a pre-existing human finalization/lock event. This
    method does not create that event. Its actor is an attribution label, not an
    authenticated signature. Only registered source copies are read.
    """
    actor = text(actor, '보관 실행자')
    if type(final) is not bool:
        raise AuditError('초안/최종 보관 모드를 명시하세요.')
    selected = Path(path)
    if not selected.is_absolute() or selected.suffix.lower() != '.zip':
        raise AuditError('사용자가 선택한 새 절대경로의 .zip 파일이 필요합니다.')
    target = _path_without_links(selected)
    if not target.parent.is_dir():
        raise AuditError('존재하는 보관 대상 폴더를 선택하세요.')
    if target.exists():
        raise AuditError('기존 보관 파일은 덮어쓸 수 없습니다. 새 파일명을 선택하세요.')
    root = _path_without_links(ws.root).resolve(strict=True)
    if target.is_relative_to(root / '원본'):
        raise AuditError('원본 폴더에는 보관 ZIP을 저장할 수 없습니다.')
    pin = _pin(ws)
    for source in pin['sources'].values():
        # Never touch an external original, including one bearing a .zip name.
        original = source.get('original')
        if original and Path(original).is_absolute() and target == Path(os.path.abspath(original)):
            raise AuditError('등록한 사용자 원본 경로에는 보관할 수 없습니다.')
    snapshot = _final_gate(ws, pin['workflow']) if final else ws.snapshot(final=False)
    snapshot = copy.deepcopy(snapshot)
    if snapshot.get('head') != pin['head']['hash'] or snapshot.get('basis') != pin['basis']:
        raise AuditError('조서 스냅샷과 고정한 이력 헤드 불일치')
    _assert_pin(ws, pin)
    mode = 'final' if final else 'draft'
    scope = _scope_record(mode, actor, now(), pin['engagement_id'], pin['head'], pin['basis'], pin['sources'], pin['workflow'])
    created_identity = None
    try:
        with tempfile.TemporaryDirectory(prefix='.audit-archive-', dir=target.parent) as staging:
            stage = Path(staging)
            _backup_database(ws, stage / DATABASE, pin)
            payloads = {DATABASE: (stage / DATABASE).read_bytes(), ANCHOR: pin['anchor'],
                        SNAPSHOT: dumps(snapshot).encode('utf-8'), WORKBOOK: workbook_bytes(snapshot, _catalog()),
                        SCOPE: dumps(scope).encode('utf-8')}
            manifest = {key: scope[key] for key in ('schema', 'mode', 'created_utc', 'actor', 'engagement_id', 'head', 'basis')}
            entries = []
            staged_zip = stage / '보관.zip'
            with zipfile.ZipFile(staged_zip, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
                def add(name, raw, kind, source_id=None):
                    entry = {'path': _safe_name(name), 'bytes': len(raw), 'sha256': _hash(raw), 'kind': kind}
                    if source_id is not None:
                        entry['source_id'] = source_id
                    entries.append(entry)
                    if len(raw) > MAX_FILE_BYTES or sum(item['bytes'] for item in entries) > MAX_TOTAL_BYTES or len(entries) >= MAX_FILES:
                        raise AuditError('보관 크기/파일 개수 한도 초과')
                    archive.writestr(name, raw)
                for name, raw in sorted(payloads.items()):
                    add(name, raw, 'core')
                for sid, source in sorted(pin['sources'].items()):
                    raw, identity = _read_regular(root / source['file'], MAX_SOURCE_BYTES)
                    if identity != pin['source_state'][sid] or _hash(raw) != source['sha256'] or len(raw) != source['bytes']:
                        raise AuditError('취합 중 등록 원본 변경')
                    add(source['file'], raw, 'source', sid)
                manifest['files'] = sorted(entries, key=lambda item: item['path'])
                archive.writestr(MANIFEST, dumps(manifest).encode('utf-8'))
            report = verify_archive(staged_zip)
            _assert_pin(ws, pin)
            _path_without_links(target)
            # Exclusive create, including when another process wins after validation.
            try:
                with target.open('xb') as output, staged_zip.open('rb') as source:
                    created_identity = os.fstat(output.fileno())
                    while chunk := source.read(1024 * 1024):
                        output.write(chunk)
                    output.flush()
                    os.fsync(output.fileno())
                    created_identity = os.fstat(output.fileno())
            except FileExistsError as exc:
                raise AuditError('대상 파일이 먼저 생성되었습니다. 덮어쓰지 않았습니다.') from exc
            _path_without_links(target)
            if _identity(target.lstat()) != _identity(created_identity):
                raise AuditError('보관 대상 경로가 저장 도중 변경되었습니다.')
            if _hash(_read_regular(target)[0]) != _hash(_read_regular(staged_zip)[0]):
                raise AuditError('저장한 보관 ZIP의 해시 불일치')
            _assert_pin(ws, pin)
            report.update(path=str(target), zip_sha256=_hash(_read_regular(target)[0]),
                          bytes=target.stat().st_size, head=pin['head'], basis=pin['basis'])
            return report
    except Exception as exc:
        # Remove only our own incomplete output; never delete a replacement file.
        if created_identity is not None:
            try:
                current = target.lstat()
                if current.st_dev == created_identity.st_dev and current.st_ino == created_identity.st_ino:
                    target.unlink()
            except OSError:
                pass
        if isinstance(exc, AuditError):
            raise
        raise AuditError('보관 취합 실패: ' + str(exc)) from exc


class _MemoryAnchor:
    def __init__(self, raw):
        self.raw = raw

    def read_text(self, encoding='utf-8'):
        return self.raw.decode(encoding)


class _ArchivedWorkspace(Workspace):
    """Read fixed app code against a validated in-memory DB, never extract it."""
    def __init__(self, raw, anchor):
        self._archive_database = raw
        self.anchor = _MemoryAnchor(anchor)
        self.root = Path('__archive_memory_only__')
        self.refresh()

    def _connect(self):
        connection = sqlite3.connect(':memory:')
        try:
            connection.deserialize(self._archive_database)
            connection.execute('PRAGMA trusted_schema=OFF')
            connection.execute('PRAGMA query_only=ON')
            return connection
        except Exception:
            connection.close()
            raise

    def verify_sources(self):
        # The caller first verifies every registered byte against this DB.
        return []

    def _append(self, kind, data):
        raise AuditError('보관 검증은 이력을 변경할 수 없습니다.')


def _verify_workbook(raw, snapshot):
    """Parse static data only; no Excel process, macros, links, or formulas."""
    from defusedxml.ElementTree import fromstring
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        names = set()
        if len(archive.infolist()) > 3000 or sum(item.file_size for item in archive.infolist()) > MAX_FILE_BYTES:
            raise AuditError('보관 조서 XLSX 크기/구조 한도 초과')
        for info in archive.infolist():
            name = _zip_name(info)
            file_mode = info.external_attr >> 16
            if name.casefold() in names or info.is_dir() or info.flag_bits & 1 or stat.S_IFMT(file_mode) not in (0, stat.S_IFREG) or info.external_attr & 0x400:
                raise AuditError('보관 조서 XLSX 중복/암호화 구조 오류')
            names.add(name.casefold())
            lower = name.lower()
            if any(token in lower for token in ('vba', 'externallinks/', 'embeddings/', 'connections.xml')):
                raise AuditError('보관 조서의 매크로·외부연결·첨부는 허용하지 않습니다.')
            if lower.endswith(('.xml', '.rels')):
                element = fromstring(archive.read(info))
                if any(node.tag.rsplit('}', 1)[-1] == 'f' for node in element.iter()):
                    raise AuditError('보관 조서에 실행 가능한 수식이 있습니다.')
                if any(node.attrib.get('TargetMode', '').lower() == 'external' for node in element.iter()):
                    raise AuditError('보관 조서에 외부 연결이 있습니다.')
        if archive.testzip() is not None:
            raise AuditError('보관 조서 XLSX CRC 오류')
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        if '증적JSON' not in workbook.sheetnames or '조서표지' not in workbook.sheetnames:
            raise AuditError('보관 조서의 증적/표지가 없습니다.')
        proof = ''.join(str(row[1] or '') for row in workbook['증적JSON'].iter_rows(min_row=2, values_only=True))
        if dumps(_json(proof, 'XLSX 증적')) != dumps(snapshot):
            raise AuditError('보관 조서와 DB 스냅샷의 증적 불일치')
        cover = {row[0]: row[1] for row in workbook['조서표지'].iter_rows(min_row=2, values_only=True)}
        if cover.get('업무ID') != snapshot['engagement']['id'] or cover.get('출력상태') != snapshot['export_state']:
            raise AuditError('보관 조서의 업무/초안·최종 표시 불일치')
    finally:
        workbook.close()


def verify_archive(path):
    """Check allowlisted composition, CRC, SHA/size, DB head and all sources.

    This checks internal consistency, not external authorship or the truth of
    evidence. It neither extracts members nor opens documents in their apps.
    """
    try:
        selected = _path_without_links(path)
        if selected.suffix.lower() != '.zip':
            raise AuditError('검증할 .zip 보관 파일을 선택하세요.')
        raw, identity = _read_regular(selected)
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > MAX_FILES or sum(item.file_size for item in infos) > MAX_TOTAL_BYTES:
                raise AuditError('보관 ZIP 크기/파일 개수 한도 초과')
            names = {}
            folded = set()
            for info in infos:
                name = _zip_name(info)
                mode = info.external_attr >> 16
                if info.is_dir() or (stat.S_IFMT(mode) not in (0, stat.S_IFREG)) or info.flag_bits & 1 or info.external_attr & 0x400:
                    raise AuditError('보관 ZIP의 링크/폴더/암호화 항목은 허용하지 않습니다.')
                if name.casefold() in folded or info.file_size > MAX_FILE_BYTES:
                    raise AuditError('보관 ZIP의 중복 파일/크기 오류')
                folded.add(name.casefold())
                names[name] = info
            if MANIFEST not in names or names[MANIFEST].file_size > MAX_MANIFEST_BYTES:
                raise AuditError('보관 manifest 없음/크기 오류')
            manifest = _json(archive.read(MANIFEST), MANIFEST)
            if not isinstance(manifest, dict) or manifest.get('schema') != SCHEMA or manifest.get('mode') not in ('draft', 'final'):
                raise AuditError('지원하지 않는 보관 manifest/모드')
            entries = manifest.get('files')
            if not isinstance(entries, list) or not entries:
                raise AuditError('보관 manifest 파일목록 오류')
            expected = {}
            payloads = {}
            for entry in entries:
                if not isinstance(entry, dict):
                    raise AuditError('보관 manifest 파일 항목 오류')
                name = _safe_name(entry.get('path'))
                if name in expected or name == MANIFEST or name not in names:
                    raise AuditError('보관 manifest의 중복/누락/고아 경로')
                if type(entry.get('bytes')) is not int or entry['bytes'] != names[name].file_size or not _HEX.fullmatch(str(entry.get('sha256', ''))):
                    raise AuditError('보관 파일 크기/해시 선언 오류')
                item = archive.read(names[name])
                if len(item) != entry['bytes'] or _hash(item) != entry['sha256']:
                    raise AuditError('보관 파일의 크기/해시 불일치: ' + name)
                expected[name] = entry
                payloads[name] = item
            if set(names) != set(expected) | {MANIFEST} or not CORE <= set(expected):
                raise AuditError('보관 ZIP에 누락 또는 허용되지 않은 고아 파일이 있습니다.')
            if archive.testzip() is not None:
                raise AuditError('보관 ZIP CRC 오류')
        snapshot = _json(payloads[SNAPSHOT], SNAPSHOT)
        scope = _json(payloads[SCOPE], SCOPE)
        head = _json(payloads[ANCHOR], ANCHOR)
        if not isinstance(snapshot, dict) or not isinstance(scope, dict) or not isinstance(head, dict):
            raise AuditError('보관 스냅샷/범위/헤드 구조 오류')
        for key in ('schema', 'mode', 'created_utc', 'actor', 'engagement_id', 'head', 'basis'):
            if scope.get(key) != manifest.get(key):
                raise AuditError('보관 manifest와 범위 기록 불일치: ' + key)
        if head != manifest.get('head') or snapshot.get('head') != head.get('hash') or snapshot.get('basis') != manifest.get('basis'):
            raise AuditError('보관 DB·헤드·스냅샷 버전 불일치')
        memory = _ArchivedWorkspace(payloads[DATABASE], payloads[ANCHOR])
        if memory.engagement['id'] != manifest.get('engagement_id') or memory.basis() != manifest.get('basis'):
            raise AuditError('보관 업무/자료 버전 불일치')
        sources = _source_records(memory.sources.values())
        source_paths = {item['file'] for item in sources.values()}
        if set(expected) != CORE | source_paths or scope.get('source_ids') != sorted(sources):
            raise AuditError('다른 업무/미등록·고아 원본 파일은 허용하지 않습니다.')
        for name in CORE:
            if expected[name].get('kind') != 'core' or set(expected[name]) != {'path', 'bytes', 'sha256', 'kind'}:
                raise AuditError('보관 핵심파일 구성 오류')
        for sid, source in sources.items():
            entry = expected[source['file']]
            if set(entry) != {'path', 'bytes', 'sha256', 'kind', 'source_id'} or entry.get('kind') != 'source' or entry.get('source_id') != sid:
                raise AuditError('보관 원본의 업무 등록 연결 오류')
            if entry['bytes'] != source['bytes'] or entry['sha256'] != source['sha256']:
                raise AuditError('DB에 등록한 원본과 보관 사본의 해시/크기 불일치')
        workflow = _workflow(memory)
        expected_scope = _scope_record(manifest['mode'], text(manifest.get('actor'), '보관 실행자'),
                                       text(manifest.get('created_utc'), '보관 시각'), manifest['engagement_id'],
                                       head, manifest['basis'], sources, workflow)
        if dumps(scope) != dumps(expected_scope):
            raise AuditError('보관 범위/초안·최종 표시 또는 미완료 기록 불일치')
        reconstructed = _final_gate(memory, workflow) if manifest['mode'] == 'final' else memory.snapshot(final=False)
        if dumps(reconstructed) != dumps(snapshot):
            raise AuditError('보관 DB 이력에서 재구성한 스냅샷과 조서 기록 불일치')
        _verify_workbook(payloads[WORKBOOK], snapshot)
        _path_without_links(selected)
        if _identity(selected.lstat()) != identity:
            raise AuditError('검증 중 보관 ZIP이 변경되었습니다.')
        return {'valid': True, 'path': str(selected), 'mode': manifest['mode'],
                'state': scope.get('state'), 'manifest': manifest, 'scope': scope,
                'head': head, 'basis': manifest['basis'], 'source_count': len(sources),
                'file_count': len(names), 'bytes': len(raw), 'zip_sha256': _hash(raw)}
    except AuditError:
        raise
    except Exception as exc:
        raise AuditError('보관 검증 실패: ' + str(exc)) from exc
