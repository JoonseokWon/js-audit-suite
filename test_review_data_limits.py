"""ACC-02 exact row boundaries; limit errors must identify the integrated module."""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest,threading
from openpyxl import Workbook
from common import AuditError
from adapters import run,MODULES
from data_controls.reader import read_module_xlsx
from input_limits import module_limits,module_limit_text,EXCEL_DATA_ROWS,MAX_SOURCE_BYTES
from engines.sales.data_io import IMPORT_FIELDS

CONTEXT={'회사코드':'REVIEW_ONLY_A','시작일':'2026-01-01','기준일':'2026-01-31'}


def sales_file(path,n):
    book=Workbook(write_only=True);sheets={k:book.create_sheet(k) for k in IMPORT_FIELDS}
    for name,headers in IMPORT_FIELDS.items():sheets[name].append(headers)
    for index in range(n):
        i=str(index+1);common=dict(order_id='O'+i,order_line_id='1',item_code='I',customer_id='C')
        rows={
          'orders':dict(common,order_date='2026-01-05',quantity=1,unit='EA',unit_price=100,price_basis='NET',tax_rate=0,currency='KRW'),
          'deliveries':dict(common,delivery_id='D'+i,delivery_line_id='1',delivery_date='2026-01-05',acceptance_date='2026-01-05',quantity=1,accepted_quantity=1,unit='EA'),
          'invoices':dict(common,invoice_id='I'+i,invoice_line_id='1',invoice_date='2026-01-05',quantity=1,unit='EA',unit_price=100,price_basis='NET',tax_rate=0,currency='KRW',supply_amount=100,tax_amount=0,total_amount=100),
          'ledger':dict(journal_id='J'+i,journal_line_id='1',company_id=CONTEXT['회사코드'],account_code='4000',posting_date='2026-01-05',invoice_id='I'+i,invoice_line_id='1',customer_id='C',currency='KRW',revenue_debit=0,revenue_credit=100)}
        for name,row in rows.items():sheets[name].append([row[h] for h in IMPORT_FIELDS[name]])
    ctl=dict(company_id=CONTEXT['회사코드'],period_start=CONTEXT['시작일'],period_end=CONTEXT['기준일'],account_codes='4000',expected_rows=n,debit_total=0,credit_total=100*n,source_reference='SYNTHETIC ONLY independent n x 100',extracted_at='2026-01-31T23:00:00+09:00',prepared_by='SYNTHETIC ONLY')
    sheets['control'].append([ctl[h] for h in IMPORT_FIELDS['control']]);book.save(path)


class ReviewLimitsTests(unittest.TestCase):
    def test_sales_999_and_1000_accept_all_rows_and_1001_rejects_without_truncation(self):
        with TemporaryDirectory() as folder:
            for n in (999,1000,1001):
                path=Path(folder)/f'{n}.xlsx';sales_file(path,n)
                if n<=1000:
                    r=run('매출',path,CONTEXT)
                    self.assertEqual(r['blocking']+r['review'],0);self.assertEqual(r['metrics']['순매출'],100*n)
                else:
                    with self.assertRaises(AuditError) as caught:run('매출',path,CONTEXT)
                    self.assertIn('매출 / orders',str(caught.exception));self.assertIn('1,001',str(caught.exception))
                    self.assertIn('1,000',str(caught.exception));self.assertIn('헤더 1행 제외',str(caught.exception))

    def test_all_four_small_module_read_gates_share_exact_999_1000_1001_bounds(self):
        with TemporaryDirectory() as folder:
            for n in (999,1000,1001):
                path=Path(folder)/f'gate{n}.xlsx';book=Workbook(write_only=True);sheet=book.create_sheet('자료')
                sheet.append(['ID']);[sheet.append([str(i)]) for i in range(n)];book.save(path)
                for module in ('매출','금융상품','감가상각','확장검사'):
                    with self.subTest(n=n,module=module):
                        if n<=1000:self.assertEqual(len(read_module_xlsx(module,path)[0]['자료']),n)
                        else:
                            with self.assertRaisesRegex(AuditError,module+' / 자료'):read_module_xlsx(module,path)

    def test_large_ledger_limit_is_kept_separate_from_other_worksheet_limits(self):
        limits=module_limits('원장·리드')
        self.assertEqual(limits['row_limits'],{'총계정원장':1048575})
        self.assertEqual(limits['max_source_bytes'],256*1024**2)
        with TemporaryDirectory() as folder:
            for title in ('총계정원장','시산표'):
                p=Path(folder)/(title+'.xlsx');b=Workbook(write_only=True);s=b.create_sheet(title);s.append(['ID'])
                for i in range(10001):s.append([str(i)])
                b.save(p)
                if title=='총계정원장':self.assertEqual(len(read_module_xlsx('원장·리드',p)[0][title]),10001)
                else:
                    with self.assertRaisesRegex(AuditError,'10,000'):read_module_xlsx('원장·리드',p)

    def test_each_integrated_module_has_format_specific_limit_text(self):
        for module in MODULES:self.assertIn(module,module_limit_text(module))
        self.assertEqual(module_limits('매출')['data_rows'],1000)
        self.assertEqual(module_limits('실무자료·원장')['raw_source_rows'],EXCEL_DATA_ROWS)
        self.assertEqual(module_limits('실무자료·원장')['raw_source_bytes'],MAX_SOURCE_BYTES)

    def test_cancelled_read_never_returns_a_partial_population(self):
        with TemporaryDirectory() as folder:
            path=Path(folder)/'cancel.xlsx';sales_file(path,999);event=threading.Event();event.set()
            with self.assertRaisesRegex(AuditError,'취소'):read_module_xlsx('매출',path,event=event)

    def test_empty_styled_and_hidden_rows_obey_declared_physical_position_limit(self):
        from openpyxl.styles import Font
        with TemporaryDirectory() as folder:
            for hidden in (False,True):
                path=Path(folder)/f'empty{hidden}.xlsx';b=Workbook();s=b.active;s.title='자료';s.append(['ID']);s.append(['one'])
                s.cell(1002,1).font=Font(bold=True);s.row_dimensions[1002].hidden=hidden;b.save(path);b.close()
                with self.assertRaisesRegex(AuditError,'Excel 행 1,002'):read_module_xlsx('매출',path)


if __name__=='__main__':unittest.main()
