"""Same v1.5 accounting observations in native source, frozen and CI runs."""
from pathlib import Path
import json,re
from common import AuditError,dumps,digest
from workspace import Workspace
import exporter

CONTEXT={'회사코드':'SYNTH01','회사명':'합성 분석 검증','시작일':'2026-01-01','기준일':'2026-12-31','회계체계':'K-IFRS / 합성','감사기준':'합성 절차 검증','작성자':'합성작성자','더미업무':True}
BASE_INPUTS=[('원장·리드','통합_v1.3_원장_리드.xlsx'),('금융상품','통합_v1.3_금융상품.xlsx'),('매출','통합_v1.3_매출.xlsx'),('감가상각','통합_v1.3_감가상각.xlsx'),('확장검사','통합_v1.3_확장검사.xlsx'),('리스','통합_v1.3_리스.xlsx'),('대손충당금','통합_v1.3_대손충당금.xlsx'),('재고평가','통합_v1.3_재고평가.xlsx'),('종업원급여','통합_v1.3_종업원급여.xlsx'),('미기록부채·비용이연','비용이연_01_정상.xlsx'),('원천증빙','통합_v1.4_원천증빙.auditpack'),('재무제표·주석','통합_v1.4_재무제표.json'),('GLA','통합_v1.5_GLA.xlsx'),('FIA','통합_v1.5_FIA.xlsx'),('현금분석','통합_v1.5_현금분석.xlsx')]

def exercise(root,output):
    root=Path(root);output=Path(output);output.mkdir(parents=True,exist_ok=True)
    checks=[];observations=[]
    def check(ok,label):
        if not ok:raise AssertionError(label)
        checks.append(label)
    def fresh(path):return Workspace(path,create=CONTEXT)
    manifest=json.loads((root/'samples/감사분석/cases.json').read_text(encoding='utf-8'))
    for case in manifest['cases']:
        folder=root/'samples/감사분석'/case['folder'];ws=fresh(output/'cases'/(case['module']+'_'+case['id']))
        current=folder/case['current_ledger'] if case.get('current_ledger') else root/'samples/통합/통합_v1.3_원장_리드.xlsx'
        gl=ws.execute('원장·리드',current);r=ws.execute(case['module'],folder/case['file'])['result']
        check(bool(r['blocking'])==case['blocking'],case['id']+'_blocking')
        check(bool(r['review'])==case['review'],case['id']+'_review')
        for code in case.get('required_codes',[]):check(code in [i['code'] for i in r['issues']],case['id']+'_'+code)
        if not r['blocking']:
            check(r['detail']['현재원장연결'] is True,case['id']+'_current_binding')
            check(all(x['상태']=='일치' for x in r['tables']['현재원장연결']),case['id']+'_exact_GL_TB')
        candidates=json.loads(re.sub(r'[0-9a-f]{32}\.xlsx','<registered-input.xlsx>',json.dumps(r['tables'].get('검토후보',[]),ensure_ascii=False)))
        observations.append({'module':case['module'],'id':case['id'],'blocking':r['blocking'],'review':r['review'],'codes':[i['code'] for i in r['issues']],'metrics':r['metrics'],'candidate_rows':candidates})
    ws=fresh(output/'normal_workspace')
    for module,name in BASE_INPUTS:
        r=ws.execute(module,root/'samples/통합'/name)['result'];check(r['blocking']==0,module+'_normal_input_not_blocked')
    check(len(ws.latest())==15,'fifteen_modules');check(len(ws.cross_checks())==35,'thirty_five_current_ledger_links')
    check(all(r['상태']=='일치' for r in ws.cross_checks()),'normal_current_bindings_equal')
    check(ws.latest()['GLA']['result']['metrics'].get('원장행수')==82,'GLA_82_input_lines')
    check(ws.latest()['GLA']['result']['metrics'].get('전표수')==34,'GLA_34_journals')
    candidates=ws.risk_register();check(bool(candidates),'normal_rule_candidates_visible')
    row=candidates[0];draft=ws.candidate_request_draft(row['후보키']);check(not ws.requests,'request_draft_not_automatically_saved')
    check(draft['requested_from']=='' and draft['state']=='요청' and draft['source_ids']==[row['자료ID']],'draft_preserves_source_and_needs_user_request_recipient')
    ws.save_request({**draft,'requested_from':'합성 담당자'})
    from review_support import save_context
    context_path=output/'검토맥락.json';save_context(ws,context_path)
    check(json.loads(context_path.read_text(encoding='utf-8'))['uses_ai_inference'] is False,'rule_engine_not_mislabeled_AI')
    try:save_context(ws,ws.anchor)
    except AuditError:check(True,'context_cannot_overwrite_history')
    else:raise AssertionError('history overwrite permitted')
    for pid,module in [('G01','GLA'),('I01','FIA'),('H01','현금분석')]:
        refs=ws.calculation_refs(pid);ws.save_procedure({'id':pid,'title':module+' 합성 검토 초안','preparer':'합성작성자','purpose':'선택 규칙과 원문 검토','risk':'후보·미확인 범위의 추가 확인','scope':dumps(refs['계산지표']),'evidence_ids':[refs['자료ID'],ws.latest()['원장·리드']['source_id']]})
    catalog=json.loads((root/'procedure_catalog.json').read_text(encoding='utf-8'))
    path=output/'통합감사_미검토초안_v1.5.xlsx';snapshot=exporter.export(ws,path,catalog)
    check(all(not p['performed'] and not p['conclusion'] and not p['reviewer'] for p in snapshot['procedures']),'human_fields_blank')
    from openpyxl import load_workbook
    w=load_workbook(path,data_only=False)
    try:
        check({'조서_G01','조서_I01','조서_H01','통합검토후보','분석필요자료'}<=set(w.sheetnames),'new_workpaper_and_context_tabs')
        check(all(any(c.hyperlink for row in w['조서_'+pid] for c in row) for pid in ('G01','I01','H01')),'all_analysis_workpapers_link_details')
        check(all(c.data_type not in ('f','e') for sh in w for row in sh for c in row),'no_formula_or_error_cells')
        proof=json.loads(''.join(row[1] for row in w['증적JSON'].iter_rows(min_row=2,values_only=True)))
        check(proof['risk_register']==candidates and proof['head']==ws.head,'export_embeds_actual_rules_and_proof')
    finally:w.close()
    try:ws.snapshot(True)
    except AuditError:check(True,'draft_cannot_be_completed_export')
    else:raise AssertionError('draft exported as complete')
    original_candidate=row['후보키'];ws.execute('원장·리드',root/'samples/통합/통합_v1.3_원장_리드.xlsx')
    check(all(r['연결상태']=='차단·재실행 필요' for r in ws.risk_register()),'rerun_invalidates_candidates')
    try:ws.candidate_request_draft(original_candidate)
    except AuditError:check(True,'stale_candidate_cannot_generate_request')
    else:raise AssertionError('stale request permitted')
    for pid in ('G01','I01','H01'):
        p=dict(ws.procedures[pid]);p.update(performed='SYNTHETIC ONLY / 재검토 통제검사',conclusion='합성 상태검사',performed_date='2026-12-31');ws.save_procedure(p)
        try:ws.request_review(pid)
        except AuditError:check(True,pid+'_stale_review_blocked_after_resave')
        else:raise AssertionError('stale review permitted')
    return {'ok':True,'case_count':len(observations),'checks':checks,'observations':observations,'normal_modules':15,'current_links':35,'candidate_count':len(candidates),'draft':str(path),'limitations':['Rule-based offline analytics, no model inference','Synthetic inputs and declared source populations only','Risk candidates require human investigation and are not fraud conclusions']}
