"""One coherent synthetic transaction cycle and deliberately altered cases.

This is a shared whole-ledger exercise, separate from the account workers'
independent calculation fixtures.  Values are fixed integers, not balanced by
plugging a generated difference to equity. All entities and evidence are fake.
"""
from copy import deepcopy
from common import digest

CONTEXT = {'회사코드': 'SYNTH01', '시작일': '2026-01-01', '기준일': '2026-12-31', '통화': 'KRW', '연결구분': '별도'}
CASES = ('normal', 'balanced_journal_missing', 'duplicate_allocation', 'partial_collection',
         'split_payment', 'wrong_company', 'wrong_period', 'blank_source_amount', 'wrong_sign', 'overallocated')


def coherent_case(name='normal'):
    if name not in CASES:
        raise ValueError('unknown synthetic case: ' + name)
    tables = []
    def table(dataset, sheet, headers, rows):
        out = {'dataset_id': dataset, 'source_file': 'sources/공통거래_합성자료.xlsx', 'sheet': sheet,
               'headers': headers, 'rows': [{'row': i + 2, 'values': row} for i, row in enumerate(rows)]}
        tables.append(out)
        return out
    sales = table('SYN-SALES', '매출증빙', ['회사코드', '송장ID', '일자', '금액'], [
        ['SYNTH01', 'SI1', '2026-06-01', 100], ['SYNTH01', 'SI2', '2026-06-02', 200], ['SYNTH01', 'CN1', '2026-06-03', -20]])
    ar = table('SYN-AR', '수금대상', ['회사코드', '채권ID', '일자', '금액'], [
        ['SYNTH01', 'AR1', '2026-06-01', 100], ['SYNTH01', 'AR2', '2026-06-03', 180]])
    purchases = table('SYN-PURCHASE', '매입증빙', ['회사코드', '송장ID', '일자', '금액'], [
        ['SYNTH01', 'PI1', '2026-06-04', 60], ['SYNTH01', 'PI2', '2026-06-05', 90]])
    bank = table('SYN-BANK', '은행거래', ['회사코드', '거래ID', '일자', '금액'], [
        ['SYNTH01', 'R1', '2026-06-10', 60], ['SYNTH01', 'R2', '2026-06-11', 40], ['SYNTH01', 'R3', '2026-06-12', 180],
        ['SYNTH01', 'D1', '2026-06-15', -100], ['SYNTH01', 'D2', '2026-06-16', -50]])
    # Explicit full journal, ten balanced postings. Cash 130 + stock 150 = sales 280.
    journal = [
        ('J-S1', '2026-06-01', '1300', 100, 0), ('J-S1', '2026-06-01', '4000', 0, 100),
        ('J-S2', '2026-06-02', '1300', 200, 0), ('J-S2', '2026-06-02', '4000', 0, 200),
        ('J-CN1', '2026-06-03', '4000', 20, 0), ('J-CN1', '2026-06-03', '1300', 0, 20),
        ('J-R1', '2026-06-10', '1000', 60, 0), ('J-R1', '2026-06-10', '1300', 0, 60),
        ('J-R2', '2026-06-11', '1000', 40, 0), ('J-R2', '2026-06-11', '1300', 0, 40),
        ('J-R3', '2026-06-12', '1000', 180, 0), ('J-R3', '2026-06-12', '1300', 0, 180),
        ('J-P1', '2026-06-04', '1400', 60, 0), ('J-P1', '2026-06-04', '2400', 0, 60),
        ('J-P2', '2026-06-05', '1400', 90, 0), ('J-P2', '2026-06-05', '2400', 0, 90),
        ('J-D1', '2026-06-15', '2400', 100, 0), ('J-D1', '2026-06-15', '1000', 0, 100),
        ('J-D2', '2026-06-16', '2400', 50, 0), ('J-D2', '2026-06-16', '1000', 0, 50),
    ]
    counts = {}
    gl_rows = []
    for jid, dt, account, debit, credit in journal:
        counts[jid] = counts.get(jid, 0) + 1
        gl_rows.append(['SYNTH01', jid, str(counts[jid]), dt, account, debit, credit, 'SYNTHETIC ONLY / 공통거래', '합성작성자'])
    gl = table('SYN-GL', '총계정원장', ['회사코드', '전표ID', '행ID', '전기일', '계정코드', '차변', '대변', '적요', '작성자'], gl_rows)
    tb_rows = [
        ['SYNTH01', '1000', '현금', 0, 280, 150, 130],
        ['SYNTH01', '1300', '매출채권', 0, 300, 300, 0],
        ['SYNTH01', '1400', '재고자산', 0, 150, 0, 150],
        ['SYNTH01', '2400', '매입채무', 0, 150, 150, 0],
        ['SYNTH01', '4000', '매출', 0, 20, 300, -280],
    ]
    tb = table('SYN-TB', '시산표', ['회사코드', '계정코드', '계정명', '기초잔액', '당기차변', '당기대변', '기말잔액'], tb_rows)
    control_rows = [
        ['SYNTH01', '2026-01-01', '2026-12-31', '1000', 5, 280, 150, 0, 130, 'SYNTHETIC ONLY / 별도 고정 기대값', '2027-01-02T00:00:00+09:00', '합성작성자'],
        ['SYNTH01', '2026-01-01', '2026-12-31', '1300', 6, 300, 300, 0, 0, 'SYNTHETIC ONLY / 별도 고정 기대값', '2027-01-02T00:00:00+09:00', '합성작성자'],
        ['SYNTH01', '2026-01-01', '2026-12-31', '1400', 2, 150, 0, 0, 150, 'SYNTHETIC ONLY / 별도 고정 기대값', '2027-01-02T00:00:00+09:00', '합성작성자'],
        ['SYNTH01', '2026-01-01', '2026-12-31', '2400', 4, 150, 150, 0, 0, 'SYNTHETIC ONLY / 별도 고정 기대값', '2027-01-02T00:00:00+09:00', '합성작성자'],
        ['SYNTH01', '2026-01-01', '2026-12-31', '4000', 3, 20, 300, 0, -280, 'SYNTHETIC ONLY / 별도 고정 기대값', '2027-01-02T00:00:00+09:00', '합성작성자'],
    ]
    control = table('SYN-CTRL', '독립통제', ['회사코드', '시작일', '기준일', '계정코드', '원장행수', '차변합계', '대변합계', '기초잔액', '기말잔액', '추출근거', '추출일시', '작성자'], control_rows)
    catalog = {'schema_version': 'raw-catalog-1', 'context': deepcopy(CONTEXT), 'tables': tables}
    populations, allocations = [], []
    def pop(pid, source, side, rows, formula):
        populations.append({'population_id': pid, 'company_code': 'SYNTH01', 'dataset_id': source['dataset_id'],
                            'rows': rows, 'amount_columns': formula, 'side': side, 'label': source['sheet'] + ' 명시 선택'})
    def alloc(aid, source, sr, target, tr, amount, tid):
        allocations.append({'allocation_id': aid, 'company_code': 'SYNTH01', 'transaction_id': tid,
                            'source': {'population_id': source, 'row': sr}, 'target': {'population_id': target, 'row': tr}, 'amount': amount})
    pop('SALES-S', sales, 'source', [2, 3, 4], {'금액': 1})
    pop('SALES-T', gl, 'target', [3, 5, 6], {'대변': 1, '차변': -1})
    alloc('A-S1', 'SALES-S', 2, 'SALES-T', 3, 100, 'SI1')
    alloc('A-S2', 'SALES-S', 3, 'SALES-T', 5, 200, 'SI2')
    alloc('A-CN1', 'SALES-S', 4, 'SALES-T', 6, -20, 'CN1')
    pop('AR-S', ar, 'source', [2, 3], {'금액': 1})
    pop('AR-T', bank, 'target', [2, 3, 4], {'금액': 1})
    alloc('A-AR1a', 'AR-S', 2, 'AR-T', 2, 60, 'AR1')
    alloc('A-AR1b', 'AR-S', 2, 'AR-T', 3, 40, 'AR1')
    alloc('A-AR2', 'AR-S', 3, 'AR-T', 4, 180, 'AR2')
    pop('AP-S', purchases, 'source', [2, 3], {'금액': -1})
    pop('AP-T', bank, 'target', [5, 6], {'금액': 1})
    # One payment settles parts of two invoices; PI2 is settled by two payments.
    alloc('A-PI1', 'AP-S', 2, 'AP-T', 5, -60, 'D1')
    alloc('A-PI2a', 'AP-S', 3, 'AP-T', 5, -40, 'D1')
    alloc('A-PI2b', 'AP-S', 3, 'AP-T', 6, -50, 'D2')
    pop('BANK-S', bank, 'source', [2, 3, 4, 5, 6], {'금액': 1})
    pop('BANK-T', gl, 'target', [8, 10, 12, 19, 21], {'차변': 1, '대변': -1})
    for i, (bn, gn, amount, tid) in enumerate([(2, 8, 60, 'R1'), (3, 10, 40, 'R2'), (4, 12, 180, 'R3'), (5, 19, -100, 'D1'), (6, 21, -50, 'D2')]):
        alloc('A-B' + str(i), 'BANK-S', bn, 'BANK-T', gn, amount, tid)
    links = {'schema_version': 'raw-links-1', 'context': deepcopy(CONTEXT), 'populations': populations,
             'allocations': allocations, 'unconfirmed': [{'범위': '선택 외 GL행·거래 실질·전체 감사 결론', '상태': '미검토', '사유': '전체 원장 산술통제는 ledger 모듈에서 별도 실행하며 외부증거와 판단은 필요합니다.'}]}
    if name == 'balanced_journal_missing':
        gl['rows'] = [r for r in gl['rows'] if r['values'][1] != 'J-R3']
    elif name == 'duplicate_allocation':
        clone = deepcopy(allocations[0]); clone['allocation_id'] = 'DUPLICATE-PAIR'; allocations.append(clone)
    elif name == 'partial_collection':
        allocations[:] = [r for r in allocations if r['allocation_id'] != 'A-AR1b']
    elif name == 'wrong_company':
        sales['rows'][0]['values'][0] = 'SYNTH02'
    elif name == 'wrong_period':
        sales['rows'][0]['values'][2] = '2027-01-01'
    elif name == 'blank_source_amount':
        sales['rows'][0]['values'][3] = None
    elif name == 'wrong_sign':
        allocations[0]['amount'] = -100
    elif name == 'overallocated':
        allocations[0]['amount'] = 101
    # native ledger input, no difference to equity or auto-approved adjustment.
    def native(table):
        return [dict(zip(table['headers'], r['values'])) | {'_row': r['row'], '_sheet': table['sheet']} for r in table['rows']]
    native_ledger = {
        '설정': [{'항목': k, '값': v, '_row': i + 2, '_sheet': '설정'} for i, (k, v) in enumerate([
            ('회사코드', 'SYNTH01'), ('시작일', '2026-01-01'), ('기준일', '2026-12-31'), ('통화', 'KRW'), ('추출범위', '1000;1300;1400;2400;4000')])],
        '시산표': native(tb), '총계정원장': native(gl), '독립통제': native(control),
        '계정매핑': [{'계정코드': a, '리드코드': a, '리드명': label, '표시구분': 'PL' if a == '4000' else 'BS', '_row': i + 2, '_sheet': '계정매핑'} for i, (a, label) in enumerate([('1000', '현금'), ('1300', '매출채권'), ('1400', '재고자산'), ('2400', '매입채무'), ('4000', '매출')])],
        '감사조정': [], '모듈연결': [],
    }
    return {'catalog': catalog, 'links': links, 'manifest': {'schema_version': 'synthetic-in-memory-1'},
            'source': {'filename': '공통거래_' + name + '.auditpack', 'bytes': 0, 'sha256': digest({'catalog': catalog, 'links': links}),
                       '자료성격': 'SYNTHETIC ONLY / 같은 생성기·독립증거 아님'}, 'native_ledger': native_ledger}


def write_coherent_case(target, name='normal'):
    """Create importable original XLSX + auditpack + complete ledger XLSX.

    The target must be a new directory. Missing journal rows keep their physical
    Excel row numbers, so removal cannot silently shift the allocation pointers.
    """
    import json
    from pathlib import Path
    from openpyxl import Workbook
    from raw_sources import build_pack, load_pack
    from raw_accounts import analyze
    import ledger
    from common import dumps, atomic
    root = Path(target)
    root.mkdir(parents=True, exist_ok=False)
    sample = coherent_case(name)
    source_root = root / '원천'
    folder = '01_현재프로그램입력'
    book_name = '01_공통거래_합성자료'
    rel_file = folder + '/' + book_name + '.xlsx'
    path = source_root / '배포' / folder / (book_name + '.xlsx')
    path.parent.mkdir(parents=True)
    book = Workbook()
    book.remove(book.active)
    datasets, headers, remap = [], {}, {}
    for i, table in enumerate(sample['catalog']['tables'], 1):
        dataset = f'D{i:03}'
        remap[table['dataset_id']] = dataset
        datasets.append({'id': dataset, 'file': rel_file, 'name': table['sheet']})
        headers[table['sheet']] = table['headers']
        sheet = book.create_sheet(table['sheet'])
        sheet.append(table['headers'])
        for record in table['rows']:
            for col, value in enumerate(record['values'], 1):
                sheet.cell(record['row'], col, value)
        sheet.freeze_panes = 'A2'
    book.save(path)
    spec = {'context': deepcopy(CONTEXT), 'books': [{'name': book_name, 'folder': folder, 'headers': headers}], 'datasets': datasets,
            '자료성격': 'SYNTHETIC ONLY / 동일 생성기의 거래·원장·통제값. 실제 기업·금융기관·회계증빙 아님.'}
    atomic(source_root / 'spec.json', dumps(spec).encode('utf-8'))
    links = deepcopy(sample['links'])
    for pop in links['populations']:
        pop['dataset_id'] = remap[pop['dataset_id']]
    package_path = build_pack(root / ('공통거래_' + name + '.auditpack'), source_root, links=links)
    native_path = root / '원장_공통거래.xlsx'
    book = Workbook()
    book.remove(book.active)
    for sheet_name, columns in ledger.SCHEMA.items():
        sheet = book.create_sheet(sheet_name)
        sheet.append(columns)
        for row in sample['native_ledger'][sheet_name]:
            for col, key in enumerate(columns, 1):
                sheet.cell(row['_row'], col, row.get(key))
        sheet.freeze_panes = 'A2'
    book.save(native_path)
    loaded = load_pack(package_path, CONTEXT)
    raw_result = analyze(loaded, CONTEXT)
    native_result = ledger.analyze(sample['native_ledger'])
    outcome = {'case': name, 'source_ledger': '원장_공통거래.xlsx', 'auditpack': package_path.name,
               'raw_status': raw_result['status'], 'raw_codes': [i['code'] for i in raw_result['issues']],
               'raw_metrics': raw_result['metrics'], 'ledger_blocking': native_result['blocking'],
               'ledger_codes': [i['code'] for i in native_result['issues']],
               'fixed_normal_oracle': {'현금': 130, '재고자산': 150, '매출': 280, 'GL행수': 20, '전표수': 10},
               'review_status': '작성·검토·결론 미확정'}
    atomic(root / '검증결과.json', json.dumps(outcome, ensure_ascii=False, indent=2).encode('utf-8'))
    return {'root': root, 'auditpack': package_path, 'ledger': native_path, 'outcome': outcome}
