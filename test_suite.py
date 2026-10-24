import runtime
ROOT=runtime.configure()
import unittest,tempfile,sqlite3,threading,copy,io,zipfile,re,shutil
from pathlib import Path
from common import *
import ledger,extensions,adapters,exporter
from workspace import Workspace
from verify import check_book,regression,synthetic_procedure
from tools.sample_specs import ledger_base,ext_base
CATALOG=json.loads((ROOT/'procedure_catalog.json').read_text(encoding='utf-8'))
SPEC=json.loads((ROOT/'samples_specs.json').read_text(encoding='utf-8'))
def tables(schema,matrix):return {k:[dict(zip(schema[k],r))|{'_row':i+2,'_sheet':k} for i,r in enumerate(rows)] for k,rows in matrix.items()}
class CoreTests(unittest.TestCase):
    def test_legacy_55_expected(self):self.assertEqual(len(regression(ROOT)),55)
    def test_je_identity_prevents_date_netting(self):
        t=tables(ledger.SCHEMA,ledger_base());t['총계정원장'][0]['전표ID']='OTHER'
        r=ledger.analyze(t);self.assertEqual(sum(i['code']=='JE_UNBALANCED' for i in r['issues']),2)
    def test_extraction_scope_detects_both_sides_missing(self):
        t=tables(ledger.SCHEMA,ledger_base())
        for k in ('시산표','계정매핑','독립통제'):t[k]=[r for r in t[k] if r['계정코드'] not in ('1100','3000')]
        self.assertIn('EXTRACTION_SCOPE',[i['code'] for i in ledger.analyze(t)['issues']])
    def test_gl_debit_negative_rejected(self):
        t=tables(ledger.SCHEMA,ledger_base());t['총계정원장'][0]['차변']=-1
        with self.assertRaises(AuditError):ledger.analyze(t)
    def test_control_timezone_required(self):
        t=tables(ledger.SCHEMA,ledger_base());t['독립통제'][0]['추출일시']='2027-01-05T09:00:00'
        self.assertIn('CONTROL_TIME',[i['code'] for i in ledger.analyze(t)['issues']])
    def test_adjustment_duplicate_not_hidden(self):
        t=tables(ledger.SCHEMA,ledger_base());r={'조정ID':'1','행ID':'1','계정코드':'1500','차변':1,'대변':0,'근거':'test','상태':'승인','_row':2};t['감사조정']=[r,r.copy()]
        self.assertIn('AJE_DUP',[i['code'] for i in ledger.analyze(t)['issues']])
    def test_bank_equal_net_missing_gross(self):
        t=tables(extensions.SCHEMA,ext_base());t['은행장부'][0]['금액']=800;t['은행장부'][1]['금액']=0
        r=extensions.analyze(t);self.assertIn('EXT_CONTROL',[i['code'] for i in r['issues']])
    def test_inventory_unit_mismatch(self):
        t=tables(extensions.SCHEMA,ext_base());t['실사'][0]['단위']='BOX'
        r=extensions.analyze(t);self.assertEqual(r['tables']['재고실사대사'][0]['상태'],'차이/누락')
    def test_recovery_not_capped_silently(self):
        t=tables(extensions.SCHEMA,ext_base());t['후속회수'][0]['금액']=120000
        r=extensions.analyze(t);self.assertIn('AR_OVERRECOVERY',[i['code'] for i in r['issues']]);self.assertEqual(r['tables']['채권연령후속회수'][0]['잔여'],-10000)
    def test_sources_are_not_mutated(self):
        t=tables(ledger.SCHEMA,ledger_base());original=copy.deepcopy(t);ledger.analyze(t);self.assertEqual(t,original)
    def test_lead_group_rollup(self):
        t=tables(ledger.SCHEMA,ledger_base())
        for r in t['계정매핑']:
            if r['계정코드'] in ('1100','1200'):r.update(리드코드='FIN',리드명='금융자산',표시구분='BS')
        r=ledger.analyze(t);s=next(s for s in r['tables']['리드집계'] if s['리드코드']=='FIN');self.assertEqual((s['계정수'],s['기초잔액'],s['조정전']),(2,1000000,1036400))
    def test_lead_group_cannot_mix_bs_pl(self):
        t=tables(ledger.SCHEMA,ledger_base())
        for r in t['계정매핑']:
            if r['계정코드'] in ('1100','4000'):r.update(리드코드='MIX',리드명='같은 리드')
        self.assertIn('LEAD_CLASS_CONFLICT',[i['code'] for i in ledger.analyze(t)['issues']])

class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(dir=ROOT/'검증');self.addCleanup(self.tmp.cleanup);self.folder=Path(self.tmp.name)/'업무'
        self.ws=Workspace(self.folder,create={'회사코드':'SYNTH01','회사명':'합성','시작일':'2026-01-01','기준일':'2026-12-31','회계체계':'K-IFRS','감사기준':'test','작성자':'합성작성자','더미업무':True})
        self.good=ROOT/'samples/통합/통합_01_원장정상.xlsx';self.asset=ROOT/'samples/통합/통합업무_감가상각정상.xlsx'
    def seed(self):
        self.ws.execute('원장·리드',self.good)
        for m,f in [('금융상품','통합업무_금융정상.xlsx'),('매출','통합업무_매출정상.xlsx'),('감가상각','통합업무_감가상각정상.xlsx')]:self.ws.execute(m,ROOT/'samples/통합'/f)
    def ready(self):
        self.seed()
        for pid in self.ws.required_procedures():
            self.ws.save_procedure(synthetic_procedure(pid,list(self.ws.sources)));self.ws.request_review(pid);self.ws.review(pid,'검토자','합성 근거 범위 검토')
    def test_cross_checks_hand_amounts(self):
        self.seed();rows=self.ws.cross_checks();self.assertEqual(len(rows),6);self.assertTrue(all(r['상태']=='일치' for r in rows))
        self.assertEqual(next(r['실제'] for r in rows if r['지표']=='당기상각'),240000)
    def test_new_source_stales_reviews(self):
        self.ready();self.ws.snapshot(True);self.ws.add_source(self.good,'추가 근거')
        self.assertTrue(all(self.ws.procedure_state(p)=='자료변경·재검토' for p in self.ws.procedures.values()))
        with self.assertRaises(AuditError):self.ws.snapshot(True)
    def test_missing_mandatory_procedure_blocks(self):
        self.seed();self.ws.save_procedure(synthetic_procedure('C01',list(self.ws.sources)));self.ws.request_review('C01');self.ws.review('C01','검토자','test')
        with self.assertRaises(AuditError):self.ws.snapshot(True)
    def test_no_evidence_review_request(self):
        self.ws.save_procedure(synthetic_procedure('A12',[]))
        with self.assertRaises(AuditError):self.ws.request_review('A12')
    def test_reviewer_separation(self):
        self.seed();self.ws.save_procedure(synthetic_procedure('A12',list(self.ws.sources)));self.ws.request_review('A12')
        with self.assertRaises(AuditError):self.ws.review('A12','합성작성자','test')
    def test_return_requires_edit_resubmit(self):
        self.seed();p=synthetic_procedure('A12',list(self.ws.sources));self.ws.save_procedure(p);self.ws.request_review('A12');self.ws.review('A12','검토자','보완 필요',False)
        with self.assertRaises(AuditError):self.ws.review('A12','검토자','바로 완료')
        self.ws.save_procedure(p);self.ws.request_review('A12');self.ws.review('A12','검토자','보완후 검토');self.assertEqual(self.ws.procedures['A12']['state'],'검토기록완료')
    def test_nonexistent_evidence_rejected(self):
        with self.assertRaises(AuditError):self.ws.save_procedure(synthetic_procedure('A12',['ghost']))
    def test_incomplete_fields_block_request(self):
        sid=self.ws.add_source(self.good,'자료');p=synthetic_procedure('A12',[sid]);p['performed']='';self.ws.save_procedure(p)
        with self.assertRaises(AuditError):self.ws.request_review('A12')
    def test_edit_clears_review(self):
        self.ready();p=dict(self.ws.procedures['A12']);p['conclusion']='개정';self.ws.save_procedure(p);self.assertEqual(self.ws.procedures['A12']['state'],'작성중');self.assertEqual(self.ws.procedures['A12']['reviewer'],'')
    def test_note_resolution_history(self):
        self.ready();nid=self.ws.add_note('의견','보완','검토자')
        with self.assertRaises(AuditError):self.ws.snapshot(True)
        self.ws.add_note('의견','해결근거','작성자',True,nid);self.ws.snapshot(True);self.assertEqual(sum(e['kind']=='note' for e in self.ws.events),2)
    def test_source_tamper_blocks(self):
        sid=self.ws.add_source(self.good,'자료');p=self.folder/self.ws.sources[sid]['file'];p.write_bytes(b'changed')
        with self.assertRaises(AuditError):self.ws.snapshot()
    def test_immutable_event_trigger(self):
        db=sqlite3.connect(self.ws.dbpath)
        try:
            with self.assertRaises(sqlite3.IntegrityError):db.execute('DELETE FROM events')
        finally:db.close()
    def test_schema_tamper_detected(self):
        db=sqlite3.connect(self.ws.dbpath);db.execute('DROP TRIGGER no_event_update');db.commit();db.close()
        with self.assertRaises(AuditError):Workspace(self.folder)
    def test_anchor_tamper_detected(self):
        self.ws.anchor.write_text('{"seq":0,"hash":""}',encoding='utf-8')
        with self.assertRaises(AuditError):Workspace(self.folder)
    def test_restart_and_relocation(self):
        self.seed();other=Path(self.tmp.name)/'이동';shutil.copytree(self.folder,other);w=Workspace(other);self.assertEqual(len(w.runs),4);self.assertEqual(w.verify_sources(),[])
    def test_failed_latest_does_not_reuse_previous(self):
        self.ws.execute('원장·리드',self.good);self.ws.execute('원장·리드',ROOT/'samples/통합/통합_03_전표한쪽누락.xlsx');self.assertGreater(self.ws.latest()['원장·리드']['result']['blocking'],0);self.assertEqual(len(self.ws.runs),2)
    def test_cancellation_saved_as_blocked(self):
        event=threading.Event();event.set();r=self.ws.execute('원장·리드',self.good,event);self.assertTrue(r['cancelled']);self.assertGreater(r['result']['blocking'],0)
    def test_no_workspace_overwrite(self):
        with self.assertRaises(AuditError):exporter.export(self.ws,self.ws.dbpath,CATALOG)
    def test_original_xlsx_overwrite_blocked(self):
        sid=self.ws.add_source(self.good,'자료')
        with self.assertRaises(AuditError):exporter.export(self.ws,self.folder/self.ws.sources[sid]['file'],CATALOG)
    def test_proof_complete_numeric_and_safe(self):
        self.ready();p=dict(self.ws.procedures['A12']);p['conclusion']='=HYPERLINK("x")';self.ws.save_procedure(p);snap=self.ws.snapshot();data=exporter.workbook_bytes(snap,CATALOG)
        from openpyxl import load_workbook
        w=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        try:
            proof=json.loads(''.join(r[1] for r in w['증적JSON'].iter_rows(min_row=2,values_only=True)));self.assertEqual(proof['all_runs'],snap['all_runs'])
            cells=[c for s in w for row in s for c in row];self.assertTrue(all(c.data_type!='f' for c in cells));self.assertTrue(any(c.value==240000 and c.data_type=='n' for c in cells))
            lead=next(s for s in w if s.title.endswith('리드스케줄'));self.assertEqual(lead.max_row,11)
        finally:w.close()
    def test_amount_diff_requires_explicit_review(self):
        self.ready();run=self.ws.latest()['금융상품'];changed=copy.deepcopy(run['result']);changed['metrics']['원금']=999999
        data={**run,'id':'difference-run','result':changed,'result_hash':digest(changed)};self.ws._append('run',data);self.ws.refresh()
        self.assertIn('차이',[r['상태'] for r in self.ws.cross_checks()]);self.assertTrue(any(x.startswith('B-') for x in self.ws.required_procedures()))
        with self.assertRaises(AuditError):self.ws.snapshot(True)
        for pid in self.ws.required_procedures():
            self.ws.save_procedure(synthetic_procedure(pid,list(self.ws.sources)));self.ws.request_review(pid);self.ws.review(pid,'검토자','차이와 미수정 영향에 대한 합성 검토')
        snap=self.ws.snapshot(True);self.assertEqual(next(x['차이'] for x in snap['cross'] if x['지표']=='원금'),-1)

class InputTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(dir=ROOT/'검증');self.addCleanup(self.tmp.cleanup);self.p=Path(self.tmp.name)/'mutated.xlsx';self.original=ROOT/'samples/통합/통합_01_원장정상.xlsx'
    def change(self,fn):
        with zipfile.ZipFile(self.original) as z,zipfile.ZipFile(self.p,'w') as out:
            for i in z.infolist():out.writestr(i,fn(i.filename,z.read(i)))
    def test_small_dimension_does_not_omit_rows(self):
        def fn(n,b):return re.sub(rb'(<x:worksheet[^>]*>)',rb'\1<x:dimension ref="A1:A1"/>',b) if n.startswith('xl/worksheets/') else b
        self.change(fn);a,_=read_xlsx(self.original,ledger.SCHEMA);b,_=read_xlsx(self.p,ledger.SCHEMA);self.assertEqual(a,b)
    def test_formula_rejected(self):
        def fn(n,b):
            if n!='xl/worksheets/sheet2.xml':return b
            changed=re.sub(rb'(<x:c r="G2"[^>]*>)',rb'\1<x:f>1+1</x:f>',b,count=1);assert changed!=b;return changed
        self.change(fn)
        with self.assertRaises(AuditError):read_xlsx(self.p,ledger.SCHEMA)
    def test_external_link_rejected(self):
        self.change(lambda n,b:b)
        with zipfile.ZipFile(self.p,'a') as z:z.writestr('xl/externalLinks/externalLink1.xml','<a/>')
        with self.assertRaises(AuditError):read_xlsx(self.p)

class FixtureTests(unittest.TestCase):pass
for index,b in enumerate(SPEC['books']):
    if b['expected'].get('template'):continue
    def test(self,book=b):check_book(book,ROOT/'samples/통합'/(book['name']+'.xlsx'),SPEC['context'])
    setattr(FixtureTests,'test_fixture_'+str(index+1).zfill(2)+'_'+b['name'],test)
if __name__=='__main__':unittest.main(verbosity=2)
