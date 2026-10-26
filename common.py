"""Shared strict input and immutable, formula-safe result serialization."""
from pathlib import Path
from datetime import date,datetime,timezone
from decimal import Decimal,InvalidOperation
from io import BytesIO
from zipfile import ZipFile,BadZipFile
import hashlib,json,re,os,tempfile

VERSION='2.2.0'
PROCEDURE_MODULES={'A02':'금융상품','A04':'매출','A12':'감가상각','A15':'리스','A07':'대손충당금','A11':'재고평가','A16':'종업원급여','A08':'미기록부채·비용이연','A09':'미기록부채·비용이연','R01':'원천증빙','FS01':'재무제표·주석','W01':'계정별증빙연습'}
PROCEDURE_MODULES.update({'G01':'GLA','I01':'FIA','H01':'현금분석'})
from onboarding import MODULES as ONBOARDING_MODULES, PROCEDURES as ONBOARDING_PROCEDURES
PROCEDURE_MODULES.update(ONBOARDING_PROCEDURES)
class AuditError(ValueError):pass
def now():return datetime.now(timezone.utc).isoformat(timespec='seconds')
def plain(v):
    if isinstance(v,Decimal):return format(v,'f')
    if isinstance(v,(date,datetime)):return v.isoformat()
    if isinstance(v,Path):return str(v)
    if isinstance(v,dict):return {str(k):plain(x) for k,x in v.items()}
    if isinstance(v,(tuple,list)):return [plain(x) for x in v]
    return v
def dumps(v):return json.dumps(plain(v),ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False)
def digest(v):return hashlib.sha256(dumps(v).encode()).hexdigest()
def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def text(v,label='값'):
    if not isinstance(v,str) or not v.strip() or len(v)>2000:raise AuditError(label+': 텍스트 필수')
    return v.strip()
def day(v):
    if isinstance(v,datetime):v=v.date()
    if isinstance(v,date):return v
    try:d=date.fromisoformat(str(v).strip())
    except ValueError as e:raise AuditError('날짜는 YYYY-MM-DD: '+str(v)) from e
    if not 1900<=d.year<=2100:raise AuditError('날짜 범위: 1900~2100')
    return d
def num(v,signed=True):
    if isinstance(v,bool) or v is None:raise AuditError('원 단위 숫자 필수')
    if not re.fullmatch(r'[+-]?\d+(?:\.\d+)?',str(v).strip()):raise AuditError('숫자 형식 오류: '+str(v))
    try:d=Decimal(str(v))
    except InvalidOperation as e:raise AuditError('숫자 오류') from e
    if not d.is_finite() or d!=d.to_integral_value() or abs(d)>10**13 or (not signed and d<0):raise AuditError('원 단위 정수/범위 오류')
    return int(d)
def cancelled(event):
    if event and event.is_set():raise AuditError('사용자가 취소했습니다. 결과는 저장하지 않았습니다.')
def atomic(path,data):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(dir=path.parent,prefix='.audit-',suffix='.tmp')
    try:
        with os.fdopen(fd,'wb') as f:f.write(data);f.flush();os.fsync(f.fileno())
        os.replace(tmp,path)
    finally:
        if Path(tmp).exists():Path(tmp).unlink()
def read_xlsx(path,schema=None,event=None,max_rows=10000):
    from openpyxl import load_workbook
    from defusedxml.ElementTree import iterparse
    p=Path(path)
    if p.suffix.lower()!='.xlsx' or p.stat().st_size>20*1024**2:raise AuditError('20MB 이하 .xlsx만 지원합니다.')
    data=p.read_bytes()
    try:
        with ZipFile(BytesIO(data)) as z:
            if len(z.infolist())>500 or sum(i.file_size for i in z.infolist())>100*1024**2:raise AuditError('Excel 압축해제 한도 초과')
            if any(any(x in i.filename.lower() for x in ('externallinks/','vba','embeddings/','connections.xml')) for i in z.infolist()):raise AuditError('외부연결·매크로·첨부 Excel은 지원하지 않습니다.')
            for item in z.infolist():
                if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',item.filename):
                    last=0
                    for _,el in iterparse(BytesIO(z.read(item)),events=('end',)):
                        if el.tag.rsplit('}',1)[-1]=='row':
                            n=int(el.get('r','0'))
                            if n<=last or n>max_rows+1:raise AuditError('물리 행 순서/한도 오류')
                            last=n;col=0
                            for c in el:
                                a=c.get('r','');m=re.fullmatch(r'([A-Z]+)(\d+)',a)
                                if not m or int(m[2])!=n:raise AuditError('셀 좌표 오류')
                                cc=0
                                for ch in m[1]:cc=cc*26+ord(ch)-64
                                if cc<=col or cc>80:raise AuditError('중복·역순·범위밖 셀')
                                col=cc
                            el.clear()
        w=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            if schema is not None and set(w.sheetnames)!=set(schema):raise AuditError('시트 목록이 입력양식과 다릅니다: '+', '.join(schema))
            tables={}
            for ws in w:
                ws.reset_dimensions();out=[];headers=None
                for n,cells in enumerate(ws.iter_rows(),1):
                    cancelled(event)
                    if n>max_rows+1 or len(cells)>80:raise AuditError('입력 한도 초과')
                    vals=[]
                    for c in cells:
                        if c.data_type in ('f','e'):raise AuditError(f'{ws.title}!{n}: 수식·오류 셀은 확정값으로 바꾸세요.')
                        v=c.value
                        if isinstance(v,datetime):v=v.date().isoformat() if v.time()==datetime.min.time() else v.isoformat()
                        elif isinstance(v,date):v=v.isoformat()
                        if isinstance(v,str) and (len(v)>2000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',v)):raise AuditError('입력 텍스트 오류')
                        vals.append(v)
                    while vals and vals[-1] is None:vals.pop()
                    if n==1:
                        headers=vals
                        if not headers or len(set(headers))!=len(headers):raise AuditError('중복/빈 헤더')
                        if schema is not None and tuple(headers)!=tuple(schema[ws.title]):raise AuditError(ws.title+': 열 이름/순서 오류')
                    elif vals:
                        if len(vals)>len(headers):raise AuditError('양식 밖 열')
                        out.append(dict(zip(headers,vals+[None]*(len(headers)-len(vals))))|{'_row':n,'_sheet':ws.title})
                if headers is None:raise AuditError('빈 시트')
                tables[ws.title]=out
        finally:w.close()
    except (BadZipFile,KeyError) as e:raise AuditError('Excel 구조 오류') from e
    if hashlib.sha256(p.read_bytes()).hexdigest()!=hashlib.sha256(data).hexdigest():raise AuditError('입력 파일이 변경되었습니다.')
    return tables,{'filename':p.name,'sha256':hashlib.sha256(data).hexdigest(),'rows':{k:len(v) for k,v in tables.items()}}

def issue(issues,code,message,ref='',blocking=True):issues.append({'code':code,'severity':'차단' if blocking else '검토','message':message,'reference':ref})
def result(module,context,tables,issues,metrics=None,detail=None):
    return plain({'module':module,'version':VERSION,'context':context,'tables':tables,'issues':issues,'metrics':metrics or {},'detail':detail or {},'blocking':sum(i['severity']=='차단' for i in issues),'review':sum(i['severity']=='검토' for i in issues)})
