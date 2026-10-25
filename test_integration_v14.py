from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from copy import deepcopy
import json,unittest
import adapters,exporter,financial_bridge
from workspace import Workspace
from common import AuditError
from app import FULL_DEMO_FILES,CATALOG
from raw_sources import load_pack
from raw_binding import analyze as bind
ROOT=Path(__file__).resolve().parent
C={'회사코드':'SYNTH01','시작일':'2026-01-01','기준일':'2026-12-31'}
def workspace(folder):return Workspace(folder,create={**C,'회사명':'합성 검증','회계체계':'K-IFRS','감사기준':'학습용','작성자':'합성작성자','더미업무':True})
class IntegratedV14Tests(unittest.TestCase):
    def setUp(self):self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name);self.ws=workspace(self.root/'engagement')
    def tearDown(self):self.tmp.cleanup()
    def normal(self):
        for module,name in FULL_DEMO_FILES:self.ws.execute(module,ROOT/'samples/통합'/name)
    def test_twelve_modules_actual_ledger_and_all_cross_links(self):
        self.normal();self.assertEqual(len(self.ws.latest()),12)
        for m,r in self.ws.latest().items():self.assertEqual(r['result']['blocking'],0,(m,r['result']['issues']))
        cross=self.ws.cross_checks();self.assertEqual(len(cross),32);self.assertTrue(all(r['상태']=='일치' for r in cross),cross)
        fs=self.ws.latest()['재무제표·주석']['result'];self.assertEqual(fs['metrics']['매핑계정수'],59);self.assertEqual(fs['metrics']['대사항목수'],94)
        self.assertTrue({'FS01','R01'}<=self.ws.required_procedures())
    def test_stale_financial_run_cannot_be_reviewed_by_resaving_procedure(self):
        self.normal();fs=self.ws.latest()['재무제표·주석']
        self.ws.execute('원장·리드',ROOT/'samples/통합/통합_v1.3_원장_리드.xlsx')
        p={'purpose_code':'financial_statement_reconciliation','id':'FS01','title':'합성 재무제표 대사','preparer':'합성작성자','purpose':'합성','risk':'합성','scope':'합성','performed':'합성 기록','conclusion':'합성 판단','performed_date':'2026-12-31','evidence_ids':[fs['source_id'],self.ws.latest()['원장·리드']['source_id']]}
        self.ws.save_procedure(p)
        with self.assertRaises(AuditError):self.ws.request_review('FS01')
        self.assertEqual(next(x for x in self.ws.cross_checks() if x['모듈']=='재무제표·주석')['상태'],'미확인')
        with self.assertRaises(AuditError):self.ws.snapshot(True)
    def test_raw_package_different_current_ledger_is_blocked(self):
        gl=adapters.run('원장·리드',ROOT/'samples/통합/통합_v1.3_원장_리드.xlsx',C)
        changed=deepcopy(gl);changed['tables']['입력_총계정원장'][0]['차변']+=1
        out=bind(ROOT/'samples/통합/통합_v1.4_원천증빙.auditpack',C,changed)
        self.assertGreater(out['blocking'],0);self.assertEqual(out['status'],'INPUT_BLOCKED')
    def test_composite_gl_identity_allows_repeated_line_numbers(self):
        gl=adapters.run('원장·리드',ROOT/'samples/통합/통합_v1.3_원장_리드.xlsx',C)
        pack=load_pack(ROOT/'samples/통합/통합_v1.4_원천증빙.auditpack',{**C,'통화':'KRW','연결구분':'별도'})
        table=next(t for t in pack['catalog']['tables'] if t['sheet']=='총계정원장');ix=table['headers'].index('행ID');group={}
        for row,raw in zip(gl['tables']['입력_총계정원장'],table['rows']):
            key=row['전표ID'];group[key]=group.get(key,0)+1;row['행ID']=str(group[key]);raw['values'][ix]=row['행ID']
        with patch('raw_sources.load_pack',return_value=pack):out=bind('ignored',C,gl)
        self.assertEqual(out['blocking'],0,out['issues'])
    def test_common_transaction_normal_and_missing_balanced_journal(self):
        folder=ROOT/'samples/공통거래/normal';self.ws.execute('원장·리드',folder/'원장_공통거래.xlsx');r=self.ws.execute('원천증빙',folder/'공통거래_normal.auditpack')
        self.assertEqual(r['result']['blocking'],0,r['result']['issues']);self.assertEqual(r['result']['review'],0)
        folder=ROOT/'samples/공통거래/balanced_journal_missing';self.ws.execute('원장·리드',folder/'원장_공통거래.xlsx');r=self.ws.execute('원천증빙',folder/'공통거래_balanced_journal_missing.auditpack')
        self.assertGreater(r['result']['blocking'],0)
    def test_request_receipt_requires_evidence_and_retains_history(self):
        entry={'title':'계약 원문 요청','requested_from':'합성 담당','note':'미수령 합성 요청','due_date':'2026-12-31','state':'요청','source_ids':[]}
        row=self.ws.save_request(entry)
        with self.assertRaises(AuditError):self.ws.save_request({**row,'state':'수령'})
        source=self.ws.add_source(ROOT/'samples/통합/통합_v1.4_재무제표.json','합성 자료')
        self.ws.save_request({**row,'state':'검토완료','note':'합성 검토 이력','source_ids':[source]})
        self.assertEqual(len([e for e in self.ws.events if e['kind']=='evidence_request']),2)
        self.assertEqual(self.ws.snapshot()['requests'][0]['state'],'검토완료')
    def test_workpaper_exports_actual_links_and_draft_states(self):
        self.normal()
        for pid,module in [('FS01','재무제표·주석'),('R01','원천증빙')]:
            refs=self.ws.calculation_refs(pid);self.ws.save_procedure({'purpose_code':'financial_statement_reconciliation' if pid=='FS01' else pid,'id':pid,'title':module,'preparer':'합성작성자','purpose':'합성 검증 초안','scope':str(refs['계산지표']),'evidence_ids':[refs['자료ID'],self.ws.latest()['원장·리드']['source_id']]})
        target=self.root/'draft.xlsx';snap=exporter.export(self.ws,target,CATALOG)
        from openpyxl import load_workbook
        w=load_workbook(target,data_only=False)
        try:
            self.assertTrue({'조서_FS01','조서_R01','자료요청·수령'}<=set(w.sheetnames))
            self.assertTrue(any(row[1].hyperlink for row in w['조서_FS01'].iter_rows(min_row=2)))
            proof=json.loads(''.join(row[1] for row in w['증적JSON'].iter_rows(min_row=2,values_only=True)))
            self.assertEqual(proof['basis'],self.ws.basis())
            self.assertTrue(all(not p['performed'] and not p['conclusion'] and not p['reviewer'] for p in proof['procedures']))
            self.assertTrue(all(c.data_type!='f' for sh in w for row in sh for c in row))
        finally:w.close()
    def test_changed_registered_source_prevents_further_calculation(self):
        run=self.ws.execute('원장·리드',ROOT/'samples/통합/통합_v1.3_원장_리드.xlsx')
        (self.ws.root/self.ws.sources[run['source_id']]['file']).write_bytes(b'changed')
        with self.assertRaises(AuditError):self.ws.execute('재무제표·주석',ROOT/'samples/통합/통합_v1.4_재무제표.json')
