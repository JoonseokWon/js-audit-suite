"""Native company selector checks, also executed in an EXE-only directory."""
from pathlib import Path
import json,sys,time,traceback
from common import atomic,sha,AuditError

def self_test(report_path):
    out=Path(report_path).parent;out.mkdir(parents=True,exist_ok=True);root=None;checks=[]
    def check(ok,name):
        if not ok:raise AssertionError(name)
        checks.append(name)
    try:
        from verify_v15 import self_test as legacy
        check(legacy(out/'legacy_v15/report.json')==0,'legacy_332_cases_and_native_checks')
        old=json.loads((out/'legacy_v15/report.json').read_text(encoding='utf-8'))
        from app import App,resource,CATALOG
        from workspace import Workspace
        from sector_demo import STORE_FILE
        from window_capture import capture_window,destroy_test_window
        from unittest.mock import patch
        import tkinter as tk
        import exporter
        root=tk.Tk();app=App(root,sector_home=out/'회사별저장')
        def settle():
            for _ in range(6):root.update();time.sleep(.05)
        def finished():
            deadline=time.monotonic()+180
            while app.busy and time.monotonic()<deadline:root.update();time.sleep(.02)
            check(not app.busy,'background_preparation_finished');settle()
        def choose(code):
            i=next(i for i,c in enumerate(app.sector_options) if c['code']==code)
            app.sector_box.current(i);app.sector_box.event_generate('<<ComboboxSelected>>');root.update()
        check([c['name'] for c in app.sector_options]==['JS','JS전자','JS자동차','JS반도체','JS바이오'],'five_named_company_choices')
        observations=[];folders={}
        for code in ('JS','JS_ELEC','JS_AUTO','JS_SEMI','JS_BIO'):
            choose(code)
            check(app.proc_text['performed'].cget('state')=='disabled',code+'_preparation_blocks_workpaper_edit')
            check(app.request_vars['title'].get()=='',code+'_no_previous_company_request')
            finished();ws=app.ws
            check(ws.engagement['회사코드']==code and app.active_sector==code,code+'_native_selection_loaded')
            check(len(ws.latest())==15 and len(ws.cross_checks())==35 and all(r['상태']=='일치' for r in ws.cross_checks()),code+'_fifteen_modules_thirty_five_ledger_links')
            check(not ws.verify_sources(),code+'_registered_source_integrity')
            check(all(not p['performed'] and not p['conclusion'] and not p['reviewer'] for p in ws.procedures.values()),code+'_human_judgment_blank')
            check(not ws.requests,code+'_requests_not_automatically_created_or_sent')
            check(app.proc_text['performed'].cget('state')=='normal',code+'_edit_restored_after_prepare')
            folders[code]=ws.root
            app.nb.select(app.tabs['원장·리드']);settle();capture_window(root,out/(code+'_원장대사.png'))
            app.load_procedure('A10');app.nb.select(app.tabs['절차·조서']);settle();capture_window(root,out/(code+'_업종조서.png'))
            check(ws.engagement['회사명'] in app.proc_heading.get(),code+'_company_workpaper_title_visible')
            target=out/(code+'_미검토초안.xlsx');exporter.export(ws,target,CATALOG)
            from openpyxl import load_workbook
            book=load_workbook(target,read_only=True,data_only=True)
            check(any(n.endswith('리드스케줄') for n in book.sheetnames) and '모듈TB연결' in book.sheetnames,code+'_xlsx_ledger_sheets')
            book.close()
            try:exporter.export(ws,out/(code+'_검토완료본_차단.xlsx'),CATALOG,True)
            except AuditError:check(True,code+'_unreviewed_final_export_blocked')
            else:raise AssertionError(code+'_final_export_was_not_blocked')
            observations.append({'company':code,'name':ws.engagement['회사명'],'modules':list(ws.latest()),'links':ws.cross_checks(),'procedures':len(ws.procedures),'risk_candidates':len(ws.risk_register()),'draft_bytes_positive':target.stat().st_size>0})
        choose('JS');finished();check(app.ws.root==folders['JS'],'reselect_uses_saved_company_workspace')
        app.load_procedure('G01');app.proc_text['performed'].insert('end','미저장 선택 취소 검증')
        previous=app.ws.engagement['id']
        with patch('app.messagebox.askyesnocancel',return_value=None):choose('JS_ELEC')
        check(not app.busy and app.ws.engagement['id']==previous and app.sector_name.get()=='JS','unsaved_cancel_keeps_old_company_and_choice')
        check(app.dirty(),'cancel_keeps_unsaved_text')
        app.proc_text['performed'].delete('1.0','end');app.proc_text['performed'].insert('1.0','합성 검증용 수동 입력 보존');app.save_procedure()
        app.request_vars['title'].set('합성 추가 증빙');app.request_vars['requested_from'].set('가상 회계담당');app.request_vars['note'].set('합성 검증용 미수령 추가 증빙 요청');app.request_vars['due_date'].set('2027-01-31');app.save_request()
        request_id=app.current_request;head=app.ws.head
        app.file.set('old-company-input.xlsx');app.search.set('old-company-search');choose('JS_ELEC');finished()
        check(not app.file.get() and not app.search.get() and not app.reviewer.get() and not app.review_note.get(),'switch_clears_previous_company_fields')
        check(not app.current_pid and not app.current_request and not app.request_vars['title'].get(),'switch_clears_previous_forms')
        choose('JS');finished()
        check(app.ws.head==head and app.ws.procedures['G01']['performed']=='합성 검증용 수동 입력 보존' and request_id in app.ws.requests,'reselect_preserves_human_workpaper_request_and_history')
        old_folder=app.ws.root;old_head=app.ws.head;app.sector_new.invoke();finished()
        check(app.ws.root!=old_folder and Workspace(old_folder).head==old_head,'fresh_button_keeps_previous_session')
        check(not app.ws.requests and not app.ws.procedures['G01']['performed'],'fresh_button_creates_blank_human_fields')
        old_current=app.ws.root;bad=folders['JS_BIO']/'업무.sqlite3';data=bad.read_bytes()
        try:
            bad.write_bytes(b'corrupt QA copy');choose('JS_BIO');finished()
            check(app.ws.root==old_current and app.sector_name.get()=='JS','failed_selection_keeps_current_company')
            check('중단' in app.status.get(),'failed_selection_is_visible')
        finally:bad.write_bytes(data)
        check(Workspace(folders['JS_BIO']).engagement['회사코드']=='JS_BIO','failure_did_not_repair_or_overwrite_other_company')
        # Cancellation is queued while the worker waits at a deterministic progress barrier.
        import sector_demo,threading
        original=sector_demo.open_company;started=threading.Event();release=threading.Event()
        def delayed(*args,**kwargs):
            started.set();release.wait(10);return original(*args,**kwargs)
        with patch('sector_demo.open_company',side_effect=delayed):
            choose('JS_AUTO');check(started.wait(2),'cancellation_worker_started')
            app.sector_cancel.invoke();release.set();finished()
        check(app.ws.root==old_current and app.sector_name.get()=='JS','cancel_preparation_keeps_previous_workspace')
        copied=app.copy_input_bundle(out/'자료복사');check((copied/'업종별자료/manifest.json').read_bytes()==resource('sector_resources/manifest.json').read_bytes(),'copy_button_includes_exact_sector_resources')
        result={'ok':True,'version':'1.6.0','frozen':bool(getattr(sys,'frozen',False)),'legacy_cases':332,'observations':old['observations'],'sector_observations':observations,'native_checks':checks,'legacy_native_checks':old['native_checks'],'resource_manifest_sha256':sha(resource('sector_resources/manifest.json'))}
        atomic(report_path,json.dumps(result,ensure_ascii=False,indent=2).encode());return 0
    except Exception as e:
        atomic(report_path,json.dumps({'ok':False,'error':str(e),'traceback':traceback.format_exc(),'native_checks':checks},ensure_ascii=False,indent=2).encode());return 1
    finally:
        if root:
            try:
                from window_capture import destroy_test_window
                destroy_test_window(root)
            except Exception:pass
