"""Shared source/frozen expected-value and native UI checks."""
from pathlib import Path
import json,sys,time,traceback,shutil
from common import *
import adapters

def check_book(book,path,context):
    r=adapters.run(book['module'],path,context);e=book['expected'];codes=[i['code'] for i in r['issues']]
    if 'code' in e:assert e['code'] in codes,(e,codes)
    for k in ('blocking','review'):
        if k in e:assert r[k]==e[k],(book['name'],k,r[k],e[k],r['issues'])
    if 'metric' in e:assert r['metrics']==e['metric'],r['metrics']
    tables=r['tables']
    if 'checks' in e:assert len(tables['원장통제'])==e['checks']
    if 'lead_rows' in e:assert len(tables['리드스케줄'])==e['lead_rows']
    if 'adjusted_asset' in e:assert next(x for x in tables['리드스케줄'] if x['계정코드']=='1500')['조정후']==e['adjusted_asset']
    if 'ar_recovered' in e:
        a=tables['채권연령후속회수'][0];assert a['후속회수']==60000 and a['잔여']==50000 and a['연체일수']==91
        assert len([x for x in tables['은행거래대사'] if x['상태']=='정확대사'])==2
    return r

def regression(root):
    """Run pinned old fixtures against the packaged engines and their independent expectations."""
    reports=[]
    from engines.finance.practice_engine import reconcile
    from engines.finance.v2_checks import verify_result as fin_check
    from engines.sales.data_io import read_table,IMPORT_FIELDS
    from engines.sales.ledger_engine import analyze_with_ledger
    from engines.sales.v2_checks import verify_result as sales_check
    from engines.assets.data_io import run_file
    from engines.assets.checks import verify as asset_check
    for kind in ('금융상품','매출 자동화'):
        folder=root/'samples/기존'/kind;spec=json.loads((folder/'v2_cases.json').read_text(encoding='utf-8'))
        for c in spec['cases']:
            p=folder/('v2_'+c['name']+'.xlsx');before=sha(p)
            if kind=='금융상품':
                if 'input_error' in c['expected']:
                    try:reconcile(p)
                    except ValueError as exc:
                        assert c['expected']['input_error'] in str(exc);actual={'expected_input_rejection':str(exc)}
                    else:raise AssertionError('Expected input rejection '+c['name'])
                else:actual=fin_check(reconcile(p),c['expected'])
            else:
                raw=p.read_bytes();t={k:read_table(raw,p.name,k,k) for k in IMPORT_FIELDS};r=analyze_with_ledger(t['orders'],t['deliveries'],t['invoices'],t['ledger'],t['control'],cutoff_date='2026-09-30',amount_tolerance='1',quantity_tolerance='0',unit_price_tolerance='0.01');actual=sales_check(r,c['expected'])
            assert before==sha(p);reports.append({'module':kind,'name':c['name'],'actual':actual,'ok':True})
    folder=root/'samples/기존/감가상각';spec=json.loads((folder/'cases.json').read_text(encoding='utf-8'))
    if isinstance(spec,dict):spec=spec['cases']
    for c in spec:
        p=folder/(c['name']+'.xlsx');before=sha(p)
        if c.get('expected',{}).get('input_error'):
            try:run_file(p)
            except ValueError as exc:
                assert c['expected']['input_error'] in str(exc);actual={'expected_failure':True}
            else:raise AssertionError('Expected input rejection '+c['name'])
        else:actual=asset_check(run_file(p),c['expected'])
        assert before==sha(p);reports.append({'module':'감가상각','name':c['name'],'actual':actual,'ok':True})
    return reports

def synthetic_procedure(pid,source_ids):return {'id':pid,'title':'합성 수행·검토 연습 '+pid,'purpose':'합성 자료의 검사·조서 연결 연습','risk':'SYNTHETIC ONLY: 실재·권리·추정 판단은 실제 미수행','scope':'등록된 합성 입력 '+str(len(source_ids))+'개 및 자동검사 결과','performed':'SYNTHETIC ONLY / 프로그램 검증에서 수행·검토 상태를 모의 입력. 실제 감사증거 검사 아님.','conclusion':'합성 시나리오의 조서 상태 확인. 실제 감사 결론 없음.','performed_date':'2026-10-22','preparer':'합성작성자','evidence_ids':source_ids,'excluded':False}

def self_test(report_path):
    from app import App,CATALOG,resource,DEMO_FILES
    import tkinter as tk
    from workspace import Workspace
    from window_capture import capture_window
    import exporter
    report_path=Path(report_path).resolve();report_path.parent.mkdir(parents=True,exist_ok=True);checks=[];cases=[];root=None
    def check(ok,name):
        if not ok:raise AssertionError(name)
        checks.append(name)
    try:
        spec=json.loads(resource('samples_specs.json').read_text(encoding='utf-8'))
        for book in spec['books']:
            if book['expected'].get('template'):continue
            r=check_book(book,resource('samples','통합',book['name']+'.xlsx'),spec['context']);cases.append({'name':book['name'],'module':book['module'],'blocking':r['blocking'],'review':r['review'],'ok':True})
        extra=json.loads(resource('new_specs.json').read_text(encoding='utf-8'))
        for book in extra['books']:
            if book['expected'].get('template'):continue
            if 'input_error' in book['expected']:
                try:adapters.run(book['module'],resource('samples','통합',book['name']+'.xlsx'),extra['context'])
                except AuditError as exc:assert book['expected']['input_error'] in str(exc),(book['name'],str(exc))
                else:raise AssertionError('Expected input rejection: '+book['name'])
                cases.append({'name':book['name'],'module':book['module'],'expected_rejection':True,'ok':True})
            else:
                r=check_book(book,resource('samples','통합',book['name']+'.xlsx'),extra['context']);cases.append({'name':book['name'],'module':book['module'],'blocking':r['blocking'],'review':r['review'],'ok':True})
        extra12=json.loads(resource('v12_specs.json').read_text(encoding='utf-8'))
        for book in extra12['books']:
            if book['expected'].get('template'):continue
            path=resource('samples','통합',book['name']+'.xlsx')
            if 'input_error' in book['expected']:
                try:adapters.run(book['module'],path,extra12['context'])
                except AuditError as exc:assert book['expected']['input_error'] in str(exc),(book['name'],str(exc))
                else:raise AssertionError('Expected input rejection: '+book['name'])
                cases.append({'name':book['name'],'module':book['module'],'expected_rejection':True,'ok':True})
            else:
                r=check_book(book,path,extra12['context']);cases.append({'name':book['name'],'module':book['module'],'blocking':r['blocking'],'review':r['review'],'ok':True})
        extra13=json.loads(resource('v13_specs.json').read_text(encoding='utf-8'))
        for book in extra13['books']:
            if book['expected'].get('template'):continue
            path=resource('samples','통합',book['name']+'.xlsx')
            if 'input_error' in book['expected']:
                try:adapters.run(book['module'],path,extra13['context'])
                except AuditError as exc:assert book['expected']['input_error'] in str(exc),(book['name'],str(exc))
                else:raise AssertionError('Expected input rejection: '+book['name'])
                cases.append({'name':book['name'],'module':book['module'],'expected_rejection':True,'ok':True})
            else:
                r=check_book(book,path,extra13['context']);cases.append({'name':book['name'],'module':book['module'],'blocking':r['blocking'],'review':r['review'],'ok':True})
        old=regression(resource());check(len(old)==55,'old_55_cases_expected')
        root=tk.Tk();app=App(root);root.update();check(len(CATALOG['modules'])==48,'48_procedure_catalog')
        folder=report_path.parent/('더미업무_'+str(time.time_ns()));app.start_demo(folder)
        deadline=time.monotonic()+90
        while app.busy and time.monotonic()<deadline:root.update();time.sleep(.02)
        check(not app.busy,'native_worker_completion');app.refresh();ws=app.ws
        check(len(ws.latest())==10,'ten_native_integrated_modules');check(all(r['result']['blocking']==r['result']['review']==0 for r in ws.latest().values()),'ten_normal_results')
        check(len(app.lead_tree.get_children())==59,'native_lead_rows');check(len(ws.cross_checks())==30 and all(x['상태']=='일치' for x in ws.cross_checks()),'thirty_independent_tb_links')
        check(len(app.proc_tree.get_children())==len(CATALOG['modules'])+3,'native_catalog_and_three_integrated_procedures');check(len(app.source_tree.get_children())==10,'ten_immutable_sources')
        ws2=Workspace(folder);check(len(ws2.runs)==10,'restart_persistence')
        app.result_module.set('감가상각');app.choose_result_module();app.result_table.set('월별계산');app.load_table();check(len(app.table_rows)==24,'monthly_detail_all_24')
        app.search.set('2026-12');app.load_table();check(len(app.table_rows)==1,'search_filter');app.search.set('');app.load_table()
        p=synthetic_procedure('A12',list(ws.sources));app.load_procedure('A12')
        for key,var in app.proc_vars.items():var.set(p[key])
        for key,t in app.proc_text.items():t.delete('1.0','end');t.insert('1.0',p[key])
        app.evidence_list.selection_set(0,4);app.save_procedure();app.request_review();app.reviewer.set('합성검토자');app.review_note.set('SYNTHETIC ONLY / UI 검토 연습');app.review(True)
        check(ws.procedures['A12']['state']=='검토기록완료','native_procedure_review_workflow')
        for module,pid,table in [('리스','A15','월별리스'),('대손충당금','A07','채권별ECL'),('재고평가','A11','재고평가요약'),('종업원급여','A16','월별급여계산'),('미기록부채·비용이연','A08','비용기간배분'),('미기록부채·비용이연','A09','비용기간배분')]:
            app.result_module.set(module);app.choose_result_module();app.result_table.set(table);app.load_table()
            app.nb.select(app.tabs['결과 상세']);root.update();capture_window(root,report_path.parent/(module+'_계산화면.png'))
            check(len(app.table_rows)==(24 if module=='리스' else 2),'native_'+module+'_detail')
            app.load_procedure(pid);check('자동계산 연결' in app.proc_text['scope'].get('1.0','end'),'native_'+pid+'_calculation_refs')
            check(len(app.evidence_list.curselection())==2,'native_'+pid+'_evidence_links')
            app.nb.select(app.tabs['절차·조서']);root.update();capture_window(root,report_path.parent/(module+('_'+pid if module=='미기록부채·비용이연' else '')+'_조서화면.png'))
        for pid in ws.required_procedures()-{'A12'}:
            ws.save_procedure(synthetic_procedure(pid,list(ws.sources)));ws.request_review(pid);ws.review(pid,'합성검토자','SYNTHETIC ONLY / 상태검사')
        nid=ws.add_note('합성 검토의견','SYNTHETIC ONLY / 예외 해결 연습','합성검토자')
        try:ws.snapshot(True)
        except AuditError:check(True,'unresolved_note_blocks_final')
        else:raise AssertionError('Unresolved note bypass')
        ws.add_note('합성 검토의견','SYNTHETIC ONLY / 해결기록','합성작성자',True,nid);snap=ws.snapshot(True);check(len(snap['procedures'])==15,'fifteen_required_manual_workpapers')
        path=report_path.parent/'통합감사_합성검토완료.xlsx';exporter.export(ws,path,CATALOG,True)
        from openpyxl import load_workbook
        w=load_workbook(path,read_only=False,data_only=False)
        try:
            proof=''.join(r[1] for r in w['증적JSON'].iter_rows(min_row=2,values_only=True));restored=json.loads(proof);check(restored['head']==snap['head'],'embedded_full_proof_reopens')
            check(all(any(str(row[0]).startswith('계산값: ') for row in w['조서_'+pid].iter_rows(values_only=True)) for pid in ('A15','A07','A11','A16','A08','A09')),'estimate_workpapers_include_calculation_and_tb')
            check(all(any(str(row[0].value).startswith('상세표: ') and row[1].hyperlink for row in w['조서_'+pid].iter_rows()) for pid in ('A15','A07','A11','A16','A08','A09')),'estimate_workpaper_detail_hyperlinks')
            check(w['수행검토조서'].max_row==16,'fifteen_complete_export_rows');check(len([n for n in w.sheetnames if n.startswith('조서_')])==15,'fifteen_readable_individual_workpapers');check(any(str(n).endswith('리드스케줄') for n in w.sheetnames),'lead_exported');check(all(c.data_type!='f' for sheet in w for row in sheet for c in row),'no_formulas_in_workpaper')
        finally:w.close()
        app.refresh();app.nb.select(app.tabs['원장·리드']);root.update();capture_window(root,report_path.parent/'통합_리드화면.png');checks.append('native_lead_capture')
        app.nb.select(app.tabs['절차·조서']);app.load_procedure('A12');root.update();capture_window(root,report_path.parent/'통합_조서화면.png');checks.append('native_workpaper_capture')
        app.ws.execute('원장·리드',resource('samples','통합','통합_02_상계되는원장누락.xlsx'));app.refresh();check(app.ws.latest()['원장·리드']['result']['blocking']>0,'bad_latest_replaces_good');check(app.ws.procedure_state(app.ws.procedures['A12'])=='자료변경·재검토','new_input_reopens_review')
        app.nb.select(app.tabs['자료·검사']);root.update();capture_window(root,report_path.parent/'통합_오류화면.png');checks.append('native_exception_capture')
        out={'ok':True,'frozen':bool(getattr(sys,'frozen',False)),'checks':checks,'new_cases':cases,'legacy_cases':old,'dummy_count':len(cases)+len(old),'workspace':str(folder),'proof':str(path)}
        atomic(report_path,json.dumps(out,ensure_ascii=False,indent=2).encode());return 0
    except Exception as e:
        atomic(report_path,json.dumps({'ok':False,'frozen':bool(getattr(sys,'frozen',False)),'error':str(e),'traceback':traceback.format_exc(),'checks':checks,'new_cases':cases},ensure_ascii=False,indent=2).encode());return 1
    finally:
        if root:
            try:
                from window_capture import destroy_test_window
                destroy_test_window(root)
            except Exception:pass
