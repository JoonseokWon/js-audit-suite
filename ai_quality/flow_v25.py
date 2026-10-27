"""Explicit synthetic state transitions after one accepted real response.

These isolated inputs are never presented as actual human performance, approval,
or review. The received model text is retained and never repaired or auto-merged.
"""
from pathlib import Path
import json
from copy import deepcopy
from common import AuditError,digest


def complete_flow(ws,packet,response,folder,check):
    import exporter,runtime
    from workspace import Workspace
    from openpyxl import load_workbook
    from demo_case import load_manifest
    pid=packet['procedure']['id'];folder=Path(folder)
    original=response['raw_response'];proposals=json.loads(original)['proposals']
    before=len(ws.events);invalid=deepcopy(response)
    parsed=json.loads(original);parsed['proposals']['scope']+=' 근거없는 금액 987654321원.'
    invalid['raw_response']=json.dumps(parsed,ensure_ascii=False)
    invalid['fields']=parsed['proposals'];invalid['unverified_proposals']=parsed['proposals']
    try:ws.record_ai_draft(packet,invalid,'SYNTHETIC ONLY negative probe')
    except AuditError:pass
    else:raise AssertionError('unsupported_number_response_was_saved')
    check(len(ws.events)==before,'unsupported_number_is_rejected_without_history_write')
    did=ws.record_ai_draft(packet,response,'SYNTHETIC ONLY request')
    check(not ws.procedures.get(pid,{}).get('ai_draft_id'),'generation_does_not_automatically_apply')
    check(not ws.ai_approvals,'generation_does_not_automatically_approve')
    ws.apply_ai_draft(did,'SYNTHETIC ONLY author')
    check(all(ws.procedures[pid][key]==value for key,value in response['verified_fields'].items()),
          'only_program_rendered_records_are_applied')
    check(not ws.ai_status(pid)['approved'],'application_is_not_content_approval')
    check(not ws.procedures[pid].get('performed_date'),'application_preserves_unperformed_human_state')
    check(all(value not in ws.procedures[pid][key] for key,value in proposals.items()),
          'unverified_model_proposals_are_not_automatically_merged')
    marker='SYNTHETIC ONLY / 직접 입력 모사: 실제 감사판단·승인 아님. 차이 원인과 증빙은 사람이 추가 확인해야 함.'
    p=dict(ws.procedures[pid]);p['conclusion']+='\n'+marker;ws.save_procedure(p)
    check(not ws.ai_status(pid)['approved'],'explicit_synthetic_edit_stays_unapproved')
    ws.approve_ai_content(pid,'SYNTHETIC ONLY state transition',
        '가상 상태전이 검증. 실제 사람 승인 아님. 프로그램 기록과 직접 입력을 모사한 수정의 버전 연결만 확인함.')
    check(ws.ai_status(pid)['approved'],'synthetic_content_approval_is_bound_to_current_content')
    check(not ws.procedures[pid].get('reviewed_at'),'content_approval_does_not_create_separate_review')
    reopened=Workspace(ws.root)
    check(reopened.ai_drafts[did]['raw_response']==original,'reopening_preserves_exact_model_original')
    check(reopened.ai_drafts[did]['unverified_proposals']==proposals,'reopening_preserves_quarantined_proposals')
    check(reopened.ai_status(pid)['approved'],'synthetic_content_approval_survives_reopen')
    catalog=json.loads((runtime.configure()/'procedure_catalog.json').read_text(encoding='utf-8'))
    export=folder/'JS전자_실제모델_프로그램근거_가상수정승인.xlsx'
    exporter.export(reopened,export,catalog,False)
    book=load_workbook(export,data_only=False)
    try:
        values=[cell.value for sheet in book for row in sheet for cell in row]
        check(any(marker in str(value) for value in values),'explicit_synthetic_edit_is_exported')
        check('AI근거참조' in book.sheetnames and 'AI최종내용승인' in book.sheetnames,
              'Excel_keeps_separate_reference_and_content_approval_records')
        for title in ('AI프로그램근거','AI미검증제안','AI구조화주장'):
            check(title in book.sheetnames,'Excel_separates_'+title)
        check(not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row),
              'OOXML_contains_no_executable_formula_cells')
        restored=json.loads(''.join(row[1] for row in book['증적JSON'].iter_rows(min_row=2,values_only=True)))
        check(next(d for d in restored['ai_drafts'] if d['id']==did)['raw_response']==original,
              'OOXML_evidence_JSON_restores_exact_original_model_response')
    finally:book.close()
    p=dict(reopened.procedures[pid]);p['conclusion']+='\nSYNTHETIC ONLY / 재수정';reopened.save_procedure(p)
    check(not reopened.ai_status(pid)['approved'],'editing_invalidates_previous_synthetic_approval')
    reopened.approve_ai_content(pid,'SYNTHETIC ONLY second state transition',
        '실제 감사 승인 아님. 데이터 변경 무효화 시험만을 위한 가상 입력.')
    prior=digest(packet);drafts=len(reopened.ai_drafts)
    entry=load_manifest(reopened.root)['inputs']['ledger_clean'];reopened.execute(entry['module'],entry['path'])
    check(reopened.ai_status(pid)['stale'] and not reopened.ai_status(pid)['approved'],
          'new_ledger_run_invalidates_previous_synthetic_approval')
    check(digest(reopened.ai_packet(pid,limits=packet['limits']))!=prior,'new_ledger_changes_the_evidence_packet')
    try:reopened.record_ai_draft(packet,response,'SYNTHETIC ONLY stale probe')
    except AuditError:pass
    else:raise AssertionError('stale_response_was_saved')
    check(len(reopened.ai_drafts)==drafts,'stale_original_response_cannot_be_saved_again')
    approvals=sum(e['kind']=='ai_content_approval' for e in reopened.events)
    check(approvals==2,'exactly_two_explicit_synthetic_approval_events')
    return {'draft_id':did,'export':str(export),'workspace':str(reopened.root),
        'synthetic_approval_events':approvals,'actual_human_approvals':0,
        'human_approval_scope':'synthetic state transition only; no real approval'}
