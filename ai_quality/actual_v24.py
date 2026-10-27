"""One opt-in local call for v2.4. Never run by test discovery or CI.

The response is either rejected with its evidence preserved or stored as an
unverified proposal. There is no repair, automatic retry, or best-of selection.
"""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import json
import time
from common import VERSION,atomic,digest


def self_test(report_path,runtime_root):
    from local_ai_runtime import LocalAIRuntime,MODEL,MODEL_DIGEST
    from ai_provider import generate_draft
    from ai_contract import validate_response,CONTRACT_VERSION
    from demo_case import create_demo
    from workspace import Workspace
    from openpyxl import load_workbook
    from .capture import ActualCapture
    import exporter,runtime
    path=Path(report_path)
    if path.exists():raise ValueError('Preserve the existing report; use a new private directory.')
    out=path.parent;out.mkdir(parents=True,exist_ok=True)
    checks=[];manager=LocalAIRuntime(runtime_root);capture=ActualCapture(out)
    started=time.monotonic();packet=None;response=None
    def check(ok,name):
        if not ok:raise AssertionError(name)
        checks.append(name)
    try:
        ws,manifest=create_demo(out/'JS전자_7원_구조화계약검증')
        clean=manifest['inputs']['ledger_clean'];ws.execute(clean['module'],clean['path'])
        book=load_workbook(manifest['inputs']['sales_clean']['path'])
        try:
            sheet=book['invoices'];headers=[c.value for c in sheet[1]]
            ident=headers.index('invoice_id')+1;amount=headers.index('supply_amount')+1
            rows=[i for i in range(2,sheet.max_row+1) if sheet.cell(i,ident).value=='SI01']
            check(len(rows)==1,'independent_invoice_SI01_is_unique')
            sheet.cell(rows[0],amount).value=90007
            source=out/'JS전자_송장_7원차이.xlsx';book.save(source)
        finally:book.close()
        run=ws.execute('매출',source)
        check(ws.latest()['원장·리드']['result']['blocking']==0,'clean_ledger_passes_its_independent_controls')
        check(run['result']['review']>0 or run['result']['blocking']>0,'independent_seven_won_invoice_difference_is_flagged')
        check(str(run['result']['metrics'].get('순매출'))=='100000','independently_expected_revenue_remains_100000')
        check(not ws.ai_approvals,'case_has_no_actual_or_synthetic_content_approval')
        packet=ws.ai_packet('A04',limits={'max_chars':12000,'max_references':40,
            'rows_per_table':3,'max_tables':8,'cell_chars':300})
        check(packet['contract']['version']==CONTRACT_VERSION,'new_typed_contract_is_explicit')
        check('calculation_inventory' in packet,'final_table_inclusion_inventory_is_present')
        atomic(out/'evidence-packet.json',json.dumps(packet,ensure_ascii=False,indent=2).encode())
        state=manager.ensure_server()
        check(state['ready'] and state['runtime_verified'] and state['model_digest']==MODEL_DIGEST,
            'pinned_local_runtime_and_model_are_ready')
        with capture:
            response=generate_draft(packet,{'provider':'ollama','model':MODEL,'endpoint':manager.base_url,
                'num_gpu':0,'num_ctx':16384,'timeout_seconds':300})
        check(len(capture.requests)==1,'exactly_one_actual_loopback_chat_call_without_retry')
        parsed=json.loads(response['raw_response']);derived=validate_response(packet,parsed)
        check(response['fields']==parsed['proposals'],'model_proposal_is_preserved_without_program_rewriting')
        check(derived['typed_claims']==response['typed_claims'],'typed_reference_kind_status_path_value_and_unit_match_exactly')
        check(response['acceptance_state']=='quarantined_proposal','free_model_narrative_is_quarantined_not_verified')
        atomic(out/'response-derived-and-original.json',json.dumps(response,ensure_ascii=False,indent=2).encode())
        original=response['raw_response'];did=ws.record_ai_draft(packet,response,'SYNTHETIC ONLY request')
        check(not ws.procedures.get('A04',{}).get('ai_draft_id'),'generation_does_not_automatically_apply')
        ws.apply_ai_draft(did,'SYNTHETIC ONLY author')
        check(all(ws.procedures['A04'][field]==value for field,value in response['verified_fields'].items()),
            'application_uses_only_program_rendered_metadata_and_typed_records')
        check(not ws.ai_status('A04')['approved'],'applying_program_records_does_not_approve')
        edited=dict(ws.procedures['A04']);marker='SYNTHETIC ONLY / 자동검증용 수정, 실제 감사판단·승인 아님.'
        edited['conclusion']+='\n'+marker;ws.save_procedure(edited)
        reopened=Workspace(ws.root);saved=reopened.ai_drafts[did]
        check(saved['raw_response']==original and saved['unverified_proposals']==parsed['proposals'],
            'reopening_preserves_original_model_text_and_quarantined_proposals')
        check(not reopened.ai_approvals and not reopened.procedures['A04']['reviewed_at'],
            'no_human_or_synthetic_approval_or_separate_review_created')
        catalog=json.loads((runtime.configure()/'procedure_catalog.json').read_text(encoding='utf-8'))
        export=out/'JS전자_구조화기록_AI격리_미승인.xlsx';exporter.export(reopened,export,catalog,False)
        book=load_workbook(export,data_only=False)
        try:
            values=[cell.value for sheet in book for row in sheet for cell in row]
            check(any(marker in str(value) for value in values),'explicit_synthetic_edit_reaches_excel')
            check(not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row),'OOXML_export_has_no_executable_formula_cells')
            check('AI근거참조' in book.sheetnames,'OOXML_export_retains_reference_links')
        finally:book.close()
        report={'status':'passed','version':VERSION,'checks':checks,'contract_version':CONTRACT_VERSION,
            'acceptance_state':'quarantined_proposal','actual_model_calls':len(capture.requests),
            'actual_openai_calls':0,'actual_human_approvals':0,'synthetic_approval_events':0,
            'model':MODEL,'model_digest':MODEL_DIGEST,'packet_hash':digest(packet),
            'original_model_response_hash':digest(original),'derived_record_hash':digest(response),
            'claim_catalog_scope':response['claim_catalog_scope'],'transport':capture.requests,
            'independent_expected':{'normal_revenue':100000,'invoice_supply':90007,'invoice_difference':7},
            'scope':'Typed equality and isolation verified; model prose semantics remain unverified. A citation failure is not an incorrect accounting amount.',
            'workspace':str(ws.root),'export':str(export),'seconds':round(time.monotonic()-started,2)}
    except Exception as exc:
        report={'status':'failed','version':VERSION,'error':str(exc),'checks':checks,
            'actual_model_calls':len(capture.requests),'actual_openai_calls':0,'actual_human_approvals':0,
            'synthetic_approval_events':0,'response_received':any('response_hash' in c for c in capture.requests),
            'scope':'One attempt preserved without response repair, retry, or success selection.',
            'seconds':round(time.monotonic()-started,2)}
    finally:manager.stop()
    atomic(path,json.dumps(report,ensure_ascii=False,indent=2).encode())
    return 0 if report['status']=='passed' else 1
