"""Synthetic editing/approval/export checks following an actual model response."""
from copy import deepcopy
from common import AuditError, digest


def complete_flow(ws, packet, response, folder, check):
    import exporter, json, runtime
    from openpyxl import load_workbook
    from workspace import Workspace
    pid=packet['procedure']['id'];before=len(ws.events)
    invalid=deepcopy(response)
    invalid['fields']['scope']+=' 근거없는 금액 987,654,321원.'
    raw=json.loads(invalid['raw_response']);raw['fields']=invalid['fields']
    invalid['raw_response']=json.dumps(raw,ensure_ascii=False)
    try:ws.record_ai_draft(packet,invalid,'SYNTHETIC ONLY / invalid-response probe')
    except AuditError:pass
    else:raise AssertionError('unsupported_number_response_was_saved')
    check(len(ws.events)==before,'unsupported_number_rejected_without_history_write')
    invalid=deepcopy(response);invalid['citations']['scope']=['REF99999']
    raw=json.loads(invalid['raw_response']);raw['citations']=invalid['citations']
    invalid['raw_response']=json.dumps(raw,ensure_ascii=False)
    try:ws.record_ai_draft(packet,invalid,'SYNTHETIC ONLY / invalid-citation probe')
    except AuditError:pass
    else:raise AssertionError('unknown_citation_response_was_saved')
    check(len(ws.events)==before,'unknown_citation_rejected_without_history_write')
    did=ws.record_ai_draft(packet,response,'SYNTHETIC ONLY / 실제모델 시험 요청자')
    check(not ws.ai_approvals and not any(p.get('reviewed_at') for p in ws.procedures.values()),
          'actual_generation_does_not_approve_or_review')
    check(not ws.procedures.get(pid,{}).get('ai_draft_id'),'actual_generation_does_not_auto_apply')
    ws.apply_ai_draft(did,'SYNTHETIC ONLY / 가상 적용자')
    check(not ws.ai_status(pid)['approved'],'applied_actual_draft_stays_unapproved')
    check(not ws.procedures[pid].get('performed_date'),'AI_does_not_create_actual_performance_date')
    marker='SYNTHETIC ONLY / 가상 수정: 청구누락과 원장 통제 차이를 실제 증빙으로 조사해야 하며 실제 감사판단은 미수행이다.'
    p=dict(ws.procedures[pid]);p['conclusion']+='\n'+marker;ws.save_procedure(p)
    ws.approve_ai_content(pid,'SYNTHETIC ONLY / 승인 상태전이 시험',
        '실제 사람의 감사 승인이 아닌 합성 상태전이 검증. 수정 내용·근거·미수행 범위를 확인하는 입력을 모사함.')
    check(ws.ai_status(pid)['approved'],'explicit_synthetic_content_approval_is_version_bound')
    check(not any(p.get('reviewed_at') for p in ws.procedures.values()),
          'content_approval_does_not_create_separate_review')
    reopened=Workspace(ws.root)
    check(reopened.ai_status(pid)['approved'] and marker in reopened.procedures[pid]['conclusion'],
          'edited_actual_draft_and_synthetic_approval_survive_reopen')
    catalog=json.loads((runtime.configure()/'procedure_catalog.json').read_text(encoding='utf-8'))
    target=folder/'JS전자_실제모델_수정승인_검토중.xlsx';exporter.export(reopened,target,catalog,False)
    book=load_workbook(target,data_only=False)
    try:
        values=[cell.value for sheet in book for row in sheet for cell in row]
        check(any(marker in str(value) for value in values),'edited_actual_content_is_exported_to_excel')
        check('AI최종내용승인' in book.sheetnames and 'AI근거참조' in book.sheetnames,
              'excel_contains_separate_AI_approval_and_evidence_tables')
        check(not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row),
              'actual_model_export_contains_no_executable_formulas')
    finally:book.close()
    p=dict(reopened.procedures[pid]);p['conclusion']+='\nSYNTHETIC ONLY / 재수정';reopened.save_procedure(p)
    check(not reopened.ai_status(pid)['approved'],'post_approval_edit_invalidates_content_approval')
    reopened.approve_ai_content(pid,'SYNTHETIC ONLY / 재승인 상태전이 시험',
        '실제 감사 승인 아님. 원장 변경 무효화 검증을 위한 가상 상태 입력.')
    prior=digest(packet);before=len(reopened.ai_drafts)
    entry=__import__('demo_case').load_manifest(reopened.root)['inputs']['ledger_clean']
    reopened.execute(entry['module'],entry['path'])
    check(reopened.ai_status(pid)['stale'] and not reopened.ai_status(pid)['approved'],
          'new_ledger_run_invalidates_prior_synthetic_approval')
    check(digest(reopened.ai_packet(pid,limits=packet['limits']))!=prior,'new_ledger_changes_evidence_packet')
    try:reopened.record_ai_draft(packet,response,'SYNTHETIC ONLY / stale response probe')
    except AuditError:pass
    else:raise AssertionError('stale_response_was_saved')
    check(len(reopened.ai_drafts)==before,'stale_actual_response_is_not_saved_again')
    approvals=sum(e['kind']=='ai_content_approval' for e in reopened.events)
    check(approvals==2,'only_two_explicit_synthetic_content_approval_events_exist')
    return {'draft_id':did,'export_path':str(target),'approval_events':approvals,
            'approval_scope':'synthetic state transition only; no real approval'}
