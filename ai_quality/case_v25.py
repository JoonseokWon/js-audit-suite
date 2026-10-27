"""Isolated synthetic accounting case with independently specified amounts."""
from pathlib import Path
from openpyxl import load_workbook
from demo_case import create_demo


def create_case(out,check):
    out=Path(out);ws,manifest=create_demo(out/'JS전자_7원_생성계약검증')
    clean=manifest['inputs']['ledger_clean'];ws.execute(clean['module'],clean['path'])
    book=load_workbook(manifest['inputs']['sales_clean']['path'])
    try:
        sheet=book['invoices'];headers=[c.value for c in sheet[1]]
        ident=headers.index('invoice_id')+1;amount=headers.index('supply_amount')+1
        rows=[i for i in range(2,sheet.max_row+1) if sheet.cell(i,ident).value=='SI01']
        check(len(rows)==1,'independent_invoice_SI01_is_unique')
        sheet.cell(rows[0],amount).value=90007
        source=out/'JS전자_송장_7원차이.xlsx';book.save(source)
    finally:book.close()
    run=ws.execute('매출',source)
    check(ws.latest()['원장·리드']['result']['blocking']==0,'clean_ledger_independent_controls_pass')
    check(run['result']['review']>0 or run['result']['blocking']>0,'independent_seven_won_invoice_difference_is_flagged')
    check(str(run['result']['metrics'].get('순매출'))=='100000','independently_expected_revenue_remains_100000')
    check(not ws.ai_approvals,'case_has_no_actual_or_synthetic_content_approval')
    check(not ws.procedures.get('A04',{}).get('performed_date'),'no_actual_human_performance_is_invented')
    return ws,manifest
