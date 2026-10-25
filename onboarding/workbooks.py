"""Human-readable, formula-free Excel inputs and existing seven-sheet GL output."""
from io import BytesIO
from pathlib import Path
from common import AuditError,atomic,read_xlsx,sha
from hashlib import sha256
from onboarding import SCHEMA
from onboarding.workbook_schema import descriptors,encode,decode,metadata

META='_양식정보'
META_COLUMNS=['모듈','버전','시트','자료','형태','키','표시','자료형']


def _workbook():
    from openpyxl import Workbook
    workbook=Workbook();workbook.remove(workbook.active)
    return workbook


def _cell(sheet,row,column,value):
    cell=sheet.cell(row,column,value)
    if isinstance(value,str):cell.data_type='s'
    return cell


def _style(workbook):
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    for sheet in workbook:
        sheet.freeze_panes='A2';sheet.sheet_view.showGridLines=False
        for cell in sheet[1]:
            cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='263F59')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,(int,float)):cell.number_format='#,##0;[Red](#,##0);0'
        for column in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(column)].width=23


def _save(workbook,path):
    path=Path(path)
    if path.exists():raise AuditError('기존 파일을 보존하도록 새 파일 이름으로 저장하세요.')
    _style(workbook);buffer=BytesIO();workbook.save(buffer);workbook.close();atomic(path,buffer.getvalue())
    return path


def save_input(module,payload,path):
    workbook=_workbook();meta=workbook.create_sheet(META);meta.append(META_COLUMNS)
    for desc in descriptors(module,payload):
        sheet=workbook.create_sheet(desc['sheet'])
        if desc['kind']=='dict':
            sheet.append(['항목','값'])
            for row,key in enumerate(desc['keys'],2):
                _cell(sheet,row,1,desc['headers'][row-2]);_cell(sheet,row,2,encode(payload[desc['key']][key]))
        else:
            sheet.append(desc['headers'])
            for n,row in enumerate(payload[desc['key']],2):
                for j,key in enumerate(desc['keys'],1):_cell(sheet,n,j,encode(row.get(key)))
        for key,label in zip(desc['keys'],desc['headers']):
            meta.append([module,SCHEMA,desc['sheet'],desc['key'],desc['kind'],key,label,desc['types'][key]])
    meta.sheet_state='hidden'
    return _save(workbook,path)


def load_input(module,path,context):
    from onboarding.integration import sample
    if Path(path).suffix.lower()!='.xlsx' or Path(path).stat().st_size>20*1024**2:
        raise AuditError('20MB 이하 Excel 입력 양식을 사용하세요.')
    captured=Path(path).read_bytes()
    tables,source_metadata=read_xlsx(path,max_rows=10000)
    if sha256(captured).hexdigest()!=source_metadata['sha256']:raise AuditError('읽는 중 입력자료가 변경됐습니다.')
    rows=tables.pop(META,None)
    if not rows:raise AuditError('이 절차의 입력 양식이 아닙니다.')
    template=sample(module,context)
    allowed={d['key']:d for d in descriptors(module,template)}
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(captured),read_only=True,data_only=False,keep_links=False)
    try:actual_headers={sheet.title:tuple(c.value for c in next(sheet.iter_rows(min_row=1,max_row=1))) for sheet in book}
    finally:book.close()
    if actual_headers.get(META)!=tuple(META_COLUMNS):raise AuditError('입력 양식 구조의 열이 다릅니다.')
    groups={}
    for row in rows:
        if row['모듈']!=module or row['버전']!=SCHEMA:raise AuditError('입력 양식 모듈·버전이 다릅니다.')
        key=row['자료'];field=row['키'];kind=row['형태'];sheet=row['시트']
        if key not in allowed or kind!=allowed[key]['kind'] or row['자료형'] not in ('value','bool','list'):
            raise AuditError('허용되지 않은 입력 양식 구조입니다.')
        canonical=allowed[key]
        if sheet!=canonical['sheet']:raise AuditError('입력 시트 이름이 양식과 다릅니다.')
        if field not in allowed[key]['keys'] and key!='source_rows':raise AuditError('알 수 없는 입력 항목: '+str(field))
        if key!='source_rows':
            n=canonical['keys'].index(field)
            if row['표시']!=canonical['headers'][n] or row['자료형']!=canonical['types'][field]:
                raise AuditError('입력 열 표시·자료형이 양식과 다릅니다.')
        group=groups.setdefault(key,{'sheet':sheet,'kind':kind,'columns':[]})
        if (sheet,kind)!=(group['sheet'],group['kind']):raise AuditError('입력 양식 시트 중복')
        group['columns'].append((field,row['표시'],row['자료형']))
    if set(groups)!=set(allowed) or {x['sheet'] for x in groups.values()}!=set(tables):raise AuditError('입력 시트가 누락·추가됐습니다.')
    output={'schema':SCHEMA,'area':module,'approval':{}}
    for key,group in groups.items():
        columns=group['columns'];labels=[c[1] for c in columns];fields=[c[0] for c in columns]
        if len(set(labels))!=len(labels) or len(set(fields))!=len(fields):raise AuditError('중복 입력 열')
        if key!='source_rows' and set(fields)!=set(allowed[key]['keys']):raise AuditError('필수 입력 열이 누락됐습니다.')
        expected_headers={'항목','값'} if group['kind']=='dict' else set(labels)
        if set(actual_headers[group['sheet']])!=expected_headers:raise AuditError('빈 표를 포함해 입력 열이 양식과 다릅니다.')
        source=tables[group['sheet']]
        if group['kind']=='dict':
            values={r['항목']:r['값'] for r in source}
            if len(values)!=len(source) or set(values)!=set(labels):raise AuditError('입력 설정 항목 누락·중복')
            output[key]={field:decode(values[label],kind) for field,label,kind in columns}
        else:
            if any(set(r)-{'_row','_sheet'}!=set(labels) for r in source):raise AuditError('입력 표의 열이 다릅니다.')
            output[key]=[{field:decode(r[label],kind) for field,label,kind in columns} for r in source]
    if output.get('context')!=context:raise AuditError('입력 회사·기간이 현재 업무와 다릅니다.')
    if sha(path)!=source_metadata['sha256']:raise AuditError('읽는 중 입력자료가 변경됐습니다.')
    return output


def save_ledger(tables,path):
    from ledger import SCHEMA as LEDGER_SCHEMA
    workbook=_workbook()
    for name,columns in LEDGER_SCHEMA.items():
        sheet=workbook.create_sheet(name);sheet.append(columns)
        for n,row in enumerate(tables[name],2):
            for j,key in enumerate(columns,1):_cell(sheet,n,j,row.get(key))
    return _save(workbook,path)
