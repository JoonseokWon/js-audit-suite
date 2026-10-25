"""Read value-only ERP Excel without macros, formulas, caches or external connections."""
from copy import deepcopy
from common import AuditError, digest, read_xlsx, sha
from .schema import PROVENANCE


def read_original(path, event=None):
    tables,metadata=read_xlsx(path,schema=None,event=event)
    # The generic reader adds these metadata keys. Reject a colliding original header before it is mistaken for provenance.
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        for sheet in book:
            headers=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())
            if any(key in ('_row','_sheet') for key in headers):
                raise AuditError(sheet.title+': _row·_sheet는 원본 보존용 필드이므로 원본 열명으로 사용할 수 없습니다.')
    finally:book.close()
    if sha(path)!=metadata['sha256']:raise AuditError('읽는 동안 원본 파일이 바뀌었습니다.')
    for sheet,rows in tables.items():
        for row in rows:
            if any(not isinstance(key,str) or not key.strip() for key in row):
                raise AuditError(sheet+': 빈 열명·문자가 아닌 열명은 지원하지 않습니다.')
            if any(key in row for key in PROVENANCE):
                raise AuditError(sheet+': 원본 열명이 보존용 필드와 충돌합니다. '+', '.join(PROVENANCE))
    return tables,metadata


def stage_original(payload, path, tables, metadata, sheet, mapping, config):
    """An edited mapping remains unconfirmed. The caller must preserve/register the source file."""
    if sha(path)!=metadata['sha256']:raise AuditError('선택한 원본이 바뀌었습니다. 파일을 다시 불러오세요.')
    if sheet not in tables or not tables[sheet]:raise AuditError('원본 시트에 자료 행이 없습니다.')
    staged=deepcopy(payload)
    raw=[]
    for row in tables[sheet]:
        raw.append({k:v for k,v in row.items() if k not in ('_row','_sheet')} | {
            'source_file':metadata['filename'],'source_sha256':metadata['sha256'],
            'source_sheet':sheet,'source_row':row['_row']})
    staged['source_rows']=raw
    staged['column_mapping']=[{'target':target,'source':source} for target,source in mapping.items() if source]
    staged['config']=dict(config,source_origin='value_only_xlsx',source_filename=metadata['filename'],source_file_sha256=metadata['sha256'],
                          source_sheet=sheet,source_rows_sha256=digest(raw))
    staged['approval']={}
    return staged
