"""Independent UI-01/02 evidence: XML safety, bounded rows and exact raw recovery."""
import copy,json,tempfile,unittest,zipfile,xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from common import AuditError
from export_support.text import require_xml_text,display_text,validate_structure
from export_support.builder import create_builder
from windows_text import validate_text


class ReviewExportTests(unittest.TestCase):
    def test_new_text_rejects_every_forbidden_C0_noncharacter_and_surrogate_with_location(self):
        for cp in list(set(range(32))-{9,10,13})+[0xD800,0xDFFF,0xFFFE,0xFFFF]:
            value='가😀\n나'+chr(cp)
            with self.subTest(codepoint=cp),self.assertRaisesRegex(AuditError,f'위치 5.*2행 2열.*U\\+{cp:04X}'):
                require_xml_text(value,'결론')
            with self.subTest(editor=cp),self.assertRaises(ValueError):validate_text(value)

    def test_allowed_unicode_tabs_and_newlines_are_not_changed(self):
        value='한글😀\tA\nB\r\n끝\n'
        self.assertEqual(require_xml_text(value),value)
        self.assertEqual(display_text(value),value)
        self.assertEqual(validate_text(value),'한글😀\tA\nB\n끝\n')
        with self.assertRaisesRegex(AuditError,'기록.내용'):
            validate_structure({'내용':'x\x01'})

    def build(self,text,streaming):
        wb,add=create_builder({}, {'금액'},set(),set(),write_only=streaming)
        add('SYNTHETIC_ONLY',[{'코드':'0001','금액':125,'기록':text}])
        out=BytesIO();wb.save(out);wb.close();data=out.getvalue()
        with zipfile.ZipFile(BytesIO(data)) as archive:
            for name in archive.namelist():
                if name.endswith('.xml'):ET.fromstring(archive.read(name))
        return load_workbook(BytesIO(data),data_only=False)

    def test_25_40_253_254_line_boundaries_preserve_text_and_only_one_amount_in_both_builders(self):
        for count in (25,40,253,254):
            original='\n'.join(f'{i}: 검토😀\t원본' for i in range(count))
            for streaming in (False,True):
                with self.subTest(lines=count,streaming=streaming):
                    wb=self.build(original,streaming)
                    try:
                        rows=list(wb['SYNTHETIC_ONLY'].iter_rows(min_row=2,values_only=True))
                        self.assertEqual(''.join(row[2] or '' for row in rows),original)
                        self.assertEqual(sum(row[1] or 0 for row in rows),125)
                        self.assertEqual(sum(row[1] is not None for row in rows),1)
                        self.assertTrue(all((d.height or 0)<=409 for s in wb for d in s.row_dimensions.values()))
                        if count>25:self.assertIn('출력문단연결',wb.sheetnames)
                        self.assertFalse(any(c.data_type=='f' for s in wb for row in s for c in row))
                    finally:wb.close()

    def test_30000_character_boundary_emoji_formula_text_and_legacy_controls_are_visible_not_truncated(self):
        for length in (29999,30000,30001):
            original=('=한글😀\t'*6001)[:length]+'\x01\ufffe\n끝'
            for streaming in (False,True):
                with self.subTest(length=length,streaming=streaming):
                    wb=self.build(original,streaming)
                    try:
                        cells=[row[2] for row in wb['SYNTHETIC_ONLY'].iter_rows(min_row=2)]
                        self.assertEqual(''.join(c.value or '' for c in cells),display_text(original))
                        self.assertTrue(all(c.data_type=='s' for c in cells))
                        self.assertTrue(all(len((c.value or '').encode('utf-16-le'))//2<=30000 for c in cells))
                    finally:wb.close()

    def test_streaming_sheet_rollover_preserves_continuation_and_does_not_repeat_amount(self):
        wb,add=create_builder({}, {'금액'},set(),set(),write_only=True,max_rows=2)
        text='한글😀\n'*130
        add('연결',[{'금액':90000,'기록':text}]);out=BytesIO();wb.save(out)
        read=load_workbook(BytesIO(out.getvalue()))
        try:
            tables=[s for s in read if s.title=='연결' or s.title.startswith('연결_')]
            rows=[r for sheet in tables for r in sheet.iter_rows(min_row=2,values_only=True)]
            self.assertEqual(''.join(r[1] or '' for r in rows),text)
            self.assertEqual(sum(r[0] or 0 for r in rows),90000)
            self.assertTrue(all(sheet.max_row<=3 for sheet in tables))
        finally:read.close()

    def test_writeonly_serializes_height_without_retaining_every_ledger_row_dimension(self):
        wb,add=create_builder({}, {'금액'},set(),set(),write_only=True)
        sheet=add('원장',[{'금액':i,'근거':'a'*64} for i in range(1002)])
        self.assertLessEqual(len(sheet.row_dimensions),1)
        out=BytesIO();wb.save(out);read=load_workbook(BytesIO(out.getvalue()))
        try:
            self.assertGreater(read['원장'].row_dimensions[2].height,25)
            self.assertLessEqual(read['원장'].row_dimensions[2].height,409)
            self.assertEqual(read['원장'].max_row,1003)
        finally:read.close()

    def snapshot(self,folder):
        from workspace import Workspace
        context={'회사코드':'SYNTH01','회사명':'SYNTHETIC ONLY','시작일':'2026-01-01','기준일':'2026-12-31',
                 '회계체계':'K-IFRS','감사기준':'SYNTHETIC ONLY','작성자':'SYNTHETIC ONLY','더미업무':True}
        ws=Workspace(folder,create=context)
        ws.save_procedure({'id':'A04','title':'SYNTHETIC ONLY','preparer':'SYNTHETIC ONLY','conclusion':'before'})
        return ws,ws.snapshot()

    def test_complete_evidence_json_restores_legacy_long_body_AI_raw_and_references_in_both_modes(self):
        import exporter
        for streaming in (False,True):
            with self.subTest(streaming=streaming),tempfile.TemporaryDirectory() as folder:
                ws,snapshot=self.snapshot(Path(folder)/'SYNTHETIC_ONLY');head=ws.head
                original='검토😀\n'*254+'\x01\ufffe'+('긴 문장'*8000)
                snapshot['procedures'][0]['conclusion']=original
                snapshot['ai_drafts']=[dict(id='SYNTHETIC_DRAFT',procedure_id='A04',actor='SYNTHETIC ONLY',
                    provider='ollama',model='NO MODEL CALL',prompt_version='fixture',packet_hash='a'*64,basis='b'*64,at='fixture',
                    company_code='SYNTH01',raw_response=original,fields={},citations={},packet={'references':[dict(id='REF001',kind='raw_json',data={'문장':original})]} )]
                if streaming:
                    snapshot['latest']=[dict(module='원장·리드',id='R',source_id='S',result_hash='a'*64,
                        result=dict(detail={},metrics={},issues=[],tables={'SYNTHETIC_ROWS':[{'금액':1}]*20001}))]
                before=copy.deepcopy(snapshot);data=exporter.workbook_bytes(snapshot,{'modules':[]})
                self.assertEqual(snapshot,before);self.assertEqual(ws.head,head)
                read=load_workbook(BytesIO(data),data_only=False)
                try:
                    proof=json.loads(''.join(r[1] for r in read['증적JSON'].iter_rows(min_row=2,values_only=True)))
                    self.assertEqual(proof,before)
                    self.assertTrue(all((d.height or 0)<=409 for s in read for d in s.row_dimensions.values()))
                    self.assertTrue(any('⟦U+0001⟧' in str(c.value) for s in read for row in s for c in row))
                    self.assertFalse(any(c.data_type=='f' for s in read for row in s for c in row))
                finally:read.close()
