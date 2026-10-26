"""Application runtime exporter: complete, static, traceable Excel workpapers."""
from common import *
from io import BytesIO
from display_schema import TABLE_COLUMNS,MODULE_TABLES
RATES={'예상정산비율','연차자산배부율','급여자산배부율','DC자산배부율','DC율','과거손실률','가산조정','조정후손실률','확률','가중손실률','월유효할인율'}
QUANTITIES={'기말수량','원가단가','단위NRV','기초일수','발생일수','사용일수','소멸일수','잔여일수'}
NUMERIC={'총공급가액','기초부채','당기발생비용','당기부채결제','기말부채','기초선급','당기선급지급','당기선급상각','기말선급','전기까지누적귀속','기말까지누적귀속','기말전귀속참고액','공급가액','공급가액지급','후속지급','계산금액','앞금액','뒤금액','차이','허용오차','기초잔액','차변','대변','조정전','승인조정','조정후','기대','실제','금액','기대금액','원금','미수이자','기말원가','기말누계','당기상각','left','right','difference','tolerance','ordered_qty','delivered_qty','accepted_qty','billed_qty','billed_supply_amount','billed_total_amount','invoice_supply','ledger_net_revenue','기대행수','실제행수','기대수치합계','실제수치합계','기대절대합계','실제절대합계','기말잔액','연체일수','후속회수','잔여','발주수량','입고수량','송장수량','송장금액','발주단가','장부수량','실사수량'}
def workbook_bytes(snapshot,catalog):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    w=Workbook();w.remove(w.active)
    def regular_add(name,rows,headers=None):
        safe=re.sub(r'[\\/*?:\[\]]','_',name)[:31] or '자료';base=safe;n=1
        while safe in w.sheetnames:n+=1;safe=base[:27]+'_'+str(n)
        priority={'리드스케줄':['계정코드','계정명','리드코드','리드명','표시구분','기초잔액','차변','대변','조정전','승인조정','조정후','원본행'],'리드집계':['리드코드','리드명','표시구분','계정수','기초잔액','차변','대변','조정전','승인조정','조정후']}
        priority.update(TABLE_COLUMNS)
        headers=headers or next((v for k,v in priority.items() if name.endswith(k)),None) or list(dict.fromkeys(k for r in rows for k in r)) or ['내용']
        ws=w.create_sheet(safe);ws.append(headers)
        for r in rows:
            for j,k in enumerate(headers,1):
                v=r.get(k)
                if isinstance(v,(dict,list)):v=dumps(v)
                if k in NUMERIC|RATES|QUANTITIES|{'전망배수'} and v is not None and isinstance(v,str):
                    try:v=float(Decimal(v)) if '.' in v else int(v)
                    except (ValueError,InvalidOperation):pass
                if isinstance(v,str) and len(v)>30000:v=v[:29970]+' [전체 내용은 증적JSON 참조]'
                cell=ws.cell(ws.max_row+1 if j==1 else ws.max_row,j,v)
                if isinstance(v,str):cell.data_type='s'
                cell.font=Font(name='맑은 고딕',size=10);cell.alignment=Alignment(vertical='center')
                if isinstance(v,(int,float)):cell.number_format='0.00####%' if k in RATES else '#,##0;[Red](#,##0);0' if isinstance(v,int) else '#,##0.00####;[Red](#,##0.00####);0'
        for c in ws[1]:c.font=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='263F59')
        ws.sheet_view.showGridLines=False;ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
        for j,k in enumerate(headers,1):ws.column_dimensions[get_column_letter(j)].width=min(55,max(14,len(k)*1.7+2, max((len(str(ws.cell(i,j).value or '')) for i in range(2,min(ws.max_row,35)+1)),default=0)*1.1))
        if name in ('감사계획·중요성','감사절차 적용범위','위험주장대응','절차계획·수행연결','완료점검표','기준서식참고','구형F01보존','구형F01승인보존','조서표지','수행검토조서','검토의견','자료요청·수령','통합검토후보','분석필요자료','AI초안목록','AI최종내용승인','AI적용이력','AI근거참조'):
            import math,unicodedata
            for row in ws.iter_rows(min_row=2):
                lines=1
                for cell in row:
                    cell.alignment=Alignment(wrap_text=True,vertical='top');width=ws.column_dimensions[cell.column_letter].width
                    lines=max(lines,sum(max(1,math.ceil(sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in line)/max(8,width-2))) for line in str(cell.value or '').split('\n')))
                ws.row_dimensions[row[0].row].height=max(25,lines*16+6)
        return ws
    large=sum(len(rows) for run in snapshot['latest'] for rows in run['result']['tables'].values())>20000
    if large:
        from export_streaming import create_builder
        priority={'리드스케줄':['계정코드','계정명','리드코드','리드명','표시구분','기초잔액','차변','대변','조정전','승인조정','조정후','원본행'],'리드집계':['리드코드','리드명','표시구분','계정수','기초잔액','차변','대변','조정전','승인조정','조정후']}|TABLE_COLUMNS
        w.close();w,add=create_builder(priority,NUMERIC,RATES,QUANTITIES)
    else:add=regular_add
    eng=snapshot['engagement'];cover=[{'항목':k,'값':eng[k]} for k in ('회사명','회사코드','시작일','기준일','작성자','회계체계','감사기준')]+[{'항목':'업무ID','값':eng['id']},{'항목':'출력상태','값':snapshot['export_state']},{'항목':'범위','값':snapshot['scope']},{'항목':'더미표시','값':'SYNTHETIC ONLY' if eng.get('더미업무') else '실제업무 기록'},{'항목':'검토정책','값':'작성자와 별도 검토자 이름 입력. 인증된 서명 아님.'}]
    ai_statuses=[p.get('ai_status',{}) for p in snapshot['procedures'] if p.get('ai_status',{}).get('required')]
    if ai_statuses:cover.append({'항목':'AI 최종내용 승인','값':f"AI 적용 {len(ai_statuses)}개 / 현재 내용 승인 {sum(bool(s.get('approved')) for s in ai_statuses)}개 / 나머지는 미승인·변경 초안. 전체 감사완료·감사의견과 다름."})
    add('조서표지',cover)
    workflow=snapshot.get('workflow',{})
    if workflow.get('active'):
        profile=workflow['profile'];materiality=profile['materiality']
        plan_labels={'audit_framework':'적용 감사체계','audit_framework_reason':'감사체계 선택·적용가능성 판단 근거','form_framework':'서식계열','actor':'계획 입력자','at':'기록시각 UTC','revision':'계획 기록ID','benchmark_name':'중요성 기준지표','benchmark_amount':'사람 입력 기준지표 금액','materiality_rate':'사람 선택 중요성 비율','performance_rate':'사람 선택 수행중요성 비율','trivial_rate':'사람 선택 사소한 금액 비율','overall_materiality':'계산된 전체 중요성','performance_materiality':'계산된 수행중요성','clearly_trivial_threshold':'계산된 명백히 사소한 금액','calculation_policy':'중요성 계산식·반올림 정책'}
        link_labels={'procedure_id':'절차ID','applicable':'적용 여부','reason':'사람 입력 적용·제외 사유','id':'위험ID','title':'위험명','accounts':'관련 계정','assertions':'관련 주장','procedure_ids':'대응 절차ID','basis':'자료·실행 근거해시','evidence_ids':'등록 원본 증거ID','nature':'성격','timing':'시기','extent':'범위','assessment':'사람 위험평가·판단','risk_ids':'연결 위험ID','plan':'절차 계획','performed':'실제 수행내용','performed_date':'실제 수행일','performer':'실제 수행자','actor':'입력자','at':'기록시각 UTC','revision':'기록ID'}
        def labeled(rows,labels):return [{labels.get(k,k):v for k,v in row.items()} for row in rows]
        add('감사계획·중요성',[{'항목':'실제 회계체계(업무 기본정보)','기록':eng['회계체계']}]+[{'항목':plan_labels.get(k,k),'기록':v} for k,v in profile.items() if k not in ('materiality','scope')]+[{'항목':plan_labels.get(k,k),'기록':v} for k,v in materiality.items()]+[{'항목':'서식 참고의 범위','기록':'서식계열은 참고 ZIP의 분류이며 실제 회계체계 또는 감사기준의 적용 판단을 대신하지 않습니다. 실제 회계체계는 업무 기본정보에 기록된 사람의 선택을 표시합니다.'}])
        add('감사절차 적용범위',labeled(profile['scope'],link_labels))
        add('위험주장대응',labeled(workflow['risks'],link_labels))
        add('절차계획·수행연결',labeled(workflow['procedure_links'],link_labels))
        add('완료점검표',labeled(workflow['rows'],{'code':'점검코드','procedure_id':'절차ID','risk_id':'위험ID','status':'상태','message':'점검내용'}),['점검코드','절차ID','위험ID','상태','점검내용'])
        add('기준서식참고',snapshot.get('form_references',[]))
    if snapshot.get('legacy_f01_documents'):
        add('구형F01보존',snapshot['legacy_f01_documents'])
        add('구형F01승인보존',snapshot.get('legacy_f01_approvals',[]))
    add('검사실행목록',[{'모듈':r['module'],'실행ID':r['id'],'시각':r['at'],'입력자료ID':r['source_id'],'차단':r['result']['blocking'],'검토':r['result']['review'],'결과해시':r['result_hash']} for r in snapshot['all_runs']])
    add('모듈TB연결',snapshot['cross']);add('증거등록부',snapshot['sources'])
    # Nested rule detail remains in module tables and the complete evidence JSON.
    # Keep the review register readable without expanding each row by that JSON.
    add('통합검토후보',snapshot.get('risk_register',[]),['모듈','규칙ID','후보ID','우선순위','사유','원본행','후속절차','연결상태','조서ID','판단상태','실행ID','자료ID','원문SHA256','후보키'])
    add('분석필요자료',snapshot.get('analytics_source_requirements',[]))
    request_labels={'id':'요청번호','title':'요청자료','requested_from':'요청처','due_date':'기한','state':'상태','note':'요청범위·처리내용','source_ids':'연결 증거번호','updated_at':'기록시각 UTC'}
    add('자료요청·수령',[{label:r.get(key,'') for key,label in request_labels.items()} for r in snapshot.get('requests',[])],list(request_labels.values()))
    add('수행검토조서',[{'절차ID':p['id'],'제목':p['title'],'상태':p['current_state'],'AI내용상태':p.get('ai_status',{}).get('state','AI 미적용'),'작성자':p['preparer'],'실제수행일':p['performed_date'],'검토자':p['reviewer'],'검토기록시각UTC':p['reviewed_at']} for p in snapshot['procedures']]);add('검토의견',snapshot['notes'])
    if snapshot.get('ai_drafts'):
        legacy_drafts=set(workflow.get('legacy_f01',{}).get('draft_ids',[]))
        def draft_purpose(d):
            return '구형 F01 보존 / '+str(workflow.get('legacy_f01',{}).get('purpose_code') or '사람 목적분류 대기') if d['id'] in legacy_drafts else d.get('packet',{}).get('procedure',{}).get('purpose_code','기존 절차 정의')
        add('AI초안목록',[{'초안ID':d['id'],'절차ID':d['procedure_id'],'절차목적':draft_purpose(d),'요청자':d['actor'],'모델제공자':d['provider'],'모델':d['model'],'프롬프트버전':d['prompt_version'],'근거패킷SHA256':d['packet_hash'],'자료실행버전':d['basis'],'생성시각UTC':d['at'],'원문응답':d['raw_response'],'적용정책':'AI 응답 보존은 적용·사람 승인·검토완료와 다름. 전체 원문은 증적JSON 참조.'} for d in snapshot['ai_drafts']])
        add('AI근거참조',[{'초안ID':d['id'],'절차ID':d['procedure_id'],'절차목적':draft_purpose(d),'회사코드':d['company_code'],'참조ID':r['id'],'참조종류':r['kind'],'자료ID':r.get('source_id',''),'원문SHA256':r.get('sha256',''),'실행ID':r.get('run_id',''),'결과SHA256':r.get('result_hash',''),'내용':r['data']} for d in snapshot['ai_drafts'] for r in d['packet']['references']])
        draft_index={d['id']:d for d in snapshot['ai_drafts']}
        add('AI적용이력',[{**a,'절차목적':draft_purpose(draft_index[a['draft_id']]) if a.get('draft_id') in draft_index else '초안 근거 연결 미확인'} for a in snapshot.get('ai_applications',[])])
        add('AI최종내용승인',[{'절차ID':p['id'],'현재상태':p.get('ai_status',{}).get('state','AI 미적용'),'승인필요':p.get('ai_status',{}).get('required',False),'현재승인유효':p.get('ai_status',{}).get('approved',False),'현재내용SHA256':p.get('ai_status',{}).get('content_hash',''),'현재자료실행버전':snapshot['basis'],'승인기록':p.get('ai_status',{}).get('approval') or '미승인','범위':'사람이 저장한 AI 적용 내용의 승인. 별도 감사조서 검토통제·감사의견을 대신하지 않음.'} for p in snapshot['procedures'] if p.get('ai_status',{}).get('required')])
    labels={'id':'조서·절차ID','title':'제목','current_state':'현재상태','purpose':'목적·주장','risk':'위험·관련 기준','scope':'모집단·선정·검사 범위','performed':'실제 수행내용·근거·예외 조사','conclusion':'판단·결론·추가절차','preparer':'작성자','performed_date':'실제 수행일','reviewer':'검토자','review_note':'검토 내용·범위','reviewed_at':'검토 기록시각 UTC','excluded_reason':'범위 제외 사유','basis':'연결된 자료·실행 버전 해시'}
    source_names={s['id']:s['name']+' / '+s['file']+' / SHA256 '+s['sha256'] for s in snapshot['sources']}
    import math,unicodedata
    for p in snapshot['procedures']:
        rows=[{'항목':label,'기록':p.get(key,'')} for key,label in labels.items()]
        rows.append({'항목':'연결 증거 사본','기록':'\n'.join(source_names.get(s,s) for s in p.get('evidence_ids',[]))})
        ai=p.get('ai_status',{})
        if ai.get('required'):
            rows.extend([{'항목':'AI 최종내용 승인 상태','기록':ai['state']},{'항목':'AI 초안·모델','기록':ai['draft_id']+' / '+ai.get('provider','')+' / '+ai.get('model','')},{'항목':'AI 적용 내용 SHA256','기록':ai.get('content_hash','')},{'항목':'사람 최종내용 승인 기록','기록':dumps(ai.get('approval')) if ai.get('approval') else '미승인 초안. 실제 수행·검토·감사완료로 사용하지 마세요.'}])
        if eng.get('더미업무'):rows.insert(0,{'항목':'더미 표시','기록':'SYNTHETIC ONLY / 가상 수행·검토 상태. 실제 감사증거·감사 결론 아님.'})
        module=PROCEDURE_MODULES.get(p['id']) if p['id']!='F01' else None
        linked=next(((i,r) for i,r in enumerate(snapshot['latest'],1) if r['module']==module),None)
        if linked:
            index,r=linked
            rows.extend([{'항목':'자동계산 실행·자료','기록':r['id']+' / '+r['source_id']+' / 결과 SHA256 '+r['result_hash']},{'항목':'자동계산 정책·가정','기록':'\n'.join(k+': '+str(v) for k,v in r['result']['detail'].items())},{'항목':'상세 계산·원장대사표','기록':' / '.join((str(index)+'_'+name)[:31] for name in r['result']['tables'] if not name.startswith('입력_'))}])
            rows.extend({'항목':'계산값: '+k,'기록':str(v)} for k,v in r['result']['metrics'].items())
            rows.extend({'항목':'상세표: '+name,'기록':(str(index)+'_'+name)[:31]} for name in MODULE_TABLES.get(module,[]) if name in r['result']['tables'])
            rows.extend({'항목':'독립 TB: '+x['지표'],'기록':x['상태']+' / 원장 '+str(x['기대'])+'원 / 계산 '+str(x['실제'])+'원 / 차이 '+str(x['차이'])+'원 / 계정 '+x['사유']} for x in snapshot['cross'] if x['모듈']==module)
        ws=add('조서_'+p['id'],rows,['항목','기록']);ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=105
        for row in ws.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
            line_count=sum(max(1,math.ceil(sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in line)/95)) for line in str(row[1].value or '').split('\n'))
            ws.row_dimensions[row[0].row].height=max(28,line_count*16+8)
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_title_rows='1:1'
    ps={p['id']:p for p in snapshot['procedures']}
    automated={'C04','C05','P07','A01','A02','A04','A06','A08','A10','A12','A15','A07','A11','A16','A09'}
    add('전체절차목록',[{'절차ID':m['id'],'절차명':m['name'],'그룹':m['group'],'현재지원':'제한된 자동검사 + 수동조서' if m['id'] in automated else '수동 수행·증거·검토 조서','향후설계참고':m['automatable'],'수동절차':m['human_required'],'기준참조':m['standards'],'조서상태':ps.get(m['id'],{}).get('current_state','미착수')} for m in catalog['modules']])
    add('전체예외',[{'실행ID':r['id'],'모듈':r['module'],'예외조서ID':'X-'+r['id']+'-'+str(i),**x} for r in snapshot['latest'] for i,x in enumerate(r['result']['issues'])])
    for index,r in enumerate(snapshot['latest'],1):
        preferred=MODULE_TABLES.get(r['module'],[])
        for name in sorted(r['result']['tables'],key=lambda n:preferred.index(n) if n in preferred else len(preferred)):add(str(index)+'_'+name,r['result']['tables'][name])
    add('변경이력',[{'순번':i+1,**e} for i,e in enumerate(snapshot['events'])])
    if large:
        from export_streaming import chunks
        add('증적JSON',({'순서':i,'JSON조각':part} for i,part in enumerate(chunks(snapshot),1)),['순서','JSON조각'])
    else:
        proof=dumps(snapshot);add('증적JSON',[{'순서':i//28000+1,'JSON조각':proof[i:i+28000]} for i in range(0,len(proof),28000)])
    from openpyxl.utils import quote_sheetname
    for pid in PROCEDURE_MODULES:
        if large:break
        if '조서_'+pid not in w.sheetnames:continue
        for row in w['조서_'+pid].iter_rows(min_row=2):
            if str(row[0].value).startswith('상세표: ') and row[1].value in w.sheetnames:
                row[1].hyperlink='#'+quote_sheetname(row[1].value)+'!A1';row[1].font=Font(name='맑은 고딕',size=10,color='0563C1',underline='single')
    out=BytesIO();w.save(out);w.close();return out.getvalue()
def export(workspace,path,catalog,final=False):
    p=Path(path).resolve()
    if p.suffix.lower()!='.xlsx':raise AuditError('출력 확장자는 .xlsx여야 합니다.')
    protected=[(workspace.root/s['file']).resolve() for s in workspace.sources.values()]+[Path(s['original']).resolve() for s in workspace.sources.values() if s.get('original')]+[workspace.dbpath,workspace.anchor]
    if p in protected or p.parent==workspace.root/'원본':raise AuditError('원본/업무 DB는 덮어쓸 수 없습니다.')
    before=workspace.head;snap=workspace.snapshot(final);data=workbook_bytes(snap,catalog)
    workspace.refresh()
    if workspace.head!=before or workspace.verify_sources():raise AuditError('내보내기 중 자료·이력 변경')
    atomic(p,data);return snap
