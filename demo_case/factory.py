"""Create a new workbook-only walkthrough, preserving all original resources."""
from io import BytesIO
from pathlib import Path
import hashlib
from openpyxl import load_workbook
from common import AuditError, atomic, digest, dumps
from .definition import BASES, EXPECTED, INPUTS, SCHEMA, TARGET_PID, steps


def _bytes(module):
    import runtime
    relative, expected = BASES[module]
    path = runtime.configure()/'sector_resources'/relative
    if not path.is_file() or path.stat().st_size > 25*1024**2:
        raise AuditError('시연용 내장 자료가 없거나 크기가 맞지 않습니다.')
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != expected:
        raise AuditError('시연용 내장 자료의 원본 해시가 다릅니다.')
    return raw


def _changed(raw, key):
    book = load_workbook(BytesIO(raw))
    try:
        sheet = book['총계정원장' if key == 'ledger_error' else 'invoices']
        headers = [cell.value for cell in sheet[1]]
        column = headers.index('전표ID' if key == 'ledger_error' else 'invoice_id')+1
        ident = 'J-S1' if key == 'ledger_error' else 'SI01'
        selected = [n for n in range(2, sheet.max_row+1) if sheet.cell(n, column).value == ident]
        if len(selected) != (3 if key == 'ledger_error' else 1):
            raise AuditError('고정 시연의 원본 거래 행 범위가 다릅니다.')
        if key == 'sales_amount_error':
            sheet.cell(selected[0], headers.index('supply_amount')+1).value = 90001
        else:
            for row in reversed(selected):sheet.delete_rows(row)
        output = BytesIO();book.save(output)
        return output.getvalue()
    finally:book.close()


def create_demo(folder):
    """Return (Workspace, manifest); no people, approvals, AI calls, or export."""
    from workspace import Workspace
    from .manifest import checked_root, load_manifest
    root = checked_root(folder)
    if root.exists() and any(root.iterdir()):
        raise AuditError('시연은 새 빈 업무 폴더에 만드세요. 기존 업무는 보존합니다.')
    raw = {module:_bytes(module) for module in BASES}
    ws = Workspace(root, create={
        '회사코드':'JS_ELEC', '회사명':'JS전자 · AI 조서 시연',
        '시작일':'2026-01-01', '기준일':'2026-12-31', '회계체계':'K-IFRS / 합성',
        '감사기준':'합성 시연 · 실제 감사수행 아님',
        '작성자':'예제 자료 생성기 · 실제 감사작성자 아님', '더미업무':True,
    })
    entries = {}
    for key, (module, name) in INPUTS.items():
        data = raw[module] if key.endswith('_clean') else _changed(raw[module], key)
        relative = '시연입력/'+name
        atomic(root/relative, data)
        entries[key] = dict(module=module, path=relative,
            bytes=len(data), sha256=hashlib.sha256(data).hexdigest())
    for key in ('ledger_error', 'sales_error'):
        entry = entries[key];ws.execute(entry['module'], root/entry['path'])
    payload = dict(schema=SCHEMA, workspace_id=ws.engagement['id'],
        company_code='JS_ELEC', synthetic=True, target_pid=TARGET_PID,
        inputs=entries, expected=dict(EXPECTED), steps=steps(),
        scope='SYNTHETIC ONLY / 사람 판단·승인·검토와 실제 모델 호출은 자동 생성하지 않음')
    atomic(root/'시연_자료목록.json', dumps(payload).encode('utf-8'))
    ws._append('demo_case_setup', dict(schema=SCHEMA, workspace_id=ws.engagement['id'],
        manifest_hash=digest(payload)))
    ws.refresh()
    guide = ('JS전자 로컬 AI 조서 시연 — SYNTHETIC ONLY\n\n'
        '입력 원장에는 균형전표 J-S1(차변·대변 각각 99,000원) 전체가 누락되어 있습니다.\n'
        '매출 입력에는 SI01 청구행(공급가액 90,000원, 세액 9,000원)이 누락되어 있습니다.\n'
        '정상 매출은 100,000원입니다. 이 값은 프로그램 결과에서 역산한 정답이 아닙니다.\n'
        '실제 고객 자료·감사증거·감사결론이 아닌 가상 사례입니다.\n'
        '사람 확인·승인·수행·판단·검토는 비어 있습니다. AI는 별도로 실행해야 합니다.\n'
        '시연입력은 고정 대조자료입니다. 입력을 고칠 때는 사본을 저장해 새 자료로 등록하세요.\n\n'
        +'\n\n'.join(s['title']+'\n'+s['detail'] for s in payload['steps'])+'\n')
    atomic(root/'시연_안내.txt', guide.encode('utf-8'))
    return ws, load_manifest(root)
