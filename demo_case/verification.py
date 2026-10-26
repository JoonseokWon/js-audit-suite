"""Opt-in deterministic walkthrough test; no provider, network, or actual model."""
from pathlib import Path
from openpyxl import load_workbook
from common import AuditError, digest, dumps
from . import create_demo, case_status


def _mock(packet):
    from ai_workpapers import FIELDS, PROMPT_VERSION
    ref = next(r['id'] for r in packet['references'] if r['kind']=='run_summary')
    fields = {key:'SYNTHETIC ONLY 모의 초안입니다. 원문·자동검사 범위를 사람이 확인해야 합니다. ['+ref+']'
              for key in FIELDS}
    fields['performed'] = ('프로그램 자동검사와 발췌 원문에 관한 모의 기록입니다. '
        '사람의 실사·조회·증빙 확인은 미수행이며 별도 확인이 필요합니다. ['+ref+']')
    fields['conclusion'] = '사람 검토 전 모의 초안입니다. 차이·미확인 범위는 추가확인과 사람의 판단이 필요합니다. ['+ref+']'
    citations = {key:[ref] for key in FIELDS}
    exceptions = [r for r in packet['references'] if r['kind']=='exception']
    if exceptions:
        fields['performed'] += ' 선택된 예외코드: '+', '.join(str(r['data']['code']) for r in exceptions)+'. '
        fields['performed'] += ''.join('['+r['id']+']' for r in exceptions)
        citations['performed'] += [r['id'] for r in exceptions]
    context = dict(company_code=packet['engagement']['company_code'],
        procedure_id=packet['procedure']['id'], basis=packet['basis'])
    return dict(fields=fields, citations=citations, provider='ollama',
        model='SYNTHETIC-ONLY-NO-MODEL-CALL', prompt_version=PROMPT_VERSION,
        packet_hash=digest(packet), **context,
        raw_response=dumps(dict(fields=fields,citations=citations,context=context)))


def verify(directory):
    """Run only in a fresh directory; synthetic transitions never reach live data."""
    import json, runtime, exporter
    from workspace import Workspace
    folder = Path(directory)
    if folder.exists():raise AuditError('흐름 검증은 새 폴더를 사용하세요.')
    folder.mkdir(parents=True)
    checks = []
    def check(condition,name):
        if not condition:raise AssertionError(name)
        checks.append('v22_demo_'+name)
    ws, manifest = create_demo(folder/'업무')
    initial = ws.latest();initial_ids = {r['id'] for r in initial.values()}
    check(ws.context['회사코드']=='JS_ELEC' and ws.engagement['더미업무'],'isolated_JS_electronics')
    check(not ws.procedures and not ws.ai_drafts and not ws.ai_approvals and not ws.requests,
          'user_fixture_has_no_human_content_or_approval')
    check(initial['원장·리드']['result']['blocking']==9,'balanced_missing_journal_is_detected')
    check({'UNBILLED_ACCEPTED','MISSING_INVOICE','LEDGER_WITHOUT_INVOICE'} ==
          {i['code'] for i in initial['매출']['result']['issues']},'missing_invoice_is_detected')
    check(initial['매출']['result']['metrics']['순매출']==100000,'independent_revenue_100000')
    packet = ws.ai_packet('A04')
    check({'raw_xlsx_row','exception','run_summary','ledger_reconciliation'} <=
          {r['kind'] for r in packet['references']},'packet_contains_raw_results_and_exceptions')
    did = ws.record_ai_draft(packet,_mock(packet),'SYNTHETIC ONLY 요청자')
    check(not ws.procedures and not ws.ai_approvals,'mock_generation_does_not_apply_or_approve')
    ws.apply_ai_draft(did,'SYNTHETIC ONLY 작성자')
    check(not ws.ai_status('A04')['approved'] and ws.procedures['A04']['performed_date']==''
          and ws.procedures['A04']['reviewer']=='','mock_application_preserves_pending_human_fields')
    edited = dict(ws.procedures['A04'])
    note = 'SYNTHETIC ONLY 수정 시험: 원장·청구 누락을 추가확인해야 하며 실제 감사판단은 수행하지 않았습니다.'
    edited['conclusion'] += '\n'+note;ws.save_procedure(edited)
    ws.approve_ai_content('A04','SYNTHETIC ONLY 승인 모사',
        '가상 상태전이만 시험. 실제 사람의 내용 승인·감사결론이 아님.')
    check(ws.ai_status('A04')['approved'] and ws.procedures['A04']['state']=='작성중',
          'synthetic_content_approval_does_not_complete_review')
    catalog = json.loads((runtime.configure()/'procedure_catalog.json').read_text(encoding='utf-8'))
    target = folder/'JS전자_합성_검토중_조서.xlsx'
    snapshot = exporter.export(ws,target,catalog,False)
    book = load_workbook(target,data_only=False)
    try:
        check(snapshot['export_state']=='검토중 초안' and 'AI근거참조' in book.sheetnames
              and 'AI최종내용승인' in book.sheetnames,'excel_has_current_draft_and_trace_sheets')
        check(any(note in str(r[1]) for r in book['조서_A04'].iter_rows(values_only=True)),
              'excel_preserves_edited_conclusion')
        check(all(c.data_type!='f' for s in book for row in s for c in row),'excel_contains_no_formulas')
    finally:book.close()
    ws = Workspace(ws.root)
    check(ws.ai_status('A04')['approved'] and note in ws.procedures['A04']['conclusion'],
          'reopen_preserves_edit_and_synthetic_approval')
    changed = dict(ws.procedures['A04']);changed['conclusion']+='\nSYNTHETIC ONLY 추가 문장 수정'
    ws.save_procedure(changed)
    check(not ws.ai_status('A04')['approved'] and not case_status(ws)['needs_rerun'],
          'text_edit_needs_reapproval_without_unnecessary_new_model')
    ws.approve_ai_content('A04','SYNTHETIC ONLY 승인 모사','가상 수정 후 재승인 상태전이 시험')
    for key in ('ledger_clean','sales_amount_error'):
        entry=manifest['inputs'][key];run=ws.execute(entry['module'],entry['path'])
    check('LEDGER_AMOUNT_MISMATCH' in {i['code'] for i in run['result']['issues']},
          'one_won_difference_remains_visible')
    for key in ('ledger_clean','sales_clean'):
        entry=manifest['inputs'][key];ws.execute(entry['module'],entry['path'])
    check(all(r['result']['blocking']+r['result']['review']==0 for r in ws.latest().values()),
          'corrected_inputs_remove_current_calculation_exceptions')
    check(initial_ids <= set(ws.runs),'correction_preserves_old_error_runs')
    check(not ws.ai_status('A04')['approved'] and case_status(ws)['needs_rerun'],
          'input_and_run_changes_require_fresh_evidence_and_approval')
    try:ws.apply_ai_draft(did,'SYNTHETIC ONLY')
    except AuditError:checks.append('v22_demo_old_mock_draft_rejected_after_rerun')
    else:raise AssertionError('Old evidence must not be reapplied')
    fresh = ws.ai_packet('A04')
    check(fresh['basis']!=packet['basis'],'fresh_packet_has_new_data_basis')
    check(not ws.workflow_locked and not ws.procedures['A04']['reviewer'],
          'audit_finalization_and_independent_review_are_not_created')
    from procedures import task_definitions
    pending,_ = create_demo(folder/'미적용_수동변경_업무')
    old_packet = pending.ai_packet('A04')
    old_id = pending.record_ai_draft(old_packet,_mock(old_packet),'SYNTHETIC ONLY 요청자')
    manual = next(p for p in task_definitions(pending,catalog) if p['id']=='A04')
    manual.update(preparer='SYNTHETIC ONLY 작성자',scope='SYNTHETIC ONLY 수동 범위 변경')
    pending.save_procedure(manual)
    check(case_status(pending)['needs_rerun'],'unapplied_draft_manual_edit_requires_fresh_packet')
    try:pending.apply_ai_draft(old_id,'SYNTHETIC ONLY')
    except AuditError:checks.append('v22_demo_unapplied_old_fingerprint_cannot_be_applied')
    else:raise AssertionError('Manual edit must invalidate the unapplied draft fingerprint')
    return dict(ok=True,checks=checks,actual_model_calls=0,actual_openai_calls=0,
        workspace=str(ws.root),export=str(target),target_pid='A04',
        human_approval_scope='SYNTHETIC ONLY mock state transitions; no actual person approval or audit judgment',
        scope='Deterministic JS전자 workflow with explicit mock response. Actual inference must be separately tested.')
