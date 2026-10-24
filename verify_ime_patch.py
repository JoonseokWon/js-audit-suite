"""Source/frozen owned-window integration check. Never activate user windows."""
from pathlib import Path
import json,sys,shutil

def self_test(report_path,workspace_path):
    from common import atomic,VERSION
    from workspace import Workspace
    from app import App,CATALOG
    from windows_text import WindowsText,NATIVE_CLASS
    from openpyxl import load_workbook
    import tkinter as tk
    import exporter
    out=Path(report_path).resolve().parent
    out.mkdir(parents=True,exist_ok=True)
    target=out/'JS전자_한글편집_시험업무'
    if target.exists():raise RuntimeError('Preserve previous test evidence')
    shutil.copytree(workspace_path,target)
    checks=[];root=tk.Tk();root.withdraw()
    def check(value,name):
        if not value:raise AssertionError(name)
        checks.append(name)
    app=None
    try:
        app=App(root,sector_home=out/'unused_sectors')
        app.ws=Workspace(target);app.reset_workspace_view();app.refresh()
        app.load_procedure('A04');root.update_idletasks()
        original=app.ws.procedures['A04']['conclusion']
        draft_ids=list(app.ws.ai_drafts)
        check(len(app.ws.latest())==15,'15_existing_module_results_preserved')
        check(len(draft_ids)==1,'previous_actual_local_AI_draft_preserved')
        check(all(isinstance(w,WindowsText) for w in app.proc_text.values()),'all_five_workpaper_fields_use_native_control')
        check(all(w.get('1.0','end').strip()==app.ws.procedures['A04'][k] for k,w in app.proc_text.items()),'all_five_saved_fields_load_exactly')
        w=app.proc_text['conclusion']
        suffix='\n한글 편집 저장검사: 검토·수정·받침 값·공백·emoji 😀\n다음 줄의 원장 대사 근거를 확인한다.'
        w.insert('end',suffix)
        check(app.dirty(),'native_edit_marks_workpaper_dirty')
        app.save_procedure()
        check(app.ws.procedures['A04']['conclusion']==original+suffix,'native_edit_saved_without_loss')
        check(not app.ws.ai_status('A04')['approved'],'content_remains_unapproved_until_person_approves')
        app.set_busy(True)
        before=w.get('1.0','end');w.insert('end','금지된 변경')
        check(w.get('1.0','end')==before,'busy_native_input_is_readonly')
        app.set_busy(False)
        for scale in (100,125,150):
            app.set_ui_zoom(scale);root.update_idletasks()
            check(w.get('1.0','end').strip()==original+suffix,'zoom_'+str(scale)+'_preserves_saved_text')
        app.open_ai_approval();win=app.ai_approval_win;win.withdraw()
        check(isinstance(win.note,WindowsText),'approval_note_uses_native_control')
        memo='더미 검토 메모: 한글·영문과 줄바꿈\n실제 감사 승인 아님.'
        win.note.insert('1.0',memo)
        check(win.note.get('1.0','end').strip()==memo,'approval_memo_unicode_roundtrip')
        check(not app.ws.ai_status('A04')['approved'],'opening_and_editing_memo_does_not_approve')
        win.destroy()
        xlsx=out/'JS전자_한글편집검사_미승인초안.xlsx'
        exporter.export(app.ws,xlsx,CATALOG,False)
        book=load_workbook(xlsx,read_only=True,data_only=True)
        check({'조서_A04','AI초안목록','AI근거참조','AI적용이력','AI최종내용승인'}<=set(book.sheetnames),'excel_contains_workpaper_and_AI_lineage')
        check(any(cell.value==original+suffix for row in book['조서_A04'] for cell in row),'excel_preserves_edited_Korean_exactly')
        book.close()
        reopened=Workspace(target)
        check(reopened.procedures['A04']['conclusion']==original+suffix,'reopened_workspace_preserves_Korean')
        check(list(reopened.ai_drafts)==draft_ids,'AI_original_draft_not_replaced_by_test_text')
        record={'ok':True,'version':VERSION,'frozen':bool(getattr(sys,'frozen',False)),
                'native_class':NATIVE_CLASS,'checks':checks,'actual_hardware_ime_test':False,
                'actual_model_calls':0,'human_approval_simulated':False,'excel':str(xlsx),
                'window_visibility':'all owned windows withdrawn; no foreground activation'}
        atomic(report_path,json.dumps(record,ensure_ascii=False,indent=2).encode('utf-8'))
        print(json.dumps(record,ensure_ascii=False));return 0
    finally:
        if app is not None:app.shutdown_ai()
        root.destroy()
