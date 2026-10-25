"""Local xlsx IO: immutable snapshot, strict schema, safe atomic exports."""
from io import BytesIO
from pathlib import Path
import hashlib
import os
import tempfile
import zipfile
from decimal import Decimal
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .engine import SCHEMA, Row, InputBook, Issue, Result, Cancelled


class InputError(Exception):
    def __init__(self, issues, source='', book=None):
        self.issues = issues
        self.source = source
        self.book = book
        super().__init__('; '.join(x.message for x in issues[:4]))


def read_input(path, cancel=None):
    p = Path(path)
    if p.suffix.lower() != '.xlsx':
        raise InputError([Issue('FILE_TYPE', 'Workbook', '', '.xlsx 입력만 지원')], str(p))
    if p.stat().st_size > 32 * 1024 * 1024:
        raise InputError([Issue('FILE_SIZE', 'Workbook', '', '입력은 32MiB 이하')], str(p))
    data = p.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    issues = []
    with zipfile.ZipFile(BytesIO(data)) as z:
        infos = z.infolist()
        if len(infos) > 2000 or sum(i.file_size for i in infos) > 128 * 1024 * 1024:
            raise InputError([Issue('ZIP_LIMIT', 'Workbook', '', '압축 해제 크기/항목 제한 초과')], str(p))
        if any(i.filename.startswith('xl/externalLinks/') or 'vbaProject' in i.filename for i in infos):
            raise InputError([Issue('EXTERNAL_LINK', 'Workbook', '', '매크로/외부링크 입력 금지')], str(p))
    wb = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    sheets = {}
    total = 0
    try:
        for ws in wb.worksheets:
            if cancel and cancel.is_set():
                raise Cancelled()
            if ws.title not in SCHEMA:
                issues.append(Issue('EXTRA_SHEET', ws.title, '', '정의되지 않은 시트'))
                continue
            if ws.sheet_state != 'visible' or ws.merged_cells.ranges:
                issues.append(Issue('SHEET_LAYOUT', ws.title, '', '숨김 시트/병합 셀 입력 금지'))
            columns = SCHEMA[ws.title]
            if ws.max_column != len(columns) or tuple(c.value for c in ws[1]) != columns:
                issues.append(Issue('HEADER', f'{ws.title}!1', '', '정확한 헤더/열 순서를 입력하세요'))
                continue
            if ws.max_row > 100001:
                issues.append(Issue('ROW_LIMIT', ws.title, '', '행 제한 100000 초과'))
                continue
            rows = []
            for cells in ws.iter_rows(min_row=2):
                if cancel and cancel.is_set():
                    raise Cancelled()
                total += 1
                if total > 100000:
                    raise InputError([Issue('ROW_LIMIT', 'Workbook', '', '전체 행 제한 100000 초과')], str(p))
                if all(c.value is None for c in cells):
                    continue
                for c in cells:
                    if c.data_type in ('f', 'e'):
                        issues.append(Issue('FORMULA_OR_ERROR', f'{ws.title}!{c.row}:{c.column}', '', '수식/오류 셀 금지: 검토된 값으로 입력'))
                    if isinstance(c.value, str) and len(c.value) > 8192:
                        issues.append(Issue('TEXT_LIMIT', f'{ws.title}!{c.row}', '', '셀 텍스트 8192자 초과'))
                rows.append(Row(ws.title, cells[0].row, {k: c.value for k, c in zip(columns, cells)}))
            sheets[ws.title] = rows
        for name in SCHEMA.keys() - sheets.keys():
            issues.append(Issue('MISSING_SHEET', name, '', '필수 시트 누락 또는 헤더 오류'))
        if issues:
            raise InputError(issues, str(p.resolve()), InputBook(sheets, str(p.resolve()), digest))
        return InputBook(sheets, str(p.resolve()), digest)
    finally:
        wb.close()


def safe_value(value):
    if isinstance(value, Decimal):
        return format(value, 'f')
    if isinstance(value, str):
        # Escape formula-trigger characters even after whitespace/control prefixes.
        plain = value.lstrip(' \t\r\n\x00\x0b\x0c')
        value = ''.join(c if ord(c) >= 32 or c in '\t\r\n' else '�' for c in value)
        if plain.startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
            return "'" + value
    return value


def append_safe(ws, values):
    ws.append([safe_value(v) for v in values])
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'
            cell.quotePrefix = True


def style(ws):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for c in ws[1]:
        c.font = Font(name='맑은 고딕', bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='15395B')
    for column in ws.columns:
        letter = column[0].column_letter
        width = min(75, max(14, max((len(str(c.value or '')) for c in column), default=0) + 2))
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name='맑은 고딕', size=10)
            c.alignment = Alignment(vertical='top', wrap_text=True)


def write_atomic(wb, path, source=None):
    p = Path(path)
    if p.suffix.lower() != '.xlsx':
        raise ValueError('저장은 .xlsx 확장자로 지정하세요')
    if source:
        original = Path(source)
        if p.resolve() == original.resolve() or p.exists() and original.exists() and os.path.samefile(p, original):
            raise ValueError('입력 원본은 덮어쓸 수 없습니다')
    if p.exists():
        raise FileExistsError('기존 파일을 보존합니다. 새 파일명을 지정하세요')
    p.parent.mkdir(parents=True, exist_ok=True)
    handle, temp = tempfile.mkstemp(prefix='.cashflow-', suffix='.xlsx', dir=p.parent)
    os.close(handle)
    try:
        wb.save(temp)
        # Windows rename fails if a destination appears concurrently; no overwrite.
        os.rename(temp, p)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
        wb.close()


def save_input(book, path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, columns in SCHEMA.items():
        ws = wb.create_sheet(name)
        append_safe(ws, columns)
        for row in book.sheets.get(name, []):
            values = []
            for k in columns:
                value = row.values.get(k)
                numeric = k in ('opening', 'closing', 'profit_effect', 'balance_effect', 'amount') or name == 'Settings' and row.values.get('key') in ('opening_cash_control', 'closing_cash_control', 'net_income', 'net_income_control', 'fx_cash_control') and k == 'value'
                # Exact signed monetary text remains data. Decimal serialization and
                # explicit string cell/quotePrefix avoid formula reinterpretation.
                if numeric and isinstance(value, str):
                    try:
                        n = Decimal(value)
                        if n.is_finite() and n == n.to_integral_value() and ',' not in value and 'e' not in value.lower():
                            value = n
                    except Exception:
                        pass
                values.append(value)
            append_safe(ws, values)
        style(ws)
    write_atomic(wb, path)


def export_result(result, path):
    wb = Workbook()
    summary = wb.active
    summary.title = '현금흐름초안'
    append_safe(summary, ('항목키', '분류', '계산금액_원_정확텍스트', '상태'))
    for lid, (cat, amount) in result.lines.items():
        append_safe(summary, (lid, cat, amount, result.status))
    if not result.calculable:
        append_safe(summary, ('계산중단', '', '', '입력 구조 오류: 오류목록을 확인하세요'))
    issues = wb.create_sheet('미해결차이')
    append_safe(issues, ('심각도', '오류코드', '원행', '키', '내용', '차이_원'))
    for i in result.issues:
        append_safe(issues, (i.severity, i.code, i.ref, i.key, i.message, i.difference))
    if not result.issues:
        append_safe(issues, ('INFO', 'MATCHED', '', '', '대사 일치. 회계정책·분류 최종 검토는 사용자가 수행', '0'))
    trace = wb.create_sheet('계산근거')
    append_safe(trace, ('항목키', '분류', '기여금액_원', '원행참조', '수식_설명'))
    for row in result.trace:
        append_safe(trace, row)
    checks = wb.create_sheet('대사내역')
    append_safe(checks, ('검증코드', '원행', '키', '실제금액_원', '기대금액_원', '차이_원'))
    for row in result.checks:
        append_safe(checks, row)
    metadata = wb.create_sheet('실행정보')
    append_safe(metadata, ('키', '값'))
    for row in [('상태', result.status), ('범위', 'KRW 원 단일기업 제한 간접법 더미 모델'), ('원본경로', result.book.source), ('SHA256', result.book.digest), ('금액', 'Decimal 정수 원. 출력은 정확한 텍스트. 합산 시 Excel 정밀도 제한 주의'), ('미지원', '종합 회계기준 적정성·감사결론·연결·RAG·조서 자동작성'), ('수식방어', '위험 접두문자는 apostrophe 처리. 원문은 원본 파일 SHA256로 식별')]:
        append_safe(metadata, row)
    for name, columns in SCHEMA.items():
        ws = wb.create_sheet('원본_' + name)
        append_safe(ws, ('원행',) + columns)
        for row in result.book.sheets.get(name, []):
            append_safe(ws, (row.number,) + tuple(row.values.get(c) for c in columns))
    for ws in wb.worksheets:
        style(ws)
    export_source = result.book.source if Path(result.book.source).exists() else None
    write_atomic(wb, path, export_source)


def input_error_result(error):
    return Result(error.book or InputBook({}, error.source), list(error.issues))
