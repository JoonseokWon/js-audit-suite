"""JS전자 walkthrough counterexamples; mock responses, zero actual inference."""
from io import BytesIO
from pathlib import Path
import json
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from common import AuditError
from demo_case import create_demo, load_manifest, case_status, is_demo_workspace
from demo_case.definition import BASES, EXPECTED
from test_ai_workpapers import mock_response, CATALOG
from workspace import Workspace
import exporter


class DemoCaseTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)/'새 JS전자 시연'
        with patch('ai_provider.generate_draft',side_effect=AssertionError('No inference in unit fixtures')):
            self.ws, self.manifest = create_demo(self.root)

    def execute(self, key):
        entry = self.manifest['inputs'][key]
        return self.ws.execute(entry['module'],entry['path'])

    def apply_mock(self):
        packet = self.ws.ai_packet('A04')
        did = self.ws.record_ai_draft(packet,mock_response(packet),'SYNTHETIC ONLY 요청자')
        self.ws.apply_ai_draft(did,'SYNTHETIC ONLY 작성자')
        return packet, did

    def test_new_user_demo_has_no_person_fields_approval_model_or_export(self):
        self.assertEqual(self.ws.context['회사코드'],'JS_ELEC')
        self.assertTrue(self.ws.engagement['더미업무'])
        self.assertEqual(self.ws.procedures,{})
        self.assertEqual(self.ws.ai_drafts,{})
        self.assertEqual(self.ws.ai_approvals,{})
        self.assertEqual(self.ws.requests,{})
        self.assertEqual(len(self.ws.runs),2)
        self.assertFalse(any(e['kind'] in ('review','ai_draft','ai_content_approval','workflow_finalize') for e in self.ws.events))
        self.assertFalse(list(self.root.glob('*.xlsx')))
        self.assertTrue(is_demo_workspace(self.root))
        status = case_status(self.ws)
        self.assertFalse(status['ai_state']['approved'])
        self.assertFalse(status['needs_rerun'])
        self.assertEqual(status['target_pid'],'A04')

    def test_balanced_whole_journal_omission_is_caught_by_independent_control(self):
        run = self.ws.latest()['원장·리드']['result']
        self.assertGreater(run['blocking'],0)
        self.assertEqual({i['code'] for i in run['issues']},{'LEDGER_CONTROL'})
        def rows(key):
            book = load_workbook(self.manifest['inputs'][key]['path'],read_only=True)
            try:return list(book['총계정원장'].values)[1:]
            finally:book.close()
        clean, dirty = rows('ledger_clean'), rows('ledger_error')
        self.assertEqual(len(clean),82);self.assertEqual(len(dirty),79)
        self.assertEqual(sum(r[5]-r[6] for r in dirty),0)
        self.assertEqual(sum(r[5] for r in clean)-sum(r[5] for r in dirty),99000)
        self.assertEqual(sum(r[6] for r in clean)-sum(r[6] for r in dirty),99000)
        self.assertFalse(any(r[1]=='J-S1' for r in dirty))

    def test_missing_invoice_keeps_revenue_total_and_specific_exceptions(self):
        result = self.ws.latest()['매출']['result']
        self.assertEqual(result['metrics']['순매출'],100000)
        self.assertEqual({r['code'] for r in result['issues']},
            {'UNBILLED_ACCEPTED','MISSING_INVOICE','LEDGER_WITHOUT_INVOICE'})
        book = load_workbook(self.manifest['inputs']['sales_error']['path'],read_only=True)
        try:
            self.assertEqual(book['invoices'].max_row,2)
            self.assertEqual(book['invoices'].cell(2,1).value,'SI02')
        finally:book.close()
        self.assertEqual(self.manifest['expected'],EXPECTED)
        self.assertEqual(EXPECTED['deleted_invoice_supply'],90000)
        self.assertEqual(EXPECTED['deleted_invoice_tax'],9000)

    def test_one_won_invoice_error_is_reviewed_instead_of_erased(self):
        self.execute('ledger_clean');run=self.execute('sales_amount_error')
        codes = {r['code'] for r in run['result']['issues']}
        self.assertIn('LEDGER_AMOUNT_MISMATCH',codes)
        self.assertIn('SUPPLY_AMOUNT_MISMATCH',codes)
        self.assertGreater(run['result']['review'],0)

    def test_clean_reexecution_clears_current_errors_and_preserves_original_history(self):
        initial = {m:r['id'] for m,r in self.ws.latest().items()}
        a=self.execute('ledger_clean');b=self.execute('sales_clean')
        self.assertEqual(a['result']['blocking']+a['result']['review'],0)
        self.assertEqual(b['result']['blocking']+b['result']['review'],0)
        self.assertEqual(b['result']['metrics']['순매출'],100000)
        self.assertTrue(set(initial.values())<=set(self.ws.runs))
        self.assertEqual(len(self.ws.runs),4)
        sales_cross=[r for r in self.ws.cross_checks() if r['모듈']=='매출']
        self.assertTrue(sales_cross)
        self.assertTrue(all(r['상태']=='일치' for r in sales_cross))

    def test_ai_packet_contains_selected_actual_input_and_current_run_evidence(self):
        packet=self.ws.ai_packet('A04',limits={'max_chars':12000,'max_references':40,
            'rows_per_table':3,'max_tables':8,'cell_chars':300})
        kinds={r['kind'] for r in packet['references']}
        self.assertTrue({'run_summary','exception','ledger_reconciliation','raw_xlsx_row'}<=kinds)
        self.assertEqual(packet['engagement']['company_code'],'JS_ELEC')
        self.assertTrue(packet['engagement']['synthetic'])
        self.assertEqual(packet['procedure']['id'],'A04')
        selected={r['run_id'] for r in packet['runs']}
        self.assertEqual(selected,{r['id'] for r in self.ws.latest().values()})
        self.assertIn('MISSING_INVOICE',{r['data']['code'] for r in packet['references'] if r['kind']=='exception'})
        self.assertEqual(self.ws.procedures,{})

    def test_mock_draft_requires_explicit_application_and_synthetic_content_approval(self):
        packet=self.ws.ai_packet('A04');head=self.ws.head
        did=self.ws.record_ai_draft(packet,mock_response(packet),'SYNTHETIC ONLY 요청자')
        self.assertNotEqual(head,self.ws.head);self.assertEqual(self.ws.procedures,{})
        self.ws.apply_ai_draft(did,'SYNTHETIC ONLY 작성자')
        p=dict(self.ws.procedures['A04'])
        self.assertEqual(p['performed_date'],'');self.assertEqual(p['reviewer'],'')
        self.assertFalse(self.ws.ai_status('A04')['approved'])
        p['conclusion']+='\nSYNTHETIC ONLY 수정: 실제 고객의 조사·판단을 수행하지 않은 상태전이 시험.'
        self.ws.save_procedure(p)
        self.ws.approve_ai_content('A04','SYNTHETIC ONLY 승인 모사',
            '자동검사에서만 가상 내용 승인 상태를 모사함. 실제 사람의 승인·감사결론 아님.')
        self.assertTrue(self.ws.ai_status('A04')['approved'])
        self.assertEqual(self.ws.procedures['A04']['state'],'작성중')
        self.assertEqual(self.ws.procedures['A04']['reviewer'],'')

    def test_synthetic_edit_approval_excel_reopen_and_further_edit_invalidation(self):
        self.apply_mock();p=dict(self.ws.procedures['A04'])
        note='SYNTHETIC ONLY 수정 문구 / 실제 감사승인 아님'
        p['conclusion']+='\n'+note;self.ws.save_procedure(p)
        self.ws.approve_ai_content('A04','SYNTHETIC ONLY 승인 모사','가상 조서 내용 상태전이 검사')
        target=self.root/'합성_검토중_출력.xlsx';snap=exporter.export(self.ws,target,CATALOG,False)
        self.assertEqual(snap['export_state'],'검토중 초안')
        book=load_workbook(target,data_only=False)
        try:
            self.assertIn('AI근거참조',book.sheetnames);self.assertIn('AI최종내용승인',book.sheetnames)
            self.assertIn(note,'\n'.join(str(r[1]) for r in book['조서_A04'].iter_rows(values_only=True)))
            self.assertTrue(all(c.data_type!='f' for s in book for row in s for c in row))
            self.assertTrue(any('MISSING_INVOICE' in str(c.value) for row in book['AI근거참조'] for c in row))
        finally:book.close()
        reopened=Workspace(self.root);self.assertTrue(reopened.ai_status('A04')['approved'])
        self.assertIn(note,reopened.procedures['A04']['conclusion'])
        p=dict(reopened.procedures['A04']);p['conclusion']+='\n새 합성 수정';reopened.save_procedure(p)
        self.assertFalse(reopened.ai_status('A04')['approved'])
        self.assertFalse(case_status(reopened)['needs_rerun'])
        self.assertIn('승인 기록',case_status(reopened)['next_step'])
        self.assertFalse(reopened.workflow_locked)

    def test_source_reexecution_invalidates_old_mock_draft_and_approval(self):
        packet,did=self.apply_mock()
        self.ws.approve_ai_content('A04','SYNTHETIC ONLY 승인 모사','가상 승인 상태전이 검사')
        original=dict(self.ws.ai_approvals['A04'])
        self.execute('ledger_clean');self.execute('sales_clean')
        self.assertTrue(case_status(self.ws)['needs_rerun'])
        self.assertFalse(self.ws.ai_status('A04')['approved'])
        self.assertEqual(self.ws.ai_approvals['A04'],original)
        with self.assertRaises(AuditError):self.ws.apply_ai_draft(did,'SYNTHETIC ONLY')
        with self.assertRaises(AuditError):self.ws.record_ai_draft(packet,mock_response(packet),'SYNTHETIC ONLY')
        fresh,did2=self.apply_mock()
        self.assertNotEqual(packet['basis'],fresh['basis']);self.assertNotEqual(did,did2)
        self.assertFalse(self.ws.ai_status('A04')['approved'])

    def test_unapplied_draft_after_manual_edit_requires_fresh_fingerprint(self):
        from procedures import task_definitions
        packet=self.ws.ai_packet('A04')
        did=self.ws.record_ai_draft(packet,mock_response(packet),'SYNTHETIC ONLY 요청자')
        task=next(p for p in task_definitions(self.ws,CATALOG) if p['id']=='A04')
        task.update(preparer='SYNTHETIC ONLY 작성자',scope='SYNTHETIC ONLY 새 수동 범위')
        self.ws.save_procedure(task)
        status=case_status(self.ws)
        self.assertTrue(status['needs_rerun'])
        self.assertIn('다시 생성',status['stages'][2]['state'])
        with self.assertRaises(AuditError):self.ws.apply_ai_draft(did,'SYNTHETIC ONLY')

    def test_manifest_cannot_redirect_inputs_or_change_expected_amounts(self):
        path=self.root/'시연_자료목록.json';original=path.read_bytes();head=self.ws.head
        for key in ('path','amount'):
            payload=json.loads(original)
            if key=='path':payload['inputs']['sales_clean']['path']='../../외부.xlsx'
            else:payload['expected']['deleted_invoice_supply']=0
            path.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
            with self.assertRaises(AuditError):load_manifest(self.root)
            self.assertFalse(is_demo_workspace(self.root))
        path.write_bytes(original)
        self.assertEqual(Workspace(self.root).head,head)

    def test_modified_or_missing_template_is_rejected_without_changing_history(self):
        path=Path(self.manifest['inputs']['sales_clean']['path']);raw=path.read_bytes();head=self.ws.head
        path.write_bytes(raw+b'changed')
        with self.assertRaises(AuditError):load_manifest(self.root)
        path.unlink()
        with self.assertRaises(AuditError):load_manifest(self.root)
        self.assertEqual(Workspace(self.root).head,head)

    def test_rewriting_template_and_self_declared_hash_cannot_replace_setup_history(self):
        import hashlib
        path=Path(self.manifest['inputs']['sales_clean']['path']);changed=path.read_bytes()+b'changed'
        path.write_bytes(changed)
        manifest_path=self.root/'시연_자료목록.json'
        payload=json.loads(manifest_path.read_text(encoding='utf-8'))
        payload['inputs']['sales_clean'].update(bytes=len(changed),sha256=hashlib.sha256(changed).hexdigest())
        manifest_path.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
        with self.assertRaisesRegex(AuditError,'최초 생성 이력'):load_manifest(self.root)

    def test_existing_folder_and_old_source_bytes_are_preserved(self):
        head=self.ws.head
        with self.assertRaises(AuditError):create_demo(self.root)
        self.assertEqual(Workspace(self.root).head,head)
        import runtime,hashlib
        for relative,expected in BASES.values():
            self.assertEqual(hashlib.sha256((runtime.configure()/'sector_resources'/relative).read_bytes()).hexdigest(),expected)


if __name__=='__main__':unittest.main()
