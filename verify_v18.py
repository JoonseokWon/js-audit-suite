"""Native source/frozen AI UI checks with labelled mock transport, no model download.

Actual local inference is separately recorded in the release validation evidence.
"""
from pathlib import Path
import json,sys,time,traceback,copy,shutil
from common import atomic,digest,dumps,AuditError

def fixture(packet):
    ref=packet['references'][0]['id']
    fields={'purpose':'모의 전송 응답이며 실제 모델 호출은 없는 UI 검사이다. ['+ref+']',
            'risk':'자료 누락·오류와 독립 증거의 신뢰성은 추가확인이 필요하다. ['+ref+']',
            'scope':'선택 회사의 자동검사와 원문 발췌 범위의 합성 초안이다. 생략된 행과 수동 절차는 별도 확인이 필요하다. ['+ref+']',
            'performed':'프로그램이 실행한 자동검사 발췌만 기술한다. 사람의 조회·수령 확인은 미수행이며 별도 검토가 필요하다. ['+ref+']',
            'conclusion':'사람 검토 전 초안이다. 차이·미확인 범위에 대한 추가확인과 회계사 판단이 필요하다. ['+ref+']'}
    context={'company_code':packet['engagement']['company_code'],'procedure_id':packet['procedure']['id'],'basis':packet['basis']}
    citations={k:[ref] for k in fields}
    for item in packet['references']:
        if item['kind']=='exception':
            fields['performed']+=' '+str(item['data'].get('code',''))+' ['+item['id']+']'
            if item['id'] not in citations['performed']:citations['performed'].append(item['id'])
    from test_review_ai_helpers import response_from_fields
    return response_from_fields(packet,fields,citations,model='native-ui-mock-no-model-call')


def self_test(report_path):
    out=Path(report_path).parent;out.mkdir(parents=True,exist_ok=True);checks=[];root=None
    def check(ok,name):
        if not ok:raise AssertionError(name)
        checks.append(name)
    try:
        from verify_v171 import self_test as baseline
        baseline_path=Path(sys.argv[sys.argv.index('--baseline-v171')+1]) if '--baseline-v171' in sys.argv else out/'baseline_v171/report.json'
        reused='--baseline-v171' in sys.argv
        if not reused:check(baseline(baseline_path)==0,'previous_native_149_and_calculation_fixtures')
        old=json.loads(baseline_path.read_text(encoding='utf-8'))
        if reused:check(old['ok'] and old['frozen']==bool(getattr(sys,'frozen',False)),'previous_native_149_and_calculation_fixtures')
        import tkinter as tk
        from app import App,CATALOG
        from workspace import Workspace
        from unittest.mock import patch
        from window_capture import capture_window,destroy_test_window
        from openpyxl import load_workbook
        import exporter
        root=tk.Tk();app=App(root,sector_home=baseline_path.parent/'baseline_v17/회사별저장')
        errors=[]
        def settle():
            for _ in range(6):root.update();time.sleep(.02)
        def wait():
            deadline=time.monotonic()+90
            while app.busy and time.monotonic()<deadline:root.update();time.sleep(.02)
            check(not app.busy,'native_background_job_finishes');settle()
        def inside(w,p):
            return w.winfo_ismapped() and w.winfo_rootx()>=p.winfo_rootx()-2 and w.winfo_rooty()>=p.winfo_rooty()-2 and w.winfo_rootx()+w.winfo_width()<=p.winfo_rootx()+p.winfo_width()+2 and w.winfo_rooty()+w.winfo_height()<=p.winfo_rooty()+p.winfo_height()+2
        source_folder=next((baseline_path.parent/'baseline_v17/회사별저장/companies/JS_ELEC').glob('*/업무'))
        folder=out/'AI_UI_합성업무';shutil.copytree(source_folder,folder)
        app.ws=Workspace(folder);app.reset_workspace_view();app.refresh();app.nb.select(app.tabs['절차·조서']);app.load_procedure('A04');settle()
        check(len(app.ws.latest())==15,'all_fifteen_calculation_modules_present')
        with patch('ai_desktop.messagebox.showerror',side_effect=lambda title,text,**kw:errors.append(text)):
            app.ai_settings_button.invoke();settle();settings=app.ai_settings_win
            check(settings.provider.get()=='로컬 AI (기본)','local_provider_default')
            check(settings.key_entry.cget('show')=='●','api_key_masked')
            check(app.ai_config['num_gpu']==0 and app.ai_config['num_ctx']==16384 and app.ai_config['timeout_seconds']==300,'cpu_context_timeout_defaults')
            settings.key.set('sk-native-fixture-never-real-123456789');settings.model.set('explicit-test-model');settings.save_button.invoke();settle()
            check(app.ai_api_key.endswith('123456789') and not settings.winfo_exists(),'settings_memory_only_apply')
            check('sk-native' not in dumps(app.ws.events),'api_key_absent_from_workspace_events')
            app.ai_api_key=''
            head=app.ws.head;before_procedure=copy.deepcopy(app.ws.procedures.get('A04'))
            with patch.object(app.ai_runtime,'ensure_server',return_value={'ready':True}),patch('ai_desktop.generate_draft',side_effect=lambda p,c,e:fixture(p)):
                app.ai_generate_button.invoke();wait()
            check(not errors,'generation_no_gui_errors')
            check(app.ws.head!=head and len(app.ws.ai_drafts)==1 and app.ws.procedures.get('A04')==before_procedure,'generate_saves_draft_without_applying_or_approving')
            win=app.ai_preview_win;check(inside(win.apply_button,win),'preview_apply_button_visible');capture_window(win,out/'AI초안_근거확인.png');win.apply_button.invoke();settle()
            check(app.ws.ai_status('A04')['required'] and not app.ws.ai_status('A04')['approved'],'apply_requires_human_content_approval')
            check(app.ws.procedures['A04']['model'] if 'model' in app.ws.procedures['A04'] else app.ws.procedures['A04']['ai_draft_id'],'model_lineage_preserved')
            check(not app.ws.procedures['A04'].get('performed_date'),'ai_does_not_invent_human_date')
            app.proc_text['conclusion'].insert('end','\n사람 수정: 독립 원본 추가 확보 전 판단을 유보함.');app.save_workpaper_button.invoke();settle()
            check('사람 수정' in app.ws.procedures['A04']['conclusion'],'human_edit_saved')
            app.ai_approve_button.invoke();settle();approval=app.ai_approval_win
            check(inside(approval.approve_button,approval),'approval_button_visible');capture_window(approval,out/'AI내용_최종승인.png')
            approval.approve_button.invoke();settle();check(not app.ws.ai_status('A04')['approved'],'unchecked_approval_blocked');errors.clear()
            approval.confirmed.set(True);approval.note.insert('1.0','합성 UI 시험: 근거·미수행·예외와 사람이 수정한 문구를 검토함.');approval.approve_button.invoke();settle()
            check(app.ws.ai_status('A04')['approved'],'explicit_human_content_approval_saved')
            check(app.ws.procedures['A04']['state']!='검토완료','content_approval_does_not_complete_independent_review')
            target=out/'AI조서_내용승인_초안.xlsx';exporter.export(app.ws,target,CATALOG,False)
            book=load_workbook(target,read_only=True,data_only=True);names=book.sheetnames
            check({'AI초안목록','AI근거참조','AI적용이력','AI최종내용승인'}<=set(names),'xlsx_has_ai_trace_and_approval_sheets');book.close()
            try:exporter.export(app.ws,out/'감사검토_미완료_차단.xlsx',CATALOG,True)
            except AuditError:check(True,'content_approval_cannot_bypass_final_review_gate')
            else:raise AssertionError('final review bypass')
            for scale in (100,125,150):
                root.geometry('1280x760');app.set_ui_zoom(scale);settle()
                check(all(inside(getattr(app,k),app.ai_actions) for k in ('ai_generate_button','ai_preview_button','ai_approve_button','ai_settings_button')),str(scale)+'_ai_actions_reachable')
                capture_window(root,out/(str(scale)+'_AI조서_작은창.png'))
            app.set_ui_zoom(100);app.display.fit_window(initial=True);settle()
            app.proc_text['conclusion'].insert('end','\n추가 수정');app.save_workpaper_button.invoke();settle()
            check(not app.ws.ai_status('A04')['approved'],'saved_human_edit_invalidates_prior_approval')
            kinds={r['기록종류'] for r in app.history_rows()}
            check({'AI 초안 생성','AI 초안 적용','AI 내용 최종 승인'}<=kinds,'history_displays_ai_actions')
            check('sk-native' not in dumps(app.ws.events),'no_api_key_in_final_history')
            app.ai_config['provider']='openai';app.ai_api_key='sk-native-fixture-never-real-123456789';app.ai_openai_model='explicit-test-model'
            called=[]
            with patch('ai_desktop.generate_draft',side_effect=lambda *args:called.append(args)):
                app.ai_generate_button.invoke();settle();consent=app.ai_consent_win
                check(inside(consent.send_button,consent),'external_send_button_visible');capture_window(consent,out/'외부전송_근거동의.png')
                check(not consent.confirmed.get() and not called,'external_consent_unchecked_and_no_send')
                consent.destroy()
            app.ai_api_key=''
            app.ws.refresh();check('sk-native' not in dumps(app.ws.snapshot(False)),'export_snapshot_no_api_key')
        check(not errors,'all_ui_error_messages_expected_and_cleared')
        app.shutdown_ai();destroy_test_window(root);root=None
        record={**old,'version':'1.8.0','frozen':bool(getattr(sys,'frozen',False)),'baseline_reused':reused,'native_checks':old['native_checks']+checks,'ai_native_checks':checks,'ai_transport':'explicit mock; actual inference separately recorded','ai_export':target.name}
        atomic(report_path,json.dumps(record,ensure_ascii=False,indent=2).encode());return 0
    except Exception as exc:
        atomic(report_path,json.dumps({'ok':False,'error':str(exc),'ui_errors':locals().get('errors',[]),'traceback':traceback.format_exc(),'ai_native_checks':checks},ensure_ascii=False,indent=2).encode());return 1
    finally:
        if root:
            try:
                from window_capture import destroy_test_window
                destroy_test_window(root)
            except Exception:pass
