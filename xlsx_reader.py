"""Stream preflight XML and value rows without buffering the whole workbook twice."""
from datetime import date,datetime
from pathlib import Path
from zipfile import ZipFile,BadZipFile
import hashlib,posixpath,re
from common import AuditError,cancelled
from input_limits import MAX_SOURCE_BYTES,MAX_EXPANDED_BYTES,MAX_WORKBOOK_ROWS,EXCEL_DATA_ROWS

def file_hash(path):
    value=hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda:stream.read(1024**2),b''):value.update(chunk)
    return value.hexdigest()

def sheet_paths(archive):
    from defusedxml.ElementTree import fromstring
    def xml(name):
        if archive.getinfo(name).file_size>2*1024**2:raise AuditError('Excel 메타데이터 한도 초과')
        return fromstring(archive.read(name))
    relations={r.get('Id'):r.get('Target','') for r in xml('xl/_rels/workbook.xml.rels')}
    paths={}
    for sheet in xml('xl/workbook.xml').iter():
        if sheet.tag.rsplit('}',1)[-1]!='sheet':continue
        rid=next((v for k,v in sheet.attrib.items() if k.rsplit('}',1)[-1]=='id'),None)
        target=relations.get(rid,'')
        name=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl',target))
        if not name.startswith('xl/worksheets/') or '..' in name.split('/'):raise AuditError('워크시트 경로 오류')
        paths[name]=sheet.get('name')
    return paths

def read_xlsx(path,schema=None,event=None,max_rows=10000,row_limits=None):
    from openpyxl import load_workbook
    from defusedxml.ElementTree import iterparse
    path=Path(path)
    if path.suffix.lower()!='.xlsx' or path.stat().st_size>MAX_SOURCE_BYTES:raise AuditError('256MB 이하 .xlsx만 지원합니다.')
    limits=dict(row_limits or {})
    if any(type(n) is not int or not 1<=n<=EXCEL_DATA_ROWS for n in [max_rows,*limits.values()]):raise AuditError('워크시트 행 한도 설정 오류')
    original=file_hash(path);total=0
    try:
        with ZipFile(path) as archive:
            entries=archive.infolist()
            if len(entries)>500 or len({e.filename for e in entries})!=len(entries) or sum(e.file_size for e in entries)>MAX_EXPANDED_BYTES:raise AuditError('Excel 압축해제 한도 초과')
            if any(any(x in e.filename.lower() for x in ('externallinks/','vba','embeddings/','connections.xml')) for e in entries):raise AuditError('외부연결·매크로·첨부 Excel은 지원하지 않습니다.')
            paths=sheet_paths(archive)
            for filename,name in paths.items():
                last=0;stack=[];limit=limits.get(name,max_rows)
                with archive.open(filename) as stream:
                    for action,element in iterparse(stream,events=('start','end')):
                        if action=='start':stack.append(element);continue
                        if element.tag.rsplit('}',1)[-1]=='row':
                            cancelled(event);number=int(element.get('r','0'))
                            if number<=last or number>limit+1:raise AuditError(name+': 물리 행 순서/한도 오류')
                            last=number;column=0
                            for cell in element:
                                match=re.fullmatch(r'([A-Z]+)(\d+)',cell.get('r',''))
                                if not match or int(match[2])!=number:raise AuditError('셀 좌표 오류')
                                index=0
                                for ch in match[1]:index=index*26+ord(ch)-64
                                if index<=column or index>80:raise AuditError('중복·역순·범위밖 셀')
                                column=index
                            element.clear()
                            if len(stack)>1:stack[-2].remove(element)
                        stack.pop()
        workbook=load_workbook(path,read_only=True,data_only=False,keep_links=False)
        try:
            if schema is not None and set(workbook.sheetnames)!=set(schema):raise AuditError('시트 목록이 입력양식과 다릅니다: '+', '.join(schema))
            tables={}
            for sheet in workbook:
                sheet.reset_dimensions();rows=[];headers=None;limit=limits.get(sheet.title,max_rows)
                for number,cells in enumerate(sheet.iter_rows(),1):
                    cancelled(event)
                    if number>limit+1 or len(cells)>80:raise AuditError('입력 한도 초과')
                    values=[]
                    for cell in cells:
                        if cell.data_type in ('f','e'):raise AuditError(f'{sheet.title}!{number}: 수식·오류 셀은 확정값으로 바꾸세요.')
                        value=cell.value
                        if isinstance(value,datetime):value=value.date().isoformat() if value.time()==datetime.min.time() else value.isoformat()
                        elif isinstance(value,date):value=value.isoformat()
                        if isinstance(value,str) and (len(value)>2000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value)):raise AuditError('입력 텍스트 오류')
                        values.append(value)
                    while values and values[-1] is None:values.pop()
                    if number==1:
                        headers=values
                        if not headers or len(set(headers))!=len(headers):raise AuditError('중복/빈 헤더')
                        if schema is not None and tuple(headers)!=tuple(schema[sheet.title]):raise AuditError(sheet.title+': 열 이름/순서 오류')
                    elif values:
                        if len(values)>len(headers):raise AuditError('양식 밖 열')
                        rows.append(dict(zip(headers,values+[None]*(len(headers)-len(values))))|{'_row':number,'_sheet':sheet.title});total+=1
                        if total>MAX_WORKBOOK_ROWS:raise AuditError('워크북 전체 2,000,000행 한도 초과')
                if headers is None:raise AuditError('빈 시트')
                tables[sheet.title]=rows
        finally:workbook.close()
    except (BadZipFile,KeyError) as exc:raise AuditError('Excel 구조 오류') from exc
    if file_hash(path)!=original:raise AuditError('입력 파일이 변경되었습니다.')
    return tables,{'filename':path.name,'sha256':original,'rows':{key:len(rows) for key,rows in tables.items()}}
