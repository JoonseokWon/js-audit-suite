"""Deterministic AI lifecycle counterexamples; no real inference/network calls."""
import copy,io,json,tempfile,threading,unittest
from pathlib import Path
import runtime
ROOT=runtime.configure()
from common import AuditError,digest,dumps
from workspace import Workspace
from ai_workpapers import FIELDS,PROMPT_VERSION,build_packet,content_hash
from verify import synthetic_procedure
import exporter

CATALOG=json.loads((ROOT/'procedure_catalog.json').read_text(encoding='utf-8'))
CONTEXT={'회사코드':'SYNTH01','회사명':'SYNTHETIC ONLY','시작일':'2026-01-01','기준일':'2026-12-31','회계체계':'K-IFRS','감사기준':'합성 검사','작성자':'합성작성자','더미업무':True}

def mock_response(packet):
    """A clearly labelled fixed JSON fixture, never a production AI fallback."""
    ref=next((r['id'] for r in packet['references'] if r['kind']=='run_summary'),packet['references'][0]['id'])
    fields={
        'purpose':'모의 AI 응답 / 모델호출 없음. 합성 자료의 자동검사 범위를 살펴볼 조서 초안이다. ['+ref+']',
        'risk':'합성 발췌 자료의 누락·오류와 추출 신뢰성을 사람이 확인해야 한다. ['+ref+']',
        'scope':'현재 회사·기간의 등록 사본과 자동검사 발췌이다. 절단·미확인 범위와 독립 증거의 한계가 있다. ['+ref+']',
        'performed':'프로그램 자동검사 결과와 원문 발췌를 기록할 초안이다. 사람이 실사·조회·수령 확인은 미수행이며 별도 확인이 필요하다. ['+ref+']',
        'conclusion':'사람 검토 전 초안이다. 차이·미확인 범위에 대한 추가확인과 회계사 판단이 필요하다. ['+ref+']'}
    citations={key:[ref] for key in FIELDS}
    exceptions=[item for item in packet['references'] if item['kind']=='exception']
    if exceptions:
        codes=list(dict.fromkeys(str(item['data'].get('code') or item['id']) for item in exceptions))
        fields['performed']+=' 선택된 예외코드: '+', '.join(codes)+'. '+''.join('['+item['id']+']' for item in exceptions)
        citations['performed']=list(dict.fromkeys([ref]+[item['id'] for item in exceptions]))
        fields['conclusion']+=' 선택된 예외는 사람이 추가조사·추가확인하여 판단해야 한다.'
    context={'company_code':packet['engagement']['company_code'],'procedure_id':packet['procedure']['id'],'basis':packet['basis']}
    return {'fields':fields,'citations':citations,'provider':'ollama','model':'unit-fixture-no-model-call','prompt_version':PROMPT_VERSION,'packet_hash':digest(packet),**context,'raw_response':dumps({'fields':fields,'citations':citations,'context':context})}

def sync_raw(response):
    response['raw_response']=dumps({'fields':response['fields'],'citations':response['citations'],'context':{k:response[k] for k in ('company_code','procedure_id','basis')}})
    return response

class AiWorkpaperTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.ws=Workspace(Path(self.temp.name)/'업무',create=CONTEXT)
        self.ledger=ROOT/'samples/통합/통합_01_원장정상.xlsx'
        self.sales=ROOT/'samples/통합/통합업무_매출정상.xlsx'
        self.ws.execute('원장·리드',self.ledger);self.ws.execute('매출',self.sales)

    def packet(self,pid='A04',**kwargs):return self.ws.ai_packet(pid,**kwargs)
    def draft(self,pid='A04'):
        packet=self.packet(pid);return self.ws.record_ai_draft(packet,mock_response(packet),'합성요청자')
    def apply(self):
        did=self.draft();self.ws.apply_ai_draft(did,'합성작성자');return did
    def approve(self):return self.ws.approve_ai_content('A04','사람승인자','합성 고정응답·근거·미수행 범위를 수동 확인한 상태 시험')
    def assert_response_blocked(self,packet,response):
        before=self.ws.head
        with self.assertRaises(AuditError):self.ws.record_ai_draft(packet,response,'합성요청자')
        self.assertEqual(self.ws.head,before)

    def test_packet_contains_real_rows_calculation_and_current_ledger_hashes(self):
        packet=self.packet();kinds={r['kind'] for r in packet['references']}
        self.assertTrue({'raw_xlsx_row','calculation_row','run_summary','ledger_reconciliation'}<=kinds)
        self.assertEqual(packet['basis'],self.ws.basis())
        self.assertEqual(packet['procedure']['id'],'A04')
        self.assertTrue(all(r['company_code']=='SYNTH01' for r in packet['references']))
        row=next(r for r in packet['references'] if r['kind']=='raw_xlsx_row')
        self.assertGreaterEqual(row['data']['row'],2)
        self.assertEqual(row['sha256'],self.ws.sources[row['source_id']]['sha256'])
        run=next(r for r in packet['references'] if r['kind']=='run_summary')
        self.assertEqual(run['result_hash'],self.ws.runs[run['run_id']]['result_hash'])

    def test_packet_has_no_original_absolute_path_or_credentials(self):
        path=Path(self.temp.name)/'metadata.json'
        path.write_text(json.dumps({'path':'C:/private/data.xlsx','value':'sk-'+'a'*30}),encoding='utf-8')
        sid=self.ws.add_source(path,'API key sk-'+'b'*30)
        procedure=synthetic_procedure('C02',[sid]);self.ws.save_procedure(procedure)
        data=dumps(self.packet('C02'))
        self.assertNotIn(str(self.ws.root),data);self.assertNotIn('C:/private',data)
        self.assertNotIn('sk-'+'a'*30,data);self.assertNotIn('sk-'+'b'*30,data)
        self.assertNotIn('original',next(iter(self.packet()['sources'])))

    def test_bounded_packet_records_omitted_rows_and_hashes(self):
        packet=self.packet(limits={'rows_per_table':1,'max_tables':2})
        self.assertLessEqual(len(dumps(packet)),24000)
        self.assertTrue(packet['truncation'])
        self.assertTrue(any(t.get('omitted',0)>0 and t.get('content_hash') for t in packet['truncation']))

    def test_compact_packet_excludes_unselected_history_but_pins_full_basis(self):
        path=Path(self.temp.name)/'extra.json';path.write_text('{"synthetic":true}',encoding='utf-8')
        for index in range(18):self.ws.add_source(path,'미선택 합성자료 '+str(index))
        before=self.ws.head;basis=self.ws.basis()
        packet=self.packet(limits={'max_chars':12000,'max_references':40,'rows_per_table':3,'max_tables':8,'cell_chars':300})
        self.assertLessEqual(len(dumps(packet)),12000)
        self.assertEqual(packet['basis'],basis);self.assertEqual(self.ws.head,before)
        self.assertEqual(len(packet['sources']),2);self.assertEqual(len(packet['runs']),2)
        self.assertEqual(packet['unselected_sources']['count'],18)
        self.assertEqual(len(packet['unselected_sources']['manifest_hash']),64)
        self.assertTrue(any(r['kind']=='raw_xlsx_row' for r in packet['references']))
        self.assertTrue(any(r['kind']=='ledger_reconciliation' for r in packet['references']))
        self.assertTrue(all('content_hash' in item or 'details_hash' in item for item in packet['truncation']))

    def test_packet_deterministic_despite_unrelated_note_event(self):
        first=self.packet();self.ws.add_note('합성','근거를 변경하지 않는 로컬 의견','검토자')
        self.assertEqual(first,self.packet())

    def test_missing_or_changed_original_blocks_packet_without_event(self):
        source=next(iter(self.ws.sources.values()));(self.ws.root/source['file']).write_bytes(b'changed')
        before=self.ws.head
        with self.assertRaisesRegex(AuditError,'원본'):self.packet()
        self.assertEqual(before,self.ws.head)

    def test_cancelled_packet_has_no_side_effect(self):
        event=threading.Event();event.set();before=self.ws.head
        with self.assertRaises(AuditError):self.packet(event=event)
        self.assertEqual(before,self.ws.head)

    def test_unknown_procedure_is_rejected(self):
        with self.assertRaises(AuditError):self.packet('UNKNOWN')

    def test_generation_is_preserved_without_applying_or_approving(self):
        before_basis=self.ws.basis();did=self.draft()
        self.assertIn(did,self.ws.ai_drafts);self.assertNotIn('A04',self.ws.procedures)
        self.assertEqual(self.ws.ai_status('A04')['state'],'AI 미적용')
        self.assertEqual(before_basis,self.ws.basis())
        self.assertEqual(self.ws.events[-1]['kind'],'ai_draft')

    def test_changed_packet_row_cannot_be_saved_as_actual_evidence(self):
        packet=self.packet();packet['references'][0]['data']['company_name']='조작'
        self.assert_response_blocked(packet,mock_response(packet))

    def test_response_wrong_company_procedure_basis_or_prompt_is_blocked(self):
        packet=self.packet()
        for key,value in [('company_code','OTHER'),('procedure_id','A12'),('basis','0'*64),('prompt_version','other'),('packet_hash','0'*64)]:
            with self.subTest(key=key):
                response=mock_response(packet);response[key]=value
                self.assert_response_blocked(packet,response)

    def test_missing_or_unknown_citations_are_blocked(self):
        packet=self.packet()
        for value in ([],['REF999'],['REF001','REF001'],[{}]):
            with self.subTest(value=value):
                response=mock_response(packet);response['citations']['purpose']=value
                self.assert_response_blocked(packet,sync_raw(response))

    def test_raw_response_and_applied_fields_cannot_disagree(self):
        packet=self.packet();response=mock_response(packet);response['fields']['purpose']='원문과 다른 내용'
        self.assert_response_blocked(packet,response)

    def test_raw_extra_keys_and_private_paths_or_keys_are_rejected(self):
        packet=self.packet()
        for value in ('C:/private/original.xlsx','sk-'+'q'*25):
            with self.subTest(value=value):
                response=mock_response(packet);response['fields']['purpose']=value
                self.assert_response_blocked(packet,sync_raw(response))
        response=mock_response(packet);raw=json.loads(response['raw_response']);raw['unexpected']='extra';response['raw_response']=dumps(raw)
        self.assert_response_blocked(packet,response)

    def test_provider_config_or_unknown_key_metadata_is_never_persisted(self):
        packet=self.packet();response=mock_response(packet)
        secret='sk-'+'z'*30;response['api_key']=secret;response['config']={'api_key':secret}
        self.ws.record_ai_draft(packet,response,'합성요청자')
        self.assertNotIn(secret,dumps(self.ws.events))
        self.assertNotIn('api_key',dumps(list(self.ws.ai_drafts.values())))

    def test_ai_claiming_receipt_inspection_or_clean_opinion_is_rejected(self):
        packet=self.packet()
        for value in ('실사를 수행했다.','조회서를 회수했다.','부정은 없습니다.','적정의견을 표명한다.','감사를 완료했다.'):
            with self.subTest(value=value):
                response=mock_response(packet);response['fields']['performed']=value
                self.assert_response_blocked(packet,sync_raw(response))

    def test_procedure_edit_during_inference_blocks_record_and_apply(self):
        packet=self.packet();self.ws.save_procedure(synthetic_procedure('A04',list(self.ws.sources)))
        self.assert_response_blocked(packet,mock_response(packet))
        did=self.draft();p=dict(self.ws.procedures['A04']);p['scope']='수정 범위';self.ws.save_procedure(p)
        with self.assertRaises(AuditError):self.ws.apply_ai_draft(did,'작성자')

    def test_apply_does_not_fabricate_performance_date_review_or_approval(self):
        did=self.apply();p=self.ws.procedures['A04']
        self.assertEqual(p['ai_draft_id'],did);self.assertEqual(p['state'],'작성중')
        self.assertEqual(p['performed_date'],'');self.assertEqual(p['reviewer'],'')
        self.assertFalse(self.ws.ai_status('A04')['approved'])
        self.assertTrue(set(p['evidence_ids'])<=set(self.ws.sources))
        self.assertEqual([e['kind'] for e in self.ws.events[-3:]],['ai_draft','procedure','ai_apply'])

    def test_human_edit_then_approval_pins_saved_content_and_evidence(self):
        self.apply();p=dict(self.ws.procedures['A04']);p['conclusion']+=' 사람이 추가 판단을 기록한 합성 시험.';self.ws.save_procedure(p)
        approval=self.approve();state=self.ws.ai_status('A04')
        self.assertTrue(state['approved']);self.assertFalse(state['stale'])
        self.assertEqual(approval['content_hash'],content_hash(self.ws.procedures['A04']))
        self.assertEqual(approval['basis'],self.ws.basis());self.assertTrue(approval['citations'])
        self.assertEqual(self.ws.procedures['A04']['state'],'작성중')

    def test_empty_actor_or_note_and_unapplied_draft_cannot_be_approved(self):
        self.draft()
        with self.assertRaises(AuditError):self.approve()
        self.apply()
        for actor,note in (('','확인'),('사람','')):
            with self.subTest(actor=actor),self.assertRaises(AuditError):self.ws.approve_ai_content('A04',actor,note)

    def test_form_save_preserves_ai_lineage_and_invalidates_approval(self):
        did=self.apply();self.approve()
        p={k:v for k,v in self.ws.procedures['A04'].items() if k not in ('ai_draft_id','edit_revision')}
        p['conclusion']='사람이 수정한 새로운 결론. 추가 검토 필요.';self.ws.save_procedure(p)
        self.assertEqual(self.ws.procedures['A04']['ai_draft_id'],did)
        self.assertFalse(self.ws.ai_status('A04')['approved']);self.assertTrue(self.ws.ai_status('A04')['stale'])

    def test_even_edit_then_restore_is_a_new_unapproved_revision(self):
        self.apply();original=dict(self.ws.procedures['A04']);self.approve()
        changed=dict(original);changed['scope']='잠시 수정';self.ws.save_procedure(changed);self.ws.save_procedure(original)
        self.assertEqual(content_hash(original),content_hash(self.ws.procedures['A04']))
        self.assertFalse(self.ws.ai_status('A04')['approved'])

    def test_source_change_stales_approval_and_requires_new_ai_packet(self):
        self.apply();self.approve();self.ws.add_source(self.ledger,'새 자료')
        self.assertFalse(self.ws.ai_status('A04')['approved'])
        with self.assertRaises(AuditError):self.approve()

    def test_rerun_stales_approval_even_when_numbers_are_unchanged(self):
        self.apply();self.approve();self.ws.execute('매출',self.sales)
        self.assertFalse(self.ws.ai_status('A04')['approved'])
        with self.assertRaises(AuditError):self.approve()

    def test_unused_new_draft_does_not_change_applied_content_approval(self):
        self.apply();self.approve();did=self.draft()
        self.assertTrue(self.ws.ai_status('A04')['approved']);self.assertIn(did,self.ws.ai_drafts)

    def test_restart_preserves_ai_raw_response_approval_and_audit_events(self):
        did=self.apply();self.approve();head=self.ws.head
        reopened=Workspace(self.ws.root)
        self.assertEqual(reopened.head,head);self.assertTrue(reopened.ai_status('A04')['approved'])
        self.assertEqual(reopened.ai_drafts[did]['raw_response'],self.ws.ai_drafts[did]['raw_response'])
        self.assertEqual(reopened.events[-1]['kind'],'ai_content_approval')

    def test_draft_excel_marks_unapproved_and_includes_trace_sheets(self):
        self.apply();snap=self.ws.snapshot();data=exporter.workbook_bytes(snap,CATALOG)
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(data),data_only=False)
        try:
            self.assertTrue({'AI초안목록','AI근거참조','AI적용이력','AI최종내용승인'}<=set(book.sheetnames))
            content='\n'.join(str(c.value) for row in book['조서_A04'] for c in row)
            self.assertIn('AI 최종내용 미승인',content)
            self.assertIn('unit-fixture-no-model-call',content)
            self.assertEqual(snap['export_state'],'검토중 초안')
        finally:book.close()

    def test_ai_text_starting_equals_is_exported_as_literal_not_formula(self):
        packet=self.packet();response=mock_response(packet)
        response['fields']['purpose']='=1+1 [REF001]';response['citations']['purpose']=['REF001']
        did=self.ws.record_ai_draft(packet,sync_raw(response),'합성요청자');self.ws.apply_ai_draft(did,'합성작성자')
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(exporter.workbook_bytes(self.ws.snapshot(),CATALOG)),data_only=False)
        try:
            cell=next(row[1] for row in book['조서_A04'] if row[0].value=='목적·주장')
            self.assertEqual(cell.value,'=1+1 [REF001]');self.assertEqual(cell.data_type,'s')
        finally:book.close()

    def test_final_needs_ai_approval_and_existing_independent_review(self):
        for module,filename in [('금융상품','통합업무_금융정상.xlsx'),('감가상각','통합업무_감가상각정상.xlsx')]:self.ws.execute(module,ROOT/'samples/통합'/filename)
        self.apply()
        for pid in self.ws.required_procedures():
            if pid=='A04':
                p=dict(self.ws.procedures[pid]);p['performed_date']='2026-10-24';self.ws.save_procedure(p)
            else:self.ws.save_procedure(synthetic_procedure(pid,list(self.ws.sources)))
            self.ws.request_review(pid);self.ws.review(pid,'다른합성검토자','별도 검토 상태 시험')
        with self.assertRaisesRegex(AuditError,'AI 적용'):self.ws.snapshot(True)
        self.approve();self.assertEqual(self.ws.snapshot(True)['export_state'],'선택조서 검토기록완료')
        p=dict(self.ws.procedures['A04']);p['risk']+=' 수정';self.ws.save_procedure(p)
        self.approve()
        with self.assertRaisesRegex(AuditError,'조서 검토'):self.ws.snapshot(True)

    def test_cross_company_response_cannot_be_recorded(self):
        packet=self.packet();other=Workspace(Path(self.temp.name)/'다른업무',create={**CONTEXT,'회사코드':'OTHER'})
        before=other.head
        with self.assertRaises(AuditError):other.record_ai_draft(packet,mock_response(packet),'다른작성자')
        self.assertEqual(other.head,before)

    def test_all_module_and_parent_manual_packets_validate(self):
        import sector_demo
        from ai_provider import validate_packet
        opened=sector_demo.open_company('JS_ELEC',Path(self.temp.name)/'업종보관',ROOT/'sector_resources')
        child=Workspace(opened['folder'])
        pids={'원장·리드':'C04','금융상품':'A02','매출':'A04','감가상각':'A12','확장검사':'A01','리스':'A15','대손충당금':'A07','재고평가':'A11','종업원급여':'A16','미기록부채·비용이연':'A09','원천증빙':'R01','재무제표·주석':'F01','GLA':'G01','FIA':'I01','현금분석':'H01'}
        self.assertEqual(set(child.latest()),set(pids))
        for module,pid in pids.items():
            with self.subTest(module=module):
                packet=child.ai_packet(pid);validate_packet(packet)
                self.assertEqual(packet['procedure']['module'],module)
                self.assertTrue(any(r['kind']=='run_summary' and r.get('run_id')==child.latest()[module]['id'] for r in packet['references']))
                for ref in packet['references']:
                    if ref['kind']=='raw_xlsx_row':self.assertEqual(len(ref['workbook_sha256']),64)
                compact=child.ai_packet(pid,limits={'max_chars':12000,'max_references':40,'rows_per_table':3,'max_tables':8,'cell_chars':300})
                validate_packet(compact);self.assertLessEqual(len(dumps(compact)),12000)
                self.assertLessEqual(len(compact['sources']),2);self.assertLessEqual(len(compact['runs']),2)
                self.assertTrue(any(r['kind']=='raw_xlsx_row' for r in compact['references']))
        parent_result=sector_demo.open_company('JS',Path(self.temp.name)/'업종보관',ROOT/'sector_resources')
        parent=Workspace(parent_result['folder'])
        packet=parent.ai_packet('C02');validate_packet(packet)
        self.assertEqual(packet['engagement']['company_code'],'JS')
        self.assertEqual({r['module'] for r in packet['runs']},{'원장·리드'})
        self.assertIn('투자',dumps(packet));self.assertNotIn('지분율":100',dumps(packet))

if __name__=='__main__':unittest.main()
