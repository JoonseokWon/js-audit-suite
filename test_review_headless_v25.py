"""Headless oracle understands continuation rows but never permits cell/row loss."""
from io import BytesIO
from copy import deepcopy
import unittest
from openpyxl import load_workbook
from export_support.builder import create_builder
from tools.verify_headless_v14 import exported_rows_match


class HeadlessContinuationOracleTests(unittest.TestCase):
    def setUp(self):
        self.expected = [{'식별자': 'SYNTHETIC ONLY A', '금액': 125, '본문': '합성 문장\n' * 40, '위치': ['원본/2', '원본/3']},
                         {'식별자': 'SYNTHETIC ONLY B', '금액': 250, '본문': '짧은 본문', '위치': ['원본/4']}]

    def book(self):
        book, add = create_builder({}, {'금액'}, set(), set())
        add('검증표', self.expected)
        stream = BytesIO(); book.save(stream); book.close(); stream.seek(0)
        result = load_workbook(stream, data_only=False); self.addCleanup(result.close)
        return result

    def test_exact_reconstruction_with_extra_physical_rows_and_numeric_once(self):
        book = self.book(); sheet = book['검증표']
        self.assertGreater(sheet.max_row, len(self.expected) + 1)
        self.assertTrue(exported_rows_match(book, sheet, self.expected))
        header = [cell.value for cell in sheet[1]]; number = header.index('금액') + 1
        self.assertEqual(375, sum(sheet.cell(row, number).value or 0 for row in range(2, sheet.max_row + 1)))
        self.assertFalse(exported_rows_match(book, sheet, self.expected[:-1]))

    def test_missing_or_duplicate_mapping_or_uncovered_row_is_rejected(self):
        book = self.book(); index = book['출력문단연결']; sheet = book['검증표']
        row = [cell.value for cell in index[2]]; index.delete_rows(2)
        self.assertFalse(exported_rows_match(book, sheet, self.expected))
        index.append(row); index.append(row)
        self.assertFalse(exported_rows_match(book, sheet, self.expected))
        index.delete_rows(3); sheet.append(['SYNTHETIC ONLY extra row', None, None, None])
        self.assertFalse(exported_rows_match(book, sheet, self.expected))

    def test_cell_tamper_and_missing_continuation_fragment_are_rejected(self):
        book = self.book(); sheet = book['검증표']; header = [cell.value for cell in sheet[1]]
        text = header.index('본문') + 1; original = sheet.cell(2, text).value
        sheet.cell(2, text).value = 'SYNTHETIC ONLY changed'
        self.assertFalse(exported_rows_match(book, sheet, self.expected))
        sheet.cell(2, text).value = original; sheet.cell(3, text).value = None
        self.assertFalse(exported_rows_match(book, sheet, self.expected))

    def test_numeric_duplicate_on_continuation_is_rejected_even_with_split_label(self):
        book = self.book(); sheet = book['검증표']; header = [cell.value for cell in sheet[1]]
        number = header.index('금액') + 1; sheet.cell(3, number).value = 125
        self.assertFalse(exported_rows_match(book, sheet, self.expected))
        index = book['출력문단연결']; split = [cell.value for cell in index[1]].index('분할열') + 1
        index.cell(2, split).value += ', 금액'
        self.assertFalse(exported_rows_match(book, sheet, self.expected))


if __name__ == '__main__':
    unittest.main()
