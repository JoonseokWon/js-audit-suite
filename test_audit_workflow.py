"""Synthetic, offline workflow counterexamples and immutable lifecycle tests."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from common import AuditError, digest, dumps, sha
from workspace import Workspace
from procedures import task_definitions
from verify import synthetic_procedure
import exporter

ROOT = Path(__file__).resolve().parent
CATALOG = json.loads((ROOT/'procedure_catalog.json').read_text(encoding='utf-8'))
CONTEXT = {'회사코드':'SYNTH01','회사명':'SYNTHETIC ONLY','시작일':'2026-01-01','기준일':'2026-12-31','회계체계':'K-IFRS','감사기준':'합성 검사','작성자':'합성작성자','더미업무':True}


def profile(pids=('C04','C05','P07')):
    return {'audit_framework':'KSA_200_720','audit_framework_reason':'합성 입력: 사람이 감사체계와 적용가능성을 판단함',
            'form_framework':'KIFRS','materiality':{'benchmark_name':'합성 매출','benchmark_amount':100000000,
            'materiality_rate':'0.05','performance_rate':'0.75','trivial_rate':'0.05'},
            'scope':[{'procedure_id':pid,'applicable':True,'reason':'사람의 합성 범위 선택'} for pid in pids]}


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.base=Path(self.tmp.name);self.ws=Workspace(self.base/'업무',create=CONTEXT)
        self.ledger=ROOT/'samples/통합/통합_01_원장정상.xlsx'

    def seed(self):self.ws.execute('원장·리드',self.ledger)
    def configure(self,pids=('C04','C05','P07')):
        self.ws.save_workflow_profile(profile(pids),'사람 계획자')
    def risk(self,pids=('C04','C05','P07'),**changes):
        return {'id':'RISK-1','title':'사람이 식별한 합성 위험','accounts':['현금'],'assertions':['실재성'],
                'procedure_ids':list(pids),'basis':self.ws.data_basis(),'evidence_ids':list(self.ws.sources),
                'nature':'증거 확인','timing':'기말','extent':'사람이 선정한 범위','assessment':'위험 수준은 사람 판단',**changes}
    def link(self,pid='C04',**changes):
        return {'procedure_id':pid,'risk_ids':['RISK-1'],'evidence_ids':list(self.ws.sources),'basis':self.ws.data_basis(),
                'nature':'합성 원본 확인','timing':'기말','extent':'선정 범위','plan':'사람이 계획한 수행',
                'performed':'SYNTHETIC ONLY / 사람이 모의 수행 사실을 입력','performed_date':'2026-12-31','performer':'사람 수행자',**changes}
    def seed_full(self):
        self.seed()
        for module,name in [('금융상품','통합업무_금융정상.xlsx'),('매출','통합업무_매출정상.xlsx'),('감가상각','통합업무_감가상각정상.xlsx')]:self.ws.execute(module,ROOT/'samples/통합'/name)
    def ready(self):
        self.seed_full();pids=sorted(self.ws.required_procedures());self.configure(pids);self.ws.save_workflow_risk(self.risk(pids),'사람 위험평가자')
        for pid in pids:self.ws.save_workflow_procedure(self.link(pid),'사람 계획자')
        for pid in pids:
            self.ws.save_procedure(synthetic_procedure(pid,list(self.ws.sources)))
            self.ws.request_review(pid);self.ws.review(pid,'사람 별도검토자','합성 자료/판단 수동 검토 기록')
    def codes(self):return {r['code'] for r in self.ws.workflow_status()['blockers']}
    def blocked(self,fn):
        before=self.ws.head
        with self.assertRaises(AuditError):fn()
        self.ws.refresh();self.assertEqual(before,self.ws.head)

    def test_profile_records_user_values_and_derived_amounts_without_approval(self):
        saved=self.ws.save_workflow_profile(profile(),'계획자')
        m=saved['materiality'];self.assertEqual((m['overall_materiality'],m['performance_materiality'],m['clearly_trivial_threshold']),(5000000,3750000,250000))
        self.assertFalse(self.ws.procedures);self.assertFalse(self.ws.ai_approvals)
        self.assertIn('WORKPAPER_MISSING',self.codes())
    def test_profile_change_invalidates_ai_approval_without_rewriting_approval(self):
        from test_ai_workpapers import mock_response
        self.seed_full()
        packet=self.ws.ai_packet('A04');did=self.ws.record_ai_draft(packet,mock_response(packet),'사람 요청자')
        self.ws.apply_ai_draft(did,'사람 작성자');approval=self.ws.approve_ai_content('A04','사람 승인자','합성 내용 및 미수행 범위를 직접 확인')
        old_events=deepcopy(self.ws.events)
        self.configure(sorted(self.ws.required_procedures()))
        self.assertFalse(self.ws.ai_status('A04')['approved']);self.assertTrue(self.ws.ai_status('A04')['stale'])
        self.assertEqual(self.ws.ai_approvals['A04'],approval);self.assertEqual(self.ws.events[:len(old_events)],old_events)
        self.blocked(lambda:self.ws.approve_ai_content('A04','사람 승인자','예전근거로 재승인 시도'))
    def test_ai_packet_contains_human_plan_risk_and_explicit_performance_links(self):
        from test_ai_workpapers import mock_response
        self.seed_full();pids=sorted(self.ws.required_procedures());self.configure(pids)
        self.ws.save_workflow_risk(self.risk(pids),'사람 평가자');self.ws.save_workflow_procedure(self.link('A04'),'사람 계획자')
        packet=self.ws.ai_packet('A04');kinds={r['kind'] for r in packet['references']}
        self.assertTrue({'human_workflow_profile','human_workflow_risk','human_workflow_procedure'}<=kinds)
        did=self.ws.record_ai_draft(packet,mock_response(packet),'사람 요청자');self.ws.apply_ai_draft(did,'사람 작성자')
        self.ws.approve_ai_content('A04','사람 승인자','합성 최종내용 직접 확인')
        self.ws.save_workflow_risk(self.risk(pids,assessment='사람 재평가'),'사람 평가자')
        self.assertFalse(self.ws.ai_status('A04')['approved'])
    def test_profile_strict_numeric_inputs_and_zero_trivial(self):
        for field,value in [('benchmark_amount',True),('benchmark_amount',float('nan')),('materiality_rate',float('inf')),('performance_rate',True),('trivial_rate','NaN'),('materiality_rate',0),('performance_rate',0),('trivial_rate',-1)]:
            p=profile();p['materiality'][field]=value
            with self.subTest(field=field,value=value):self.blocked(lambda:self.ws.save_workflow_profile(p,'계획자'))
        p=profile();p['materiality']['trivial_rate']=0
        self.assertEqual(self.ws.save_workflow_profile(p,'계획자')['materiality']['clearly_trivial_threshold'],0)
    def test_profile_no_automatic_framework_selection_and_required_reason(self):
        for changes in ({'audit_framework':'auto'},{'audit_framework_reason':''},{'form_framework':'KSA_1200'}):
            p=profile();p.update(changes);self.blocked(lambda:self.ws.save_workflow_profile(p,'계획자'))
        p=profile();p['audit_framework']='KSA_1200';self.ws.save_workflow_profile(p,'사람 선택자')
        self.assertEqual(self.ws.workflow_profile['audit_framework'],'KSA_1200')
    def test_form_series_is_independent_from_engagement_accounting_and_audit_frameworks(self):
        original=deepcopy(self.ws.engagement);chosen=profile();chosen['form_framework']='소규모기업'
        chosen['audit_framework']='KSA_200_720'
        saved=self.ws.save_workflow_profile(chosen,'사람 선택자')
        self.assertEqual(saved['form_framework'],'소규모기업');self.assertNotIn('accounting_framework',saved)
        self.assertEqual(self.ws.engagement,original);self.assertEqual(self.ws.engagement['회계체계'],'K-IFRS')
        snap=self.ws.snapshot();self.assertEqual(snap['engagement']['회계체계'],'K-IFRS')
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(exporter.workbook_bytes(snap,CATALOG)),read_only=True)
        try:
            plan=dict(book['감사계획·중요성'].iter_rows(min_row=2,values_only=True))
            self.assertEqual(plan['실제 회계체계(업무 기본정보)'],'K-IFRS')
            self.assertEqual(plan['서식계열'],'소규모기업');self.assertEqual(plan['적용 감사체계'],'KSA_200_720')
            self.assertIn('적용 판단을 대신하지',plan['서식 참고의 범위'])
        finally:book.close()
    def test_obsolete_accounting_framework_profile_key_is_rejected(self):
        chosen=profile();chosen['accounting_framework']=chosen.pop('form_framework')
        self.blocked(lambda:self.ws.save_workflow_profile(chosen,'사람 선택자'))
    def test_duplicate_unknown_oversized_and_boolean_scope_inputs_rejected(self):
        self.blocked(lambda:self.ws.save_workflow_profile('{"audit_framework":"a","audit_framework":"b"}','사람'))
        self.blocked(lambda:self.ws.save_workflow_profile('{"x":"'+'a'*200001+'"}','사람'))
        for mutate in (lambda p:p.update(unknown='x'),lambda p:p['scope'].append(deepcopy(p['scope'][0])),lambda p:p['scope'][0].update(procedure_id='OTHER-WORKSPACE'),lambda p:p['scope'][0].update(applicable=1)):
            p=profile();mutate(p);self.blocked(lambda:self.ws.save_workflow_profile(p,'사람'))
    def test_risk_rejects_foreign_duplicate_and_stale_links(self):
        self.seed();self.configure()
        for changes in ({'procedure_ids':['F01']},{'evidence_ids':['foreign']},{'evidence_ids':list(self.ws.sources)*2},{'basis':'not-current'},{'accounts':['현금','현금']}):
            self.blocked(lambda:self.ws.save_workflow_risk(self.risk(**changes),'위험 평가자'))
    def test_link_rejects_orphans_missing_actual_fields_and_unregistered_sources(self):
        self.seed();self.configure();self.ws.save_workflow_risk(self.risk(),'위험 평가자')
        for changes in ({'risk_ids':['foreign']},{'risk_ids':['RISK-1','RISK-1']},{'procedure_id':'F01'},{'performed_date':''},{'performer':''},{'evidence_ids':[]},{'basis':'stale'}):
            self.blocked(lambda:self.ws.save_workflow_procedure(self.link(**changes),'계획자'))
    def test_plan_only_never_records_performance_or_completion(self):
        self.seed();self.configure();self.ws.save_workflow_risk(self.risk(),'위험 평가자')
        saved=self.ws.save_workflow_procedure(self.link(performed='',performed_date='',performer=''),'계획자')
        self.assertFalse(saved['performed']);self.assertIn('PROCEDURE_NOT_PERFORMED',self.codes())
        self.assertTrue(self.ws.workflow_status()['human_pending']);self.assertFalse(self.ws.workflow_locked)
    def test_profile_and_risk_changes_stale_review_and_keep_original_events(self):
        self.ready();old=deepcopy(self.ws.events);before=self.ws.basis()
        self.ws.save_workflow_risk(self.risk(assessment='사람이 변경한 평가'),'평가자')
        self.assertNotEqual(before,self.ws.basis());self.assertEqual(self.ws.events[:len(old)],old)
        self.assertTrue(all(self.ws.procedure_state(p)=='자료변경·재검토' for p in self.ws.procedures.values()))
        self.assertIn('WORKPAPER_REVIEW_PENDING',self.codes());self.blocked(lambda:self.ws.snapshot(True))
    def test_retiring_a_changed_plan_cannot_reactivate_an_old_review_basis(self):
        self.ready();before=self.ws.basis()
        self.ws.save_workflow_risk(self.risk(sorted(self.ws.required_procedures()),id='RISK-2'),'사람 평가자')
        self.ws.retire_workflow_risk('RISK-2','사람 평가자','추가 평가 기록 사용 해제')
        self.assertNotEqual(self.ws.basis(),before)
        self.assertTrue(all(self.ws.procedure_state(p)=='자료변경·재검토' for p in self.ws.procedures.values()))
    def test_data_change_blocks_plans_risks_and_review(self):
        self.ready();self.ws.add_source(self.ledger,'추가 자료')
        self.assertTrue({'RISK_STALE','PROCEDURE_PLAN_STALE','WORKPAPER_REVIEW_PENDING'}<=self.codes())
    def test_original_missing_blocks_status_and_new_link(self):
        self.seed();self.configure();self.ws.save_workflow_risk(self.risk(),'평가자')
        source=next(iter(self.ws.sources.values()));(self.ws.root/source['file']).unlink()
        self.assertIn('SOURCE_INTEGRITY',self.codes());self.blocked(lambda:self.ws.save_workflow_procedure(self.link(),'계획자'))
    def test_scope_change_creates_explicit_orphan_until_human_retires_links(self):
        self.seed();self.configure(('F01',));self.ws.save_workflow_risk(self.risk(('F01',)),'평가자');self.ws.save_workflow_procedure(self.link('F01'),'계획자')
        self.configure(('C04','C05','P07'))
        self.assertTrue({'RISK_PROCEDURE_ORPHAN','PROCEDURE_LINK_ORPHAN'}<=self.codes())
        old=deepcopy(self.ws.events)
        self.ws.retire_workflow_risk('RISK-1','평가자','사람이 변경한 범위');self.ws.retire_workflow_procedure('F01','계획자','범위 변경')
        self.assertEqual(self.ws.events[:len(old)],old);self.assertFalse(self.ws.workflow_risks);self.assertFalse(self.ws.workflow_procedure_links)
    def test_unexecuted_selected_module_is_a_blocker(self):
        self.seed();self.configure(('A02',));self.assertIn('MODULE_NOT_EXECUTED',self.codes())
    def test_final_snapshot_and_lock_require_all_current_human_records(self):
        self.seed();self.configure();self.blocked(lambda:self.ws.finalize_workflow('담당자','미완료 잠금 시도'))
        self.assertFalse(self.ws.workflow_locked)
    def test_successful_lock_blocks_all_mutations_and_reopen_preserves_export(self):
        self.ready();self.assertFalse(self.ws.workflow_status()['blockers'])
        self.ws.finalize_workflow('사람 잠금자','선택 범위 및 증거 검토 후 보관')
        self.assertTrue(self.ws.workflow_locked);self.ws.snapshot(True)
        output=self.base/'보관.xlsx';exporter.export(self.ws,output,CATALOG,True);old_sha=sha(output)
        files=set((self.ws.root/'원본').iterdir())
        for fn in (lambda:self.ws.add_source(self.ledger,'새 자료'),lambda:self.ws.execute('원장·리드',self.ledger),lambda:self.ws.add_note('제목','내용','기록자'),lambda:self.ws.save_workflow_profile(profile(),'계획자'),lambda:self.ws.save_procedure(synthetic_procedure('C04',list(self.ws.sources))),lambda:self.ws.review('C04','검토자','기록'),lambda:self.ws._append('fake',{})):
            self.blocked(fn)
        self.assertEqual(files,set((self.ws.root/'원본').iterdir()))
        self.blocked(lambda:self.ws.reopen_workflow('사람',''))
        self.ws.reopen_workflow('사람 재개자','추가 증거를 확인하기 위해 재개')
        self.assertFalse(self.ws.workflow_locked);self.ws.add_note('추가 질문','사람 검토 필요','담당자')
        self.assertEqual(sha(output),old_sha)
    def test_workflow_apis_reject_stale_explicit_heads_before_saving(self):
        self.ready();stale=self.ws.head;Workspace(self.ws.root).add_note('다른 화면','추가 정보','다른 담당자')
        actions=(lambda:self.ws.save_workflow_profile(profile(),'계획자',expected_head=stale),
                 lambda:self.ws.save_workflow_risk(self.risk(),'평가자',expected_head=stale),
                 lambda:self.ws.save_workflow_procedure(self.link(),'계획자',expected_head=stale),
                 lambda:self.ws.retire_workflow_risk('RISK-1','평가자','사용 해제',expected_head=stale),
                 lambda:self.ws.retire_workflow_procedure('C04','계획자','사용 해제',expected_head=stale),
                 lambda:self.ws.finalize_workflow('담당자','보관',expected_head=stale))
        for action in actions:
            with self.assertRaisesRegex(AuditError,'이력'):action()
    def test_workflow_apis_cas_without_explicit_head_reject_between_validation_and_append(self):
        self.ready()
        actions=(lambda:self.ws.save_workflow_profile(profile(),'계획자'),
                 lambda:self.ws.save_workflow_risk(self.risk(),'평가자'),
                 lambda:self.ws.save_workflow_procedure(self.link(),'계획자'),
                 lambda:self.ws.retire_workflow_risk('RISK-1','평가자','사용 해제'),
                 lambda:self.ws.retire_workflow_procedure('C04','계획자','사용 해제'))
        for action in actions:
            self.ws.refresh();count=len(self.ws.events);original=self.ws._append
            def raced(kind,data,expected_head=None):
                Workspace(self.ws.root).add_note('동시 변경','검증과 저장 사이 새 정보','다른 담당자')
                return original(kind,data,expected_head=expected_head)
            with patch.object(self.ws,'_append',side_effect=raced):
                with self.assertRaisesRegex(AuditError,'이력'):action()
            self.ws.refresh();self.assertEqual(len(self.ws.events),count+1);self.assertEqual(self.ws.events[-1]['kind'],'note')
    def test_reopen_checks_explicit_head_and_internal_compare_and_swap(self):
        self.ready();self.ws.finalize_workflow('담당자','보관');stale=self.ws.head
        Workspace(self.ws.root).reopen_workflow('다른 담당자','동시 재개')
        with self.assertRaisesRegex(AuditError,'이력'):self.ws.reopen_workflow('담당자','늦은 화면 재개',expected_head=stale)
        self.assertEqual(self.ws.events[-1]['data']['actor'],'다른 담당자')
    def test_finalization_checks_head_concurrency(self):
        self.ready();original=self.ws.snapshot
        def concurrent(final=False):
            snap=original(final);other=Workspace(self.ws.root);other.add_note('동시 변경','새 질문','다른 담당자');return snap
        with patch.object(self.ws,'snapshot',side_effect=concurrent):self.blocked_no_head(lambda:self.ws.finalize_workflow('담당자','동시 변경 차단'))
        self.ws.refresh();self.assertFalse(self.ws.workflow_locked)
    def blocked_no_head(self,fn):
        with self.assertRaises(AuditError):fn()
    def test_export_contains_workflow_crosswalk_and_complete_event_json(self):
        self.ready();snap=self.ws.snapshot();data=exporter.workbook_bytes(snap,CATALOG)
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(data),data_only=False)
        try:
            self.assertTrue({'감사계획·중요성','위험주장대응','절차계획·수행연결','기준서식참고','완료점검표'}<=set(book.sheetnames))
            proof=json.loads(''.join(row[1] for row in book['증적JSON'].iter_rows(min_row=2,values_only=True)))
            self.assertEqual(proof['events'],self.ws.events);self.assertEqual(proof['workflow']['profile'],self.ws.workflow_profile)
            self.assertTrue(all(c.data_type!='f' for sheet in book for row in sheet for c in row))
        finally:book.close()
    def test_dynamic_exception_and_reconciliation_ids_export_with_explicit_unmapped_form_rows(self):
        self.seed_full()
        latest=self.ws.latest()['금융상품'];changed=deepcopy(latest['result'])
        changed['metrics']['원금']-=1
        changed['issues'].append({'code':'SYNTHETIC_REVIEW','severity':'검토','message':'합성 수동 조사','reference':'SYNTHETIC ONLY'})
        self.ws._append('run',{**latest,'id':'synthetic-exception-run','result':changed,'result_hash':digest(changed)})
        self.ws.refresh();self.configure(sorted(self.ws.required_procedures()))
        snap=self.ws.snapshot();rows={r['절차ID']:r for r in snap['form_references']}
        dynamic=[pid for pid in self.ws.required_procedures() if pid.startswith(('X-','B-'))]
        self.assertTrue(any(pid.startswith('X-') for pid in dynamic));self.assertTrue(any(pid.startswith('B-') for pid in dynamic))
        self.assertTrue(all(rows[pid]['연결상태']=='공식 참고 미지정' and not rows[pid]['공식서식ID'] for pid in dynamic))
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(exporter.workbook_bytes(snap,CATALOG)),read_only=True)
        try:self.assertEqual(book['기준서식참고'].max_row,len(snap['form_references'])+1)
        finally:book.close()
    def test_export_cannot_overwrite_registered_external_original(self):
        original=self.base/'수령원본.xlsx';original.write_bytes(self.ledger.read_bytes());self.ws.add_source(original,'수령 원본')
        before=original.read_bytes()
        with self.assertRaises(AuditError):exporter.export(self.ws,original,CATALOG)
        self.assertEqual(original.read_bytes(),before)
    def test_legacy_opt_out_keeps_original_basis_and_final_rules(self):
        self.seed_full()
        for pid in self.ws.required_procedures():
            self.ws.save_procedure(synthetic_procedure(pid,list(self.ws.sources)));self.ws.request_review(pid);self.ws.review(pid,'사람 별도검토자','합성 검토')
        self.assertEqual(self.ws.basis(),self.ws.data_basis());self.assertFalse(self.ws.workflow_status()['active']);self.ws.snapshot(True)


class LegacyF01Tests(unittest.TestCase):
    setUp=WorkflowTests.setUp
    seed=WorkflowTests.seed
    codes=WorkflowTests.codes
    blocked=WorkflowTests.blocked
    def old_records(self):
        p=synthetic_procedure('F01',[]);p.pop('purpose_code');p.update(basis=self.ws.basis(),state='검토기록완료',reviewer='과거검토자',review_note='과거 원문',reviewed_at='2026-01-01',edit_revision='old')
        self.ws._append('procedure',p)
        approval={'procedure_id':'F01','draft_id':'old-draft','basis':self.ws.basis(),'actor':'과거승인자','note':'과거승인 원문'}
        self.ws._append('ai_content_approval',approval)
        self.ws._append('ai_draft',{'id':'old-draft','procedure_id':'F01','basis':self.ws.basis(),'packet':{'procedure':{'id':'F01'}}})
        self.ws.refresh();return deepcopy(self.ws.events),p,approval
    def test_ambiguous_old_f01_is_never_inferred_from_title_or_runs(self):
        events,p,approval=self.old_records();self.assertIn('LEGACY_F01_UNCLASSIFIED',self.codes())
        for fn in (lambda:self.ws.ai_packet('F01'),lambda:self.ws.request_review('F01'),lambda:self.ws.approve_ai_content('F01','사람','확인'),lambda:self.ws.snapshot(True)):
            self.blocked(fn)
        self.assertEqual(self.ws.events,events)
    def test_classification_keeps_raw_documents_and_approvals_and_creates_independent_ids(self):
        events,old,approval=self.old_records()
        self.ws.classify_legacy_f01('financial_statement_reconciliation','사람 분류자','옛 문서와 근거를 직접 확인')
        self.assertEqual(self.ws.events[:len(events)],events);self.assertEqual(self.ws.legacy_f01_documents,[old]);self.assertEqual(self.ws.legacy_f01_approvals,[approval])
        self.assertNotIn('F01',self.ws.procedures);self.assertNotIn('F01',self.ws.ai_approvals)
        for pid in ('F01','FS01'):self.ws.save_procedure(synthetic_procedure(pid,[]))
        self.assertEqual(self.ws.procedures['F01']['purpose_code'],'subsequent_events');self.assertEqual(self.ws.procedures['FS01']['purpose_code'],'financial_statement_reconciliation')
        self.assertEqual(self.ws.legacy_f01_documents,[old]);self.assertFalse(self.ws.ai_status('F01')['required'])
        self.blocked(lambda:self.ws.apply_ai_draft('old-draft','사람'))
        self.blocked(lambda:self.ws.classify_legacy_f01('subsequent_events','사람','변경 시도'))
    def test_legacy_classification_cas_preserves_other_window_event(self):
        self.old_records();before=self.ws.head;Workspace(self.ws.root).add_note('다른 화면','원본 목적 확인 중','담당자')
        with self.assertRaisesRegex(AuditError,'이력'):self.ws.classify_legacy_f01('financial_statement_reconciliation','사람','확인',expected_head=before)
        original=self.ws._append;count=len(self.ws.events)
        def raced(kind,data,expected_head=None):
            Workspace(self.ws.root).add_note('다른 화면','새 근거','담당자')
            return original(kind,data,expected_head=expected_head)
        with patch.object(self.ws,'_append',side_effect=raced):
            with self.assertRaisesRegex(AuditError,'이력'):self.ws.classify_legacy_f01('financial_statement_reconciliation','사람','확인')
        self.ws.refresh();self.assertIsNone(self.ws.workflow_classification);self.assertEqual(len(self.ws.events),count+1)
    def test_f01_and_fs01_definitions_and_packets_have_distinct_semantics(self):
        self.seed();self.ws.execute('재무제표·주석',ROOT/'samples/통합/통합_v1.4_재무제표.json')
        tasks=task_definitions(self.ws,CATALOG);self.assertEqual(len([p for p in tasks if p['id']=='F01']),1)
        packet=self.ws.ai_packet('F01');fs=self.ws.ai_packet('FS01')
        self.assertIsNone(packet['procedure']['module']);self.assertEqual(packet['procedure']['purpose_code'],'subsequent_events')
        self.assertEqual(fs['procedure']['module'],'재무제표·주석');self.assertEqual(fs['procedure']['purpose_code'],'financial_statement_reconciliation')
        self.assertNotIn('재무제표·주석',{r['module'] for r in packet['runs']})
    def test_duplicate_procedure_definitions_are_rejected_instead_of_overwritten(self):
        catalog=deepcopy(CATALOG);catalog['modules'].append(deepcopy(catalog['modules'][0]))
        with self.assertRaises(AuditError):task_definitions(self.ws,catalog)
    def test_missing_or_wrong_explicit_purpose_cannot_bind_new_f01_workpaper(self):
        p=synthetic_procedure('F01',[]);p.pop('purpose_code');self.blocked(lambda:self.ws.save_procedure(p))
        p['purpose_code']='financial_statement_reconciliation';self.blocked(lambda:self.ws.save_procedure(p))


if __name__=='__main__':unittest.main()
